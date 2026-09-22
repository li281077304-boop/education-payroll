from __future__ import annotations

import csv
import copy
import hashlib
import io
import json
import math
import re
import shutil
import uuid
from dataclasses import asdict, replace
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from payroll_core.excel.check_workbook import read_check_workbook_schedule
from payroll_core.excel.package import discover_payroll_package
from payroll_core.period import coverage_for, dominant_month, month_from_filename, normalize_period_window, previous_period
from payroll_core.mapping import SCHEDULE_AC_REQUIREMENT, analyze_mapping, resolve_schedule_import
from payroll_core.mapping.schedule import read_schedule_with_mapping
from payroll_core.models.evidence import AdapterIssue
from payroll_core.models.class_type_rules import (
    SPECIAL_RULES_KEY,
    default_rule_versions,
    normalize_special,
    rule_version_for_period,
    structured_rules,
)
from payroll_core.payroll_generation import build_legacy_generated_payroll, generated_from_calculation
from payroll_core.final_fields import renewal_snapshot_entry, renewal_fields_from_snapshot
from payroll_core.excel.standard_payroll_render import is_payroll_template, render_generated_payroll
from payroll_core.excel.output_paths import default_output_dir, describe_location, safe_output_path
from payroll_core.excel.inspect import inspect_workbook
from payroll_core.excel.payroll import read_payroll_excel
from payroll_core.excel.schedule import read_schedule_excel
from payroll_core.adapters.payroll_sheet import read_payroll_csv
from payroll_core.adapters.schedule import read_schedule_csv_result
from payroll_core.excel.common import grade_from_class_name, load_student_grade_lookup
from payroll_core.adapters import read_refund_report, read_renewal_report
from payroll_core.grade_inference import (
    CourseExportSnapshot,
    StudentGradeEvidence,
    course_export_grade_override,
    evidence_identity,
    export_timestamp_from_source,
    normalize_lesson_start_time,
    remove_export_pollution,
    split_student_names,
)
from payroll_core.formula_audit import audit_payroll_formulas
from payroll_core.rules.authority import TeacherCompensationProfile, TeacherRating, default_compensation_bands, policy_fee_checks, rating_and_rate_checks
from payroll_core.reconcile.payroll_scope import HEADCOUNT_COEFFICIENTS, LESSON_HOUR_FACTOR, SMALL_GROUP_CLASS_TYPES, FieldCheck, class_value_contribution, normalize_class_rules, schedule_field_checks, total_salary_read_checks
from payroll_core.reconcile.ac_resolution import (
    ApprovedPayrollOverride,
    SourceDataCorrection,
    assess_counterfactual_cause,
    resolve_ac,
    schedule_record_id,
)
from payroll_core.excel.reconciliation_bridge import GRADE_COEFFICIENTS
from payroll_core.comments import render_comment
from payroll_core.excel.writeback import preview as preview_writeback, sha256 as workbook_sha256, write_new_workbook
from payroll_core.calculation import course_record_key
from payroll_ui.assessment_flow import AssessmentService
from payroll_ui.submissions import PayrollSubmissionService

from .storage import RunStore
from .business import build_groups, build_user_actions, invalidate as invalidate_business_decisions
from .business_inputs import BusinessInputService, is_active_business_input
from .core_flow import CoreFlow, valid_period

REQUIRED = ("schedule",)
SCOPE_ROLES = ("math", "science")
#: 内部模式名。界面上只说“核对一份工资表 / 直接生成工资表”。
MODE_AUDIT = "AUDIT"
MODE_GENERATE = "GENERATE"

# Field-check states are deliberately classified by meaning rather than by
# subtracting an ever-growing list from ``manual_review``.  In particular,
# GENERATE mode uses DETERMINED/NOT_APPLICABLE for facts that are complete even
# though they do not compare against a pre-existing payroll workbook.
COMPLETE_STATUSES = frozenset({
    "MATCH", "FORMULA_MATCH", "RATE_MATCH", "AF_POLICY_MATCH",
    "READ_ONLY", "EXPLAINED_DIFFERENCE", "DETERMINED", "NOT_APPLICABLE",
})
MANUAL_REVIEW_STATUSES = frozenset({
    "UNEXPLAINED_DIFFERENCE", "RATING_MISMATCH", "RATE_MISMATCH",
    "AF_POLICY_MISMATCH", "FORMULA_MISSING", "FORMULA_DIFFERENCE",
    "FORMULA_PATTERN_MISMATCH", "FORMULA_REGION_BREAK",
    "FORMULA_REPLACED_BY_VALUE", "ROW_REFERENCE_SHIFT", "MISSING_SOURCE",
    "MISSING_TARGET", "MISSING_PAYROLL_VALUE", "NEEDS_MANUAL_REVIEW",
    "NEEDS_INPUT", "NEEDS_CONFIRMATION", "NEEDS_RECONFIRMATION",
    "CONFLICT_NEEDS_CONFIRMATION", "GRADE_UNRESOLVED", "RULE_NOT_FOUND",
    "MULTIPLE_RULES_MATCHED", "MISSING_AUTHORITY", "IDENTITY_NOT_STABLE", "NOT_PROVIDED",
})


def _day_before(day: str) -> str:
    from datetime import date, timedelta

    try:
        return (date.fromisoformat(day) - timedelta(days=1)).isoformat()
    except ValueError:
        return day
LABELS = {"schedule": "原始排课数据", "math": "数学组提交表", "science": "理化组提交表", "baseline": "基准最终工资表", "check": "最终工资核对表"}
LAYOUTS = {"schedule": "SCHEDULE_EXPORT_V1", "math": "PAYROLL_SHEET_V1", "science": "PAYROLL_SHEET_V1", "baseline": "PAYROLL_SHEET_V1", "check": "PAYROLL_CHECK_V1"}
STATUS = {"DRAFT": "待导入", "FILES_READY": "材料已准备", "CHECKING": "正在核对", "REVIEW_REQUIRED": "需要复核", "STALE": "文件已变化", "PASS": "字段核对完成"}
ISSUE_LABELS = {
    "MATCH": "一致",
    "EXPLAINED_DIFFERENCE": "已确认",
    "UNEXPLAINED_DIFFERENCE": "差异待处理",
    "MISSING_SOURCE": "缺少原始依据",
    "MISSING_TARGET": "工资表缺少人员",
    "NEEDS_MANUAL_REVIEW": "需要人工确认",
    "FORMULA_DIFFERENCE": "公式复算不一致",
    "RATING_MISMATCH": "星级不一致",
    "RATE_MATCH": "档位金额一致",
    "RATE_MISMATCH": "档位金额不一致",
    "MISSING_AUTHORITY": "缺少权威资料",
    "IDENTITY_NOT_STABLE": "教师身份无法稳定绑定",
    "MISSING_PAYROLL_VALUE": "工资表缺少值",
    "RULE_NOT_FOUND": "未找到适用规则",
    "MULTIPLE_RULES_MATCHED": "规则冲突",
    "FORMULA_REPLACED_BY_VALUE": "公式被固定值替代",
    "FORMULA_MISSING": "公式缺失",
    "FORMULA_PATTERN_MISMATCH": "公式结构不同",
    "ROW_REFERENCE_SHIFT": "公式引用错行",
    "COLUMN_REFERENCE_SHIFT": "公式引用错列",
    "FORMULA_REGION_BREAK": "公式区域异常",
}
ACTION_LABELS = {"special": "已确认特殊情况", "payroll_error": "工资表待修改", "defer": "暂时保留", "confirm_source": "来源待核实"}

# This is a read-only projection of the existing final-payroll contract.  It
# deliberately names every component in the historical template and records
# whether the current repository has a source adapter/rule for it.  It is not
# a second calculation engine; the values and statuses below are only joined
# with the current Run's durable final_fields evidence by av_source_map().
AV_SOURCE_DEFINITIONS: dict[str, dict[str, Any]] = {
    "M": {"business_name": "实际基本工资", "relationship": "(G + H + I + J) / K × L", "source": "真实工资模板与 July/August 最终工资表", "source_type": "PAYROLL_TEMPLATE_FORMULA", "category": "SOURCE_AVAILABLE_NOT_CONNECTED", "authoritative_source": "薪资表模板.xlsx、July/August final payroll（M5 公式）", "can_auto_calculate": True, "needs_manual_input": False, "needs_business_rule": False},
    "AF": {"business_name": "总课时费", "relationship": "MAX(0, (AD - 义务课时) × AE)", "source": "Core 排课计算 + Run 级义务课时政策", "source_type": "CORE_CALCULATION", "category": "DONE", "authoritative_source": "core_calculation / af_policy_confirmation", "can_auto_calculate": True, "needs_manual_input": False, "needs_business_rule": False},
    "AG": {"business_name": "领航伴学课时费", "relationship": "模板列，当前没有已定义计算关系", "source": "未发现明确业务来源或规则", "source_type": "UNKNOWN", "category": "RULE_NOT_DEFINED", "authoritative_source": "NO EVIDENCE", "can_auto_calculate": False, "needs_manual_input": False, "needs_business_rule": True},
    "AH": {"business_name": "续费一对一课时", "relationship": "取已审核 RENEWAL_RESULT 的 one_to_one_hours", "source": "已审核续费最终结果", "source_type": "RENEWAL_RESULT", "category": "SOURCE_MISSING", "authoritative_source": "approved business input bound to Run", "can_auto_calculate": True, "needs_manual_input": True, "needs_business_rule": False},
    "AI": {"business_name": "续费班课课时", "relationship": "取已审核 RENEWAL_RESULT 的 class_hours", "source": "已审核续费最终结果", "source_type": "RENEWAL_RESULT", "category": "SOURCE_MISSING", "authoritative_source": "approved business input bound to Run", "can_auto_calculate": True, "needs_manual_input": True, "needs_business_rule": False},
    "AJ": {"business_name": "领航续费课时", "relationship": "取已审核 RENEWAL_RESULT 的 mentor_hours", "source": "已审核续费最终结果", "source_type": "RENEWAL_RESULT", "category": "SOURCE_MISSING", "authoritative_source": "approved business input bound to Run", "can_auto_calculate": True, "needs_manual_input": True, "needs_business_rule": False},
    "AK": {"business_name": "推荐续费奖", "relationship": "AH × 1 + AI × 1.5 + AJ × 0.75", "source": "AH/AI/AJ 的已审核续费结果", "source_type": "RENEWAL_RESULT", "category": "SOURCE_AVAILABLE_NOT_CONNECTED", "authoritative_source": "approved renewal result + existing AK adapter", "can_auto_calculate": True, "needs_manual_input": True, "needs_business_rule": False},
    "AL": {"business_name": "进步率奖金", "relationship": "模板列，当前没有已定义计算关系", "source": "未发现明确业务来源或规则", "source_type": "UNKNOWN", "category": "RULE_NOT_DEFINED", "authoritative_source": "NO EVIDENCE", "can_auto_calculate": False, "needs_manual_input": False, "needs_business_rule": True},
    "AM": {"business_name": "管理团队奖", "relationship": "模板列；管理考核与金额规则尚未统一", "source": "管理岗位考核资料存在入口，但没有 AM 金额权威规则", "source_type": "ASSESSMENT", "category": "RULE_NOT_DEFINED", "authoritative_source": "management assessment source + explicit amount rule required", "can_auto_calculate": False, "needs_manual_input": True, "needs_business_rule": True},
    "AN": {"business_name": "退费/拒收学员", "relationship": "汇总已审核退费结果的人头金额与业绩金额", "source": "已审核退费最终结果", "source_type": "REFUND_RESULT", "category": "SOURCE_MISSING", "authoritative_source": "approved refund result bound to Run", "can_auto_calculate": True, "needs_manual_input": True, "needs_business_rule": False},
    "AO": {"business_name": "房租", "relationship": "模板列，需明确指向 AO 的已审核金额", "source": "OTHER 业务输入（target_field=AO）适配器", "source_type": "OTHER", "category": "SOURCE_AVAILABLE_NOT_CONNECTED", "authoritative_source": "approved explicit-field business input", "can_auto_calculate": True, "needs_manual_input": True, "needs_business_rule": False},
    "AP": {"business_name": "社保", "relationship": "模板列，需明确指向 AP 的已审核金额", "source": "OTHER 业务输入（target_field=AP）适配器", "source_type": "OTHER", "category": "SOURCE_AVAILABLE_NOT_CONNECTED", "authoritative_source": "approved explicit-field business input", "can_auto_calculate": True, "needs_manual_input": True, "needs_business_rule": False},
    "AQ": {"business_name": "工装费", "relationship": "模板列，需明确指向 AQ 的已审核金额", "source": "OTHER 业务输入（target_field=AQ）适配器", "source_type": "OTHER", "category": "SOURCE_AVAILABLE_NOT_CONNECTED", "authoritative_source": "approved explicit-field business input", "can_auto_calculate": True, "needs_manual_input": True, "needs_business_rule": False},
    "AR": {"business_name": "内部推荐奖金", "relationship": "模板列，需明确指向 AR 的已审核金额", "source": "OTHER 业务输入（target_field=AR）适配器", "source_type": "OTHER", "category": "SOURCE_AVAILABLE_NOT_CONNECTED", "authoritative_source": "approved explicit-field business input", "can_auto_calculate": True, "needs_manual_input": True, "needs_business_rule": False},
    "AS": {"business_name": "月度激励", "relationship": "模板列，需明确标注 AS 的已审核金额", "source": "OTHER 业务输入（target_field=AS）适配器", "source_type": "OTHER", "category": "SOURCE_AVAILABLE_NOT_CONNECTED", "authoritative_source": "approved explicit-field business input with effective period", "can_auto_calculate": True, "needs_manual_input": True, "needs_business_rule": False},
    "AT": {"business_name": "补发工资", "relationship": "模板列，需明确指向 AT 的已审核金额", "source": "OTHER 业务输入（target_field=AT）适配器", "source_type": "OTHER", "category": "SOURCE_AVAILABLE_NOT_CONNECTED", "authoritative_source": "approved explicit-field business input", "can_auto_calculate": True, "needs_manual_input": True, "needs_business_rule": False},
    "AU": {"business_name": "考勤罚款", "relationship": "模板列，需明确指向 AU 的已审核金额", "source": "OTHER 业务输入（target_field=AU）适配器", "source_type": "OTHER", "category": "SOURCE_AVAILABLE_NOT_CONNECTED", "authoritative_source": "approved explicit-field business input", "can_auto_calculate": True, "needs_manual_input": True, "needs_business_rule": False},
    "AV": {"business_name": "总工资", "relationship": "M + AF + AG + AK + AL + AM + AN + AO + AP + AQ + AR + AS + AT + AU", "source": "模板公式汇总 14 个直接组成项", "source_type": "DERIVED_OUTPUT", "category": "DERIVED_OUTPUT", "authoritative_source": "template AV formula + 14 direct component sources", "can_auto_calculate": True, "needs_manual_input": False, "needs_business_rule": False},
}


def version(path: Path) -> dict[str, Any]:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    stat = path.stat()
    return {"sha256": digest.hexdigest(), "size": stat.st_size, "mtime_ns": stat.st_mtime_ns}


def _template_av_evidence(path: str | Path | None) -> dict[str, Any]:
    """Read AV-column evidence from the actual workbook without mutating it."""
    if not path:
        return {"path": None, "exists": False, "formula": None, "field_formulas": {}, "historical_evidence": "NO TEMPLATE BOUND"}
    template = Path(path)
    if not template.is_file():
        return {"path": str(template), "exists": False, "formula": None, "field_formulas": {}, "historical_evidence": "TEMPLATE NOT FOUND"}
    try:
        book = load_workbook(template, data_only=False, read_only=False, keep_links=True)
        sheet = book.worksheets[0]
        columns: dict[str, int] = {}
        aliases = {
            "M": ("实际基本工资",),
            "AF": ("总课时费",), "AG": ("领航伴学课时费",), "AH": ("1对1课时",),
            "AI": ("班课&1对2领航伴课次",), "AJ": ("小班领航伴学课次",), "AK": ("推荐续费奖",),
            "AL": ("进步率奖金",), "AM": ("管理团队奖",), "AN": ("退费/拒收学员",),
            "AO": ("房租",), "AP": ("社保",), "AQ": ("工装费",), "AR": ("内部推荐奖金",),
            "AS": ("月度激励",), "AT": ("补发工资",), "AU": ("考勤罚款",), "AV": ("总工资数",),
        }
        for column in range(1, sheet.max_column + 1):
            values = {str(sheet.cell(row, column).value or "").replace("\n", "").strip() for row in (3, 4)}
            for code, labels in aliases.items():
                if code not in columns and values.intersection(labels):
                    columns[code] = column
        formulas: dict[str, str | None] = {}
        for code, column in columns.items():
            value = sheet.cell(5, column).value if sheet.max_row >= 5 else None
            formulas[code] = value if isinstance(value, str) and value.startswith("=") else None
        formula = formulas.get("AV")
        return {
            "path": str(template), "exists": True, "sha256": version(template)["sha256"],
            "formula": formula, "field_formulas": formulas,
            "historical_evidence": f"真实工资模板 {template.name}；原始列位与公式按 data_only=False 读取。",
        }
    except Exception as exc:  # pragma: no cover - defensive read-only projection
        return {"path": str(template), "exists": True, "formula": None, "field_formulas": {}, "historical_evidence": f"模板读取失败：{exc}"}


def safe_csv(value: Any) -> Any:
    if not isinstance(value, str):
        return value
    return "'" + value if value.lstrip(" \t\r\n").startswith(("=", "+", "-", "@")) else value


class PayrollService(CoreFlow):
    def __init__(self, root: Path):
        self.root = Path(root)
        self.store = RunStore(self.root)
        self.inputs = BusinessInputService(self.store)
        # Teacher payroll sheets and group-leader assessments are two separate
        # flows on purpose: a personal payroll sheet is never an assessment.
        self.submissions = PayrollSubmissionService(self.store)
        self.assessments = AssessmentService(self.store)

    def import_payroll_sheets(self, paths: list[str], period: str, submitted_by: str, *, default_teacher: str = "") -> dict:
        return self.submissions.import_sheets(paths, period, submitted_by, default_teacher=default_teacher)

    def confirm_payroll_layout(self, batch_id: str, fingerprint: str, mapping: dict[str, str], confirmed_by: str) -> dict:
        return self.submissions.confirm_layout(batch_id, fingerprint, mapping, confirmed_by)

    def preview_payroll_merge(self, batch_id: str, expected_teachers: list[str] | None = None) -> dict:
        return self.submissions.preview_merge(batch_id, expected_teachers=expected_teachers or [])

    def confirm_payroll_merge(self, batch_id: str, output_path: str, reviewer: str, expected_teachers: list[str] | None = None) -> dict:
        return self.submissions.confirm_merge(batch_id, output_path, reviewer, expected_teachers=expected_teachers or [])

    def import_management_assessment(self, path: str, period: str, leader_id: str, submitted_by: str) -> dict:
        return self.assessments.import_assessment(path, period, leader_id, submitted_by)

    def assessment_findings(self, period: str = "", expected_leaders: list[str] | None = None) -> list[dict]:
        return self.assessments.findings(period, expected_leaders=expected_leaders or [])

    def confirm_management_assessment(self, record_id: str, reviewer: str, *, subjective_confirmations: dict[str, float] | None = None, amount_rule: dict | None = None) -> dict:
        return self.assessments.confirm(record_id, reviewer, subjective_confirmations=subjective_confirmations, amount_rule=amount_rule)

    def bind_management_assessment(self, run_id: str, result_id: str) -> dict:
        run = self.store.get(run_id)
        run = self.assessments.bind_to_run(result_id, run)
        self.store.save(run)
        return run

    # Business inputs are deliberately separate from payroll calculation.
    # These methods only create source-backed records, review them and bind
    # approved records to an existing Run.
    def create_teacher_access(self, teacher_id: str, display_name: str = "") -> dict:
        return self.inputs.create_teacher_access(teacher_id, display_name)

    def teacher_submit(self, access_token: str, period: str, item_type: str, note: str, link: dict | None = None) -> dict:
        return self.inputs.submit_teacher(access_token, period, item_type, note, link)

    def teacher_save_draft(self, access_token: str, period: str, item_type: str, note: str, link: dict | None = None, input_id: str = "") -> dict:
        return self.inputs.save_teacher_draft(access_token, period, item_type, note, link, input_id)

    def teacher_inputs(self, access_token: str) -> list[dict]:
        return self.inputs.list_teacher(access_token)

    def business_inputs(self, period: str = "", status: str = "") -> list[dict]:
        return self.inputs.list_admin(period=period, status=status)

    def create_manual_adjustment(
        self,
        run_id: str,
        *,
        teacher_id: str,
        field: str,
        raw_calculated: float,
        adjustment_value: float,
        final_value: float,
        reason: str,
        evidence: str,
        actor: str,
        confirmed_at: str = "",
    ) -> dict:
        """Persist and bind one explicitly confirmed Run-scoped adjustment."""
        run = self._load(run_id)
        teacher_id = str(teacher_id or "").strip()
        field = str(field or "").strip().upper()
        reason, evidence, actor = str(reason or "").strip(), str(evidence or "").strip(), str(actor or "").strip()
        if not teacher_id or not field or not reason or not evidence or not actor:
            raise ValueError("人工调整必须包含教师、字段、原因、证据和确认人。")
        if field not in {"AH", "AI", "AJ", "AK", "AN"}:
            raise ValueError("当前只允许对已接入的工资业务字段建立人工调整。")
        raw, delta, final = float(raw_calculated), float(adjustment_value), float(final_value)
        if abs((raw + delta) - final) > 1e-9:
            raise ValueError("人工调整的最终值必须等于原始值加调整值。")
        when = confirmed_at.strip() or datetime.now(timezone.utc).isoformat()
        item = {
            "id": uuid.uuid4().hex[:16], "input_type": "MANUAL_ADJUSTMENT",
            "source_type": "MANUAL_CONFIRMATION", "source_ref": evidence,
            "source_file_hash": "", "source_row": "", "submitted_by": actor,
            "submitted_at": when, "status": "FINAL_CONFIRMED",
            "period": run["period"], "run_id": run_id, "teacher_id": teacher_id,
            "teacher_name": teacher_id, "field": field,
            "raw_calculated": raw, "adjustment_value": delta, "final_value": final,
            "reason": reason, "evidence": evidence, "actor": actor,
            "confirmed_at": when, "created_at": when, "updated_at": when,
            "reviewed_by": actor, "reviewed_at": when,
        }
        self.store.save_business_input(item)
        run, _ = self.inputs.bind_to_run(item["id"], run)
        self.store.save(run)
        return self.store.get_business_input(item["id"])

    def import_package(self, run_id: str, package_path: str) -> dict:
        """Discover a local materials package and bind its star authority.

        This is intentionally a thin bridge into the existing Run flow: the
        package scanner owns content recognition and provenance, while the
        normal ``import_file``/``check`` path remains the calculation path.
        Source workbooks are never copied into the store.
        """
        run = self._load(run_id)
        package = discover_payroll_package(package_path, run["period"], period_start=self._period_window(run)[0], period_end=self._period_window(run)[1])
        for evidence in package.grade_evidence:
            identifier = hashlib.sha256(
                f"package-grade|{evidence.source_hash}|{evidence.coordinate}|{evidence.student}|{evidence.lesson_date}|{evidence.grade}".encode()
            ).hexdigest()[:24]
            self.store.save_student_grade_evidence({
                "id": identifier, "student": evidence.student, "lesson_date": evidence.lesson_date,
                "grade": evidence.grade, "origin": evidence.origin, "source_file": evidence.source_file,
                "source_hash": evidence.source_hash, "sheet": evidence.sheet, "coordinate": evidence.coordinate,
                "teacher": evidence.teacher, "subject": evidence.subject,
                "lesson_start_time": evidence.lesson_start_time, "class_type": evidence.class_type,
                "note": "资料包自动识别的历史年级证据。",
            })

        run["reference_ratings"] = dict(package.reference_ratings)
        run["reference_rating_sources"] = {key: list(value) for key, value in package.rating_sources.items()}
        run["authority_ratings"] = dict(package.authority_ratings)
        run["star_records"] = list(package.star_records)
        run["star_conflicts"] = list(package.star_conflicts)
        run["personnel_records"] = list(package.personnel_records)
        run["identity_conflicts"] = list(package.identity_conflicts)
        run["weekly_reports"] = list(package.weekly_reports)
        run["renewal_reports"] = list(package.renewal_reports)
        run["refund_reports"] = list(package.refund_reports)
        # Keep the user-facing material cards informative even when the
        # optional folder-package fallback is used.  The package registry
        # already contains the source-backed file metadata, so this does not
        # copy or re-read any workbook.
        run["material_inputs"] = dict(run.get("material_inputs") or {})
        for kind, report_key, label in (
            ("renewal", "renewal_reports", "续费表"),
            ("refund", "refund_reports", "退费表"),
        ):
            reports = run[report_key]
            source_type = kind.upper()
            source = next((item for item in package.source_registry if item.get("source_type") == source_type), None)
            if reports and source:
                run["material_inputs"][kind] = {
                    "name": source.get("file_name") or Path(source.get("file_path", "")).name,
                    "path": source.get("file_path", ""),
                    "sha256": source.get("file_hash", ""),
                    "records": len(reports),
                    "label": label,
                }
        run["assessment_reports"] = list(package.assessment_reports)
        run["package_root"] = str(package.root)
        run["package_inventory"] = package.inventory
        run["source_registry"] = list(package.source_registry)
        if package.template_path is not None:
            run["template_path"] = str(package.template_path)
            run["template"] = {
                "name": package.template_path.name,
                "path": str(package.template_path),
                "source": "资料包自动识别的工资模板",
            }
        if package.star_conflicts:
            # A source can be readable yet not bindable for this Run.  Mark
            # that distinction in the Run-local registry so the UI does not
            # claim a verified authority where the values conflict.
            for source in run["source_registry"]:
                if source.get("source_type") == "STAR":
                    source["status"] = "NEEDS_CONFIRMATION"
                    source.setdefault("source_evidence", {})["conflicts"] = list(package.star_conflicts)
        run["scope_teachers"] = list(package.scope_teachers)
        run["data_center"] = {
            "sources": list(package.source_registry),
            "physical_file_audit": list(package.inventory.get("physical_file_audit") or []),
            "pending_confirmation": [
                item for item in package.source_registry
                if item.get("status") == "NEEDS_CONFIRMATION"
            ],
        }

        # A package's standalone star files are system authority.  Conflicting
        # teachers are deliberately excluded, preserving the existing
        # needs-review behavior instead of choosing one value silently.
        if package.authority_ratings:
            source_hash = hashlib.sha256(
                "|".join(sorted(
                    f"{item.get('source_file', '')}:{item.get('sheet', '')}:{item.get('cell', '')}:{item.get('teacher', '')}:{item.get('rating', '')}"
                    for item in package.star_records
                )).encode()
            ).hexdigest()
            source_version = f"package-stars-{source_hash[:12]}"
            existing = next((item for item in self.store.list_rating_versions()
                             if item.get("source_hash") == source_hash
                             and item.get("source_version") == source_version
                             and item.get("effective_from") <= run["period"] <= item.get("effective_to")), None)
            if existing is None:
                self.save_rating_version(
                    run["period"], run["period"], "资料包系统权威星级", source_version,
                    [{"teacher": teacher, "rating": rating} for teacher, rating in sorted(package.authority_ratings.items())],
                    source_hash=source_hash,
                )
                existing = next(item for item in self.store.list_rating_versions()
                                if item.get("source_hash") == source_hash and item.get("source_version") == source_version)
            run["rating_version_id"] = existing["id"]
            run["star_authority_status"] = "VERIFIED_WITH_CONFLICTS" if package.star_conflicts else "VERIFIED"
        elif package.star_conflicts:
            # Clear any pre-existing auto-binding as well: retaining an older
            # rating would silently choose a value in the face of this
            # package's conflicting system sources.
            run["rating_version_id"] = None
            run["star_authority_status"] = "CONFLICT_NEEDS_CONFIRMATION"
        else:
            run["star_authority_status"] = "NOT_PROVIDED"

        # Personnel is a dated input, not a prompt-time override.  Materialize
        # the package's explicit part-time rates into the same calculation
        # version store used by ordinary Runs, while leaving identity
        # conflicts unbound until they are resolved.
        conflicted_names = {
            name for conflict in package.identity_conflicts for name in conflict.get("names", [])
        }
        personnel_profiles = [
            {
                "teacher": item["teacher"], "grade_scope": "*",
                "pricing_mode": "FIXED_GRADE_RATE", "fixed_rate": item["fixed_rate"], "rate_per_session": item["fixed_rate"],
                "notes": f"人员资料生效 {item.get('effective_from', run['period'])}～{item.get('effective_to', run['period'])}",
            }
            for item in package.personnel_records
            if item.get("employment_type") == "PART_TIME"
            and item.get("fixed_rate") is not None
            and item.get("teacher") not in conflicted_names
        ]
        # Keep employment identity even when a personnel row has no rate yet;
        # Core can then fail closed with NEEDS_INPUT instead of treating the
        # teacher as full-time and silently omitting the missing price.
        run["personnel_contexts"] = [
            {"teacher": item["teacher"], "employment_type": item["employment_type"],
             "effective_from": item["effective_from"], "effective_to": item["effective_to"],
             "source": item["source_file"]}
            for item in package.personnel_records if item.get("teacher") not in conflicted_names
        ]
        if personnel_profiles:
            source = "资料包人员资料（只读提取）"
            existing_part_time = next((item for item in self.part_time_rate_versions()
                                       if item.get("source") == source
                                       and item.get("effective_from") == run["period"]
                                       and item.get("effective_to") == run["period"]
                                       and all(
                                           any(
                                               stored.get("teacher") == expected.get("teacher")
                                               and stored.get("grade_scope", "*") == expected.get("grade_scope", "*")
                                               and float(stored.get("fixed_rate", stored.get("rate_per_session"))) == float(expected.get("fixed_rate", expected.get("rate_per_session")))
                                               for stored in (item.get("profiles") or [])
                                           ) for expected in personnel_profiles
                                       )), None)
            if existing_part_time is None:
                versions = self.save_part_time_rate_version(personnel_profiles, source, run["period"], run["period"], "系统识别")
                existing_part_time = next(item for item in versions if item.get("source") == source and item.get("effective_from") == run["period"] and item.get("effective_to") == run["period"])
            run["part_time_rate_version_id"] = existing_part_time["id"]
            # Package discovery may bind a source-backed policy after the Run
            # was created.  Freeze it before the first calculation; subsequent
            # policy edits cannot alter this snapshot.
            if run.get("status") == "DRAFT":
                run["run_policy_snapshot"] = self._build_run_policy_snapshot(run)

        if package.schedule_path is None:
            raise ValueError("资料包中没有可识别的排课表。")
        self.store.save(run)
        imported = self.import_file(run_id, "schedule", str(package.schedule_path))
        return {
            "run": imported,
            "inventory": package.inventory,
            "schedule": str(package.schedule_path),
            "reference_ratings": len(package.reference_ratings),
            "authority_ratings": len(package.authority_ratings),
            "star_conflicts": len(package.star_conflicts),
            "grade_evidence": len(package.grade_evidence),
        }

    @staticmethod
    def _read_csv(role: str, path: Path, period: str, *, period_start: str | None = None, period_end: str | None = None):
        """Read the normalized CSV shapes accepted by the material flow."""
        try:
            if role == "schedule":
                return read_schedule_csv_result(path, period, period_start=period_start, period_end=period_end)
            records = read_payroll_csv(path, period)
        except (KeyError, ValueError) as exc:
            raise ValueError(f"CSV 文件缺少必要列或数据格式不正确：{exc}") from exc
        from payroll_core.models.evidence import AdapterResult
        return AdapterResult(records=records)

    def _infer_subject_group_role(self, path: Path, run: dict) -> str:
        """Classify one submitted payroll workbook without exposing departments in UI."""
        parsed = self._read_for_run("math", path, run) if path.suffix.lower() != ".csv" else self._read_csv("math", path, run["period"])
        if parsed.errors or not parsed.records:
            raise ValueError("学科组提交表缺少可读取的教师和工资字段。")
        teachers = {str(item.teacher).strip() for item in parsed.records if getattr(item, "teacher", "").strip()}
        schedule_teachers: dict[str, set[str]] = {}
        schedule_path = run.get("files", {}).get("schedule", {}).get("path")
        if schedule_path and Path(schedule_path).is_file():
            schedule = self._read_for_run("schedule", Path(schedule_path), run)
            for item in schedule.records:
                schedule_teachers.setdefault(item.teacher, set()).add(str(item.subject or ""))
        math_words = ("数学", "math")
        science_words = ("物理", "化学", "理化", "physics", "chemistry", "science")
        math_score = sum(1 for teacher in teachers if any(word.lower() in subject.lower() for subject in schedule_teachers.get(teacher, set()) for word in math_words))
        science_score = sum(1 for teacher in teachers if any(word.lower() in subject.lower() for subject in schedule_teachers.get(teacher, set()) for word in science_words))
        stem = path.stem.lower()
        if math_score > science_score or (math_score == science_score and any(word.lower() in stem for word in math_words)):
            return "math"
        if science_score > math_score or any(word.lower() in stem for word in science_words):
            return "science"
        existing = {role for role in ("math", "science") if role in run.get("files", {})}
        missing = [role for role in ("math", "science") if role not in existing]
        if len(missing) == 1:
            return missing[0]
        raise ValueError("无法自动判断这份学科组提交表属于哪个学科，请确认文件内容后重试。")

    def import_material_file(self, run_id: str, kind: str, path: str) -> dict:
        """Import a user-facing material and keep the internal role mapping hidden."""
        run = self._load(run_id)
        source = Path(path).expanduser().resolve()
        if not source.is_file():
            raise ValueError("找不到这份材料，请重新拖入或粘贴文件。")
        if kind in {"package", "auto"}:
            if source.suffix.lower() == ".csv":
                try:
                    kind = "schedule" if self._read_csv("schedule", source, run["period"]).records else ""
                except ValueError:
                    kind = ""
                if not kind:
                    try:
                        kind = "subject_group" if self._read_csv("math", source, run["period"]).records else ""
                    except ValueError:
                        kind = ""
                if not kind:
                    renewal = read_renewal_report(source, run["period"])
                    refund = read_refund_report(source, run["period"])
                    if renewal.records and not refund:
                        kind = "renewal"
                    elif refund and not renewal.records:
                        kind = "refund"
                    else:
                        raise ValueError("无法自动识别这份 CSV。请确认包含排课、学科组提交、续费或退费字段。")
            else:
                inspection = inspect_workbook(source)
                layout = inspection.records[0].fingerprint.layout if inspection.records else ""
                if layout == LAYOUTS["schedule"]:
                    kind = "schedule"
                elif layout == LAYOUTS["math"]:
                    kind = "subject_group"
                else:
                    renewal = read_renewal_report(source, run["period"])
                    refund = read_refund_report(source, run["period"])
                    if renewal.records and not refund:
                        kind = "renewal"
                    elif refund and not renewal.records:
                        kind = "refund"
                    else:
                        raise ValueError("无法自动识别这份材料。请拖入排课、学科组提交、续费或退费表。")
        if kind == "subject_group":
            role = self._infer_subject_group_role(source, run)
            rendered = self.import_file(run_id, role, str(source))
            rendered["material_kind"] = "subject_group"
            rendered["recognized_role"] = role
            return {"run": rendered, "material_kind": "subject_group", "recognized_role": role}
        if kind == "schedule":
            rendered = self.import_file(run_id, "schedule", str(source))
            return {"run": rendered, "material_kind": "schedule", "recognized_role": "schedule"}
        if kind not in {"renewal", "refund"}:
            raise ValueError("无法识别这类材料。")
        before = version(source)
        result = read_renewal_report(source, run["period"]) if kind == "renewal" else read_refund_report(source, run["period"])
        if hasattr(result, "errors") and result.errors:
            raise ValueError("文件缺少必要列：" + "；".join(issue.message for issue in result.errors))
        records = list(result.records) if hasattr(result, "records") else list(result)
        if not records:
            raise ValueError(f"未识别到有效的{ '续费' if kind == 'renewal' else '退费' }记录，请检查工作表和表头。")
        if version(source) != before:
            raise ValueError("文件在读取期间发生变化，请关闭 Excel/WPS 后重试。")
        reports = [item.as_dict() for item in records]
        material = {
            "name": source.name, "path": str(source), **before,
            "records": len(reports), "label": "续费表" if kind == "renewal" else "退费表",
        }
        run.setdefault("material_inputs", {})[kind] = material
        run[f"{kind}_reports"] = reports
        run["last_error"] = ""
        self.store.save(run)
        return {"run": self.render(run), "material_kind": kind, "recognized_role": kind, "records": len(reports)}

    def _bind_template_from_run_files(self, run: dict) -> bool:
        """Bind an explicit package or persistent company template only."""
        if run.get("template_path") and Path(run["template_path"]).is_file():
            return True
        for item in self.store.list_company_payroll_templates():
            if item.get("status") != "ACTIVE":
                continue
            path = Path(item.get("managed_path") or item.get("path", "")) if (item.get("managed_path") or item.get("path")) else None
            if not path or not path.is_file() or version(path)["sha256"] != item.get("sha256"):
                continue
            run["template_path"] = str(path.resolve())
            run["template"] = {"name": path.name, "path": str(path.resolve()), "source": "已登记公司工资模板", "template_id": item.get("id")}
            return True
        return False

    def register_company_template(self, path: str, actor: str = "") -> dict:
        """Explicitly register a reusable company payroll template."""
        source = Path(path).expanduser().resolve()
        if not source.is_file():
            raise ValueError("找不到公司工资模板。")
        if not is_payroll_template(source):
            raise ValueError("该文件不符合公司工资模板结构，不能登记。")
        digest = version(source)
        now = datetime.now(timezone.utc).isoformat()
        managed_dir = self.root / "company-templates"
        managed_dir.mkdir(parents=True, exist_ok=True)
        managed_path = managed_dir / f"{uuid.uuid4().hex[:16]}-{source.name}"
        shutil.copy2(source, managed_path)
        if version(managed_path)["sha256"] != digest["sha256"]:
            raise ValueError("公司工资模板复制校验失败，请重试。")
        for item in self.store.list_company_payroll_templates():
            if item.get("status") == "ACTIVE":
                item["status"] = "SUPERSEDED"
                self.store.save_company_payroll_template(item)
        item = {"id": uuid.uuid4().hex[:16], "status": "ACTIVE", "name": source.name,
                "path": str(source), "managed_path": str(managed_path), "sha256": digest["sha256"], "size": digest["size"],
                "mtime_ns": digest["mtime_ns"], "registered_by": actor or "管理员",
                "created_at": now, "source": "明确登记的公司工资模板"}
        self.store.save_company_payroll_template(item)
        return item

    def import_business_results(self, input_type: str, period: str, path: str, submitted_by: str, activation_scope: str = "SUPPLEMENT", replace_input_ids: list[str] | None = None) -> list[dict]:
        return self.inputs.import_results(input_type, period, path, submitted_by, activation_scope=activation_scope, replace_input_ids=replace_input_ids)

    def review_business_input(self, input_id: str, action: str, reviewer: str, note: str = "") -> dict:
        return self.inputs.review(input_id, action, reviewer, note)

    def bind_business_input(self, run_id: str, input_id: str) -> dict:
        run = self._load(run_id)
        self._require_fresh(run)
        run, item = self.inputs.bind_to_run(input_id, run, persist=False)
        if item.get("input_type") == "RENEWAL_RESULT":
            self._bind_renewal_snapshot(run, item)
        run["status"] = "FILES_READY" if self._materials_ready(run) else "DRAFT"
        run["business_context_stale"] = True
        invalidate_business_decisions(run.setdefault("business_decisions", []))
        self.store.save_run_and_business_input(run, item)
        self.inputs._event(item, "BOUND_TO_RUN", "系统", f"绑定核算 {run['id']}")
        return self.render(run)

    def _bind_renewal_snapshot(self, run: dict, item: dict) -> None:
        """Freeze approved renewal results at the Run boundary.

        The live business-input row remains auditable, but generation reads
        this snapshot so a later source edit cannot alter an existing Run.
        Duplicate approved results for one identity are rejected rather than
        silently summed.
        """
        snapshot = run.get("run_renewal_result_snapshot")
        if not isinstance(snapshot, dict):
            snapshot = {"version": "RUN_RENEWAL_RESULT_SNAPSHOT/v1", "run_id": run["id"], "period_label": run["period"], "created_at": datetime.now(timezone.utc).isoformat(), "entries": {}}
        entries = snapshot.setdefault("entries", {})
        teacher_id = str(item.get("teacher_id", "")).strip()
        if not teacher_id:
            raise ValueError("续费结果缺少稳定教师标识，不能绑定到工资核算。")
        if teacher_id in entries and entries[teacher_id].get("source_result_id") != item.get("id"):
            raise ValueError("同一教师已有另一条已绑定续费结果；请先明确替代版本，不能静默合并。")
        entry = renewal_snapshot_entry(item, run_id=run["id"], period=run["period"], display_name=str((item.get("payload") or {}).get("teacher") or teacher_id))
        entries[teacher_id] = entry
        unsigned = {key: value for key, value in snapshot.items() if key != "sha256"}
        snapshot["sha256"] = hashlib.sha256(json.dumps(unsigned, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")).hexdigest()
        run["run_renewal_result_snapshot"] = snapshot

    def comment_candidates(self, run_id: str = "") -> list[dict]:
        values = self.store.list_comment_candidates()
        if run_id:
            values = [item for item in values if item.get("run_id") == run_id]
        for item in values:
            self._refresh_candidate(item)
        return values

    def create_refund_comment_candidate(self, run_id: str, input_id: str, target_role: str, sheet: str, cell: str) -> dict:
        run = self._load(run_id); self._require_fresh(run)
        source = self.store.get_business_input(input_id)
        if source.get("input_type") != "REFUND_RESULT" or source.get("status") != "APPROVED" or not self.inputs.current(source["id"]):
            raise ValueError("只有已审核通过的退费结果才能生成退费批注候选。")
        if source.get("period") != run["period"]:
            raise ValueError("退费结果月份与当前工资核算月份不一致。")
        if source["id"] not in {item.get("input_id") for item in run.get("business_input_bindings", [])}:
            raise ValueError("退费结果必须先绑定到当前工资核算。")
        return self._create_candidate(run, source, "REFUND_NOTE", target_role, sheet, cell, {
            "period": run["period"], "note": source.get("payload", {}).get("note") or source.get("payload", {}).get("说明") or "已审核退费结果",
            "source_label": f"退费结果表第{source.get('source_row') or '—'}行",
        })

    def create_class_comment_candidate(self, run_id: str, resolution_id: str, target_role: str, sheet: str, cell: str) -> dict:
        """Create a note only from an already confirmed structured resolution.

        The caller may only name a persisted resolution.  It cannot claim a
        browser-supplied dictionary is confirmed.
        """
        run = self._load(run_id); self._require_fresh(run)
        resolution = self.store.get_resolution(resolution_id)
        kind = str(resolution.get("kind", ""))
        state = str(resolution.get("outcome", resolution.get("status", "")))
        if resolution.get("run_id") != run_id or resolution.get("period") != run["period"]:
            raise ValueError("该班课处理不属于当前工资核算。")
        if kind not in {"SOURCE_DATA_CORRECTION", "APPROVED_PAYROLL_OVERRIDE"} or state not in {"CONFIRMED", "RESOLVED_BY_SOURCE_CORRECTION", "RESOLVED_BY_APPROVED_OVERRIDE"}:
            raise ValueError("只有已确认且可追溯的结构化班课处理，才能生成确定性批注。")
        template = "SOURCE_CORRECTION_NOTE" if kind == "SOURCE_DATA_CORRECTION" else "APPROVED_OVERRIDE_NOTE"
        return self._create_candidate(run, None, template, target_role, sheet, cell, {
            "period": run["period"], "course_label": resolution.get("course_label", "相关课程"),
            "reason": resolution.get("reason", "已确认上游事实修正"), "corrected_value": resolution.get("corrected_value", "确认后的"),
            "approved_treatment": resolution.get("approved_treatment", "批准口径"), "system_value": resolution.get("system_value", ""),
        }, resolution=resolution)

    def preview_comment_candidate(self, candidate_id: str, strategy: str = "APPEND") -> dict:
        item = self.store.get_comment_candidate(candidate_id)
        self._refresh_candidate(item)
        if item["status"] == "NEEDS_RECONFIRMATION":
            raise ValueError("批注依据已变化，请重新生成。")
        view = preview_writeback(Path(item["source_workbook"]), [item], strategy=strategy)[0]
        token = hashlib.sha256((item["id"] + view.before_comment + view.proposed_comment + strategy).encode()).hexdigest()
        item.update({"status": "PREVIEWED", "preview_token": token, "previewed_at": datetime.now(timezone.utc).isoformat(), "preview_strategy": strategy, "before_comment": view.before_comment, "after_comment": view.proposed_comment})
        self.store.save_comment_candidate(item)
        return item

    def approve_comment_candidate(self, candidate_id: str, reviewer: str, preview_token: str) -> dict:
        if not reviewer.strip() or not preview_token:
            raise ValueError("请先预览最终批注，再填写确认人。")
        item = self.store.get_comment_candidate(candidate_id)
        self._refresh_candidate(item)
        if item.get("status") != "PREVIEWED" or item.get("preview_token") != preview_token:
            raise ValueError("请先预览最终批注；若预览已失效，请重新预览后确认。")
        item.update({"status": "APPROVED", "approved_by": reviewer.strip(), "approved_at": datetime.now(timezone.utc).isoformat(), "writeback_strategy": item["preview_strategy"]})
        self.store.save_comment_candidate(item)
        return item

    def writeback_comments(self, run_id: str, source_workbook: str, candidate_ids: list[str], output_path: str, reviewer: str) -> dict:
        if not reviewer.strip():
            raise ValueError("请填写确认回填人。")
        run = self._load(run_id); self._require_fresh(run)
        source = Path(source_workbook).expanduser().resolve()
        items = [self.store.get_comment_candidate(item_id) for item_id in candidate_ids]
        if not items or any(item.get("run_id") != run_id for item in items):
            raise ValueError("请选择当前核算中已确认的批注候选。")
        if any(item.get("status") != "APPROVED" for item in items):
            raise ValueError("所有批注都必须先在预览后确认。")
        if any(Path(item["source_workbook"]).resolve() != source for item in items):
            raise ValueError("一次回填只能处理同一份原工资表。")
        for item in items:
            self._refresh_candidate(item)
            if item["status"] == "NEEDS_RECONFIRMATION":
                raise ValueError("存在依据已变化的批注候选，请重新确认。")
        strategies = {item.get("writeback_strategy", "APPEND") for item in items}
        if len(strategies) != 1:
            raise ValueError("一次回填的已有批注处理方式必须一致。")
        actual_path = safe_output_path(output_path)
        result = write_new_workbook(source, actual_path, items, strategy=strategies.pop())
        when = datetime.now(timezone.utc).isoformat()
        for item in items:
            item.update({"status": "WRITTEN", "written_at": when, "written_by": reviewer.strip(), "output_workbook": str(actual_path.resolve())})
            self.store.save_comment_candidate(item)
        run.setdefault("writeback_history", []).append({"source_workbook": str(source), "output_workbook": str(actual_path.resolve()), "candidate_ids": candidate_ids, "written_by": reviewer.strip(), "written_at": when})
        self.store.save(run)
        return {"output_path": str(actual_path.resolve()), "location": describe_location(actual_path), "written": len(result), "candidates": items}

    def _create_candidate(self, run: dict, source_input: dict | None, kind: str, target_role: str, sheet: str, cell: str, values: dict, *, resolution: dict | None = None) -> dict:
        if target_role not in run.get("files", {}):
            raise ValueError("请先导入要写入批注的工资表。")
        target = run["files"][target_role]
        source = Path(target["path"])
        if not source.is_file():
            raise ValueError("目标工资表无法读取。")
        teacher_id = (source_input or resolution or {}).get("teacher_id", (source_input or resolution or {}).get("teacher", ""))
        self._verify_comment_target(run, target_role, sheet, cell, teacher_id)
        content, template_version = render_comment(kind, values)
        item = {
            "id": uuid.uuid4().hex[:16], "run_id": run["id"], "period": run["period"], "teacher_id": teacher_id,
            "comment_type": kind, "source_workbook": str(source), "source_file_hash": workbook_sha256(source), "sheet": sheet, "cell": cell,
            "content": content, "template_version": template_version, "source_input_ids": [source_input["id"]] if source_input else [], "source_input_hashes": [source_input.get("source_file_hash", "")] if source_input else [], "source_resolution_id": resolution.get("id", "") if resolution else "", "source_resolution_fingerprint": resolution.get("fingerprint", "") if resolution else "", "source_resolution_source_hash": (resolution or {}).get("source_hash", (resolution or {}).get("source_file_hash", "")), "source_resolution_rule_version": (resolution or {}).get("rule_version", ""), "source_resolution": resolution or {},
            "before_comment": "", "after_comment": "", "status": "PROPOSED", "created_at": datetime.now(timezone.utc).isoformat(), "approved_by": "", "approved_at": "", "written_at": "",
        }
        self.store.save_comment_candidate(item)
        return item

    def _refresh_candidate(self, item: dict) -> None:
        if item.get("status") in {"WRITTEN", "REJECTED"}:
            return
        stale = False
        source = Path(item.get("source_workbook", ""))
        try:
            stale = not source.is_file() or workbook_sha256(source) != item.get("source_file_hash")
        except OSError:
            stale = True
        for input_id in item.get("source_input_ids", []):
            try:
                input_record = self.store.get_business_input(input_id)
                stale = stale or not self.inputs.current(input_id, item.get("source_input_hashes", [""])[0])
            except ValueError:
                stale = True
        resolution_id = item.get("source_resolution_id")
        if resolution_id:
            try:
                resolution = self.store.get_resolution(resolution_id)
                stale = stale or resolution.get("fingerprint") != item.get("source_resolution_fingerprint") or resolution.get("outcome") not in {"RESOLVED_BY_SOURCE_CORRECTION", "RESOLVED_BY_APPROVED_OVERRIDE"}
                # A resolution that lost its source must invalidate its note too:
                # the outcome alone still looks resolved, but its basis changed.
                stale = stale or resolution.get("status", "ACTIVE") != "ACTIVE"
                stale = stale or resolution.get("source_hash", resolution.get("source_file_hash", "")) != item.get("source_resolution_source_hash", "")
                stale = stale or resolution.get("rule_version", "") != item.get("source_resolution_rule_version", "")
            except ValueError:
                stale = True
        if stale and item.get("status") != "NEEDS_RECONFIRMATION":
            item["status"] = "NEEDS_RECONFIRMATION"
            self.store.save_comment_candidate(item)

    def _verify_comment_target(self, run: dict, role: str, sheet: str, cell: str, teacher_id: str) -> None:
        records = self._read_for_run(role, Path(run["files"][role]["path"]), run).records
        matching = [record for record in records if getattr(record, "teacher", "") == teacher_id]
        if not matching:
            raise ValueError("目标工资表没有该教师，不能写入其他教师单元格。")
        if not any(any(item.sheet == sheet and item.coordinate == cell for item in record.provenance.values()) for record in matching):
            raise ValueError("目标单元格不属于该教师的已识别工资表字段。")

    def create(self, period: str, mode: str = MODE_AUDIT, *, period_start: str = "", period_end: str = "", period_boundary_source: str = "") -> dict:
        if len(period) != 7 or period[4] != "-" or not period.replace("-", "").isdigit() or not 1 <= int(period[5:]) <= 12:
            raise ValueError("请选择有效月份。")
        if mode not in (MODE_AUDIT, MODE_GENERATE):
            raise ValueError("未知的工资任务方式。")
        versions = [item for item in self.store.list_rating_versions() if item.get("status", "ACTIVE") == "ACTIVE" and item["effective_from"] <= period <= item["effective_to"]]
        policies = [item for item in self.store.list_policy_versions() if item.get("status", "ACTIVE") == "ACTIVE" and item["effective_from"] <= period <= item["effective_to"]]
        class_rules = [item for item in self.class_type_rule_versions() if item.get("status", "ACTIVE") == "ACTIVE" and item["effective_from"] <= period <= item["effective_to"]]
        class_rules.sort(key=lambda item: item["effective_from"], reverse=True)
        window = normalize_period_window(period, period_start, period_end, period_boundary_source)
        persisted_salary = self._effective_base_salary_inputs(period)
        persisted_af = self._effective_af_policy(period)
        run = {"id": uuid.uuid4().hex[:12], "period": period, **window, "mode": mode, "created_at": datetime.now(timezone.utc).isoformat(), "status": "DRAFT", "files": {}, "issues": [], "field_records": [], "issue_groups": [], "user_actions": [], "decisions": [], "business_decisions": [], "management": [], "resolutions": [], "resolution_history": [], "rating_version_id": versions[0]["id"] if len(versions) == 1 else None, "policy_version_id": policies[0]["id"] if len(policies) == 1 else None, "class_type_rule_version_id": class_rules[0]["id"] if class_rules else None, "confirmed_hours": {}, "af_policy_confirmation": persisted_af, "base_salary_inputs": persisted_salary, "base_salary_input_snapshot": self._base_salary_snapshot(period, persisted_salary) if persisted_salary else None, "base_salary_deferred": False, "base_salary_deferred_by": "", "base_salary_deferred_at": None, "run_renewal_result_snapshot": None, "field_status": self._field_status([]), "summary": self._summary([])}
        self._bind_new_calculation(run)
        run["run_policy_snapshot"] = self._build_run_policy_snapshot(run)
        self._bind_template_from_run_files(run)
        self.store.save(run)
        return self.render(run)

    def _base_salary_snapshot(self, period: str, inputs: dict[str, dict]) -> dict:
        snapshot = {"version": "BASE_SALARY_INPUT_SNAPSHOT/v1", "run_id": "AUTO_PROFILE", "period": period, "confirmed_by": "长期教师基本工资资料", "confirmed_at": datetime.now(timezone.utc).isoformat(), "source": "已确认的长期教师基本工资资料", "inputs": inputs}
        snapshot["sha256"] = hashlib.sha256(json.dumps(snapshot, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")).hexdigest()
        return snapshot

    def _effective_base_salary_inputs(self, period: str) -> dict[str, dict]:
        grouped: dict[str, list[dict]] = {}
        for item in self.store.list_teacher_base_salary_profiles():
            if item.get("status", "ACTIVE") != "ACTIVE" or not item.get("effective_from", "") <= period <= item.get("effective_to", "9999-12"):
                continue
            grouped.setdefault(str(item.get("teacher", "")), []).append(item)
        result: dict[str, dict] = {}
        for teacher, profiles in grouped.items():
            distinct = {json.dumps(p.get("entry", {}), ensure_ascii=False, sort_keys=True) for p in profiles}
            if teacher and len(distinct) == 1:
                result[teacher] = profiles[-1]["entry"]
        return result

    def _effective_af_policy(self, period: str) -> dict | None:
        policies = [item for item in self.store.list_af_default_policies() if item.get("status", "ACTIVE") == "ACTIVE" and item.get("effective_from", "") <= period <= item.get("effective_to", "9999-12")]
        if not policies:
            return None
        selected = copy.deepcopy(sorted(policies, key=lambda item: item.get("effective_from", ""), reverse=True)[0])
        selected["exceptions"] = {}
        selected["effective_to"] = selected.get("effective_to", "9999-12")
        return selected

    def save_base_salary_inputs(self, run_id: str, inputs: list[dict], confirmed_by: str, source: str = "本次 Run 基本工资确认") -> dict:
        """Persist and freeze the source-backed G:L inputs used to calculate M."""
        run = self._load(run_id)
        self._require_fresh(run)
        actor = str(confirmed_by or "").strip()
        if not actor:
            raise ValueError("请填写基本工资输入确认人。")
        if not isinstance(inputs, list) or not inputs:
            raise ValueError("请至少填写一位教师的 G～L 基本工资输入。")
        from payroll_core.payroll_generation import BASE_SALARY_FIELDS, base_salary_field
        now = datetime.now(timezone.utc).isoformat()
        normalized: dict[str, dict] = {}
        for item in inputs:
            teacher = str(item.get("teacher") or item.get("display_name") or "").strip()
            if not teacher or teacher in normalized:
                raise ValueError("基本工资输入中的教师不能为空且不能重复。")
            raw_fields = item.get("fields") or {}
            if not isinstance(raw_fields, dict):
                raise ValueError(f"{teacher} 的基本工资输入格式无效。")
            fields: dict[str, dict] = {}
            for code in BASE_SALARY_FIELDS:
                raw = raw_fields.get(code)
                value = raw.get("value") if isinstance(raw, dict) else raw
                # Preserve an explicit blank as a blocked input so the Run can
                # show a teacher-scoped action while other teachers continue.
                # Missing G:L is never converted to zero.
                if value in (None, ""):
                    fields[code] = {"value": None, "source": str((raw or {}).get("source") if isinstance(raw, dict) else source) or source, "provenance": (raw or {}).get("provenance", {}) if isinstance(raw, dict) else {}, "confirmed_at": now, "confirmed_by": actor}
                    continue
                try:
                    number = float(value)
                except (TypeError, ValueError) as exc:
                    raise ValueError(f"{teacher} 的 {code} 不是有效数字。") from exc
                if not math.isfinite(number):
                    raise ValueError(f"{teacher} 的 {code} 不是有限数字。")
                fields[code] = {"value": number, "source": str((raw or {}).get("source") if isinstance(raw, dict) else source) or source, "provenance": (raw or {}).get("provenance", {}) if isinstance(raw, dict) else {}, "confirmed_at": now, "confirmed_by": actor}
            if ((fields["K"]["value"] is not None and fields["K"]["value"] <= 0)
                    or (fields["L"]["value"] is not None and fields["L"]["value"] < 0)):
                raise ValueError(f"{teacher} 的出勤输入无效：K 必须大于 0，L 不能为负数。")
            entry = {"teacher_id": str(item.get("teacher_id") or teacher), "display_name": teacher, "fields": fields, "source": source, "provenance": {"kind": "RUN_BASE_SALARY_INPUT", "confirmed_by": actor, "confirmed_at": now}}
            entry["m"] = base_salary_field(teacher, {teacher: entry})
            normalized[teacher] = entry
            effective_from = str(item.get("effective_from") or run["period"])
            effective_to = str(item.get("effective_to") or "9999-12")
            for old in self.store.list_teacher_base_salary_profiles():
                if old.get("status", "ACTIVE") == "ACTIVE" and old.get("teacher") == teacher:
                    old_from = str(old.get("effective_from") or "")
                    old_to = str(old.get("effective_to") or "9999-12")
                    if old_from == effective_from:
                        # Re-confirming the same effective month replaces the
                        # prior version without changing older periods.
                        old["status"] = "SUPERSEDED"
                        self.store.save_teacher_base_salary_profile(old)
                    elif old_from < effective_from <= old_to:
                        # Close the old version immediately before the new
                        # version.  It remains ACTIVE for historical months.
                        old["effective_to"] = previous_period(effective_from)
                        self.store.save_teacher_base_salary_profile(old)
                    elif effective_from < old_from <= effective_to:
                        raise ValueError(f"{teacher} 的基本工资版本生效期重叠，请先明确旧版本有效期。")
            self.store.save_teacher_base_salary_profile({
                "id": uuid.uuid4().hex[:16], "teacher": teacher,
                "effective_from": effective_from, "effective_to": effective_to,
                "source": source, "confirmed_by": actor, "version": "TEACHER_BASE_SALARY/v1",
                "status": "ACTIVE", "entry": entry, "created_at": now,
            })
        snapshot = {"version": "BASE_SALARY_INPUT_SNAPSHOT/v1", "run_id": run_id, "period": run["period"], "confirmed_by": actor, "confirmed_at": now, "source": source, "inputs": normalized}
        snapshot["sha256"] = hashlib.sha256(json.dumps(snapshot, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")).hexdigest()
        run["base_salary_inputs"] = normalized
        run["base_salary_input_snapshot"] = snapshot
        run["business_context_stale"] = True
        # Validate the complete snapshot before saving it, including the
        # derived M value for every submitted teacher.
        for teacher in normalized:
            base_salary_field(teacher, normalized)
        self.store.save(run)
        return self.render(run)

    def confirm_af_policy(self, run_id: str, confirmed_by: str, *, default_obligation_hours: float = 30, exceptions: dict | None = None, reason: str = "") -> dict:
        """Confirm this Run's default obligation-hours policy once.

        The decision is Run-scoped and never mutates the dated compensation
        policy version.  Personal exceptions override the Run default only for
        the named teacher; a later Run starts with no confirmation.
        """
        run = self._load(run_id)
        self._require_fresh(run)
        actor = str(confirmed_by or "").strip()
        try:
            default_hours = float(default_obligation_hours)
        except (TypeError, ValueError) as exc:
            raise ValueError("默认义务课时必须是非负有限数值。") from exc
        if not actor or not math.isfinite(default_hours) or default_hours < 0:
            raise ValueError("请填写确认人和有效的默认义务课时。")
        cleaned: dict[str, dict] = {}
        for teacher, item in (exceptions or {}).items():
            name = str(teacher or "").strip()
            if not name or not isinstance(item, dict):
                raise ValueError("义务课时特殊情况格式无效。")
            try:
                hours = float(item.get("obligation_hours"))
            except (TypeError, ValueError) as exc:
                raise ValueError("特殊教师义务课时必须是非负有限数值。") from exc
            if not math.isfinite(hours) or hours < 0:
                raise ValueError("特殊教师义务课时必须是非负有限数值。")
            cleaned[name] = {"obligation_hours": hours, "deduction_enabled": bool(item.get("deduction_enabled", True)), "reason": str(item.get("reason", "")).strip()}
        timestamp = datetime.now(timezone.utc).isoformat()
        run["af_policy_confirmation"] = {
            "default_obligation_hours": default_hours,
            "confirmed": True,
            "confirmed_by": actor,
            "confirmed_at": timestamp,
            "effective_from": run["period"],
            "effective_to": run["period"],
            "source": "本次核算义务课时确认",
            "reason": str(reason or "普通全职教师按默认义务课时扣除；特殊人员按例外配置。"),
            "exceptions": cleaned,
        }
        self.store.save_af_default_policy({**run["af_policy_confirmation"], "id": uuid.uuid4().hex[:16], "effective_to": "9999-12", "status": "ACTIVE", "version": "AF_DEFAULT_POLICY/v1"})
        invalidate_business_decisions(run.setdefault("business_decisions", []))
        self.store.save(run)
        return self.check(run_id) if self._materials_ready(run) else self.render(run)

    def defer_base_salary(self, run_id: str, confirmed_by: str, reason: str = "") -> dict:
        """Allow an explicit generate-now decision without inventing M=0."""
        run = self._load(run_id)
        self._require_fresh(run)
        actor = str(confirmed_by or "").strip()
        if not actor:
            raise ValueError("请填写确认人，才能暂不录入基本工资并继续生成。")
        run["base_salary_deferred"] = True
        run["base_salary_deferred_by"] = actor
        run["base_salary_deferred_at"] = datetime.now(timezone.utc).isoformat()
        run["base_salary_deferred_reason"] = str(reason or "用户选择暂不录入基本工资；后续补录后可重新生成。")
        self.store.save(run)
        return self.render(run)

    def list(self) -> list[dict]:
        return [self.render(self._load(item["id"])) for item in self.store.list()]

    def list_index(self) -> list[dict]:
        """Return the cheap, resumable projection used by the home page.

        Opening the workbench must not render every archived Run (rendering a
        Run can parse its source workbooks and build grade help).  The index
        deliberately contains only persisted metadata and the material
        readiness that can be derived without reading a workbook.  Opening a
        selected Run still goes through :meth:`get` and the authoritative
        render path.
        """
        indexed: list[dict] = []
        for stored in self.store.list():
            files = stored.get("files") or {}
            required = self._missing_materials({**stored, "files": files})
            mode = stored.get("mode", MODE_AUDIT)
            ready_count = int("schedule" in files) + int(any(key in files for key in SCOPE_ROLES))
            readiness = 100 if mode == MODE_GENERATE and "schedule" in files else round(100 * ready_count / 2)
            status = str(stored.get("status", "DRAFT"))
            summary = dict(stored.get("summary") or {})
            indexed.append({
                "id": stored["id"],
                "period": stored.get("period", ""),
                "period_label": stored.get("period_label", stored.get("period", "")),
                "created_at": stored.get("created_at"),
                "updated_at": stored.get("updated_at", stored.get("created_at")),
                "mode": mode,
                "status": status,
                "status_label": STATUS.get(status, status),
                "summary": summary,
                "issue_groups": list(stored.get("issue_groups") or []),
                "user_actions": list(stored.get("user_actions") or []),
                "health": {
                    "ready": not required and status != "STALE",
                    "missing": required,
                    "readiness": readiness,
                },
            })
        return sorted(indexed, key=lambda item: item.get("updated_at") or item.get("created_at") or "", reverse=True)

    def get(self, run_id: str) -> dict:
        return self.render(self._load(run_id))

    def rating_versions(self) -> list[dict]:
        return self._version_views("rating")

    def save_rating_version(self, effective_from: str, effective_to: str, source: str, source_version: str, ratings: list[dict], supersedes_version_id: str | None = None, source_hash: str = "") -> list[dict]:
        if len(effective_from) != 7 or len(effective_to) != 7 or effective_from > effective_to or not source.strip() or not ratings:
            raise ValueError("请填写生效期、来源和至少一位教师的星级。")
        cleaned = []
        for item in ratings:
            name, rating = str(item.get("teacher", "")).strip(), item.get("rating")
            if not name or not isinstance(rating, int) or rating not in range(1, 7):
                raise ValueError("星级资料必须包含教师姓名和一至六星。")
            cleaned.append({
                "teacher": name,
                "rating": rating,
                "role": str(item.get("role", "教师")).strip() or "教师",
                "allow_blank_payroll_rating": bool(item.get("allow_blank_payroll_rating", False)),
            })
        if supersedes_version_id:
            prior = self.store.get_rating_version(supersedes_version_id)
            if prior["effective_from"] != effective_from or prior["effective_to"] != effective_to:
                raise ValueError("修正版必须沿用原版本的生效期；请另建新的年度版本。")
        version = {"id": uuid.uuid4().hex[:12], "effective_from": effective_from, "effective_to": effective_to, "source": source.strip(), "source_hash": source_hash.strip(), "source_version": source_version.strip() or effective_from, "ratings": cleaned, "created_at": datetime.now(timezone.utc).isoformat(), "status": "ACTIVE", "supersedes_version_id": supersedes_version_id}
        self.store.save_rating_version(version)
        if supersedes_version_id:
            self._supersede("rating", supersedes_version_id, version["id"])
        return self.rating_versions()

    def policy_versions(self) -> list[dict]:
        return self._version_views("policy")

    def payroll_policy_registry(self, run_id: str = "", period: str = "") -> dict:
        """Return one auditable registry view over all compensation policies.

        Registry rows are projections, not a second calculation engine.  A
        supplied Run is the preferred scope and exposes its frozen snapshot;
        otherwise only currently stored dated policy versions are shown.
        """
        run = self._load(run_id) if run_id else None
        target_period = period or (run or {}).get("period", "")
        if target_period and not valid_period(target_period):
            raise ValueError("工资月份格式无效。")
        rows: list[dict] = []
        snapshot = (run or {}).get("run_policy_snapshot")
        if snapshot and snapshot.get("sha256") != self._policy_snapshot_hash(snapshot):
            raise ValueError("Run 政策快照校验失败，不能继续核算。")

        default = (snapshot or {}).get("default_full_time")
        if default is None:
            core = self._calculation_version(run, "core") if run else None
            candidate = (core or {}).get("rules", {}).get("af", {}).get("default_policy_candidate") or {}
            default = {"policy_type": "DEFAULT_FULL_TIME", "obligation_hours": float(candidate.get("obligation_hours", 30) or 30), "unit": "HOURS", "source": (core or {}).get("source", "Core AF 默认政策"), "source_hash": (core or {}).get("sha256", ""), "status": "ACTIVE"}
        rows.append({"policy_type": "DEFAULT_FULL_TIME", "teacher_id": "*", "display_name": "普通全职教师", "rate": None, "unit": "HOURS", "obligation_hours": default.get("obligation_hours", 30), "effective_from": target_period or "—", "effective_to": target_period or "—", "source": default.get("source", "Core AF 默认政策"), "source_hash": default.get("source_hash", ""), "provenance": {"kind": "CORE_RULE"}, "status": "ACTIVE"})

        if snapshot:
            personal_versions = [{"id": snapshot.get("policy_version_id"), "effective_from": target_period, "effective_to": target_period, "source": "RUN_POLICY_SNAPSHOT", "source_hash": "", "status": "ACTIVE", "profiles": snapshot.get("personal_policies") or []}]
            rate_versions = [{"id": snapshot.get("part_time_rate_version_id"), "effective_from": target_period, "effective_to": target_period, "source": "RUN_POLICY_SNAPSHOT", "source_hash": "", "status": "ACTIVE", "profiles": snapshot.get("part_time_rates") or [], "actor": "RUN_POLICY_SNAPSHOT", "created_at": snapshot.get("created_at", "")}]
        else:
            personal_versions = self.store.list_policy_versions()
            rate_versions = self.store.calculation_versions("part_time")

        for version in personal_versions:
            effective = bool(target_period) and version.get("effective_from", "") <= target_period <= version.get("effective_to", "") and version.get("status", "ACTIVE") == "ACTIVE"
            status = "ACTIVE" if effective else ("EXPIRED" if target_period and version.get("effective_to", "") < target_period else version.get("status", "EXPIRED"))
            for profile in version.get("profiles") or []:
                rows.append({"policy_type": profile.get("policy_type", "PERSONAL_POLICY"), "teacher_id": profile.get("teacher_id") or "", "display_name": profile.get("teacher", ""), "rate": None, "unit": "HOURS", "obligation_hours": profile.get("obligation_hours"), "effective_from": version.get("effective_from", ""), "effective_to": version.get("effective_to", ""), "source": version.get("source", ""), "source_hash": version.get("source_hash", ""), "provenance": {"version_id": version.get("id"), "role": profile.get("role", ""), **(profile.get("provenance") or {})}, "status": status})

        for version in rate_versions:
            effective = bool(target_period) and version.get("effective_from", "") <= target_period <= version.get("effective_to", "") and version.get("status", "ACTIVE") == "ACTIVE"
            status = "ACTIVE" if effective else ("EXPIRED" if target_period and version.get("effective_to", "") < target_period else version.get("status", "EXPIRED"))
            for profile in version.get("profiles") or []:
                rows.append({"policy_type": "PART_TIME_RATE", "teacher_id": profile.get("teacher_id") or "", "display_name": profile.get("teacher", ""), "rate": profile.get("rate_per_session"), "unit": profile.get("unit", "CNY_PER_LESSON"), "obligation_hours": None, "effective_from": version.get("effective_from", ""), "effective_to": version.get("effective_to", ""), "source": version.get("source", ""), "source_hash": version.get("source_hash", ""), "provenance": {"version_id": version.get("id"), "grade_scope": profile.get("grade_scope", "*"), "pricing_mode": profile.get("pricing_mode", "FIXED_GRADE_RATE"), "student_id": profile.get("student_id", ""), "class_id": profile.get("class_id", ""), **(profile.get("provenance") or {}), "actor": version.get("actor", "")}, "status": status})

        # Historical July rates are deliberately a separate projection.  They
        # are visible for audit and never enter a future Run snapshot.
        if run and target_period == "2026-07":
            try:
                reconciliation = self.historical_reconciliation(run["id"])
            except (ValueError, OSError):
                reconciliation = {}
            for teacher in (reconciliation.get("source_classifications") or {}).get("PART_TIME_RATE", []):
                detail = next((item for item in reconciliation.get("rows", []) if item.get("teacher") == teacher), {})
                rate = None
                source = "2026-07 历史工资表（仅证据）"
                for field in (detail.get("fields") or {}).values():
                    for evidence in field.get("evidence") or []:
                        if evidence.get("source_classification") == "PART_TIME_RATE":
                            rate = evidence.get("rate")
                            source = evidence.get("source_cell") or source
                rows.append({"policy_type": "PART_TIME_RATE", "teacher_id": "", "display_name": teacher, "rate": rate, "unit": "CNY_PER_LESSON", "obligation_hours": None, "effective_from": "2026-07", "effective_to": "2026-07", "source": source, "source_hash": "", "provenance": {"kind": "HISTORICAL_EVIDENCE", "run_id": run["id"]}, "status": "HISTORICAL_ONLY"})

        # Same display name with multiple identities must never silently bind.
        by_name: dict[str, set[str]] = {}
        for row in rows:
            if row["display_name"]:
                by_name.setdefault(row["display_name"], set()).add(row.get("teacher_id", ""))
        for row in rows:
            ids = {item for item in by_name.get(row["display_name"], set()) if item}
            if len(ids) > 1 and not row.get("teacher_id"):
                row["status"] = "NEEDS_CONFIRMATION"
        return {"version": "PAYROLL_POLICY_REGISTRY/v1", "period": target_period or None, "run_id": run_id or None, "priority": ["PERSONAL_POLICY", "PART_TIME_RATE", "DEFAULT_FULL_TIME"], "rows": rows}

    def av_source_map(self, run_id: str = "") -> dict[str, Any]:
        """Return a read-only AV source map backed by the current Run/template.

        The map is intentionally an evidence projection.  It does not create
        business inputs, alter a Run, or calculate any payroll value.
        """
        run = self._load(run_id) if run_id else None
        template = _template_av_evidence((run or {}).get("template_path"))
        generated_rows = ((run or {}).get("generated_payroll") or {}).get("rows") or []
        core_rows = (run or {}).get("core_calculation", {}).get("rows") or []
        rows = generated_rows or core_rows
        statuses: dict[str, set[str]] = {code: set() for code in AV_SOURCE_DEFINITIONS}
        for row in rows:
            final = row.get("final_fields") or {}
            core = row.get("fields") or {}
            for code in statuses:
                item = final.get(code) or core.get(code) or {}
                statuses[code].add(str(item.get("state", "NO_EVIDENCE")))

        fields: list[dict[str, Any]] = []
        payable = {"DETERMINED", "NOT_APPLICABLE"}
        for code, definition in AV_SOURCE_DEFINITIONS.items():
            item = copy.deepcopy(definition)
            state_set = statuses.get(code) or set()
            if not state_set:
                item["current_system_status"] = "NO_EVIDENCE"
            elif len(state_set) == 1:
                item["current_system_status"] = next(iter(state_set))
            else:
                item["current_system_status"] = "MIXED: " + ", ".join(sorted(state_set))
            item["historical_evidence"] = template.get("historical_evidence", "NO EVIDENCE")
            if code in template.get("field_formulas", {}):
                item["template_formula"] = template["field_formulas"].get(code)
            # Once this Run has a complete, source-backed G:L snapshot, M is
            # no longer merely an available-but-unconnected source: it is a
            # determined production component.  Keep the original source
            # provenance in the same record while exposing the current state.
            if code in {"M", "AK"} and state_set and state_set <= payable:
                item["category"] = "DETERMINED"
            fields.append({"column": code, **item})

        direct_components = ("M", "AF", "AG", "AK", "AL", "AM", "AN", "AO", "AP", "AQ", "AR", "AS", "AT", "AU")
        complete_count = sum(1 for code in direct_components if statuses.get(code) and statuses[code] <= payable)
        av_blockers = [code for code in direct_components if not statuses.get(code) or not statuses[code] <= payable]
        av_status = "DETERMINED" if not av_blockers else "BLOCKED_BY_COMPONENTS"
        for field in fields:
            if field["column"] == "AV":
                field["current_system_status"] = av_status
        denominator = len(direct_components)
        return {
            "version": "AV_SOURCE_MAP/v1",
            "run_id": run_id or None,
            "period": (run or {}).get("period_label") or (run or {}).get("period"),
            "template": template,
            "av_formula": template.get("formula") or AV_SOURCE_DEFINITIONS["AV"]["relationship"],
            "direct_components": list(direct_components),
            "upstream_dependencies": {"AK": ["AH", "AI", "AJ"]},
            "components": list(direct_components),
            "av_status_model": {"state": av_status, "blocked_by": av_blockers, "direct_components": list(direct_components)},
            "full_payroll_completeness": {"complete": complete_count, "total": denominator, "label": f"{complete_count} / {denominator}"},
            "fields": fields,
        }

    def save_policy_version(self, effective_from: str, effective_to: str, source: str, profiles: list[dict], supersedes_version_id: str | None = None, source_hash: str = "") -> list[dict]:
        from math import isfinite
        from .core_flow import valid_period
        if not valid_period(effective_from) or not valid_period(effective_to):
            raise ValueError("请填写有效政策月份。")
        if len(effective_from) != 7 or len(effective_to) != 7 or effective_from > effective_to or not source.strip() or not profiles:
            raise ValueError("请填写生效期、来源和至少一位教师的工资政策。")
        cleaned = []
        for item in profiles:
            teacher, role = str(item.get("teacher", "")).strip(), str(item.get("role", "")).strip()
            if not teacher or not role:
                raise ValueError("每条工资政策必须包含教师和身份。")
            if any(p["teacher"] == teacher for p in cleaned):
                raise ValueError("同一政策版本不能重复配置教师。")
            if "obligation_hours" not in item or not isinstance(item.get("obligation_hours_deduction_enabled"), bool):
                raise ValueError("必须明确填写义务小时及是否扣除，缺少政策不能默认免扣。")
            try:
                obligation = float(item["obligation_hours"])
            except (TypeError, ValueError) as exc:
                raise ValueError("义务课时必须是非负有限数值。") from exc
            if not isfinite(obligation) or obligation < 0:
                raise ValueError("义务课时必须是非负有限数值。")
            employment = item.get("employment_type", "FULL_TIME")
            if employment not in {"FULL_TIME", "PART_TIME"} or not isinstance(item.get("allow_no_teaching", False), bool):
                raise ValueError("请明确选择全职/兼职及是否允许当月无课。")
            cleaned.append({"teacher": teacher, "teacher_id": str(item.get("teacher_id", "")).strip(), "role": role, "policy_type": str(item.get("policy_type", "PERSONAL_POLICY")), "rating": item.get("rating"), "rating_override": item.get("rating_override"), "special_approval": str(item.get("special_approval", "")).strip(), "obligation_hours": float(item.get("obligation_hours", 0)), "obligation_hours_deduction_enabled": bool(item.get("obligation_hours_deduction_enabled", False)), "pricing_mode": item.get("pricing_mode"), "base_rate": item.get("base_rate"), "fixed_rate": item.get("fixed_rate"), "override_rate": item.get("override_rate"), "student_id": str(item.get("student_id", "")).strip(), "class_id": str(item.get("class_id", "")).strip(), "provenance": item.get("provenance") or {}, "note": str(item.get("note", "")).strip()})
            cleaned[-1].update(employment_type=employment, allow_no_teaching=item.get("allow_no_teaching", False))
        if supersedes_version_id:
            prior = self.store.get_policy_version(supersedes_version_id)
            if prior["effective_from"] != effective_from or prior["effective_to"] != effective_to:
                raise ValueError("修正版必须沿用原版本的生效期；请另建新的年度版本。")
        version = {"id": uuid.uuid4().hex[:12], "effective_from": effective_from, "effective_to": effective_to, "source": source.strip(), "source_hash": source_hash.strip(), "profiles": cleaned, "created_at": datetime.now(timezone.utc).isoformat(), "status": "ACTIVE", "supersedes_version_id": supersedes_version_id}
        self.store.save_policy_version(version)
        if supersedes_version_id:
            self._supersede("policy", supersedes_version_id, version["id"])
        return self.policy_versions()

    def _versions(self, kind: str) -> list[dict]:
        return self.store.list_rating_versions() if kind == "rating" else self.store.list_policy_versions()

    def _get_version(self, kind: str, version_id: str) -> dict:
        return self.store.get_rating_version(version_id) if kind == "rating" else self.store.get_policy_version(version_id)

    def _save_version(self, kind: str, version: dict) -> None:
        (self.store.save_rating_version if kind == "rating" else self.store.save_policy_version)(version)

    def _supersede(self, kind: str, prior_id: str, replacement_id: str) -> None:
        prior = self._get_version(kind, prior_id)
        prior.update({"status": "SUPERSEDED", "superseded_by": replacement_id, "superseded_at": datetime.now(timezone.utc).isoformat()})
        self._save_version(kind, prior)

    def _version_views(self, kind: str) -> list[dict]:
        binding_key = f"{kind}_version_id"
        views = []
        for version in self._versions(kind):
            view = dict(version)
            view.setdefault("status", "ACTIVE")
            view["used_by_runs"] = [
                {"id": run["id"], "period": run["period"]}
                for run in self.store.list() if run.get(binding_key) == version["id"]
            ]
            views.append(view)
        return views

    def authority_catalog(self) -> dict:
        return {
            "ratings": self.rating_versions(),
            "policies": self.policy_versions(),
            "rules": [{
                "id": "default-compensation-bands",
                "name": "现行档位金额规则",
                "source": "skill/payroll/references/ae_tier_rules.md",
                "source_version": "2025-10",
                "effective_from": "2025-10",
                "effective_to": "持续维护",
                "editable": False,
            }],
        }

    def rebind_authority(self, run_id: str, kind: str, version_id: str) -> dict:
        if kind not in {"rating", "policy"}:
            raise ValueError("未知基础资料类型。")
        run = self._load(run_id)
        self._require_fresh(run)
        version = self._get_version(kind, version_id)
        if not version["effective_from"] <= run["period"] <= version["effective_to"]:
            raise ValueError("该版本不适用于当前核算月份。")
        key = f"{kind}_version_id"
        if run.get(key) == version_id:
            return self.render(run)
        run.setdefault("authority_rebind_history", []).append({"kind": kind, "from_version_id": run.get(key), "to_version_id": version_id, "changed_at": datetime.now(timezone.utc).isoformat()})
        run[key] = version_id
        invalidate_business_decisions(run.setdefault("business_decisions", []))
        run["business_context_stale"] = True
        run["status"] = "FILES_READY" if self._materials_ready(run) else "DRAFT"
        self.store.save(run)
        return self.render(run)

    # --------------------------------------------------------- class type rules
    # 班型折算来自配置，不是 Python 分支：新增班型只需新增规则版本。
    def class_type_rule_versions(self) -> list[dict]:
        stored = self.store.list_class_type_rule_versions()
        if not stored:
            for version in default_rule_versions():
                self.store.save_class_type_rule_version({
                    "id": version.id, "source": version.source, "effective_from": version.effective_from,
                    "effective_to": version.effective_to, "rules": dict(version.rules), "status": version.status,
                    "created_at": version.created_at, "notes": version.notes,
                })
            stored = self.store.list_class_type_rule_versions()
        return sorted(stored, key=lambda item: item.get("effective_from", ""), reverse=True)

    def save_class_type_rule_version(self, *, effective_from: str, effective_to: str, rules: dict, source: str, actor: str, notes: str = "", supersedes_version_id: str | None = None) -> list[dict]:
        """Editing a coefficient always creates a new version; old ones stay."""
        special = self._validate_class_rules(rules)
        if effective_from > effective_to:
            raise ValueError("生效开始时间不能晚于结束时间。")
        # The version being superseded is about to be closed, so it must not
        # count as an overlap; every other active version must stay untouched.
        overlapping = [
            item for item in self.class_type_rule_versions()
            if item.get("status") == "ACTIVE" and item["id"] != supersedes_version_id
            and item["effective_from"] <= effective_to and effective_from <= item["effective_to"]
        ]
        if overlapping:
            raise ValueError("生效时间与已有规则版本重叠，请先确认旧版本的结束时间。")
        self.store.save_class_type_rule_version({
            "id": f"CLASS_TYPE_RULE_{uuid.uuid4().hex[:8]}", "source": source.strip(), "actor": actor.strip(),
            "effective_from": effective_from, "effective_to": effective_to,
            "rules": structured_rules(special),
            "status": "ACTIVE", "created_at": datetime.now(timezone.utc).isoformat(timespec="seconds"), "notes": notes.strip(),
            "supersedes": supersedes_version_id or "",
        })
        if supersedes_version_id:
            prior = next((item for item in self.store.list_class_type_rule_versions() if item["id"] == supersedes_version_id), None)
            if prior is not None:
                # The old version is closed, never rewritten: months it already
                # covered keep resolving to it.
                prior["status"] = "SUPERSEDED"
                prior["effective_to"] = _day_before(effective_from)
                self.store.save_class_type_rule_version(prior)
        return self.class_type_rule_versions()

    @staticmethod
    def _validate_class_rules(rules: dict) -> dict[str, dict[int, float]]:
        """Validate the legacy editor's special-class version shape.

        Only special classes remain configurable here.  One-to-one and ordinary
        small groups are stable Core rules, so accepting either as external
        coefficients would create a second authority path.
        """
        if not rules:
            raise ValueError("请至少配置一个特殊班型的实到人数系数。")
        if SPECIAL_RULES_KEY not in rules:
            raise ValueError("特殊班型必须按“班型 → 实到人数 → 系数”配置。")
        raw_special = rules.get(SPECIAL_RULES_KEY) or {}
        try:
            special = normalize_special(raw_special)
        except ValueError as exc:
            raise ValueError(str(exc)) from None
        for name, values in special.items():
            if not name:
                raise ValueError("班型名称不能为空。")
            if name in SMALL_GROUP_CLASS_TYPES:
                raise ValueError("普通小班按主核实到人数折算，不能配置外部系数。")
            if name == "1对1":
                raise ValueError("一对一走主核 AA 逻辑，不能配置班型系数。")
            if not values or any(attended < 1 or value <= 0 for attended, value in values.items()):
                raise ValueError(f"特殊班型“{name}”必须填写大于 0 的实到人数系数。")
        if not special:
            raise ValueError("请至少配置一个特殊班型的实到人数系数。")
        return special

    def _class_type_rules_for_run(self, run: dict) -> tuple[dict[str, object], str]:
        bound = run.get("class_type_rule_version_id")
        versions = self.class_type_rule_versions()
        if bound:
            version = next((item for item in versions if item["id"] == bound), None)
            if version is not None:
                return dict(version.get("rules", {})), bound
        version = rule_version_for_period(versions, run["period"])
        if version is None:
            return {}, ""
        return dict(version.get("rules", {})), version["id"]

    def rebind_class_type_rules(self, run_id: str, version_id: str) -> dict:
        run = self._load(run_id)
        if run.get("calculation_engine") == "CONFIGURED_V1":
            raise ValueError("本记录使用完整核心规则，请从核心规则面板创建并绑定新版本。")
        version = next((item for item in self.class_type_rule_versions() if item["id"] == version_id), None)
        if version is None:
            raise ValueError("未找到该班型折算规则版本。")
        run.setdefault("authority_rebind_history", []).append({"kind": "class_type_rules", "from_version_id": run.get("class_type_rule_version_id"), "to_version_id": version_id, "changed_at": datetime.now(timezone.utc).isoformat()})
        run["class_type_rule_version_id"] = version_id
        run["business_context_stale"] = True
        self.store.save(run)
        return self.render(run)

    # ----------------------------------------------------------- generate mode
    def default_export_path(self, filename: str = "标准工资表.xlsx") -> dict:
        target = default_output_dir(filename)
        return {"path": str(target), "location": describe_location(target), "exists": target.exists()}

    def generate_payroll(self, run_id: str, output_path: str, *, confirmed_hours: dict | None = None, production: bool = False) -> dict:
        """生成模式：同一套 Core 结果直接渲染成标准工资表。"""
        run = self._load(run_id)
        if not str(output_path or "").strip():
            output_path = str(default_output_dir(f"工资表-{run['period']}.xlsx"))
        if self._bind_template_from_run_files(run):
            self.store.save(run)
        if production and not run.get("template_path"):
            raise ValueError("尚未设置公司工资模板，请先到基础资料中设置一次。")
        business_inputs = self._approved_business_inputs(run)
        if run.get("calculation_engine") == "CONFIGURED_V1":
            if confirmed_hours:
                raise ValueError("AD 已由 AA + AC 独立计算，不接受手工或工资表 AD 覆盖。")
            checked = self.check(run_id)
            payroll = self._generated_from_checked(checked)
            payroll = self._apply_generation_guardrails(payroll, self.store.get(run_id))
            self._ensure_production_export_allowed(payroll, production)
            path = render_generated_payroll(payroll, safe_output_path(output_path), template_path=run.get("template_path"), manual_adjustments=business_inputs)
            run = self.store.get(run_id)
            run["generated_payroll"] = {"path": path, "status": payroll.status, "blockers": list(payroll.blockers), "rule_versions": dict(payroll.rule_versions), "created_at": datetime.now(timezone.utc).isoformat(), "rows": [asdict(row) for row in payroll.rows]}
            self.store.save(run)
            return run["generated_payroll"]
        missing = self._missing_materials(run)
        if missing:
            raise ValueError("请先导入：" + "、".join(missing))
        self._require_fresh(run)
        coefficients, rule_version_id = self._class_type_rules_for_run(run)
        reads = self._read_for_run("schedule", Path(run["files"]["schedule"]["path"]), run)
        if reads.errors:
            raise ValueError("排课数据重新读取失败，请返回材料页重新选择。")
        schedule = self._apply_schedule_grade_resolutions(reads.records, run)
        rating_version = self._rating_version_for_run(run)
        ratings = {item["teacher"]: item.get("rating", "") for item in (rating_version or {}).get("ratings", [])}
        hours = dict(confirmed_hours if confirmed_hours is not None else run.get("confirmed_hours", {}))
        if confirmed_hours is not None:
            run["confirmed_hours"] = {key: float(value) for key, value in confirmed_hours.items()}
        payroll = build_legacy_generated_payroll(
            period=run["period"], schedule_records=schedule, coefficients=coefficients,
            ratings_by_teacher=ratings, confirmed_hours=hours,
            rule_versions={"class_type_rules": rule_version_id, "rating": (rating_version or {}).get("id", "")},
            blocked=bool(run.get("business_context_stale")),
            base_salary_inputs=run.get("base_salary_inputs") if run.get("base_salary_input_snapshot") else None,
            renewal_snapshot=run.get("run_renewal_result_snapshot"),
        )
        payroll = self._apply_generation_guardrails(payroll, run)
        self._ensure_production_export_allowed(payroll, production)
        path = render_generated_payroll(payroll, safe_output_path(output_path), template_path=run.get("template_path"), manual_adjustments=business_inputs)
        run["generated_payroll"] = {
            "path": path, "status": payroll.status, "blockers": list(payroll.blockers),
            "rule_versions": dict(payroll.rule_versions), "created_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
            "rows": [asdict(row) for row in payroll.rows],
        }
        self.store.save(run)
        return {"path": path, "status": payroll.status, "blockers": list(payroll.blockers), "rows": [asdict(row) for row in payroll.rows], "rule_versions": dict(payroll.rule_versions)}

    @staticmethod
    def _ensure_production_export_allowed(payroll, production: bool) -> None:
        if not production:
            return
        if "PERIOD_COVERAGE_INCOMPLETE" in payroll.blockers:
            raise ValueError("排课周期不完整，请补齐后再生成最终工资表。")
        if "PERIOD_MATERIAL_CONFLICT" in payroll.blockers or "PERIOD_NEEDS_CONFIRMATION" in payroll.blockers:
            raise ValueError("材料所属月份尚未确认，确认后再生成最终工资表。")

    def _approved_business_inputs(self, run: dict) -> tuple[dict, ...]:
        """Return only current, explicitly bound inputs eligible for output."""
        values: list[dict] = []
        for binding in run.get("business_input_bindings", []):
            item = self.store.get_business_input(binding["input_id"])
            if not is_active_business_input(item):
                continue
            if not self.inputs.current(item["id"], binding.get("source_file_hash", "")):
                continue
            values.append(item)
        return tuple(values)

    def preview_payroll(self, run_id: str) -> dict:
        """核算并返回完整工资预览，不创建或修改任何输出工作簿."""
        run = self._load(run_id)
        if run.get("calculation_engine") != "CONFIGURED_V1":
            raise ValueError("当前核算记录尚未启用完整工资计算链。")
        checked = self.check(run_id)
        return self._with_generated_preview(checked)

    def historical_reconciliation(self, run_id: str) -> dict:
        """Build an auditable, read-only July-style historical ledger.

        The historical workbook remains an oracle only.  Current values come
        from the run's configured calculation and every difference carries
        source/formula evidence; no value is rounded into a match.
        """
        run = self._load(run_id)
        self._require_fresh(run)
        baseline_item = run.get("files", {}).get("baseline")
        core = run.get("core_calculation") or {}
        if not baseline_item or not core.get("rows"):
            raise ValueError("当前记录缺少可用于历史工资对账的基准工资表或核心计算结果。")
        baseline_result = self._read_for_run("baseline", Path(baseline_item["path"]), run)
        historical = {row.teacher: row for row in baseline_result.records if getattr(row, "teacher", "")}
        historical_comments = {}
        for comment in getattr(baseline_result, "comments", []) or []:
            if comment.target in historical:
                historical_comments.setdefault((comment.target, str(comment.field).upper()), []).append(comment)
        current = {row.get("teacher"): row for row in core.get("rows", []) if row.get("teacher")}
        common = sorted(set(historical) & set(current))
        schedule_result = self._read_for_run("schedule", Path(run["files"]["schedule"]["path"]), run) if run.get("files", {}).get("schedule") else None
        schedule = self._apply_schedule_grade_resolutions(schedule_result.records, run) if schedule_result else []
        by_key = {course_record_key(record): record for record in schedule}
        contributions = [item for item in core.get("course_contributions", []) if item.get("teacher") in common]
        fields = (("AA", "one_to_one"), ("AC", "class_value"), ("AD", "teaching_hours"), ("AE", "ae"), ("AF", "af"))
        stats = {field: {"matches": 0, "comparable": 0} for field, _ in fields}
        differences = []
        category_counts = {name: 0 for name in ("INPUT_SCOPE", "MISSING_SOURCE", "PART_TIME_RATE", "COURSE_CONTRIBUTION", "PERSONAL_EXCEPTION", "RULE_DIFFERENCE", "DATA_QUALITY", "UNEXPLAINED")}
        for teacher in common:
            old = historical[teacher]
            new = current[teacher]
            field_output = {}
            teacher_categories = []
            teacher_evidence = []
            for field, old_name in fields:
                historical_value = getattr(old, old_name, None)
                current_value = (new.get("fields", {}).get(field) or {}).get("value")
                difference = None if historical_value is None or current_value is None else round(float(current_value) - float(historical_value), 6)
                comparable = historical_value is not None and current_value is not None
                if comparable:
                    stats[field]["comparable"] += 1
                    if abs(difference or 0) <= 1e-6:
                        stats[field]["matches"] += 1
                if difference is None or abs(difference) <= 1e-6:
                    continue
                category, evidence = self._historical_difference_evidence(
                    field, old, new, difference, contributions, by_key, run,
                    historical_comments.get((teacher, field), []),
                )
                category_counts[category] += 1
                teacher_categories.append(category)
                teacher_evidence.extend(evidence)
                field_output[field] = {
                    "historical": historical_value,
                    "current": current_value,
                    "diff": difference,
                    "status": "DIFFERENCE",
                    "difference_category": category,
                    "evidence": evidence,
                }
            if field_output:
                primary = next((item for item in ("PART_TIME_RATE", "MISSING_SOURCE", "PERSONAL_EXCEPTION", "COURSE_CONTRIBUTION", "RULE_DIFFERENCE", "DATA_QUALITY", "INPUT_SCOPE", "UNEXPLAINED") if item in teacher_categories), "DATA_QUALITY")
                differences.append({"teacher": teacher, "fields": field_output, "status": "DIFFERENCE", "difference_category": primary, "evidence": teacher_evidence})
        source_classifications = {name: [] for name in ("PART_TIME_RATE", "PERSONAL_POLICY", "FIXED_HISTORICAL_VALUE", "MISSING_SOURCE")}
        for row in differences:
            for field in row["fields"].values():
                for item in field.get("evidence", []):
                    classification = item.get("source_classification")
                    if classification in source_classifications and row["teacher"] not in source_classifications[classification]:
                        source_classifications[classification].append(row["teacher"])
        for teachers in source_classifications.values():
            teachers.sort()
        return {
            "run_id": run["id"],
            "period_label": run.get("period_label", run.get("period")),
            "period_start": run.get("period_start"),
            "period_end": run.get("period_end"),
            "historical_source": baseline_item.get("path"),
            "current_source": run.get("files", {}).get("schedule", {}).get("path"),
            "teachers_compared": len(common),
            "field_stats": stats,
            "difference_teachers": len(differences),
            "category_counts": category_counts,
            "source_classifications": source_classifications,
            "unexplained": category_counts["UNEXPLAINED"],
            "rows": differences,
        }

    @staticmethod
    def _historical_obligation(formula: object) -> float | None:
        text = str(formula or "").upper().replace("$", "")
        if re.search(r"AD\d*\s*\*\s*AE", text):
            return 0.0
        match = re.search(r"AD\d*\s*[-−]\s*(\d+(?:\.\d+)?)", text)
        return float(match.group(1)) if match else None

    def _historical_difference_evidence(self, field: str, old: object, new: dict, difference: float, contributions: list[dict], by_key: dict, run: dict, comments: list[object] | None = None) -> tuple[str, list[dict]]:
        old_provenance = getattr(old, "provenance", {}) or {}
        af_fact = old_provenance.get("af")
        historical_formula = getattr(af_fact, "raw_value", "") if af_fact else ""
        current_fact = (new.get("fields", {}).get(field) or {})
        evidence: list[dict] = [{"source": getattr(old, "source", ""), "historical_formula": historical_formula, "difference": difference}]
        if field in {"AC", "AD"}:
            course_rows = []
            for item in contributions:
                if item.get("teacher") != new.get("teacher") or item.get("field") != "ac" or item.get("value") is None:
                    continue
                record = by_key.get(item.get("record_key"))
                if record is None:
                    continue
                locations = [value for value in (getattr(record, "provenance", {}) or {}).values() if getattr(value, "coordinate", "")]
                inputs = (item.get("evidence") or [{}])[0].get("inputs", {})
                course_rows.append({"record_key": item.get("record_key"), "date": getattr(record, "lesson_date", ""), "class_name": getattr(record, "class_name", ""), "student": getattr(record, "student", ""), "lesson_status": getattr(record, "lesson_status", ""), "attended": getattr(record, "attended", None), "grade": getattr(record, "grade", ""), "class_type": getattr(record, "class_type", ""), "coefficient": inputs, "contribution": item.get("value"), "source_row": ", ".join(f"{getattr(value, 'sheet', '')}!{getattr(value, 'coordinate', '')}" for value in locations)})
            evidence.append({"course_contributions": course_rows, "course_contribution_count": len(course_rows), "note": "当前 AC/AD 差异由逐课贡献与历史工资表字段对照；历史逐课公式未存入工资表，保留原始 AC 公式作为依据。"})
            return "COURSE_CONTRIBUTION", evidence
        if field == "AF":
            part_time = self._historical_part_time_policy(old, comments or [], run)
            if part_time is not None:
                evidence.append(part_time)
                return "PART_TIME_RATE", evidence
            historical_obligation = self._historical_obligation(historical_formula)
            current_obligation = None
            for item in current_fact.get("evidence", []):
                try:
                    current_obligation = float((item.get("inputs") or {}).get("obligation_hours"))
                    break
                except (TypeError, ValueError):
                    continue
            if historical_obligation is not None and current_obligation is not None and abs(historical_obligation - current_obligation) > 1e-6:
                evidence.append({"historical_obligation_hours": historical_obligation, "current_obligation_hours": current_obligation, "historical_formula": historical_formula, "current_formula": current_fact.get("reason", "")})
                return "PERSONAL_EXCEPTION", evidence
            if historical_obligation is not None and abs(float(getattr(old, "teaching_hours", 0) or 0) - float((new.get("fields", {}).get("AD") or {}).get("value") or 0)) > 1e-6:
                evidence.append({"historical_obligation_hours": historical_obligation, "current_obligation_hours": current_obligation, "reason": "AF 差异由 AD 课程贡献差异传导。"})
                return "COURSE_CONTRIBUTION", evidence
            if historical_obligation is None:
                evidence.append({"source_classification": "MISSING_SOURCE", "reason": "历史 AF 公式未引用 AD/AE，且没有可绑定的兼职/个人政策证据；该教师历史来源缺少可复算字段。", "historical_formula": historical_formula})
                return "MISSING_SOURCE", evidence
            evidence.append({"historical_obligation_hours": historical_obligation, "current_obligation_hours": current_obligation, "historical_formula": historical_formula, "current_formula": current_fact.get("reason", "")})
            return "RULE_DIFFERENCE", evidence
        evidence.append({"reason": "历史字段与当前版本化规则结果不同，保留双方来源供后续规则对账。", "current_reason": current_fact.get("reason", "")})
        return "RULE_DIFFERENCE", evidence

    @staticmethod
    def _historical_part_time_policy(old: object, comments: list[object], run: dict) -> dict | None:
        """Extract an explicit historical per-lesson policy from a payroll comment.

        A fixed AF formula alone is not enough to infer employment type.  The
        July workbook's AF comments explicitly record a per-lesson rate and
        count; require both that wording and a matching ``rate * count``
        formula before classifying the source as PART_TIME_RATE.
        """
        formula = str(getattr((getattr(old, "provenance", {}) or {}).get("af"), "raw_value", "") or "")
        factors = re.search(r"=\s*(\d+(?:\.\d+)?)\s*\*\s*(\d+(?:\.\d+)?)", formula)
        if not factors:
            return None
        rate, lessons = float(factors.group(1)), float(factors.group(2))
        for comment in comments:
            text = str(getattr(comment, "text", "") or "")
            rate_match = re.search(r"(\d+(?:\.\d+)?)\s*(?:元\s*)?/\s*节", text)
            if not rate_match or abs(float(rate_match.group(1)) - rate) > 1e-6 or "节" not in text:
                continue
            source_file = str(getattr(comment, "source_file", "") or "")
            source_sheet = str(getattr(comment, "sheet", "") or "")
            coordinate = str(getattr(comment, "coordinate", "") or "")
            return {
                "source_classification": "PART_TIME_RATE",
                "employment_type": "PART_TIME",
                "teacher": getattr(old, "teacher", ""),
                "period": run.get("period_label", run.get("period", "")),
                "rate": rate,
                "unit": "CNY_PER_LESSON",
                "lesson_count": lessons,
                "formula": formula,
                "source": source_file,
                "source_cell": f"{source_sheet}!{coordinate}" if source_sheet or coordinate else source_file,
                "source_text": text,
                "author": getattr(comment, "author", None),
            }
        return None

    def _with_generated_preview(self, checked: dict) -> dict:
        """Attach the canonical final-field preview without writing a file."""
        if not checked.get("core_calculation"):
            return checked
        payroll = self._generated_from_checked(checked)
        payroll = self._apply_generation_guardrails(payroll, checked)
        return {**checked, "generated_payroll": {"status": payroll.status, "blockers": list(payroll.blockers), "rule_versions": dict(payroll.rule_versions), "rows": [asdict(row) for row in payroll.rows]}}

    @staticmethod
    def _apply_generation_guardrails(payroll, run: dict):
        """Prevent an incomplete/mismatched period from being advertised final.

        Core values remain available for review.  The guard only changes the
        generated artifact status and blocker list; it never changes AA/AC or
        any payroll formula.
        """
        check = run.get("period_check") or {}
        blockers = set(payroll.blockers)
        coverage = check.get("coverage") or {}
        if coverage.get("incomplete_tail"):
            blockers.add("PERIOD_COVERAGE_INCOMPLETE")
        if check.get("conflict"):
            blockers.add("PERIOD_MATERIAL_CONFLICT")
        if check.get("mismatch") and check.get("decision") not in {"SWITCHED", "KEPT"}:
            blockers.add("PERIOD_NEEDS_CONFIRMATION")
        if not blockers:
            return payroll
        return replace(payroll, status="NEEDS_CONFIRMATION", blockers=tuple(sorted(blockers)))

    def _generated_from_checked(self, checked: dict):
        """Build the one canonical final-field result used by preview/export."""
        bindings = checked.get("business_input_bindings", [])
        inputs = [
            item for binding in bindings
            for item in (self.store.get_business_input(binding["input_id"]),)
            if is_active_business_input(item)
        ]
        # An empty mapping is deliberate: it asks the shared generator to
        # expose M as MISSING rather than omitting it or treating it as zero.
        if "base_salary_input_snapshot" in checked:
            base_salary = checked.get("base_salary_inputs") if checked.get("base_salary_input_snapshot") else {}
        else:
            # Minimal calculation fixtures from older callers do not include
            # production Run input metadata; preserve their historical final
            # field semantics.
            base_salary = None
        return generated_from_calculation(checked["core_calculation"], business_inputs=inputs, base_salary_inputs=base_salary, renewal_snapshot=checked.get("run_renewal_result_snapshot"))

    def writeback_to_generated(self, run_id: str, candidate_ids: list[str], output_path: str, reviewer: str, strategy: str = "APPEND") -> dict:
        """生成模式复用同一套批注机制：同样的预览与回填函数，没有第二套逻辑。"""
        run = self._load(run_id)
        generated = (run.get("generated_payroll") or {}).get("path", "")
        if not generated:
            raise ValueError("请先生成标准工资表。")
        rows = {item["teacher"]: index + 4 for index, item in enumerate(sorted((run.get("generated_payroll") or {}).get("rows", []), key=lambda row: row["teacher"]))}
        items = []
        for candidate_id in candidate_ids:
            item = self.store.get_comment_candidate(candidate_id)
            self._refresh_candidate(item)
            if item["status"] != "APPROVED":
                raise ValueError("只有已预览确认的批注候选才能写入。")
            if item["teacher_id"] not in rows:
                raise ValueError("目标教师不在本次生成的工资表里，不能写入其他教师单元格。")
            if item["sheet"] != "标准工资表" or item["cell"] != f"A{rows[item['teacher_id']]}":
                raise ValueError("生成工资表的批注只能写在该教师自己的行上。")
            items.append(item)
        actual_path = safe_output_path(output_path)
        result = write_new_workbook(generated, actual_path, items, strategy=strategy, author=reviewer.strip() or "工资核算助手")
        when = datetime.now(timezone.utc).isoformat(timespec="seconds")
        for item in items:
            item.update({"status": "WRITTEN", "written_at": when, "output_workbook": str(actual_path.resolve())})
            self.store.save_comment_candidate(item)
        run.setdefault("writeback_history", []).append({"source_workbook": generated, "output_workbook": str(actual_path.resolve()), "candidate_ids": candidate_ids, "written_by": reviewer.strip(), "written_at": when})
        self.store.save(run)
        return {"output_path": str(Path(output_path).resolve()), "written": len(result), "candidates": items}

    # ------------------------------------------------------------------ mapping
    # A layout mismatch is not an error: the engine maps business fields onto
    # whatever columns the file actually has, and only asks when it cannot tell.
    def import_requirement(self, role: str = "schedule"):
        return SCHEDULE_AC_REQUIREMENT if role in {"schedule", "check"} else None

    def preview_import_mapping(self, path: str, role: str = "schedule", period: str = "") -> dict:
        source = Path(path).expanduser().resolve()
        if not source.is_file():
            raise ValueError("找不到原始文件。请关闭 Excel/WPS 后重新选择。")
        requirement = self.import_requirement(role)
        if requirement is None:
            raise ValueError("该材料类别暂不支持字段映射。")
        analysis = analyze_mapping(source, requirement, profiles=self.store.list_import_profiles(requirement.name))
        if role == "schedule":
            known = read_schedule_excel(source, period or "0000-00")
            known_ok = known.ok and bool(known.records)
        else:
            known_ok = False
        fields = []
        for item in requirement.fields:
            header = analysis.mapped_headers.get(item.field, "")
            candidates = list(analysis.candidates.get(item.field, ()))
            fields.append({
                "field": item.field, "label": item.label, "required": item.required,
                "column": analysis.mapping.get(item.field),
                "header": header,
                "candidates": candidates,
                "needs_choice": item.required and not header,
            })
        return {
            "status": "KNOWN_LAYOUT" if known_ok else analysis.status,
            "requirement": requirement.name,
            "sheet": analysis.sheet,
            "header_row": analysis.header_row,
            "fingerprint": analysis.fingerprint,
            "detected_columns": list(analysis.detected_columns),
            "columns": [{"column": column, "header": text} for column, text in sorted(analysis.column_headers.items())],
            "fields": fields,
            "missing_labels": [requirement.field(name).label for name in analysis.missing],
            "message": analysis.message,
            "profile_id": analysis.profile_id,
            "profile_drift": analysis.profile_drift,
            "mapping": dict(analysis.mapping),
        }

    def save_import_profile(self, path: str, role: str, mapping: dict, actor: str, name: str = "") -> dict:
        requirement = self.import_requirement(role)
        if requirement is None:
            raise ValueError("该材料类别暂不支持字段映射。")
        analysis = analyze_mapping(Path(path).expanduser().resolve(), requirement)
        if not analysis.fingerprint:
            raise ValueError("无法识别这份文件的表头，不能保存格式。")
        profile = {
            "id": uuid.uuid4().hex[:12], "requirement": requirement.name, "name": name or f"{Path(path).name} 已确认格式",
            "fingerprint": analysis.fingerprint, "mapping": mapping,
            "headers": {field: analysis.column_headers.get(int(column), "") for field, column in mapping.items()},
            "sheet": analysis.sheet,
            "header_row": analysis.header_row, "created_by": actor.strip(), "created_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        }
        self.store.save_import_profile(profile)
        return profile

    def _mapping_error(self, analysis) -> str:
        """Explain exactly what is missing and what was detected, with a way forward."""
        lines: list[str] = []
        if getattr(analysis, "profile_drift", False):
            lines.append("这份文件用了已确认过的格式，但某些列已经变化，需要重新确认字段。")
        requirement = SCHEDULE_AC_REQUIREMENT
        if analysis.candidates:
            for field, options in analysis.candidates.items():
                lines.append(f"字段“{requirement.field(field).label}”有多个候选：{'、'.join(options)}，需要你确认用哪一列。")
        if analysis.missing:
            labels = "、".join(f"{requirement.field(name).label}（{name}）" for name in analysis.missing)
            lines.append(f"缺少必要字段：{labels}。")
        if analysis.detected_columns:
            lines.append("检测到的列：" + "、".join(analysis.detected_columns) + "。")
        lines.append("可以在“字段识别”里指定每一列代表哪个业务字段，确认后即可继续导入。")
        return "\n".join(lines)

    def import_file(self, run_id: str, role: str, path: str, expected_hash: str | None = None, mapping: dict | None = None, profile_name: str = "", profile_actor: str = "") -> dict:
        if role not in LABELS:
            raise ValueError("未知材料类别。")
        source = Path(path).expanduser().resolve()
        if not source.is_file():
            raise ValueError("找不到原始文件。请关闭 Excel/WPS 后重新选择。")
        try:
            before = version(source)
        except OSError as exc:
            raise ValueError(self._file_error(exc)) from exc
        if expected_hash and expected_hash != before["sha256"]:
            raise ValueError("所选原文件与刚才拖放识别的文件不一致。")
        run = self._load(run_id)
        if any(key != role and item["sha256"] == before["sha256"] for key, item in run["files"].items()):
            raise ValueError("同一份文件不能同时充当两个材料类别。")
        is_csv = source.suffix.lower() == ".csv"
        workbook = None
        if not is_csv:
            inspection = inspect_workbook(source)
            mapping_capable = self.import_requirement(role) is not None
            # An unrecognized layout is no longer a dead end for mapping-capable
            # roles: the semantic engine takes over instead of refusing the file.
            blocking = [issue for issue in inspection.errors if not (mapping_capable and issue.code == "UNSUPPORTED_LAYOUT")]
            if blocking or not inspection.records:
                raise ValueError(self._inspection_error(blocking))
            workbook = inspection.records[0]
            if workbook.fingerprint.layout != LAYOUTS[role] and not mapping_capable:
                raise ValueError(f"文件不属于“{LABELS[role]}”，请检查后重新选择。")
        try:
            if is_csv:
                result = self._read_for_run(role, source, run)
            elif role == "schedule" and mapping_capable:
                profiles = self.store.list_import_profiles(SCHEDULE_AC_REQUIREMENT.name)
                result, analysis = resolve_schedule_import(
                    source, run["period"], profiles=profiles, confirmed=mapping,
                    course_export_snapshots=self._stored_course_export_snapshots(), period_start=self._period_window(run)[0], period_end=self._period_window(run)[1],
                )
                if analysis is not None and not analysis.ready:
                    raise ValueError(self._mapping_error(analysis))
                if not result.errors and not result.records and not any(
                    issue.code == "OUT_OF_PERIOD_ROWS_EXCLUDED" for issue in result.warnings
                ):
                    raise ValueError("这份排课表没有识别到有效课程记录，请检查是否选错了工作表或文件。")
            else:
                result = self._read_for_run(role, source, run)
        except OSError as exc:
            raise ValueError(self._file_error(exc)) from exc
        if is_csv and role == "schedule" and not result.records:
            # An all-out-of-period CSV is still useful month evidence and may
            # be confirmed/switchable.  A CSV with no usable in-period rows
            # for any other reason must not masquerade as a valid schedule.
            if not any(issue.code == "OUT_OF_PERIOD_ROWS_EXCLUDED" for issue in result.warnings):
                message = next(
                    (issue.message for issue in result.warnings if issue.code == "UNPARSEABLE_LESSON_DATE"),
                    "这份排课表没有识别到有效课程记录，请检查日期和文件格式。",
                )
                raise ValueError(message)
        if result.errors:
            raise ValueError("文件缺少当前核对所需字段：" + "；".join(issue.message for issue in result.errors))
        try:
            unchanged = version(source) == before
        except OSError as exc:
            raise ValueError(self._file_error(exc)) from exc
        if not unchanged:
            raise ValueError("文件在读取期间发生变化，请关闭 Excel/WPS 后重试。")
        if role == "schedule":
            # Persist only course-version snapshots here.  The current payroll
            # source must not become student history used to justify itself.
            self._save_course_export_snapshots(result.records, before)
        if mapping and profile_name:
            # Remember the confirmed layout so next month's identical file imports
            # without asking again. Drift is still re-checked on every import.
            self.save_import_profile(str(source), role, dict(mapping.get("mapping", {})), profile_actor, profile_name)
        run["files"][role] = {"name": source.name, "path": str(source), **before, "label": LABELS[role], "records": len(result.records), "teachers": len({record.teacher for record in result.records if hasattr(record, "teacher")}), "warnings": [issue.code for issue in result.warnings], "warning_messages": [self._warning_message(issue.code) for issue in result.warnings], "sheets": [sheet.name for sheet in workbook.sheets] if workbook is not None else ["CSV"], "formula_count": sum(sheet.formula_count for sheet in workbook.sheets) if workbook is not None else 0, "missing_cache": sum(sheet.formula_cache_missing for sheet in workbook.sheets) if workbook is not None else 0, "external_references": workbook.external_link_count + sum(sheet.external_formula_count for sheet in workbook.sheets) if workbook is not None else 0}
        if role == "schedule":
            run["period_check"] = self._period_evidence(
                run, result.records, source.name, source,
                evidence_dates=result.coverage.get("all_lesson_dates", ()) if is_csv else (),
            )
            run.setdefault("material_period_evidence", []).append({"role": role, "source_file": source.name, "source_month": run["period_check"].get("source_month"), "basis": "课表实际课程日期"})
            # Grade resolutions are bound to coordinates in the imported
            # schedule workbook. Replacing that source invalidates them.
            run.pop("schedule_grade_resolutions", None)
            for item in run.get("resolutions", []):
                if item.get("status") == "ACTIVE":
                    item["status"] = "NEEDS_RECONFIRMATION"
                    item["invalidated_at"] = datetime.now(timezone.utc).isoformat()
                    run.setdefault("resolution_history", []).append({**item, "history_event": "SOURCE_CHANGED"})
        else:
            hint = self._file_period_hint(run, source, workbook)
            if hint:
                run.setdefault("material_period_evidence", []).append({"role": role, "source_file": source.name, "source_month": hint, "basis": "材料文件名或工作表月份"})
        self._refresh_material_period_check(run)
        # Legacy per-field decisions predate business review cards.  They are
        # cleared for backward compatibility; durable business decisions are
        # retained but explicitly require a fresh confirmation.
        run["issues"], run["field_records"], run["issue_groups"], run["user_actions"], run["decisions"] = [], [], [], [], []
        invalidate_business_decisions(run.setdefault("business_decisions", []))
        run.pop("last_error", None)
        run.pop("stale_files", None)
        run["status"] = "FILES_READY" if self._materials_ready(run) else "DRAFT"
        # Replacing one changed file must not silently clear changes in another.
        self._fresh(run)
        self.store.save(run)
        return self.render(run)

    def _period_evidence(self, run: dict, records, file_name: str, source_path: Path | None = None, *, evidence_dates=()) -> dict:
        # Keep source-month evidence separate from the dates that actually
        # survived the current Run's period filter.  A mixed-month source can
        # identify the file without allowing an out-of-period row to make the
        # selected month's coverage look complete.
        raw_source_dates = list(evidence_dates)
        in_period_dates = [getattr(record, "lesson_date", "") or getattr(record, "lesson_time", "") for record in records]
        if not raw_source_dates and source_path is not None:
            # A source whose rows were filtered out for the selected month
            # still needs month evidence so the UI can offer a safe switch.
            analysis = analyze_mapping(source_path, SCHEDULE_AC_REQUIREMENT)
            if analysis.ready:
                all_rows = read_schedule_with_mapping(
                    source_path, "", mapping=analysis.mapping,
                    requirement=SCHEDULE_AC_REQUIREMENT,
                    sheet_name=analysis.sheet, header_row=analysis.header_row,
                )
                raw_source_dates = [getattr(record, "lesson_date", "") or getattr(record, "lesson_time", "") for record in all_rows.records]
        if not raw_source_dates:
            raw_source_dates = list(in_period_dates)
        raw_months = sorted({month for month in (dominant_month([item]) for item in raw_source_dates) if month})
        if len(raw_months) == 1:
            source_month = raw_months[0]
        elif run["period"] in raw_months:
            # A mixed source is usable for an explicitly selected month when
            # that month has effective rows.  Keep all months visible as
            # evidence, but do not force a switch to the dominant month.
            source_month = run["period"]
        else:
            source_month = None
        file_month, file_has_year = month_from_filename(file_name)
        start, end = self._period_window(run)
        coverage = coverage_for(run["period"], in_period_dates, period_start=start, period_end=end)
        mismatch = bool(source_month and source_month != run["period"])
        filename_disagrees = bool(file_month and source_month and (file_month != source_month if file_has_year else file_month != source_month[5:7]))
        source_conflict = len(raw_months) > 1 and run["period"] not in raw_months
        return {"run_month": run["period"], "period_start": start, "period_end": end,
                "period_boundary_source": run.get("period_boundary_source", "LEGACY_CALENDAR_DEFAULT"),
                "source_month": source_month, "mismatch": mismatch,
                "source_months": raw_months, "source_conflict": source_conflict,
                "source_file": file_name,
                "file_name_month": file_month, "file_name_has_year": file_has_year,
                "filename_disagrees": filename_disagrees, "coverage": coverage.as_dict(),
                "final_generation_blocked": bool(coverage.incomplete_tail) or source_conflict,
                "conflict": source_conflict,
                "decision": "" if not mismatch else "PENDING"}

    @staticmethod
    def _file_period_hint(run: dict, source: Path, workbook=None) -> str | None:
        """Use a workbook title/file name only as secondary month evidence."""
        hinted, has_year = month_from_filename(source.name)
        if hinted:
            return hinted if has_year else f"{run['period'][:4]}-{int(hinted):02d}"
        for sheet in getattr(workbook, "sheets", ()) or ():
            match = re.search(r"(?:20(\d{2})[年-])?(0?[1-9]|1[0-2])月", str(sheet.name))
            if match:
                year = match.group(1) or run["period"][:4]
                return f"20{year}-{int(match.group(2)):02d}"
        return None

    def _refresh_material_period_check(self, run: dict) -> None:
        sources = [item for item in run.get("material_period_evidence", []) if item.get("source_month")]
        months = sorted({item["source_month"] for item in sources})
        if len(months) > 1:
            current = run.get("period_check") or {}
            run["period_check"] = {
                "run_month": run["period"], "source_month": None, "mismatch": False,
                "conflict": True, "decision": "PENDING", "sources": sources,
                "coverage": current.get("coverage", {}),
                "final_generation_blocked": True,
                "message": "不同材料判断出的工资月份不一致，不能自动选择。",
            }
            return
        if not sources:
            return
        source_month = months[0]
        current = run.get("period_check") or {}
        if current.get("source_file") and current.get("coverage"):
            # Preserve the schedule's precise coverage and only add the
            # cross-material evidence to the existing check.
            current["sources"] = sources
            current["conflict"] = False
            run["period_check"] = current
            return
        run["period_check"] = {
            "run_month": run["period"], "source_month": source_month,
            "mismatch": source_month != run["period"], "conflict": False,
            "decision": "PENDING" if source_month != run["period"] else "",
            "sources": sources, "final_generation_blocked": False,
        }

    def change_period(self, run_id: str, period: str) -> dict:
        if not valid_period(period):
            raise ValueError("请选择有效月份。")
        run = self._load(run_id)
        if period == run["period"]:
            return self.render(run)
        run["period"] = period
        window = normalize_period_window(period)
        run.update(window)
        for role, item in run.get("files", {}).items():
            path = Path(item["path"])
            if not path.is_file():
                run.setdefault("stale_files", []).append(role)
                continue
            result = self._read_for_run(role, path, run)
            if result.errors:
                run.setdefault("stale_files", []).append(role)
                continue
            item["records"] = len(result.records)
            item["teachers"] = len({record.teacher for record in result.records if hasattr(record, "teacher")})
            if role == "schedule":
                period_evidence = self._period_evidence(
                    run, result.records, path.name, path,
                    evidence_dates=result.coverage.get("all_lesson_dates", ()) if path.suffix.lower() == ".csv" else (),
                )
                run["period_check"] = period_evidence
                run["material_period_evidence"] = [
                    evidence for evidence in run.get("material_period_evidence", [])
                    if not (evidence.get("role") == "schedule" and evidence.get("source_file") == path.name)
                ]
                run["material_period_evidence"].append({
                    "role": "schedule", "source_file": path.name,
                    "source_month": period_evidence.get("source_month"),
                    "basis": "课表实际课程日期",
                })
        self._refresh_material_period_check(run)
        check = run.get("period_check") or {}
        if check and not check.get("conflict") and not check.get("mismatch"):
            check["decision"] = "SWITCHED"
            run["period_check"] = check
        run["status"] = "FILES_READY" if self._materials_ready(run) else "DRAFT"
        self.store.save(run)
        return self.render(run)

    def resolve_period_check(self, run_id: str, decision: str) -> dict:
        run = self._load(run_id)
        check = run.get("period_check") or {}
        if not check.get("mismatch"):
            return self.render(run)
        if decision == "SWITCH":
            self.change_period(run_id, check["source_month"])
            stored = self._load(run_id)
            stored_check = stored.get("period_check") or {}
            stored_check["decision"] = "SWITCHED"
            stored["period_check"] = stored_check
            self.store.save(stored)
            return self.render(stored)
        if decision == "KEEP":
            check["decision"] = "KEPT"
            run["period_check"] = check
            self.store.save(run)
            return self.render(run)
        raise ValueError("请选择切换到课表月份，或仍按当前月份核算。")

    def create_resolution(self, run_id: str, issue_id: str, kind: str, course_record_id: str, values: dict[str, Any], confirmed_by: str, expected_fingerprint: str) -> dict:
        """Persist one run-scoped AC correction or approved treatment.

        The API deliberately accepts one selected course at a time.  That keeps
        every change and its counterfactual evidence attributable to a stable
        source record.
        """
        run = self._load(run_id)
        self._require_current_audit(run)
        if run.get("schedule_grade_resolutions"):
            raise ValueError("该核算含旧版年级修正记录；请先完成历史迁移，避免重复应用。")
        group = next((item for item in run.get("issue_groups", []) if item["id"] == issue_id), None)
        if not group or "class_value" not in group.get("affected_fields", []):
            raise ValueError("只能从当前 AC 班课业务问题创建结构化处理。")
        if expected_fingerprint != group.get("fingerprint"):
            raise ValueError("问题依据已变化，请刷新后重新选择课程。")
        if kind not in {"SOURCE_DATA_CORRECTION", "APPROVED_PAYROLL_OVERRIDE"}:
            raise ValueError("请选择修正上游事实或特殊核算口径。")
        if not confirmed_by.strip() or not course_record_id:
            raise ValueError("请选择具体课程并填写确认人。")
        schedule = self._read_for_run("schedule", Path(run["files"]["schedule"]["path"]), run).records
        record = next((item for item in schedule if schedule_record_id(item) == course_record_id), None)
        if record is None or record.teacher != group["teacher"] or record.period != run["period"]:
            raise ValueError("所选课程不属于当前教师、月份或当前排课来源。")
        source_hash = run["files"]["schedule"]["sha256"]
        if any(item.get("status") == "ACTIVE" and item.get("kind") == kind and item.get("source_record_id") == course_record_id for item in run.get("resolutions", [])):
            raise ValueError("该课程已经存在同类型的有效处理记录。")
        timestamp = datetime.now(timezone.utc).isoformat()
        if kind == "SOURCE_DATA_CORRECTION":
            field = str(values.get("field", ""))
            original = getattr(record, field, object())
            correction = SourceDataCorrection(
                run_id=run["id"], period=run["period"], teacher=record.teacher, source_record_id=course_record_id,
                source_file_hash=source_hash, field=field, original_value=original,
                corrected_value=values.get("corrected_value"), reason_code=str(values.get("reason_code", "")),
                reason_text=str(values.get("reason", "")).strip(), confirmed_by=confirmed_by.strip(), confirmed_at=timestamp,
            )
            item = {**asdict(correction), "id": uuid.uuid4().hex[:16], "kind": kind, "issue_id": group["id"], "issue_fingerprint": group["fingerprint"], "created_at": timestamp, "outcome": "PENDING_RECOMPUTE"}
        else:
            default_value, calculation = self._configured_contribution(run)(record)
            if default_value is None:
                raise ValueError("该课程没有可确定的默认班课折算，不能创建特殊核算口径。")
            override = ApprovedPayrollOverride(
                run_id=run["id"], period=run["period"], teacher=record.teacher, source_record_id=course_record_id,
                source_file_hash=source_hash, affected_field="class_value", default_treatment=calculation,
                default_contribution=default_value, approved_treatment=str(values.get("approved_treatment", "")).strip(),
                approved_contribution=float(values.get("approved_contribution")), reason=str(values.get("reason", "")).strip(),
                approved_by=confirmed_by.strip(), created_at=timestamp,
            )
            item = {**asdict(override), "id": uuid.uuid4().hex[:16], "kind": kind, "issue_id": group["id"], "issue_fingerprint": group["fingerprint"], "created_at": timestamp, "outcome": "PENDING_RECOMPUTE"}
        run.setdefault("resolutions", []).append(item)
        invalidate_business_decisions(run.setdefault("business_decisions", []))
        self.store.save(run)
        return self.check(run_id)

    def active_resolution(self, run_id: str, resolution_id: str) -> dict:
        run = self._load(run_id)
        item = next((value for value in run.get("resolutions", []) if value.get("id") == resolution_id), None)
        if item is None:
            raise ValueError("未找到该结构化处理记录。")
        return item

    def resolution_recomputation(self, run_id: str, resolution_id: str) -> dict:
        item = self.active_resolution(run_id, resolution_id)
        return dict(item.get("recomputation") or {})

    def check(self, run_id: str) -> dict:
        run = self._load(run_id)
        missing = self._missing_materials(run)
        if missing:
            raise ValueError("请先导入：" + "、".join(missing))
        self._require_fresh(run)
        run["status"] = "CHECKING"
        run.pop("last_error", None)
        self.store.save(run)
        try:
            result = self._perform_check(run)
            # Keep the post-decision recheck on the same complete preview path
            # as the explicit preview action. This is read-only: workbook
            # rendering remains exclusively inside generate_payroll().
            return self._with_generated_preview(result)
        except OSError as exc:
            message = self._file_error(exc)
            self._finish_failed_check(run, message)
            raise ValueError(message) from exc
        except ValueError as exc:
            self._finish_failed_check(run, str(exc))
            raise

    def _perform_check(self, run: dict) -> dict:
        reads = {key: self._read_for_run(key, Path(item["path"]), run) for key, item in run["files"].items()}
        if not self._fresh(run):
            raise ValueError("原文件在核对时发生变化，请重新导入。")
        if any(result.errors for result in reads.values()):
            raise ValueError("材料重新读取失败，请返回材料页重新选择。")
        payroll, scope_teachers = self._payroll_records(reads)
        if run.get("mode", MODE_AUDIT) == MODE_GENERATE and not scope_teachers:
            # Generate mode has no submitted sheet: the schedule itself defines
            # who is in scope, so blockers such as an unknown class type still
            # surface instead of quietly producing an empty run.
            scope_teachers = {row.teacher for row in reads["schedule"].records}
        math_teachers = {row.teacher for row in reads["math"].records} if "math" in reads else set()
        science_teachers = {row.teacher for row in reads["science"].records} if "science" in reads else set()
        if math_teachers & science_teachers:
            raise ValueError("两张工资表含有重复教师，请确认分组后再核对。")
        # A source export can cover the whole campus.  The active payroll
        # targets define this run's population; teachers outside that target
        # set are out of scope, not missing payroll recipients.
        resolved_schedule = self._apply_schedule_grade_resolutions(reads["schedule"].records, run)
        scoped_schedule = [row for row in resolved_schedule if row.teacher in scope_teachers]
        # Same Core function both modes use; only the coefficients come from the
        # run's bound ClassTypeRule version instead of code constants.
        coefficients, _rule_version_id = self._class_type_rules_for_run(run)
        configured = run.get("calculation_engine") == "CONFIGURED_V1"
        if configured:
            core_result = self._configured_calculation(run, scoped_schedule, payroll)
            checks = self._core_checks(core_result, payroll, run.get("mode", MODE_AUDIT))
        else:
            checks = schedule_field_checks(scoped_schedule, payroll, rules=coefficients)
        checks = self._apply_ac_resolutions(run, scoped_schedule, payroll, checks)
        # M is a production input-derived component.  In GENERATE mode each
        # in-scope teacher gets one explicit, grouped check for the Run's G:L
        # snapshot; missing fields remain actionable instead of becoming zero.
        if run.get("mode", MODE_AUDIT) == MODE_GENERATE:
            from payroll_core.payroll_generation import base_salary_field
            base_inputs = run.get("base_salary_inputs") if run.get("base_salary_input_snapshot") else None
            for teacher in sorted(scope_teachers):
                m_result = base_salary_field(teacher, base_inputs)
                m_status = "DETERMINED" if m_result.get("state") == "DETERMINED" else "MISSING_SOURCE"
                checks.append(FieldCheck(teacher, "base_salary", m_result.get("value"), m_result.get("value"), m_status, str(m_result.get("reason", "M 的 G～L 输入尚未齐全。"))))
        # A bound renewal snapshot is the production authority for AH/AI/AJ.
        # Surface one grouped, teacher-scoped action only when that frozen
        # source is incomplete; determined/no-event rows remain quiet.
        renewal_snapshot = run.get("run_renewal_result_snapshot")
        if renewal_snapshot is not None:
            for teacher in sorted(scope_teachers):
                renewal = renewal_fields_from_snapshot(renewal_snapshot, teacher)
                if renewal is None:
                    checks.append(FieldCheck(teacher, "renewal_result", None, None, "MISSING_SOURCE", "本期没有可绑定的已审核续费最终结果；不能把缺少来源当作 0。"))
                    continue
                missing_codes = [code for code, field in renewal.items() if field.get("state") not in {"DETERMINED", "NOT_APPLICABLE"}]
                if missing_codes:
                    state = "IDENTITY_NOT_STABLE" if any(renewal[code].get("state") == "IDENTITY_NOT_STABLE" for code in missing_codes) else "MISSING_SOURCE"
                    checks.append(FieldCheck(teacher, "renewal_result", None, None, state, f"已审核续费结果缺少或无法稳定绑定：{', '.join(missing_codes)}。"))
        if configured:
            from payroll_core.calculation import course_record_key
            keys = {schedule_record_id(row): course_record_key(row) for row in scoped_schedule}
            effective = {}
            for resolution in run.get("resolutions", []):
                if resolution.get("status") != "ACTIVE":
                    continue
                for entry in resolution.get("recomputation", {}).get("contributions", []):
                    if entry.get("effective_contribution") is not None and entry["source_record_id"] in keys:
                        effective[keys[entry["source_record_id"]]] = entry["effective_contribution"]
            if effective:
                core_result = self._configured_calculation(run, scoped_schedule, payroll, effective)
                recalculated = self._core_checks(core_result, payroll, run.get("mode", MODE_AUDIT))
                ac_checks = [c for c in checks if c.field == "class_value"]
                checks = [c for c in recalculated if c.field != "class_value"] + ac_checks
            run["core_calculation"] = core_result
            run["calculation_context"] = self._calculation_context(run)
            if core_result.get("course_contributions"):
                # The lesson-level snapshot is produced from the exact
                # calculation result, so later registry edits cannot change
                # an already calculated Run's effective prices.
                run.setdefault("run_part_time_pricing_snapshot", self._build_part_time_pricing_snapshot(run, scoped_schedule, core_result))
        is_generate = run.get("mode", MODE_AUDIT) == MODE_GENERATE
        if not is_generate:
            # These checks describe reconciliation against an existing payroll
            # target.  GENERATE has no target by design; the schedule defines
            # the population and must not manufacture one issue per teacher.
            for teacher in sorted(scope_teachers - {row.teacher for row in payroll}):
                checks.extend((
                    FieldCheck(teacher, "one_to_one", None, None, "MISSING_TARGET", "提交范围内教师未出现在基准最终工资表。"),
                    FieldCheck(teacher, "class_value", None, None, "MISSING_TARGET", "提交范围内教师未出现在基准最终工资表。"),
                ))
            checks += total_salary_read_checks(payroll)
        rating_version = self._rating_version_for_run(run)
        ratings = [TeacherRating(item["teacher"], item["rating"], item.get("role", "教师"), rating_version["effective_from"], rating_version["effective_to"], rating_version["source"], rating_version["source_version"], allow_blank_payroll_rating=item.get("allow_blank_payroll_rating", False)) for item in rating_version.get("ratings", [])] if rating_version else []
        exempt = set()
        if not is_generate:
            rating_checks = rating_and_rate_checks(payroll, ratings, default_compensation_bands(), run["period"])
            if configured:
                exempt = {row["teacher"] for row in core_result["rows"] if row["fields"]["AE"]["state"] == "NOT_APPLICABLE"}
                checks += [c for c in rating_checks if c.field == "rating" and c.teacher not in exempt]
            else:
                checks += rating_checks
        policy_version = self._policy_version_for_run(run)
        profiles = [TeacherCompensationProfile(item["teacher"], item["role"], item.get("rating"), item.get("rating_override"), item.get("special_approval", ""), item.get("obligation_hours", 0), item.get("obligation_hours_deduction_enabled", False), policy_version["effective_from"], policy_version["effective_to"], policy_version["source"], item.get("note", "")) for item in policy_version.get("profiles", [])] if policy_version else []
        if not configured:
            checks += policy_fee_checks(payroll, profiles, default_compensation_bands(), run["period"])
        for role in (("baseline",) if "baseline" in reads else tuple(role for role in SCOPE_ROLES if role in reads)):
            audit_rows = payroll if role == "baseline" else reads[role].records
            if configured:
                audit_rows = [row for row in audit_rows if row.teacher not in exempt]
            for item in self._formula_audit_for_scope(run["files"][role]["path"], audit_rows):
                status, evidence = self._formula_status_with_policy(item, audit_rows, profiles)
                checks.append(FieldCheck("工作簿", "formula", None, None, status, f"{Path(item.workbook).name} / {item.sheet} / {item.cell}：{evidence} 正常模式：{item.expected_pattern or '待确认'}；当前公式：{item.formula or '空白/固定值'}。"))
        raw_checks = list(checks)
        # Only historical, field-level actions retain their old display
        # behaviour.  Business decisions never alter a payroll audit status.
        checks = self._apply_decisions(checks, run["decisions"])
        visible_checks = [check for check in checks if check.status not in {"MATCH", "FORMULA_MATCH", "RATE_MATCH", "AF_POLICY_MATCH", "READ_ONLY", "NOT_APPLICABLE", "DETERMINED"}]
        run["field_records"] = self._annotate_decisions([self._issue(check) for check in raw_checks], run["decisions"])
        run["issues"] = self._annotate_decisions([self._issue(check) for check in visible_checks], run["decisions"])
        run["issues"].sort(key=lambda item: (item["severity_rank"], item["title"], item["teacher"]))
        run["audit_context"] = self._business_context(run)
        run.pop("business_context_stale", None)
        run["issue_groups"] = self._business_groups(run)
        run["user_actions"] = build_user_actions(run, run["issue_groups"])
        run["field_status"] = self._field_status(checks)
        run["summary"] = self._summary(checks)
        if configured:
            run["field_status"] = self._core_field_status(core_result, checks) + [f for f in run["field_status"] if f["field"] in {"rating", "formula", "av"}]
            run["summary"]["core_calculation_complete"] = bool(core_result["rows"]) and all(v["state"] in {"DETERMINED", "NOT_APPLICABLE"} for row in core_result["rows"] for v in row["fields"].values())
            run["summary"]["scope_note"] = "AA/AC/AD 来自独立排课；AE/AF 按版本规则与个人政策逐层计算。估算不算已核对；总工资外围项目不在本轮计算范围。"
        run["status"] = "PASS" if run["summary"]["full_scope_complete"] else "REVIEW_REQUIRED"
        run.pop("last_error", None)
        self.store.save(run)
        return self.render(run)

    def _rating_version_for_run(self, run: dict) -> dict | None:
        version_id = run.get("rating_version_id")
        if version_id:
            return self.store.get_rating_version(version_id)
        if run.get("star_authority_status") == "CONFLICT_NEEDS_CONFIRMATION":
            # A package conflict is an explicit stop for automatic authority
            # selection; do not let the generic single-version fallback pick
            # an older unrelated rating behind the user's back.
            return None
        matches = [item for item in self.store.list_rating_versions() if item.get("status", "ACTIVE") == "ACTIVE" and item["effective_from"] <= run["period"] <= item["effective_to"]]
        if len(matches) == 1:
            run["rating_version_id"] = matches[0]["id"]
            return matches[0]
        return None

    @staticmethod
    def _payroll_records(reads: dict) -> tuple[list, set[str]]:
        submissions = [
            row
            for role in SCOPE_ROLES
            for row in (reads[role].records if role in reads else [])
        ]
        scope_teachers = {row.teacher for row in submissions}
        if "baseline" not in reads:
            return submissions, scope_teachers
        return [row for row in reads["baseline"].records if row.teacher in scope_teachers], scope_teachers

    @staticmethod
    def _formula_audit_for_scope(path: str, payroll: list) -> list:
        """Keep final-workbook structural checks within this group leader's scope.

        A base payroll workbook can contain teachers a group leader did not
        submit.  Those rows are useful as formula context, but must not become
        issues in the group's own run.
        """
        cells = {
            (item.provenance[field].sheet, item.provenance[field].coordinate)
            for item in payroll
            for field in ("ae", "af")
            if field in item.provenance
        }
        if Path(path).suffix.lower() == ".csv":
            return []
        return [item for item in audit_payroll_formulas(path) if (item.sheet, item.cell) in cells]

    @staticmethod
    def _formula_status_with_policy(item: Any, payroll: list, profiles: list[TeacherCompensationProfile]) -> tuple[str, str]:
        """Recognize an intentional AF formula only when a dated policy proves it.

        A no-deduction teacher legitimately uses ``AD * AE`` rather than the
        usual ``(AD - obligation) * AE``.  This is not a blanket exception for
        a role: it requires that teacher's bound compensation profile.
        """
        if item.status != "FORMULA_PATTERN_MISMATCH" or item.field != "AF 总课时费":
            return item.status, item.evidence
        record_by_cell = {
            (record.provenance["af"].sheet, record.provenance["af"].coordinate): record
            for record in payroll if "af" in record.provenance
        }
        record = record_by_cell.get((item.sheet, item.cell))
        profile = next((profile for profile in profiles if record and profile.teacher == record.teacher), None)
        if profile and not profile.obligation_hours_deduction_enabled and item.normalized_formula == "=AD{row}*AE{row}":
            return "FORMULA_MATCH", "该教师的有效工资政策明确不扣义务课时，公式结构与个人政策一致。"
        return item.status, item.evidence

    @staticmethod
    def _missing_materials(run: dict) -> list[str]:
        missing = [LABELS[key] for key in REQUIRED if key not in run["files"]]
        # A submitted payroll sheet is the cross-check target in audit mode only.
        # Generate mode computes the payroll itself, so it must never be forced.
        if run.get("mode", MODE_AUDIT) == MODE_AUDIT and not any(key in run["files"] for key in SCOPE_ROLES):
            missing.append("至少一张教师提交表")
        return missing

    @classmethod
    def _materials_ready(cls, run: dict) -> bool:
        return not cls._missing_materials(run)

    @staticmethod
    def _apply_schedule_grade_resolutions(records: list, run: dict) -> list:
        """Apply run-bound, source-cited grade facts without changing the source workbook.

        A resolution is keyed to the original worksheet cell that supplied the
        class name.  It remains tied to the imported schedule hash through the
        run's ordinary stale-file gate.
        """
        if run.get("resolutions"):
            # New correction records use stable course identities.  Never apply
            # the older coordinate layer as well.
            return records
        resolutions = {
            str(item.get("class_name_cell", "")): str(item.get("grade", "")).strip()
            for item in run.get("schedule_grade_resolutions", [])
            if str(item.get("class_name_cell", "")).strip() and str(item.get("grade", "")).strip()
        }
        if not resolutions:
            return records
        return [
            replace(record, grade=resolutions.get(record.provenance["class_name"].coordinate, record.grade))
            if "class_name" in record.provenance and record.provenance["class_name"].coordinate in resolutions
            else record
            for record in records
        ]

    def _apply_ac_resolutions(self, run: dict, schedule: list, payroll: list, checks: list[FieldCheck]) -> list[FieldCheck]:
        """Replace AC checks only for teachers with active, verified resolutions."""
        active = [item for item in run.get("resolutions", []) if item.get("status") == "ACTIVE"]
        if not active:
            return checks
        by_teacher: dict[str, list[dict]] = {}
        for item in active:
            by_teacher.setdefault(str(item.get("teacher", "")), []).append(item)
        payroll_by_teacher = {item.teacher: item for item in payroll}
        output = [item for item in checks if not (item.field == "class_value" and item.teacher in by_teacher)]
        for teacher, stored in by_teacher.items():
            corrections: list[SourceDataCorrection] = []
            overrides: list[ApprovedPayrollOverride] = []
            for item in stored:
                common = {key: value for key, value in item.items() if key not in {"id", "kind", "issue_id", "issue_fingerprint", "outcome", "recomputation", "invalidated_at", "history_event"}}
                if item["kind"] == "SOURCE_DATA_CORRECTION":
                    common.pop("created_at", None)
                    corrections.append(SourceDataCorrection(**common))
                else:
                    overrides.append(ApprovedPayrollOverride(**common))
            teacher_schedule = [row for row in schedule if row.teacher == teacher and row.class_type != "1对1"]
            result = resolve_ac(
                teacher_schedule, run_id=run["id"], period=run["period"], corrections=corrections, overrides=overrides,
                source_hashes={row.source: run["files"]["schedule"]["sha256"] for row in teacher_schedule},
                contribution_calculator=self._configured_contribution(run),
            )
            target = payroll_by_teacher.get(teacher)
            actual = target.class_value if target else None
            baseline = result.original_total
            if result.effective_total is None:
                check = FieldCheck(teacher, "class_value", None, actual, "NEEDS_MANUAL_REVIEW", "已存在结构化处理，但仍有课程没有可确定的班课折算，不能完成重算。")
                assessment = {"status": "UNEXPLAINED", "blockers": list(result.blockers)}
            elif actual is None:
                check = FieldCheck(teacher, "class_value", result.effective_total, None, "NEEDS_MANUAL_REVIEW", "工资表 AC 没有可靠可读值，不能确认结构化处理结果。")
                assessment = {"status": "UNEXPLAINED", "blockers": []}
            else:
                cause = assess_counterfactual_cause(baseline_total=baseline, recomputed_total=result.effective_total, payroll_total=actual)
                status = "MATCH" if cause.status == "EXACT_CAUSE" else "UNEXPLAINED_DIFFERENCE"
                reason = "结构化处理重算后与工资表 AC 精确一致。" if status == "MATCH" else "结构化处理已参与重算，但仍未精确闭合工资表 AC 差异。"
                check = FieldCheck(teacher, "class_value", result.effective_total, actual, status, reason)
                assessment = asdict(cause)
            for item in stored:
                item["recomputation"] = {"original_total": result.original_total, "corrected_total": result.corrected_total, "effective_total": result.effective_total, "assessment": assessment, "contributions": [
                    {"source_record_id": entry.source_record_id, "default_contribution": entry.default_contribution, "effective_contribution": entry.effective_contribution, "calculation": entry.calculation}
                    for entry in result.contributions
                ]}
                if check.status == "MATCH" and assessment.get("status") == "EXACT_CAUSE":
                    item["outcome"] = "RESOLVED_BY_SOURCE_CORRECTION" if item["kind"] == "SOURCE_DATA_CORRECTION" else "RESOLVED_BY_APPROVED_OVERRIDE"
                else:
                    item["outcome"] = "PENDING_REVIEW"
            output.append(check)
        return output

    def _policy_version_for_run(self, run: dict) -> dict | None:
        snapshot = run.get("run_policy_snapshot")
        if snapshot:
            if snapshot.get("sha256") != self._policy_snapshot_hash(snapshot):
                raise ValueError("Run 政策快照校验失败，不得继续核算。")
            if not snapshot.get("personal_policies"):
                return None
            return {
                "id": snapshot.get("policy_version_id") or "RUN_POLICY_SNAPSHOT",
                "effective_from": run["period"], "effective_to": run["period"],
                "source": "RUN_POLICY_SNAPSHOT", "source_hash": snapshot.get("sha256", ""),
                "profiles": copy.deepcopy(snapshot.get("personal_policies") or []),
            }
        version_id = run.get("policy_version_id")
        if version_id:
            return self.store.get_policy_version(version_id)
        matches = [item for item in self.store.list_policy_versions() if item.get("status", "ACTIVE") == "ACTIVE" and item["effective_from"] <= run["period"] <= item["effective_to"]]
        if len(matches) == 1:
            run["policy_version_id"] = matches[0]["id"]
            return matches[0]
        return None

    def _finish_failed_check(self, run: dict, message: str) -> None:
        if self._fresh(run):
            run["status"] = "FILES_READY"
            run["last_error"] = message
            self.store.save(run)

    def decide(self, run_id: str, issue_id: str, action: str, person: str, reason: str, expected_fingerprint: str | None = None) -> dict:
        if action not in {"special", "payroll_error", "defer", "confirm_source", "CONFIRMED_ERROR", "ACCEPTED_EXCEPTION", "DEFERRED"} or not person.strip() or not reason.strip():
            raise ValueError("请选择处理方式，并填写确认人和理由。")
        run = self._load(run_id)
        self._require_fresh(run)
        if action in {"CONFIRMED_ERROR", "ACCEPTED_EXCEPTION", "DEFERRED"}:
            self._require_current_audit(run)
            group = next((item for item in run.get("issue_groups", []) if item["id"] == issue_id), None)
            # API callers from the pre-card UI can still submit a field id.
            if group is None:
                field = next((item for item in run.get("issues", []) if item["id"] == issue_id), None)
                if field:
                    group = next((item for item in run.get("issue_groups", []) if field["id"] in item.get("field_record_ids", [])), None)
            if group is None:
                raise ValueError("该问题已不存在，请重新核对。")
            if expected_fingerprint is not None and expected_fingerprint != group["fingerprint"]:
                raise ValueError("问题依据已变化，请刷新详情并重新确认。")
            context = self._business_context(run)
            decision = {"group_id": group["id"], "root_cause_key": group["root_cause_key"], "fingerprint": group["fingerprint"], "affected_fields": group["affected_fields"], "context": context, "action": action, "status": "ACTIVE", "person": person.strip(), "reason": reason.strip(), "created_at": datetime.now(timezone.utc).isoformat()}
            previous = [item for item in run.setdefault("business_decisions", []) if item.get("group_id") == group["id"]]
            if previous:
                run.setdefault("business_decision_history", []).extend({**item, "superseded_at": decision["created_at"], "superseded_by": action} for item in previous)
            run["business_decisions"] = [item for item in run["business_decisions"] if item.get("group_id") != group["id"]] + [decision]
            self._refresh_business_groups(run)
            self.store.save(run)
            return self.render(run)
        issue = next((item for item in run["issues"] if item["id"] == issue_id), None)
        if issue is None:
            raise ValueError("该问题已不存在，请重新核对。")
        if action == "special" and issue["status"] != "UNEXPLAINED_DIFFERENCE":
            raise ValueError("只有已有明确系统值与工资表值的差异，才能记录为特殊情况。")
        run["decisions"] = [item for item in run["decisions"] if item["issue_id"] != issue_id] + [{"issue_id": issue_id, "action": action, "person": person.strip(), "reason": reason.strip(), "teacher": issue["teacher"], "field": issue["field"], "expected": issue["expected"], "actual": issue["actual"], "versions": {key: value["sha256"] for key, value in run["files"].items()}}]
        run["issues"] = self._annotate_decisions(run["issues"], run["decisions"])
        run["status"] = "REVIEW_REQUIRED"
        self.store.save(run)
        return self.render(run)

    def evidence(self, run_id: str, issue_id: str) -> dict:
        run = self._load(run_id); self._require_fresh(run)
        self._require_current_audit(run)
        group = next((item for item in run.get("issue_groups", []) if item["id"] == issue_id), None)
        issue = next((item for item in run["issues"] if item["id"] == issue_id), None)
        if group is None and issue is not None:
            group = next((item for item in run.get("issue_groups", []) if issue["id"] in item.get("field_record_ids", [])), None)
        if group is None:
            raise ValueError("未找到问题。")
        records = [item for item in run.get("field_records", []) if item["id"] in group["field_record_ids"]]
        schedule = self._apply_schedule_grade_resolutions(
            self._read_for_run("schedule", Path(run["files"]["schedule"]["path"]), run).records,
            run,
        )
        reads = {
            role: self._read_for_run(role, Path(item["path"]), run)
            for role, item in run["files"].items()
            if role in {"math", "science", "baseline"}
        }
        payroll, _ = self._payroll_records(reads)
        target = next((item for item in payroll if item.teacher == group["teacher"]), None)
        lines = self._course_evidence(schedule, group, target, run)
        ac_calculation = self._class_value_calculation(schedule, group["teacher"]) if "class_value" in group["affected_fields"] else None
        sections = self._evidence_sections(run, group, records, target)
        if run.get("calculation_engine") == "CONFIGURED_V1":
            for line in lines:
                if "计算依据" in line:
                    line["计算依据"] = "见本页版本化逐课计算；历史常量公式不参与当前结果。"
            core_row = next((row for row in run.get("core_calculation", {}).get("rows", []) if row["teacher"] == group["teacher"]), None)
            if core_row:
                sections = [{"title": "独立工资核心计算链", "items": [{"字段": key, "系统值": value["value"], "确定性": value["state"], "计算依据": value["reason"], "规则证据": value.get("evidence", [])} for key, value in core_row["fields"].items()]}, {"title": "工资表只作为对照目标", "items": [{"项目": r["field_label"], "系统值": r["expected"], "工资表值": r["actual"], "差额": r["difference"]} for r in records]}, {"title": "独立性边界", "items": [{"AD": "由独立排课 AA + AC 得到，不读取提交表 AD", "AE/AF": "仅 DETERMINED 为确定结果；参考星级或默认候选为估算，不冒充权威", "版本": run["core_calculation"]["rule_versions"]}]}]
            if ac_calculation is not None:
                from payroll_core.calculation import course_record_key
                by_key = {course_record_key(row): row for row in schedule}
                lessons = []
                for c in run.get("core_calculation", {}).get("course_contributions", []):
                    if c["teacher"] != group["teacher"] or c.get("field") not in {"ac", ""}:
                        continue
                    source_row = by_key.get(c["record_key"])
                    if source_row is None:
                        continue
                    locations = [v for name, v in source_row.provenance.items() if name != "student"]
                    lessons.append({"日期": source_row.lesson_date, "班级": source_row.class_name, "课程": source_row.course_name, "班型": source_row.class_type, "年级": source_row.grade, "实到人数": source_row.attended, "折算规则": c["reason"], "本条折算值": c["value"], "状态": c["state"], "来源文件": Path(source_row.source).name, "来源工作表": locations[0].sheet if locations else "未提供", "来源位置": ", ".join(v.coordinate for v in locations)})
                ac_calculation = {"records": lessons, "system_total": core_row["fields"]["AC"]["value"] if core_row else None, "formula": "系统 AC 合计 = Σ 有效课程贡献；未知不计为0"}
        comments = [c for read in reads.values() for c in read.comments if c.target == group["teacher"] and c.field in set(group["affected_fields"]) | {"ae", "af"}]
        if "class_value" in group["affected_fields"]:
            sections.append(self._class_value_comparison(records, comments, run))
        elif set(group["affected_fields"]) & {"one_to_one"}:
            sections.append({"title": "特殊处理线索（不代表已获批准）", "items": [{"工资表批注": c.text, "来源文件": Path(c.source_file).name, "工作表": c.sheet, "单元格": c.coordinate} for c in comments] or [{"线索状态": "当前文件没有可读取的相关批注。历史特殊班型须由负责人核实，不自动改变折算规则。"}]})
        if not self._fresh(run):
            raise ValueError("文件在查看证据时发生变化，请重新导入。")
        resolution_courses = []
        if "class_value" in group["affected_fields"]:
            for item in schedule:
                if item.teacher == group["teacher"] and item.class_type != "1对1":
                    resolution_courses.append({"id": schedule_record_id(item), "label": " / ".join(value for value in (item.lesson_date, item.class_name, item.course_name, item.class_type, item.grade) if value)})
        return {"issue": group, "field_records": records, "sections": sections, "evidence": lines, "ac_calculation": ac_calculation, "resolution_courses": resolution_courses, "note": "课程级明细仅适用于 AC/AA 排课核对；其余依据见分段证据。", "boundary": {"run_id": run["id"], "fingerprint": group["fingerprint"], "source_hashes": {key: item["sha256"] for key, item in run["files"].items()}, "audit_context": run.get("audit_context")}}

    def export_csv(self, run_id: str) -> str:
        run = self._load(run_id); self._require_fresh(run)
        out = io.StringIO(); writer = csv.writer(out)
        writer.writerow(["核算月份", "项目", "核对状态", "教师", "排课计算值", "工资表值", "差异", "处理情况", "说明"])
        for field in run["field_status"]:
            writer.writerow([run["period"], field["label"], field["state"], "", "", "", "", "", field["note"]])
        for issue in run["issues"]:
            writer.writerow([run["period"], issue["title"], issue["status_label"], safe_csv(issue["teacher"]), issue["expected"], issue["actual"], issue["difference"], issue["decision_label"], safe_csv(issue["reason"])])
        return out.getvalue()

    def save_management(self, run_id: str, values: dict[str, str], person: str) -> dict:
        """Store human confirmations; the UI never generates a suggested value."""
        if not person.strip():
            raise ValueError("请填写确认人。")
        run = self._load(run_id); self._require_fresh(run)
        allowed = {"平均课时", "续推人次", "退费人次", "管理考核说明"}
        run["management"] = [
            {"field": key, "value": str(value).strip(), "confirmed_by": person.strip()}
            for key, value in values.items() if key in allowed and str(value).strip()
        ]
        self.store.save(run)
        return self.render(run)

    def _load(self, run_id: str) -> dict:
        run = self.store.get(run_id)
        # A stale source is terminal for this read: never downgrade STALE to
        # FILES_READY merely because a separately bound authority changed too.
        if not self._fresh(run):
            self._refresh_business_groups(run)
            self.store.save(run)
            return run
        if run.get("audit_context") and run["audit_context"] != self._business_context(run):
            invalidate_business_decisions(run.setdefault("business_decisions", []))
            self._refresh_business_groups(run)
            run["status"] = "FILES_READY" if self._materials_ready(run) else "DRAFT"
            run["business_context_stale"] = True
            self.store.save(run)
        elif run.get("field_records") or run.get("issues"):
            prior = [(x.get("group_id"), x.get("status")) for x in run.get("business_decisions", [])]
            self._refresh_business_groups(run)
            if prior != [(x.get("group_id"), x.get("status")) for x in run.get("business_decisions", [])]:
                self.store.save(run)
        return run

    def _fresh(self, run: dict) -> bool:
        stale_roles: list[str] = []
        for role, item in run.get("files", {}).items():
            path = Path(item["path"])
            try:
                current = version(path) if path.is_file() else None
            except OSError:
                current = None
            if current != {key: item[key] for key in ("sha256", "size", "mtime_ns")}:
                stale_roles.append(role)
        stale_inputs = [
            binding["input_id"] for binding in run.get("business_input_bindings", [])
            if not self.inputs.current(binding["input_id"], binding.get("source_file_hash", ""))
        ]
        if stale_roles or stale_inputs:
            run["status"] = "STALE"
            run["stale_files"] = stale_roles
            run["stale_business_inputs"] = stale_inputs
            run["decisions"] = []
            invalidate_business_decisions(run.setdefault("business_decisions", []))
            self.store.save(run)
        else:
            run.pop("stale_files", None)
            run.pop("stale_business_inputs", None)
        return not stale_roles and not stale_inputs

    def _require_fresh(self, run: dict) -> None:
        if run["status"] == "STALE":
            raise ValueError("原始文件已变化，请重新导入并重新核对。")

    def _business_context(self, run: dict) -> dict:
        rating = self._stable_authority_version(self._rating_version_for_run(run))
        policy = self._stable_authority_version(self._policy_version_for_run(run))
        bands = [asdict(item) for item in default_compensation_bands()]
        return {
            "run_id": run["id"],
            "period": run["period"],
            "source_hashes": {key: item.get("sha256") for key, item in sorted(run.get("files", {}).items())},
            "rating_version": rating,
            "policy_version": policy,
            "default_compensation_bands": bands,
            "schedule_grade_resolutions": run.get("schedule_grade_resolutions", []),
            "business_input_bindings": run.get("business_input_bindings", []),
            "af_policy_confirmation": run.get("af_policy_confirmation"),
            **({"calculation_versions": self._calculation_context(run)} if run.get("calculation_engine") == "CONFIGURED_V1" else {}),
        }

    @staticmethod
    def _stable_authority_version(version: dict | None) -> dict | None:
        """Exclude lifecycle labels from an audit fingerprint.

        Marking v1 as superseded preserves provenance but must not silently
        invalidate an historical Run that is still explicitly bound to v1.
        Actual contents, source hash and effective period remain fingerprinted.
        """
        if version is None:
            return None
        return {
            key: value for key, value in version.items()
            if key not in {"status", "superseded_by", "superseded_at", "used_by_runs"}
        }

    def _refresh_business_groups(self, run: dict) -> None:
        context = self._business_context(run)
        records = run.get("field_records") or run.get("issues", [])
        run["issue_groups"] = build_groups(run, records, context["rating_version"], context["policy_version"], context["default_compensation_bands"], ISSUE_LABELS)
        run["user_actions"] = build_user_actions(run, run["issue_groups"])
        if run.get("business_context_stale"):
            for group in run["issue_groups"]:
                group["status_label"] = "依据已变化，需重新核对"

    def _business_groups(self, run: dict) -> list[dict]:
        self._refresh_business_groups(run)
        return run["issue_groups"]

    def _require_current_audit(self, run: dict) -> None:
        self._require_fresh(run)
        if run.get("audit_context") != self._business_context(run):
            invalidate_business_decisions(run.setdefault("business_decisions", []))
            self._refresh_business_groups(run)
            run["status"] = "FILES_READY" if self._materials_ready(run) else "DRAFT"
            self.store.save(run)
            raise ValueError("星级、政策或档位依据已变化，请重新核对后再记录处理意见。")

    def _course_evidence(self, schedule: list, group: dict, target: Any, run: dict | None = None) -> list[dict]:
        if not set(group["affected_fields"]) & {"one_to_one", "class_value"}:
            return []
        # Explain each row with the rule it actually used: 普通小班 只看实到人数系数，
        # 特殊班型用固定系数，两者不会互相相乘。
        rules = self._class_type_rules_for_run(run)[0] if run else None
        special_classes, headcount_coefficients = normalize_class_rules(rules)
        lines = []
        for item in schedule:
            if item.teacher != group["teacher"]:
                continue
            source = next(iter(item.provenance.values()))
            rationale = "未进入可计算范围"
            if item.class_type == "1对1":
                coefficient = GRADE_COEFFICIENTS.get(item.grade)
                rationale = f"AA：实到 {item.attended} × 2 × 年级系数 {coefficient}" if coefficient is not None and item.attended else "AA：年级或实到缺失，不能计算"
            elif item.class_type in special_classes:
                grade = GRADE_COEFFICIENTS.get(item.grade)
                rationale = f"AC（特殊班型 {item.class_type}）：年级系数 {grade} × 班型系数 {special_classes[item.class_type]} × {LESSON_HOUR_FACTOR}" if grade is not None else "AC：年级缺失，不能计算"
            elif item.class_type in SMALL_GROUP_CLASS_TYPES:
                grade, people = GRADE_COEFFICIENTS.get(item.grade), headcount_coefficients.get(item.attended)
                rationale = f"AC（普通小班）：年级系数 {grade} × 实到人数系数 {people} × {LESSON_HOUR_FACTOR}" if grade is not None and people is not None else "AC：年级或实到人数超出已确认系数，不能计算"
            if item.lesson_status.strip() != "已上课" or item.attended is None or item.attended <= 0:
                rationale = "未计入：只计已上课且实到人数大于零的记录。"
            lines.append({"来源": "原始排课", "来源文件": Path(source.source_file).name, "来源工作表": source.sheet, "课程时间": item.lesson_time, "课程状态": item.lesson_status, "班级": item.class_name, "班型": item.class_type, "年级": item.grade or "待确认", "实到": item.attended, "计算依据": rationale, "来源位置": ", ".join(value.coordinate for name, value in item.provenance.items() if name != "student")})
        if target:
            for field in group["affected_fields"]:
                value = target.provenance.get(field)
                if value:
                    lines.append({"来源": "工资表", "字段": value.source_field, "工资表值": value.normalized_value, "来源文件": Path(value.source_file).name, "来源工作表": value.sheet, "来源位置": value.coordinate, "读取情况": self._cell_state(value.state.value)})
        return lines

    @staticmethod
    def _class_value_calculation(schedule: list, teacher: str) -> dict:
        """Explain the existing AC calculation without changing its rules.

        This mirrors the already-established ``schedule_field_checks`` formula
        solely to expose each contributing row as evidence.  Rows without a
        deterministic coefficient are intentionally not included in the sum;
        their existing manual-review audit remains the source of truth.
        """
        lessons: list[dict] = []
        total = 0.0
        for item in schedule:
            if item.teacher != teacher:
                continue
            value, rationale = class_value_contribution(item)
            if value is None:
                continue
            total += value
            source = next(iter(item.provenance.values()))
            lessons.append({
                "日期": item.lesson_date or item.lesson_time or "未提供",
                "班级": item.class_name or "未提供",
                "课程": item.course_name or "未提供",
                "班型": item.class_type,
                "年级": item.grade,
                "原始时长/课次": item.duration_text or "未提供",
                "实到人数": item.attended,
                "折算规则": rationale,
                "本条折算值": round(value, 6),
                "来源文件": Path(source.source_file).name,
                "来源工作表": source.sheet,
                "来源位置": ", ".join(value.coordinate for name, value in item.provenance.items() if name != "student"),
            })
        return {
            "教师": teacher,
            "records": lessons,
            "system_total": round(total, 6),
            "formula": "系统 AC 合计 = Σ 每条折算值",
        }

    @staticmethod
    def _structured_class_comment(text: str, configured_rules: dict | None = None) -> float | None:
        """Parse only an explicit AC/班课 assertion; prose is never guessed."""
        normalized = " ".join(str(text).replace("\n", " ").split())
        matches = re.findall(
            r"(?:\bAC\b|班课(?:折算(?:小时(?:数)?)?|小时(?:数)?|值)?)\s*(?:为|是|[:：=])\s*(-?\d+(?:\.\d+)?)",
            normalized,
            flags=re.IGNORECASE,
        )
        values = {float(value) for value in matches}
        if len(values) == 1:
            return next(iter(values))
        # A labelled grade/headcount tally (for example “高二 2人班×3”)
        # is also structured.  It is only accepted when every item maps to
        # the existing small-class table; otherwise it stays natural language.
        tally = re.findall(r"(九年级|高一|高二|高三)\s*(\d+)\s*人班\s*(?:×|x|X|\*)?\s*(\d+)(?:节|次)?", normalized)
        if not tally:
            # The actual payroll comments often put the grade in a heading and
            # list only “N人班  N” on following lines.  This remains structured
            # because every count is explicitly under that named heading.
            current_grade = None
            nested = []
            for raw_line in str(text).splitlines():
                line = raw_line.strip()
                heading = re.fullmatch(r"(九年级|高一|高二|高三)班课[：:]?", line)
                if heading:
                    current_grade = heading.group(1)
                    continue
                entry = re.fullmatch(r"(\d+)\s*人班\s*(\d+)", line)
                if entry and current_grade:
                    nested.append((current_grade, entry.group(1), entry.group(2)))
            tally = nested
        if not tally:
            return None
        total = 0.0
        for grade, people_text, count_text in tally:
            if configured_rules is not None:
                from payroll_core.calculation import calculate_course
                from payroll_core.models.records import ScheduleRecord
                record = ScheduleRecord(configured_rules["effective_from"], "批注结构化对照", grade, "", "小班", int(people_text), lesson_status="已上课")
                result = calculate_course(record, configured_rules)
                if result.value is None or result.state != "DETERMINED":
                    return None
                total += float(result.value) * int(count_text)
                continue
            coefficient = GRADE_COEFFICIENTS.get(grade)
            people = HEADCOUNT_COEFFICIENTS.get(int(people_text))
            if coefficient is None or people is None:
                return None
            # 普通小班没有班型系数：只有 年级系数 × 实到人数系数 × 2。
            total += coefficient * people * LESSON_HOUR_FACTOR * int(count_text)
        return round(total, 6)

    def _class_value_comparison(self, records: list[dict], comments: list, run: dict | None = None) -> dict:
        fact = next((item for item in records if item["field"] == "class_value"), None)
        system_value = fact["expected"] if fact else None
        payroll_value = fact["actual"] if fact else None
        items: list[dict] = [{
            "排课系统计算值": system_value,
            "工资表 AC 最终值": payroll_value,
            "系统 vs 工资表": round(system_value - payroll_value, 6) if system_value is not None and payroll_value is not None else "无法比较",
        }]
        if not comments:
            items.append({"人工批注": "当前工资表没有可读取的 AC 批注。", "批注状态": "NO_COMMENT"})
        for comment in comments:
            configured = self._calculation_version(run, "core") if run and run.get("calculation_engine") == "CONFIGURED_V1" else None
            claimed = self._structured_class_comment(comment.text, configured["rules"] if configured else None)
            item = {
                "人工批注原文": comment.text,
                "批注来源文件": Path(comment.source_file).name,
                "批注工作表": comment.sheet,
                "批注单元格": comment.coordinate,
                "批注作者": comment.author or "未提供",
            }
            if claimed is None:
                item["批注状态"] = "COMMENT_NOT_STRUCTURED"
                item["说明"] = "批注不是可可靠解析的“AC/班课 = 数值”声明；只作人工说明，不参与计算。"
            else:
                item["批注主张值"] = claimed
                item["批注状态"] = "STRUCTURED_COMMENT"
                item["系统 vs 批注"] = round(system_value - claimed, 6) if system_value is not None else "无法比较"
                item["批注 vs 工资表"] = round(claimed - payroll_value, 6) if payroll_value is not None else "无法比较"
            items.append(item)
        return {"title": "排课系统、工资表与人工批注的三方对照", "items": items}

    def _evidence_sections(self, run: dict, group: dict, records: list[dict], target: Any) -> list[dict]:
        context = self._business_context(run)
        rating, policy = context["rating_version"], context["policy_version"]
        target_items = []
        if target:
            for field in ("one_to_one", "class_value", "teaching_hours", "ae", "af"):
                value = target.provenance.get(field)
                if value:
                    target_items.append({"字段": {"one_to_one": "AA 一对一折算小时", "class_value": "AC 班课折算小时", "teaching_hours": "AD 最终授课小时", "ae": "AE 课时单价", "af": "AF（总课时费）"}[field], "实际值": value.normalized_value, "来源文件": Path(value.source_file).name, "来源工作表": value.sheet, "来源位置": value.coordinate, "读取情况": self._cell_state(value.state.value)})
        audit_names = {"rate": "AE 课时单价", "ae": "AE 课时单价", "af_policy": "AF（总课时费）", "af": "AF（总课时费）"}
        audit_items = [{"字段": audit_names.get(row["field"], row["field_label"]), "状态": row["status_label"], "期望值": row["expected"], "实际值": row["actual"], "差异": row["difference"], "审计说明": row["reason"]} for row in records]
        authority = []
        if rating:
            teacher_rating = next((item for item in rating.get("ratings", []) if item.get("teacher") == group["teacher"]), None)
            authority.append({"依据": "星级权威资料", "来源": rating.get("source"), "有效期": f"{rating.get('effective_from')} 至 {rating.get('effective_to')}", "版本": rating.get("source_version"), "教师": group["teacher"], "权威星级": teacher_rating.get("rating") if teacher_rating else None, "身份": teacher_rating.get("role") if teacher_rating else None})
        if policy:
            profile = next((item for item in policy.get("profiles", []) if item.get("teacher") == group["teacher"]), None)
            authority.append({"依据": "教师工资政策", "来源": policy.get("source"), "有效期": f"{policy.get('effective_from')} 至 {policy.get('effective_to')}", "版本": policy.get("id"), "身份": profile.get("role") if profile else None, "政策星级": (profile.get("rating_override") if profile and profile.get("rating_override") is not None else profile.get("rating")) if profile else None, "义务课时": profile.get("obligation_hours") if profile else None, "扣减义务课时": profile.get("obligation_hours_deduction_enabled") if profile else None, "特殊审批": profile.get("special_approval") if profile else None})
        sections = [{"title": "原始字段审计", "items": audit_items}, {"title": "工资表目标与 AD 来源", "items": target_items}]
        if not set(group["affected_fields"]) & {"rate", "af_policy", "ae", "af", "rating"}:
            return sections
        teacher_rating = next((x for x in (rating or {}).get("ratings", []) if x["teacher"] == group["teacher"]), None)
        profile = next((x for x in (policy or {}).get("profiles", []) if x["teacher"] == group["teacher"]), None)
        hours = target.teaching_hours if target else None
        sections.append({"title": "共同输入与独立权威依据", "items": [{"AD 最终授课小时": hours, "AD 来源": "工资表自身；具体文件与单元格见上方"}] + authority})
        for field, label, basis, version_info in (("rate", "AE 规则链", teacher_rating, rating), ("af_policy", "AF 规则链", profile, policy)):
            if field not in group["affected_fields"]:
                continue
            selected_rating = basis.get("rating_override", basis.get("rating")) if basis else None
            if basis and selected_rating is None:
                selected_rating = basis.get("rating")
            valid = version_info and version_info["effective_from"] <= run["period"] <= version_info["effective_to"]
            # Use the exact Core selector, independently for AE and AF.
            matched = [b for b in default_compensation_bands() if valid and basis and hours is not None and b.applies_to(run["period"], basis.get("role", "教师"), hours) and b.rating == selected_rating]
            items = [{"适用星级": selected_rating, "命中规则数": len(matched)}]
            items.extend({"AD 所属档位": b.band, "小时下界（含）": b.min_hours, "小时上界（含）": b.max_hours if b.max_hours is not None else "无上限", "基础金额": b.base_amount, "星级加成": b.rating_bonus, "规则有效期": f"{b.effective_from} 至 {b.effective_to}", "规则版本": b.source_version, "规则来源": b.source, "规则文件": "skill/payroll/references/ae_tier_rules.md", "单价计算": f"{b.base_amount:g} + {b.rating_bonus:g}"} for b in matched)
            fact = next((x for x in records if x["field"] == field), None)
            if field == "af_policy" and basis:
                deductible = basis.get("obligation_hours", 0) if basis.get("obligation_hours_deduction_enabled", False) else 0
                items.append({"教师政策版本": (version_info or {}).get("id"), "原始星级": basis.get("rating"), "特批星级": basis.get("rating_override"), "特批说明": basis.get("special_approval") or "无", "是否扣义务课时": "是" if basis.get("obligation_hours_deduction_enabled") else "否", "本月扣减小时": deductible, "计算过程": f"({hours} - {deductible}) × ({matched[0].base_amount} + {matched[0].rating_bonus})" if len(matched) == 1 and hours is not None else "依据不足，不能唯一计算", "政策备注": basis.get("note", "")})
            if fact:
                items.append({"系统应有值": fact["expected"], "工资表值": fact["actual"], "差额": fact["difference"], "计算出处": "核算程序的原始字段结果；本页不另行计算工资"})
            sections.append({"title": label, "items": items})
        sections.append({"title": "独立性边界", "items": [{"星级": "独立权威" if teacher_rating else "权威缺失", "档位规则": "独立规则表，按生效期匹配", "教师工资政策": "独立权威" if profile else "权威缺失", "AD": "仍来自工资表自身，AE/AF 尚未完全独立闭环"}]})
        return sections

    @staticmethod
    def _period_window(run: dict) -> tuple[str, str]:
        window = normalize_period_window(run["period"], run.get("period_start"), run.get("period_end"), run.get("period_boundary_source"))
        return window["period_start"], window["period_end"]

    def _read_for_run(self, role: str, path: Path, run: dict):
        """Read a run-bound source while keeping legacy test/provider seams."""
        start, end = self._period_window(run)
        from payroll_core.period import calendar_bounds
        if (start, end) == calendar_bounds(run["period"]):
            return self._read(role, path, run["period"])
        return self._read(role, path, run["period"], period_start=start, period_end=end)

    def _read(self, role: str, path: Path, period: str, *, period_start: str | None = None, period_end: str | None = None):
        """Read one material. Schedule falls back to its confirmed mapping.

        The check phase must see the same rows the import produced, so a layout
        that needed field mapping keeps using it on every later read.
        """
        if path.suffix.lower() == ".csv":
            return self._read_csv(role, path, period, period_start=period_start, period_end=period_end)
        if role == "schedule":
            # This is deliberately a local authority file, never a repository
            # fixture.  Empty/missing means the adapter leaves such grades
            # unresolved rather than inventing a default.
            lookup = self._student_grade_lookup()
            manual, history = self._stored_grade_evidence()
            result, analysis = resolve_schedule_import(
                path, period, profiles=self.store.list_import_profiles(SCHEDULE_AC_REQUIREMENT.name), student_grades=lookup,
                manual_grade_evidence=manual, historical_grade_evidence=history,
                course_export_snapshots=self._stored_course_export_snapshots(), period_start=period_start, period_end=period_end,
            )
            if analysis is not None and not analysis.ready:
                result.errors.append(AdapterIssue("NEEDS_FIELD_CONFIRMATION", self._mapping_error(analysis)))
            return result
        return {"math": read_payroll_excel, "science": read_payroll_excel, "baseline": read_payroll_excel, "check": read_check_workbook_schedule}[role](path, period)

    def import_grade_history(self, path: str, period: str = "0000-00", expected_hash: str | None = None) -> dict:
        """Store minimum grade facts and course-export snapshots only.

        The one-time migration path keeps attended student facts separate from
        all course-version snapshots, including unstarted rows, and never
        retains a copy of the workbook.
        """
        source = Path(path).expanduser().resolve()
        if not source.is_file():
            raise ValueError("找不到历史课表。请关闭 Excel/WPS 后重试。")
        before = version(source)
        if expected_hash and expected_hash != before["sha256"]:
            raise ValueError("所选历史课表与刚才识别的文件不一致。")
        # No lookup/history is supplied here: this operation must save only
        # facts explicitly present in that historical schedule.
        result, analysis = resolve_schedule_import(source, period)
        if analysis is not None and not analysis.ready:
            raise ValueError(self._mapping_error(analysis))
        if result.errors:
            raise ValueError("历史课表缺少必要字段：" + "；".join(issue.message for issue in result.errors))
        if version(source) != before:
            raise ValueError("历史课表在读取期间发生变化，请关闭 Excel/WPS 后重试。")
        snapshots = self._save_course_export_snapshots(result.records, before)
        stored = self._save_direct_grade_evidence(result.records, before)
        return {"source": source.name, "period": period, "direct_grade_evidence": stored, "course_export_snapshots": snapshots, "ignored_export_pollution": getattr(self, "_last_grade_pollution_count", 0), "records": len(result.records)}

    def import_grade_history_for_run(self, run_id: str, path: str, expected_hash: str | None = None) -> dict:
        """Add one more past schedule and report its practical effect."""
        run = self._load(run_id)
        before = self._grade_help_for_run(run)
        imported = self.import_grade_history(path, expected_hash=expected_hash)
        after = self._grade_help_for_run(run)
        rendered = self.check(run_id) if self._materials_ready(run) else self.render(run)
        return {"import": imported, "before_count": before["count"], "resolved": max(0, before["count"] - after["count"]), "remaining": after["count"], "grade_help": after, "run": rendered}

    def save_student_grade_confirmation(self, student: str, grade: str, period: str, confirmed_by: str, note: str = "", *, effective_date: str = "") -> dict:
        """Save a human-confirmed dated fact for later, explainable reuse."""
        name, value = student.strip(), grade.strip()
        if not name or not value or not period:
            raise ValueError("请填写学生、确认年级和适用月份。")
        if not confirmed_by.strip():
            raise ValueError("请填写确认人。")
        fact_date = effective_date or f"{period}-01"
        try:
            datetime.fromisoformat(fact_date).date()
        except ValueError as exc:
            raise ValueError("人工年级确认需要有效的课程日期。") from exc
        identifier = hashlib.sha256(f"manual-grade|{name}|{fact_date}".encode()).hexdigest()[:24]
        item = {
            "id": identifier, "student": name, "lesson_date": fact_date, "grade": value,
            "origin": "MANUAL_CONFIRMATION", "confirmed_by": confirmed_by.strip(), "note": note.strip(),
            "source_file": "", "source_hash": "", "sheet": "", "coordinate": "",
        }
        self.store.save_student_grade_evidence(item)
        return item

    def save_grade_confirmations_for_run(self, run_id: str, confirmations: list[dict], confirmed_by: str, note: str = "") -> dict:
        """Save one dated student fact and re-read every affected course."""
        run = self._load(run_id)
        pending = {item["student"]: item for item in self._grade_help_for_run(run)["students"] if item.get("manual_allowed")}
        saved = []
        for item in confirmations:
            student = str(item.get("student", "")).strip()
            grade = str(item.get("grade", "")).strip()
            group = pending.get(student)
            if group is None:
                raise ValueError("只能确认当前待补齐列表中的学生。请刷新后重试。")
            if not grade:
                raise ValueError(f"请为 {student} 选择年级。")
            saved.append(self.save_student_grade_confirmation(student, grade, run["period"], confirmed_by, str(item.get("note", note)), effective_date=str(group["first_course_date"])))
        if not saved:
            raise ValueError("请至少填写一名学生的年级。")
        refreshed = self._grade_help_for_run(run)
        rendered = self.check(run_id) if self._materials_ready(run) else self.render(run)
        return {"saved": len(saved), "remaining": refreshed["count"], "grade_help": refreshed, "run": rendered}

    def grade_help(self, run_id: str) -> dict:
        return self._grade_help_for_run(self._load(run_id))

    def _grade_help_for_run(self, run: dict) -> dict:
        schedule_file = run.get("files", {}).get("schedule")
        if not schedule_file:
            return {"available": False, "count": 0, "students": [], "message": "请先导入本月排课表。"}
        try:
            result = self._read_for_run("schedule", Path(schedule_file["path"]), run)
        except (OSError, ValueError):
            return {"available": False, "count": 0, "students": [], "message": "排课表暂时无法重新读取。"}
        groups: dict[str, dict] = {}
        for row in result.records:
            if row.grade:
                continue
            students = split_student_names(row.student)
            if not students:
                students = (f"__missing_student__:{row.lesson_date}:{row.teacher}:{row.class_name}",)
            for student in students:
                missing = student.startswith("__missing_student__:")
                key = student
                group = groups.setdefault(key, {"student": "未填写学生姓名" if missing else student, "course_count": 0, "course_dates": [], "teachers": set(), "manual_allowed": not missing, "reason": row.grade_reason or "课程没有可用的年级信息。"})
                group["course_count"] += 1
                if row.lesson_date:
                    group["course_dates"].append(row.lesson_date)
                if row.teacher:
                    group["teachers"].add(row.teacher)
        students = []
        for item in groups.values():
            dates = sorted(set(item.pop("course_dates")))
            item["first_course_date"] = dates[0] if dates else f"{run['period']}-01"
            item["course_dates"] = dates
            item["crosses_grade_boundary"] = any(
                date(year, 9, 20).isoformat() > item["first_course_date"][:10] <= max(dates, default=item["first_course_date"])[:10]
                for year in range(int(item["first_course_date"][:4]), int(max(dates, default=item["first_course_date"])[:4]) + 1)
            ) if dates else False
            item["teachers"] = sorted(item["teachers"])
            students.append(item)
        students.sort(key=lambda item: item["student"])
        return {"available": True, "count": len(students), "students": students, "message": "部分赠送、换购、特批课程没有写年级，或者当前班级名称已发生变化。年级会影响课时折算，需要先确认。"}

    def _stored_grade_evidence(self) -> tuple[list[StudentGradeEvidence], list[StudentGradeEvidence]]:
        manual: list[StudentGradeEvidence] = []
        history: list[StudentGradeEvidence] = []
        for item in self.store.list_student_grade_evidence():
            if item.get("status") == "EXPORT_POLLUTION":
                continue
            try:
                evidence = StudentGradeEvidence(
                    student=str(item.get("student", "")), lesson_date=str(item.get("lesson_date", "")), grade=str(item.get("grade", "")),
                    source_file=str(item.get("source_file", "")), source_hash=str(item.get("source_hash", "")),
                    sheet=str(item.get("sheet", "")), coordinate=str(item.get("coordinate", "")), origin=str(item.get("origin", "")),
                    confirmed_by=str(item.get("confirmed_by", "")), note=str(item.get("note", "")), exported_at=str(item.get("exported_at", "")),
                    teacher=str(item.get("teacher", "")), subject=str(item.get("subject", "")), lesson_start_time=str(item.get("lesson_start_time", "")), class_type=str(item.get("class_type", "")),
                )
            except (TypeError, ValueError):
                continue
            (manual if evidence.origin == "MANUAL_CONFIRMATION" else history).append(evidence)
        return manual, history

    def _stored_course_export_snapshots(self) -> list[CourseExportSnapshot]:
        snapshots: list[CourseExportSnapshot] = []
        for item in self.store.list_course_export_snapshots():
            try:
                snapshots.append(CourseExportSnapshot(
                    lesson_date=str(item.get("lesson_date", "")),
                    lesson_start_time=str(item.get("lesson_start_time", "")),
                    teacher=str(item.get("teacher", "")),
                    subject=str(item.get("subject", "")),
                    class_name=str(item.get("class_name", "")),
                    parsed_grade=str(item.get("parsed_grade", "")),
                    lesson_status=str(item.get("lesson_status", "")),
                    attended=item.get("attended"),
                    teaching_form=str(item.get("teaching_form", "")),
                    source_file=str(item.get("source_file", "")),
                    source_hash=str(item.get("source_hash", "")),
                    exported_at=str(item.get("exported_at", "")),
                    sheet=str(item.get("sheet", "")),
                    coordinate=str(item.get("coordinate", "")),
                    student=str(item.get("student", "")),
                ))
            except (TypeError, ValueError):
                continue
        return snapshots

    def _save_course_export_snapshots(self, records: list, source: dict[str, Any]) -> int:
        """Persist only course-version evidence, including unstarted rows.

        An unstarted row is deliberately excluded from student-grade facts but
        remains useful for comparing two exports of the same scheduled lesson.
        """
        saved = 0
        for record in records:
            class_cell = record.provenance.get("class_name") or record.provenance.get("grade")
            parsed_grade = grade_from_class_name(record.class_name)
            if not parsed_grade and record.grade_origin == "DIRECT_SOURCE":
                parsed_grade = record.grade
            lesson_start_time = normalize_lesson_start_time(record.lesson_time)
            if not (record.lesson_date and lesson_start_time and record.teacher and record.subject and parsed_grade):
                continue
            coordinate = getattr(class_cell, "coordinate", "")
            sheet = getattr(class_cell, "sheet", "")
            source_file = str(record.source or source.get("path", ""))
            snapshot = {
                "id": hashlib.sha256(f"course-snapshot|{source['sha256']}|{coordinate}|{record.lesson_date}|{lesson_start_time}|{record.teacher}|{record.subject}".encode()).hexdigest()[:32],
                "lesson_date": record.lesson_date,
                "lesson_start_time": lesson_start_time,
                "teacher": record.teacher,
                "subject": record.subject,
                "class_name": record.class_name,
                "parsed_grade": parsed_grade,
                "lesson_status": record.lesson_status,
                "attended": record.attended,
                "teaching_form": str(getattr(record.provenance.get("class_type"), "normalized_value", "") or ""),
                "source_file": source_file,
                "source_hash": source["sha256"],
                "exported_at": export_timestamp_from_source(source_file),
                "sheet": sheet,
                "coordinate": coordinate,
                "student": record.student,
                "origin": "COURSE_EXPORT_SNAPSHOT",
            }
            self.store.save_course_export_snapshot(snapshot)
            saved += 1
        return saved

    def _save_direct_grade_evidence(self, records: list, source: dict[str, Any]) -> int:
        candidates = []
        snapshot_index = self._stored_course_export_snapshots()
        snapshot_pollution_count = 0
        for record in records:
            if (
                record.grade_origin != "DIRECT_SOURCE" or not record.student or not record.lesson_date or not record.grade
                or record.lesson_status != "已上课" or record.attended is None or record.attended <= 0
            ):
                continue
            # The current export can already contain a promoted class name
            # even when it was exported before 20 September.  If an earlier
            # snapshot of the same stable lesson proves the natural previous
            # grade, do not persist the later class label as a student fact.
            # The snapshot remains available for course-level audit and the
            # current read will use it through the same override function.
            if course_export_grade_override(
                teacher=record.teacher,
                subject=record.subject,
                lesson_date=record.lesson_date,
                lesson_time=record.lesson_time,
                current_grade=record.grade,
                current_source_file=record.source,
                snapshots=snapshot_index,
            ):
                snapshot_pollution_count += len(split_student_names(record.student))
                continue
            class_cell = record.provenance.get("grade") or record.provenance.get("class_name")
            coordinate = getattr(class_cell, "coordinate", "")
            sheet = getattr(class_cell, "sheet", "")
            for student in split_student_names(record.student):
                identifier = hashlib.sha256(f"schedule-grade|{source['sha256']}|{coordinate}|{student}|{record.lesson_date}|{record.grade}".encode()).hexdigest()[:24]
                item = {
                    "id": identifier, "student": student, "lesson_date": record.lesson_date, "grade": record.grade,
                    "origin": "HISTORICAL_SCHEDULE", "source_file": record.source, "source_hash": source["sha256"],
                    "sheet": sheet, "coordinate": coordinate, "confirmed_by": "", "note": "源课表直接年级。",
                    "teacher": record.teacher, "subject": record.subject, "lesson_start_time": normalize_lesson_start_time(record.lesson_time), "class_type": record.class_type,
                }
                candidates.append((item, StudentGradeEvidence(
                    student=student, lesson_date=record.lesson_date, grade=record.grade,
                    source_file=record.source, source_hash=source["sha256"], sheet=sheet, coordinate=coordinate,
                    teacher=record.teacher, subject=record.subject, lesson_start_time=normalize_lesson_start_time(record.lesson_time), class_type=record.class_type,
                )))
        existing_rows = self.store.list_student_grade_evidence()
        existing = []
        for raw in existing_rows:
            if raw.get("status") == "EXPORT_POLLUTION" or raw.get("origin") != "HISTORICAL_SCHEDULE":
                continue
            existing.append(StudentGradeEvidence(
                student=str(raw.get("student", "")), lesson_date=str(raw.get("lesson_date", "")), grade=str(raw.get("grade", "")),
                source_file=str(raw.get("source_file", "")), source_hash=str(raw.get("source_hash", "")),
                sheet=str(raw.get("sheet", "")), coordinate=str(raw.get("coordinate", "")),
                origin=str(raw.get("origin", "HISTORICAL_SCHEDULE")), exported_at=str(raw.get("exported_at", "")),
                teacher=str(raw.get("teacher", "")), subject=str(raw.get("subject", "")), lesson_start_time=str(raw.get("lesson_start_time", "")), class_type=str(raw.get("class_type", "")),
            ))
        _kept, ignored = remove_export_pollution([*existing, *(evidence for _item, evidence in candidates)])
        ignored_keys = {evidence_identity(item) for item in ignored}
        for raw in existing_rows:
            if raw.get("status") == "EXPORT_POLLUTION":
                continue
            evidence = StudentGradeEvidence(
                student=str(raw.get("student", "")), lesson_date=str(raw.get("lesson_date", "")), grade=str(raw.get("grade", "")),
                source_file=str(raw.get("source_file", "")), source_hash=str(raw.get("source_hash", "")),
                sheet=str(raw.get("sheet", "")), coordinate=str(raw.get("coordinate", "")),
                teacher=str(raw.get("teacher", "")), subject=str(raw.get("subject", "")), lesson_start_time=str(raw.get("lesson_start_time", "")), class_type=str(raw.get("class_type", "")),
            )
            if evidence_identity(evidence) in ignored_keys:
                self.store.save_student_grade_evidence({**raw, "status": "EXPORT_POLLUTION", "note": "后续导出导致班名提前升级，未作为学生历史年级事实使用。"})
        saved = 0
        ignored_count = snapshot_pollution_count
        for item, evidence in candidates:
            if evidence_identity(evidence) in ignored_keys:
                item["status"] = "EXPORT_POLLUTION"
                item["note"] = "后续导出导致班名提前升级，未作为学生历史年级事实使用。"
                ignored_count += 1
            else:
                saved += 1
            self.store.save_student_grade_evidence(item)
        self._last_grade_pollution_count = ignored_count
        return saved

    def _student_grade_lookup(self) -> dict[str, str]:
        """Return the local grade authority, with a read-only legacy bridge.

        New installations keep this file under the application data directory.
        On this Mac, the existing Skill is still the authoritative local source
        during migration; reading it is intentionally a fallback only and
        never copies sensitive student data into the repository.
        """
        configured = self.root / "config" / "student_grade_lookup.csv"
        if configured.is_file():
            return load_student_grade_lookup(configured)
        legacy = Path.home() / ".workbuddy" / "skills" / "工资核对表制作" / "references" / "学生年级查表.csv"
        return load_student_grade_lookup(legacy)

    @staticmethod
    def _issue(item: FieldCheck) -> dict:
        difference = round(item.actual - item.expected, 6) if item.actual is not None and item.expected is not None else None
        severity = {
            "MISSING_SOURCE": (0, "严重"),
            "MISSING_TARGET": (0, "严重"),
            "UNEXPLAINED_DIFFERENCE": (1, "重要"),
            "FORMULA_DIFFERENCE": (1, "重要"),
            "RATING_MISMATCH": (1, "重要"),
            "RATE_MISMATCH": (1, "重要"),
            "FORMULA_REPLACED_BY_VALUE": (0, "严重"),
            "FORMULA_MISSING": (0, "严重"),
            "FORMULA_REGION_BREAK": (0, "严重"),
            "ROW_REFERENCE_SHIFT": (0, "严重"),
            "COLUMN_REFERENCE_SHIFT": (0, "严重"),
            "FORMULA_PATTERN_MISMATCH": (1, "重要"),
            "MISSING_AUTHORITY": (1, "重要"),
            "MISSING_PAYROLL_VALUE": (1, "重要"),
            "RULE_NOT_FOUND": (1, "重要"),
            "MULTIPLE_RULES_MATCHED": (0, "严重"),
            "AF_POLICY_MISMATCH": (1, "重要"),
            "NEEDS_MANUAL_REVIEW": (2, "需确认"),
            "EXPLAINED_DIFFERENCE": (3, "已确认"),
        }.get(item.status, (2, "需确认"))
        return {
            "id": hashlib.sha256(f"{item.teacher}|{item.field}|{item.reason}".encode()).hexdigest()[:16],
            "teacher": item.teacher,
            "field": item.field,
            "field_label": {"one_to_one": "AA 一对一", "class_value": "AC 班课", "teaching_hours": "AD 授课小时合计", "part_time": "兼职按节课时费", "ae": "AE 课时单价", "af": "AF（总课时费）", "af_policy": "AF（总课时费）政策", "base_salary": "M 实际基本工资（G～L 输入）", "rating": "教师星级", "rate": "档位金额", "formula": "公式完整性"}.get(item.field, "其他项目"),
            "title": {"one_to_one": "一对一折算小时需要处理", "class_value": "班课折算小时需要处理", "ae": "课时单价需要确认", "af": "总课时费需要确认", "af_policy": "AF（总课时费）政策不一致", "base_salary": "实际基本工资输入需要补齐", "rating": "教师星级不一致", "rate": "档位金额需要处理", "formula": "工资表公式异常"}.get(item.field, "需要人工处理"),
            "difference": difference,
            "status_label": ISSUE_LABELS.get(item.status, "需要处理"),
            "severity_rank": severity[0],
            "severity_label": severity[1],
            **asdict(item),
        }

    @staticmethod
    def _group_issues(issues: list[dict]) -> list[dict]:
        """Keep field-level audits, while exposing one business cause card."""
        grouped: dict[tuple[str, str], list[dict]] = {}
        for issue in issues:
            cause = "compensation" if issue["field"] in {"rate", "af_policy"} else issue["field"]
            grouped.setdefault((issue["teacher"], cause), []).append(issue)
        output = []
        for (teacher, cause), rows in grouped.items():
            main = next((row for row in rows if row["field"] == "rate"), rows[0])
            output.append({"id": main["id"], "teacher": teacher, "title": "档位与课时费依据需要处理" if cause == "compensation" else main["title"], "fields": [row["field_label"] for row in rows], "count": len(rows), "severity_rank": min(row["severity_rank"] for row in rows), "severity_label": main["severity_label"], "expected": main["expected"], "actual": main["actual"], "difference": main["difference"]})
        return sorted(output, key=lambda item: (item["severity_rank"], item["title"], item["teacher"]))

    @staticmethod
    def _annotate_decisions(issues: list[dict], decisions: list[dict]) -> list[dict]:
        by_values = {
            (item.get("teacher"), item.get("field"), item.get("expected"), item.get("actual")): item
            for item in decisions
        }
        for issue in issues:
            decision = by_values.get((issue["teacher"], issue["field"], issue["expected"], issue["actual"]))
            issue["decision"] = decision
            issue["decision_label"] = ACTION_LABELS.get(decision.get("action"), "已记录意见") if decision else "待处理"
        return issues

    @staticmethod
    def _apply_decisions(checks: list[FieldCheck], decisions: list[dict]) -> list[FieldCheck]:
        special = {
            (item.get("teacher"), item.get("field"), item.get("expected"), item.get("actual")): item
            for item in decisions if item.get("action") == "special"
        }
        output: list[FieldCheck] = []
        for check in checks:
            decision = special.get((check.teacher, check.field, check.expected, check.actual))
            if check.status == "UNEXPLAINED_DIFFERENCE" and decision:
                output.append(FieldCheck(check.teacher, check.field, check.expected, check.actual, "EXPLAINED_DIFFERENCE", f"人工确认的特殊情况：{decision['reason']}（{decision['person']}）"))
            else:
                output.append(check)
        return output

    @staticmethod
    def _summary(checks: list[FieldCheck]) -> dict:
        automatic = [row for row in checks if row.field in {"one_to_one", "class_value"}]
        completed = [row for row in automatic if row.status in COMPLETE_STATUSES]
        field_states = PayrollService._field_status(checks)
        return {
            "automatic_required": len(automatic),
            "automatic_completed": len(completed),
            "automatic_coverage": round(100 * len(completed) / len(automatic)) if automatic else 0,
            "automatic_pass": bool(automatic) and len(completed) == len(automatic),
            "unexplained": sum(row.status == "UNEXPLAINED_DIFFERENCE" for row in checks),
            "manual_review": sum(row.status in MANUAL_REVIEW_STATUSES for row in checks),
            "full_scope_complete": bool(field_states) and all(row["state"].startswith("已核对") for row in field_states),
            "scope_note": "AA、AC 已接入独立排课源；AE、AF、AV 尚无完整独立权威源，不能判定整份工资核对通过。",
        }

    @staticmethod
    def _field_status(checks: list[FieldCheck]) -> list[dict]:
        labels = {"one_to_one": "AA 一对一折算小时", "class_value": "AC 班课折算小时", "ae": "AE 该档每小时金额", "af": "AF（总课时费）", "af_policy": "AF（总课时费）政策", "av": "AV 总工资", "rating": "教师星级", "rate": "档位金额", "formula": "公式完整性"}
        output = []
        for field in ("one_to_one", "class_value", "rating", "rate", "formula", "af_policy", "ae", "af", "av"):
            rows = [row for row in checks if row.field == field]
            if field in {"one_to_one", "class_value"}:
                state = "已核对" if rows and all(row.status in {"MATCH", "EXPLAINED_DIFFERENCE"} for row in rows) else "待处理"
                note = "原始排课已独立计算并与工资表比较。" if state == "已核对" else "存在差异、缺失或无法确定的排课规则。"
            elif field in {"ae", "af"}:
                state, note = "公式复算 / 待权威确认", "已按现有公式复算；星级或输入仍来自工资表，尚非独立权威源。"
            elif field == "rating":
                state = "已核对" if rows and all(row.status == "MATCH" for row in rows) else "待处理"
                note = "使用按生效期保存的教师星级权威资料比较工资表。" if state == "已核对" else "缺少、过期或不一致的星级权威资料会阻止通过。"
            elif field == "rate":
                state = "依据已比对 / 待AD独立来源" if rows and all(row.status == "RATE_MATCH" for row in rows) else "待处理"
                note = "已用独立星级与生效期档位规则比对 AE；AD 最终授课小时仍读取自工资表自身，尚非完整独立闭环。" if state.startswith("依据已比对") else "缺少规则、规则冲突或金额不一致会阻止通过。"
            elif field == "formula":
                state = "已核对" if rows and all(row.status == "FORMULA_MATCH" for row in rows) else "待处理"
                note = "关键公式区域按同列结构模式扫描。" if state == "已核对" else "尚无足够公式样本，或已发现公式结构异常。"
            elif field == "af_policy":
                state = "依据已比对 / 待AD独立来源" if rows and all(row.status == "AF_POLICY_MATCH" for row in rows) else "待处理"
                note = "已按独立教师工资政策计算 AF；AD 最终授课小时仍来自工资表自身，不能声称独立闭环。" if state.startswith("依据已比对") else "缺少有效政策、特殊审批或 AF 金额不一致。"
            else:
                state, note = "仅读取 / 待人工确认", "总工资包含续费、退费、激励、管理奖等未接入来源。"
            # An empty check cannot be advertised as independently verified.
            automatic_available = field in {"one_to_one", "class_value", "rating", "formula"} and bool(rows)
            formula_recomputed = field in {"ae", "af", "rate", "af_policy"} and any(row.expected is not None for row in rows)
            formula_compared = field in {"ae", "af", "rate", "af_policy"} and any(row.expected is not None and row.actual is not None for row in rows)
            output.append({"field": field, "label": labels[field], "state": state, "note": note, "read": bool(rows), "authority": automatic_available, "computed": automatic_available or formula_recomputed, "compared": automatic_available or formula_compared})
        return output

    @staticmethod
    def _file_error(exc: OSError) -> str:
        if isinstance(exc, PermissionError):
            return "文件无法读取。请在系统设置中允许本工具访问该文件所在文件夹，然后重试。"
        if exc.errno == 4:
            return "文件读取被系统中断。请检查本工具是否有权限访问该文件所在文件夹。"
        return "文件打不开。请关闭 Excel/WPS，确认文件仍在原位置后重试。"

    @staticmethod
    def _inspection_error(errors: list) -> str:
        codes = {item.code for item in errors}
        if "UNSUPPORTED_FILE_FORMAT" in codes:
            return "文件格式暂不支持。请选择当前流程使用的 Excel 工作簿。"
        if "UNSUPPORTED_LAYOUT" in codes or "UNKNOWN_LAYOUT" in codes:
            return "无法识别这张表。请确认选择了正确类别的当前工资材料。"
        return "文件无法完成结构检查，请确认工作簿完整后重试。"

    @staticmethod
    def _warning_message(code: str) -> str:
        return {
            "MISSING_CACHE": "部分公式没有保存可读取的计算结果，相关字段不能可靠核对。",
            "EXTERNAL_REFERENCE_UNRESOLVED": "部分结果依赖其他 Excel 文件，目前无法确认是否最新。",
            "GRADE_UNRESOLVED": "部分排课无法确定年级，需要人工确认。",
            "MISSING_ATTENDANCE": "部分排课缺少可读取的实到人数，需要人工确认。",
            "UNPARSEABLE_LESSON_DATE": "有已上课排课无法识别上课日期，已排除，请修正日期后重新导入。",
            "OUT_OF_PERIOD_ROWS_EXCLUDED": "部分排课不属于当前工资月份，已按月份排除。",
        }.get(code, "存在需要查看的材料提示。")

    @staticmethod
    def _cell_state(state: str) -> str:
        return {
            "RAW_VALUE": "直接读取",
            "CACHED_VALUE": "读取已保存的公式结果",
            "MISSING_CACHE": "公式结果不可读取",
            "EXTERNAL_REFERENCE": "依赖其他文件",
        }.get(state, "待确认")

    def render(self, run: dict) -> dict:
        required = PayrollService._missing_materials(run)
        window = normalize_period_window(run["period"], run.get("period_start"), run.get("period_end"), run.get("period_boundary_source"))
        summary = run.get("summary", PayrollService._summary([]))
        status_label = STATUS[run["status"]]
        if run["status"] == "REVIEW_REQUIRED":
            status_label = "排课项目已核对，仍需人工确认" if summary.get("automatic_pass") else "有问题待处理"
        mode = run.get("mode", MODE_AUDIT)
        materials = []
        for role, label in LABELS.items():
            item = run.get("files", {}).get(role)
            state = "失效" if role in run.get("stale_files", []) else "已准备" if item else "未导入"
            scope_required = role in SCOPE_ROLES and mode == MODE_AUDIT and not any(key in run.get("files", {}) for key in SCOPE_ROLES)
            materials.append({"role": role, "label": label, "required": role in REQUIRED or scope_required, "state": state, "file": item})
        rating = self._rating_version_for_run(run)
        policy = self._policy_version_for_run(run)
        schedule = run.get("files", {}).get("schedule")
        class_rules, class_rule_version_id = self._class_type_rules_for_run(run)
        return {
            **run, **window,
            "period_display": {"label": window["period_label"], "start": window["period_start"], "end": window["period_end"], "boundary_source": window["period_boundary_source"]},
            "status_label": status_label,
            "materials": materials,
            "health": {
                "ready": not required and run["status"] != "STALE",
                "missing": required,
                "readiness": 100 if mode == MODE_GENERATE and "schedule" in run["files"] else round(100 * (int("schedule" in run["files"]) + int(any(key in run["files"] for key in SCOPE_ROLES))) / 2),
                "warnings": [message for item in run.get("files", {}).values() for message in item.get("warning_messages", [])],
            },
            "authority_context": {
                "schedule": {"label": "排课权威源", "name": schedule.get("name") if schedule else "尚未导入", "sha256": schedule.get("sha256") if schedule else None},
                "rating": {
                    **self._authority_reference(rating, "教师星级"),
                    "binding_status": run.get("star_authority_status", "BOUND" if rating else "NOT_PROVIDED"),
                    "conflicts": list(run.get("star_conflicts", [])),
                },
                "policy": self._authority_reference(policy, "教师工资政策"),
                "class_type_rules": {"label": "班型折算规则", "version_id": class_rule_version_id, "coefficients": class_rules},
                "core_rules": self._calculation_version(run, "core"),
                "part_time_rates": self._calculation_version(run, "part_time"),
                "rules": {"label": "工资规则", "name": "现行档位金额规则", "source": "skill/payroll/references/ae_tier_rules.md", "source_version": "2025-10", "effective_period": "2025-10 起持续维护"},
            },
            "business_input_bindings": run.get("business_input_bindings", []),
            "grade_help": self._grade_help_for_run(run),
        }

    @staticmethod
    def _authority_reference(version: dict | None, label: str) -> dict:
        if version is None:
            return {"label": label, "name": "尚未绑定", "version_id": None}
        return {
            "label": label,
            "name": version.get("source_version") or version.get("id"),
            "version_id": version.get("id"),
            "source": version.get("source"),
            "source_hash": version.get("source_hash") or None,
            "effective_period": f"{version.get('effective_from')} ～ {version.get('effective_to')}",
            "created_at": version.get("created_at"),
            "status": version.get("status", "ACTIVE"),
        }
