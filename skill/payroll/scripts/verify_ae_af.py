#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""验证薪资表的 AE（该档每小时金额）与 AF（总课时费）列是否吻合规则。

用法:
  python3 verify_ae_af.py --file 2026数学组6月薪资表.xlsx \
                          [--sheet 工资核对] \
                          [--star-override star_override.csv]

规则（有效期至 2026-09，10 月可能调整）:
  AE = 基础档位(查 AD) + 星级加成(查 F 列教师级别)
  基础档位: 0-30→0, 31-60→30, 61-80→32, 81-100→34, 101-130→36, 131-160→37, 160+→38
  星级加成: 一二星+0, 三星+5, 四星+10, 五星+15, 六星+20
  AF: 非 TRMT = (AD-30)*AE；TRMT = AD*AE（全额，前 30h 不减免）

默认列：C=姓名, F=教师级别(含星级), AD=最终授课小时, AE=该档每小时金额, AF=总课时费。
可用 --*-col 覆盖。管理/底薪类行用 --skip-keywords 管理,底薪 跳过（不走 AE 规则）。

--star-override：当表内 F 列的星级不可信时（用户曾明确要求以口头告知的星级为准），
  传入「姓名,星级」CSV，脚本以覆盖表为准计算加成，忽略表内 F 列星级。
"""
import argparse
import os
import re
import sys

import openpyxl
from openpyxl.utils import column_index_from_string

STAR_BONUS = {1: 0, 2: 0, 3: 5, 4: 10, 5: 15, 6: 20}
# 基础档位上限 -> 基础金额（最后一个为 160 以上）
TIER = [(30, 0), (60, 30), (80, 32), (100, 34), (130, 36), (160, 37), (10**9, 38)]

CN_NUM = {"一": 1, "二": 2, "三": 3, "四": 4, "五": 5, "六": 6}


def tier_base(ad):
    for upper, base in TIER:
        if ad <= upper:
            return base
    return 38


def extract_star(text):
    if not isinstance(text, str):
        return None
    m = re.search(r"([一二三四五六])星", text)
    if m:
        return CN_NUM[m.group(1)]
    m = re.search(r"(\d)星", text)
    if m:
        return int(m.group(1))
    return None


def is_trmt(text):
    return isinstance(text, str) and "TRMT" in text.upper()


def load_override(path):
    table = {}
    if not path:
        return table
    with open(path, encoding="utf-8-sig") as f:
        for line in f:
            line = line.strip()
            if not line or line.startswith("姓名"):
                continue
            name, _, star = line.partition(",")
            table[name.strip()] = int(star.strip())
    return table


def find_header_row(ws, target="姓名", max_scan=6):
    for r in range(1, max_scan + 1):
        for c in range(1, 20):
            if ws.cell(row=r, column=c).value == target:
                return r
    return None


def main():
    ap = argparse.ArgumentParser(description="验证薪资表 AE / AF 列")
    ap.add_argument("--file", required=True, help="薪资表 xlsx")
    ap.add_argument("--sheet", default=None, help="指定 sheet（默认 active）")
    ap.add_argument("--star-override", default=None, help="姓名,星级 权威覆盖表")
    ap.add_argument("--ae-col", default="AE")
    ap.add_argument("--af-col", default="AF")
    ap.add_argument("--ad-col", default="AD")
    ap.add_argument("--name-col", default="C")
    ap.add_argument("--level-col", default="F")
    ap.add_argument("--skip-keywords", default="", help="级别含这些词则跳过，逗号分隔")
    args = ap.parse_args()

    # raw 用于读取级别/姓名文本；val 用于读取 AD/AE/AF 的计算缓存值（AF 为公式）
    wb_raw = openpyxl.load_workbook(args.file, data_only=False)
    wb_val = openpyxl.load_workbook(args.file, data_only=True)
    ws = wb_raw[args.sheet] if args.sheet else wb_raw.active
    ws_val = wb_val[args.sheet] if args.sheet else wb_val.active

    override = load_override(args.star_override)
    if override:
        print(f"[覆盖] 使用权威星级表: {len(override)} 人")

    hr = find_header_row(ws, "姓名")
    if hr is None:
        print("[错误] 找不到「姓名」表头")
        sys.exit(1)

    name_ci = column_index_from_string(args.name_col)
    level_ci = column_index_from_string(args.level_col)
    ad_ci = column_index_from_string(args.ad_col)
    ae_ci = column_index_from_string(args.ae_col)
    af_ci = column_index_from_string(args.af_col)

    skip_kw = [k.strip() for k in args.skip_keywords.split(",") if k.strip()]

    mism_ae, mism_af = [], []
    n = 0
    for r in range(hr + 1, ws.max_row + 1):
        name = ws.cell(row=r, column=name_ci).value
        level = ws.cell(row=r, column=level_ci).value
        ad = ws_val.cell(row=r, column=ad_ci).value
        ae = ws_val.cell(row=r, column=ae_ci).value
        af = ws_val.cell(row=r, column=af_ci).value
        if name in (None, ""):
            continue
        if not isinstance(ad, (int, float)):
            continue
        if skip_kw and isinstance(level, str) and any(k in level for k in skip_kw):
            continue
        n += 1

        star = override.get(name) if name in override else extract_star(level)
        base = tier_base(ad)
        bonus = STAR_BONUS.get(star, 0) if star else 0
        exp_ae = base + bonus
        if isinstance(ae, (int, float)) and abs(ae - exp_ae) > 1e-6:
            mism_ae.append((name, ad, star, base, bonus, ae, exp_ae))

        exp_af = ad * exp_ae if is_trmt(level) else (ad - 30) * exp_ae
        if isinstance(af, (int, float)) and abs(af - exp_af) > 1e-6:
            mism_af.append((name, ad, exp_ae, is_trmt(level), af, exp_af))

    print(f"[统计] 校验 {n} 名教师")
    print(f"[AE 不匹配] {len(mism_ae)} 条:")
    for name, ad, star, base, bonus, ae, exp in mism_ae:
        print(f"   {name}: AD={ad} 星级={star} 基础={base} 加成={bonus} "
              f"表AE={ae} 应AE={exp}")
    print(f"[AF 不匹配] {len(mism_af)} 条:")
    for name, ad, exp_ae, trmt, af, exp in mism_af:
        print(f"   {name}: AD={ad} AE={exp_ae} TRMT={trmt} 表AF={af} 应AF={exp}")
    if not mism_ae and not mism_af:
        print("[OK] AE 与 AF 全部吻合规则。")


if __name__ == "__main__":
    main()
