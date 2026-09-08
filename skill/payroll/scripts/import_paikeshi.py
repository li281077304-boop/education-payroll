#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""工资核对表制作：排课记录源文件 -> 工资核对表.xlsx（排课记录 sheet）。

用法:
  python3 import_paikeshi.py \
      --source  排课列表_06月01日到06月28日.xls \
      --template 工资核对表模板.xlsx \
      --output  工资核对表6月.xlsx

流程:
  1. 读取源文件（xls/xlsx），自动识别列名（兼容新旧两版格式）
  2. 过滤：上课状态 == 有效状态（默认「已上课」）且 实到 > 0
  3. 逐行提取 年级 / 学科 / 班型 / 实到（见 _grade_extractor）
  4. 写入模板的「排课记录」sheet（A-H），F/G/H 写 SWITCH 公式
  5. 清理多余行 + fullCalcOnLoad（解决 Excel 打开公式显示 0 的问题）

依赖：pandas + openpyxl + xlrd（读 .xls 用 xlrd）。
"""
import argparse
import os
import shutil
import sys

import pandas as pd
import openpyxl

sys.path.insert(0, os.path.dirname(__file__))
from _grade_extractor import (extract_grade, extract_subject,
                              map_class_type, build_switch_formulas)


def detect_columns(df):
    """返回 (teacher, class_name, form, status, att, subject) 列名，缺失则报错。"""
    cols = list(df.columns)

    def find(candidates):
        for c in candidates:
            if c in cols:
                return c
        return None

    teacher = find(["任课老师"])
    class_name = find(["上课班级", "班级名称"])
    form = find(["教学形式", "课程所属班型"])
    status = find(["上课状态"])
    att = find(["实到", "实到人数"])
    subject = find(["上课科目", "课程所属学科"])
    missing = [n for n, v in dict(teacher=teacher, class_name=class_name,
                                  form=form, status=status, att=att).items()
               if v is None]
    if missing:
        raise ValueError(f"源文件缺少必要列: {missing}。实际列: {cols}")
    return teacher, class_name, form, status, att, subject


def read_source(path):
    if path.lower().endswith(".xls"):
        try:
            return pd.read_excel(path, engine="xlrd")
        except Exception:
            return pd.read_excel(path)
    return pd.read_excel(path, engine="openpyxl")


def main():
    ap = argparse.ArgumentParser(description="排课记录 -> 工资核对表（排课记录 sheet）")
    ap.add_argument("--source", required=True, help="排课记录源文件 xls/xlsx")
    ap.add_argument("--template", required=True, help="工资核对表模板 xlsx")
    ap.add_argument("--output", required=True, help="输出文件 xlsx")
    ap.add_argument("--status-ok", default="已上课", help="视为有效上课状态的值")
    ap.add_argument("--no-backup", action="store_true", help="不备份输出文件（默认会备份）")
    args = ap.parse_args()

    if os.path.exists(args.output) and not args.no_backup:
        bak = args.output + ".bak"
        shutil.copy2(args.output, bak)
        print(f"[备份] {bak}")

    print(f"[读取源] {args.source}")
    df = read_source(args.source)
    teacher, class_name, form, status, att, subject = detect_columns(df)
    print(f"[列映射] 老师={teacher} 班级={class_name} 班型={form} "
          f"状态={status} 实到={att} 科目={subject}")

    before = len(df)
    df = df[df[status].astype(str).str.strip() == args.status_ok]
    df = df[df[att].notna()]
    df = df[df[att] > 0]
    print(f"[过滤] {before} -> {len(df)} 行（保留「{args.status_ok}」且 实到>0）")

    wb = openpyxl.load_workbook(args.template)
    if "排课记录" not in wb.sheetnames:
        raise ValueError("模板缺少「排课记录」sheet")
    ws = wb["排课记录"]

    # 定位表头行（含「任课老师」）
    header_row = 1
    for r in range(1, 6):
        if ws.cell(row=r, column=1).value == "任课老师":
            header_row = r
            break
    data_start = header_row + 1

    # 清空旧数据（含公式）
    for r in range(data_start, ws.max_row + 1):
        for c in range(1, 9):
            ws.cell(row=r, column=c).value = None

    ridx = data_start
    unknown_grades = []
    for _, row in df.iterrows():
        tname = str(row[teacher]).strip()
        cname = row[class_name]
        subj = extract_subject(cname, row[subject] if subject else None)
        ctype = map_class_type(row[form])
        att_n = int(row[att])
        grade = extract_grade(cname)
        if not grade:
            unknown_grades.append((tname, cname))
        f, g, h = build_switch_formulas(ridx)
        ws.cell(row=ridx, column=1, value=tname)
        ws.cell(row=ridx, column=2, value=grade)
        ws.cell(row=ridx, column=3, value=subj)
        ws.cell(row=ridx, column=4, value=ctype)
        ws.cell(row=ridx, column=5, value=att_n)
        ws.cell(row=ridx, column=6, value=f)
        ws.cell(row=ridx, column=7, value=g)
        ws.cell(row=ridx, column=8, value=h)
        ridx += 1

    # 清理多余行（模板可能带了上月的旧行）
    last = ridx - 1
    if ws.max_row > last:
        for r in range(ws.max_row, last, -1):
            ws.delete_rows(r)
        print(f"[清理] 删除多余行，保留至 {last} 行")

    wb.calculation.fullCalcOnLoad = True
    wb.save(args.output)
    print(f"[完成] 写入 {last - data_start + 1} 条记录 -> {args.output}")
    if unknown_grades:
        print(f"[警告] {len(unknown_grades)} 条无法识别年级，需人工核对（查表或补充 references/学生年级查表.csv）：")
        for t, c in unknown_grades[:60]:
            print(f"   - {t}: {c}")


if __name__ == "__main__":
    main()
