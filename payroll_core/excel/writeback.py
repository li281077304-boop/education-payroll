"""Safe, preview-first Excel comment writeback.

The source workbook is read only. Approved candidates are applied to a newly
created workbook and then re-opened to verify both comments and formulas.
"""
from __future__ import annotations

import hashlib
import shutil
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook
from openpyxl.comments import Comment


@dataclass(frozen=True)
class WritebackPreview:
    sheet: str
    cell: str
    before_comment: str
    proposed_comment: str
    strategy: str


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def preview(source: str | Path, candidates: Iterable[dict], *, strategy: str = "APPEND") -> list[WritebackPreview]:
    if strategy not in {"KEEP_EXISTING", "APPEND", "REPLACE_CONFIRMED"}:
        raise ValueError("未知已有批注处理方式。")
    workbook = load_workbook(source, data_only=False, read_only=False, keep_links=True)
    views: list[WritebackPreview] = []
    for candidate in candidates:
        sheet = workbook[candidate["sheet"]]
        existing = sheet[candidate["cell"]].comment
        before = existing.text if existing else ""
        proposed = str(candidate["content"])
        if before and strategy == "KEEP_EXISTING":
            proposed = before
        elif before and strategy == "APPEND":
            proposed = before.rstrip() + "\n\n" + proposed
        views.append(WritebackPreview(candidate["sheet"], candidate["cell"], before, proposed, strategy))
    return views


def write_new_workbook(source: str | Path, output: str | Path, candidates: Iterable[dict], *, strategy: str = "APPEND", author: str = "工资核算助手") -> list[WritebackPreview]:
    source_path, output_path = Path(source), Path(output)
    if source_path.resolve() == output_path.resolve():
        raise ValueError("回填必须输出为新文件，不能覆盖原工资表。")
    if not source_path.is_file():
        raise ValueError("找不到原工资表。")
    items = list(candidates)
    if not items:
        raise ValueError("没有已确认的批注候选。")
    preview_rows = preview(source_path, items, strategy=strategy)
    original_hash = sha256(source_path)
    for item in items:
        expected = item.get("source_file_hash", "")
        if expected and expected != original_hash:
            raise ValueError("原工资表已变化，批注候选需要重新确认。")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source_path, output_path)
    source_book = load_workbook(source_path, data_only=False, read_only=False, keep_links=True)
    book = load_workbook(output_path, data_only=False, read_only=False, keep_links=True)
    for row in preview_rows:
        book[row.sheet][row.cell].comment = Comment(row.proposed_comment, author)
    book.save(output_path)
    verified = load_workbook(output_path, data_only=False, read_only=False, keep_links=True)
    for row in preview_rows:
        actual = verified[row.sheet][row.cell].comment
        if actual is None or actual.text != row.proposed_comment:
            raise ValueError("回填后的批注无法验证。")
    for source_sheet, result_sheet in zip(source_book.worksheets, verified.worksheets, strict=True):
        for source_row in source_sheet.iter_rows():
            for cell in source_row:
                if isinstance(cell.value, str) and cell.value.startswith("=") and result_sheet[cell.coordinate].value != cell.value:
                    raise ValueError(f"回填意外改变了公式：{source_sheet.title}!{cell.coordinate}")
    return preview_rows
