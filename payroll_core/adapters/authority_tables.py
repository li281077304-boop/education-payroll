"""Read the established renewal and refund authority workbook layouts.

These adapters preserve the layouts staff already use.  They do not decide
eligibility, responsibility, or a refund-to-payroll conversion rule.
"""
from __future__ import annotations

from dataclasses import dataclass
from hashlib import sha256
from pathlib import Path
from typing import Any, Iterable

import xlrd
from openpyxl import load_workbook

from .business_results import ImportedBusinessResult


def _text(value: object) -> str:
    return str(value or "").strip().replace("\n", "")


def _number(value: object) -> float:
    if value in (None, ""):
        return 0.0
    if not isinstance(value, (int, float)) or isinstance(value, bool):
        raise ValueError("续费推荐数据中的合计必须是数字。")
    return float(value)


def _month_sheet(workbook, period: str):
    month = f"{int(period[-2:])}月"
    candidates = [sheet for sheet in workbook.worksheets if _text(sheet.title).startswith(month)]
    if len(candidates) != 1:
        raise ValueError(f"未找到唯一的“{month}”工作表。")
    return candidates[0]


def _find_column(sheet, row: int, predicate) -> int:
    matches = [column for column in range(1, sheet.max_column + 1) if predicate(_text(sheet.cell(row, column).value))]
    if len(matches) != 1:
        raise ValueError("续费推荐数据的表头不完整或存在歧义。")
    return matches[0]


def _group_total_column(sheet, label: str, *, start: int, end: int) -> int:
    totals = [column for column in range(start, end + 1) if _text(sheet.cell(2, column).value) == "合计"]
    if len(totals) != 1:
        raise ValueError(f"未找到唯一的“{label}合计”列。")
    return totals[0]


def _cached_or_sum(raw_sheet, cached_sheet, row: int, start: int, total_column: int) -> float:
    cached = cached_sheet.cell(row, total_column).value
    if isinstance(cached, (int, float)) and not isinstance(cached, bool):
        return float(cached)
    # The official table uses a total formula.  A newly created template may
    # not yet have an Excel formula cache, so summing the displayed input
    # buckets is the transparent equivalent, not a guessed payroll rule.
    return sum(_number(raw_sheet.cell(row, column).value) for column in range(start, total_column))


def read_renewal_2026(path: str | Path, period: str) -> list[ImportedBusinessResult]:
    """Read the monthly-sheet 2026 renewal/recommendation workbook.

    Header groups, not absolute column letters, locate the three payroll
    inputs: one-to-one, class, and lead-study sessions.
    """
    source = Path(path)
    # Some newly created workbooks omit OOXML dimension metadata.  Openpyxl's
    # streaming reader then reports ``max_column=None`` even though the
    # visible Chinese headers are present.  Opening in normal mode is still
    # read-only from this application's perspective (we never call save), and
    # lets semantic header discovery work for both these templates and the
    # established monthly workbooks.
    raw_book = load_workbook(source, data_only=False, read_only=False, keep_links=True)
    cached_book = load_workbook(source, data_only=True, read_only=False, keep_links=True)
    raw = _month_sheet(raw_book, period)
    cached = cached_book[raw.title]
    teacher_column = _find_column(raw, 1, lambda value: value == "教师")
    one_start = _find_column(raw, 1, lambda value: value in {"1V1课时", "1对1课时"})
    class_start = _find_column(raw, 1, lambda value: value == "班课")
    mentor_start = _find_column(raw, 1, lambda value: "领航伴学" in value)
    groups = sorted(((one_start, "1对1"), (class_start, "班课"), (mentor_start, "领航伴学")), key=lambda item: item[0])
    bounds = {name: (start, (groups[index + 1][0] - 1 if index + 1 < len(groups) else raw.max_column)) for index, (start, name) in enumerate(groups)}
    one_total = _group_total_column(raw, "1对1", start=bounds["1对1"][0], end=bounds["1对1"][1])
    class_total = _group_total_column(raw, "班课", start=bounds["班课"][0], end=bounds["班课"][1])
    mentor_total = _group_total_column(raw, "领航伴学", start=bounds["领航伴学"][0], end=bounds["领航伴学"][1])
    results: list[ImportedBusinessResult] = []
    for row in range(3, raw.max_row + 1):
        teacher = _text(raw.cell(row, teacher_column).value)
        if not teacher:
            continue
        one = _cached_or_sum(raw, cached, row, one_start, one_total)
        classes = _cached_or_sum(raw, cached, row, class_start, class_total)
        mentor = _cached_or_sum(raw, cached, row, mentor_start, mentor_total)
        award = one + classes * 1.5 + mentor * 0.75
        payload = {
            "teacher": teacher,
            "renewal_one_to_one": one,
            "renewal_class": classes,
            "renewal_mentor": mentor,
            "payroll_fields": {"AH": one, "AI": classes, "AJ": mentor, "AK": award},
            "award_formula": "AH×1 + AI×1.5 + AJ×0.75",
        }
        results.append(ImportedBusinessResult(teacher, str(row), payload, {
            "source_file": source.name, "sheet": raw.title, "row": str(row),
            "semantic_columns": {"教师": teacher_column, "1对1合计": one_total, "班课合计": class_total, "领航伴学合计": mentor_total},
        }))
    if not results:
        raise ValueError("续费推荐数据表没有可导入的教师记录。")
    return results


def _header_row_values(rows: list[list[object]]) -> int:
    for row_index, row in enumerate(rows[:10]):
        labels = {_text(value) for value in row}
        has_student = bool({"姓名", "学生", "学生姓名"} & labels)
        has_teacher_impact_group = any(label.replace(" ", "").startswith("教师") for label in labels)
        if has_student and "退费总金额" in labels and "学科教师" in labels and has_teacher_impact_group:
            return row_index
    raise ValueError("未找到退费事实表表头。")


def _first_column(headers: list[str], label: str) -> int:
    matches = [index for index, value in enumerate(headers) if value == label]
    if len(matches) != 1:
        raise ValueError(f"退费表缺少唯一的“{label}”列。")
    return matches[0]


def _first_matching_column(headers: list[str], labels: tuple[str, ...]) -> int:
    for label in labels:
        matches = [index for index, value in enumerate(headers) if value == label]
        if len(matches) == 1:
            return matches[0]
    expected = " / ".join(f"“{label}”" for label in labels)
    raise ValueError(f"退费表缺少唯一的学生识别列（{expected}）。")


def _teacher_groups(headers: list[str]) -> list[tuple[int, int, int]]:
    groups: list[tuple[int, int, int]] = []
    for index, value in enumerate(headers):
        normalized = value.replace(" ", "")
        if normalized.startswith("教师") and index + 2 < len(headers):
            headcount, performance = headers[index + 1].replace(" ", ""), headers[index + 2].replace(" ", "")
            if headcount.startswith("人头") and (performance.startswith("业绩") or performance.startswith("绩效")):
                groups.append((index, index + 1, index + 2))
    if not groups:
        raise ValueError("未找到连续的“教师 / 人头 / 绩效”填写组。")
    return groups


def _cell(row: Iterable[object], index: int) -> object:
    values = list(row)
    return values[index] if index < len(values) else ""


def read_refund_2026(path: str | Path, period: str) -> list[ImportedBusinessResult]:
    """Read one-refund-per-row data and expand teacher impact groups safely."""
    source = Path(path)
    month = f"{int(period[-2:])}月"
    if source.suffix.lower() == ".xls":
        workbook = xlrd.open_workbook(source)
        sheets = [sheet for sheet in workbook.sheets() if _text(sheet.name).startswith(month)]
        if len(sheets) != 1:
            raise ValueError(f"未找到唯一的“{month}”退费工作表。")
        sheet_name = sheets[0].name
        rows = [sheets[0].row_values(row) for row in range(sheets[0].nrows)]
    elif source.suffix.lower() in {".xlsx", ".xlsm"}:
        workbook = load_workbook(source, data_only=True, read_only=True, keep_links=True)
        sheets = [sheet for sheet in workbook.worksheets if _text(sheet.title).startswith(month)]
        if len(sheets) != 1:
            raise ValueError(f"未找到唯一的“{month}”退费工作表。")
        sheet_name = sheets[0].title
        rows = [list(row) for row in sheets[0].iter_rows(values_only=True)]
    else:
        raise ValueError("退费绩效数据只支持 .xls、.xlsx 或 .xlsm。")
    header_row = _header_row_values(rows)
    headers = [_text(value) for value in rows[header_row]]
    student_column = _first_matching_column(headers, ("姓名", "学生", "学生姓名"))
    teacher_groups = _teacher_groups(headers)
    group_columns = {column for group in teacher_groups for column in group}
    fact_columns = {label: index for index, label in enumerate(headers) if label and index not in group_columns}
    results: list[ImportedBusinessResult] = []
    for row_index in range(header_row + 1, len(rows)):
        row = rows[row_index]
        if not any(value not in (None, "") for value in row):
            continue
        student = _text(_cell(row, student_column))
        if not student:
            continue
        fact_id = sha256(f"{source.name}|{sheet_name}|{row_index + 1}".encode()).hexdigest()[:20]
        facts = {label: _cell(row, index) for label, index in fact_columns.items() if _cell(row, index) not in (None, "")}
        for group_index, (teacher_column, headcount_column, performance_column) in enumerate(teacher_groups, start=1):
            teacher = _text(_cell(row, teacher_column))
            if not teacher:
                continue
            payload = {
                **facts,
                "teacher": teacher,
                "student": student,
                "refund_fact_id": fact_id,
                "refund_group_index": group_index,
                "refund_headcount": _cell(row, headcount_column),
                "refund_performance": _cell(row, performance_column),
            }
            results.append(ImportedBusinessResult(teacher, str(row_index + 1), payload, {
                "source_file": source.name, "sheet": sheet_name, "row": str(row_index + 1),
                "refund_fact_id": fact_id, "teacher_group_index": group_index,
                "teacher_group_headers": [headers[teacher_column], headers[headcount_column], headers[performance_column]],
            }))
    if not results:
        raise ValueError("退费绩效表没有可导入的教师绩效记录。")
    return results
