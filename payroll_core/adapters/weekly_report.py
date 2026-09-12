"""Read-only adapter for already-produced weekly operating reports."""
from __future__ import annotations

from dataclasses import asdict, dataclass, field
from functools import lru_cache
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from typing import Any, Iterable
from zipfile import BadZipFile

from ..models.evidence import AdapterIssue, AdapterResult


@dataclass(frozen=True)
class WeeklyReportRecord:
    period: str
    campus: str
    group: str
    teacher: str
    one_to_one_students: float | None
    class_students: float | None
    total_students: float | None
    one_to_one_weekly_average: float | None
    average_class_frequency: float | None
    source_file: str
    sheet: str
    cell: str
    range: str
    evidence: dict[str, Any] = field(default_factory=dict)

    def as_dict(self) -> dict[str, Any]:
        return asdict(self)


def _text(value: Any) -> str:
    return "" if value is None else str(value).replace("\n", " ").strip()


def _number(value: Any) -> float | None:
    if value in (None, "") or (isinstance(value, str) and value.strip().startswith("=")):
        return None
    try:
        return float(str(value).replace(",", ""))
    except (TypeError, ValueError):
        return None


def _period_from_value(value: Any, default: str) -> str:
    text = _text(value)
    if len(text) == 7 and text[4] == "-":
        return text
    return default


def _load_rows(path: Path) -> list[tuple[str, list[Any]]]:
    """Return one shared row snapshot per physical file content.

    Content hashing makes the cache independent of a copied filename and
    avoids stale rows when a workbook is replaced without a size/mtime change.
    The package scanner and all business adapters therefore consume the same
    decoded snapshot for one file hash.
    """
    payload = path.read_bytes()
    digest = sha256(payload).hexdigest()
    return _load_rows_cached(digest, payload)


@lru_cache(maxsize=128)
def _load_rows_cached(digest: str, payload: bytes) -> list[tuple[str, list[Any]]]:
    """Return (sheet, row) pairs for OOXML and legacy .xls-with-OOXML files."""
    try:
        import openpyxl
        _raw = openpyxl.load_workbook(BytesIO(payload), data_only=False, read_only=False, keep_links=True)
        cached = openpyxl.load_workbook(BytesIO(payload), data_only=True, read_only=False, keep_links=True)
        return [(ws.title, list(row)) for ws in cached.worksheets for row in ws.iter_rows(values_only=True)]
    except (ImportError, ValueError, OSError, BadZipFile):
        try:
            import xlrd
        except ImportError as exc:  # pragma: no cover - dependency is in the project test command
            raise ValueError("读取周报需要 openpyxl 或 xlrd。") from exc
        book = xlrd.open_workbook(file_contents=payload, on_demand=True)
        return [(sheet.name, sheet.row_values(row)) for sheet in book.sheets() for row in range(sheet.nrows)]


def _matrix(path: Path) -> list[tuple[str, list[Any]]]:
    try:
        return _load_rows(path)
    except Exception as exc:
        raise ValueError(f"无法读取周报 {path.name}: {exc}") from exc


def _find_header(rows: list[tuple[str, list[Any]]]) -> tuple[int, str, list[str]] | None:
    # Prefer a two-row grouped header when present.  This matters for both
    # the .xlsx and legacy .xls weekly templates.
    # When a workbook contains both raw course aggregates and a teacher
    # summary, the latter is the authoritative already-produced weekly
    # result.  Prefer it before scanning the raw-data sheet.
    for index, (sheet, row) in enumerate(rows):
        if "教师" not in sheet:
            continue
        texts = [_text(value) for value in row]
        joined = " ".join(texts)
        exact_teacher = any(text.replace(" ", "") in {"教师", "教师姓名", "姓名", "任课老师"} for text in texts)
        if not exact_teacher:
            continue
        if index + 1 < len(rows) and rows[index + 1][0] == sheet:
            second = [_text(value) for value in rows[index + 1][1]]
            merged = [" ".join(value for value in (texts[col] if col < len(texts) else "", second[col] if col < len(second) else "") if value) for col in range(max(len(texts), len(second)))]
            if any("一对一" in value for value in merged) and any("班课" in value for value in merged):
                return index, sheet, merged
        if "一对一" in joined and "班课" in joined:
            return index, sheet, texts
    for index in range(len(rows) - 1):
        sheet, first = rows[index]
        next_sheet, second = rows[index + 1]
        if sheet != next_sheet:
            continue
        first_texts, second_texts = [_text(value) for value in first], [_text(value) for value in second]
        exact_teacher = any(text.replace(" ", "") in {"教师", "教师姓名", "姓名", "任课老师"} for text in second_texts)
        if exact_teacher and ("一对一" in first_texts or "1对1" in first_texts) and "班课" in first_texts:
            merged = [" ".join(value for value in (first_texts[col] if col < len(first_texts) else "", second_texts[col] if col < len(second_texts) else "") if value) for col in range(max(len(first_texts), len(second_texts)))]
            return index, sheet, merged
    for index, (sheet, row) in enumerate(rows):
        texts = [_text(value) for value in row]
        joined = " ".join(texts)
        exact_teacher = any(text.replace(" ", "") in {"教师", "教师姓名", "姓名", "任课老师"} for text in texts)
        if exact_teacher and ("一对一" in joined or "1对1" in joined) and ("班课" in joined or "班级数" in joined):
            if index + 1 < len(rows) and rows[index + 1][0] == sheet and any("生数" in _text(value) or "周均" in _text(value) for value in rows[index + 1][1]):
                second = [_text(value) for value in rows[index + 1][1]]
                texts = [" ".join(value for value in (texts[col] if col < len(texts) else "", second[col] if col < len(second) else "") if value) for col in range(max(len(texts), len(second)))]
            return index, sheet, texts
        if exact_teacher and "续费" in joined:
            return index, sheet, texts
    return None


def _column(headers: list[str], *names: str) -> int | None:
    normalized = [_text(value).replace(" ", "") for value in headers]
    for name in names:
        target = name.replace(" ", "")
        for index, value in enumerate(normalized):
            if value == target or target in value:
                return index
    return None


def _is_summary_teacher(value: str) -> bool:
    normalized = value.replace(" ", "")
    return normalized in {"教师", "合计", "总计", "月总", "科组", "科组汇总", "校区", "学校"} or normalized.endswith(("组汇总", "科组"))


def read_weekly_report(path: str | Path, period: str) -> AdapterResult[WeeklyReportRecord]:
    source = Path(path).expanduser().resolve()
    result: AdapterResult[WeeklyReportRecord] = AdapterResult()
    try:
        rows = _matrix(source)
    except ValueError as exc:
        result.errors.append(AdapterIssue("UNREADABLE_WEEKLY_REPORT", str(exc)))
        return result
    found = _find_header(rows)
    if found is None:
        result.errors.append(AdapterIssue("UNRECOGNIZED_WEEKLY_REPORT", "没有找到含教师、一对一和班课字段的周报表头。"))
        return result
    header_index, sheet, headers = found
    one_students = _column(headers, "一对一生数", "一对一学生", "1对1学生", "1V1生数")
    class_students = _column(headers, "班课生数", "班课学生", "小班生数")
    total_students = _column(headers, "单科总数", "总学生数", "总学员数", "总人数")
    weekly_average = _column(headers, "一对一周均", "周平均", "周均")
    frequency = _column(headers, "平均课次", "平均班级频次", "平均频次")
    course_count = _column(headers, "总课次", "课次")
    teacher_col = next((index for index, value in enumerate(headers) if value.replace(" ", "") in {"教师", "教师姓名", "姓名", "任课老师"}), None)
    if teacher_col is None or one_students is None or class_students is None:
        result.errors.append(AdapterIssue("WEEKLY_REQUIRED_FIELD_MISSING", "周报缺少教师、一对一生数或班课生数字段。", sheet=sheet))
        return result
    title = " ".join(_text(value) for _sheet, row in rows[max(0, header_index - 3):header_index] for value in row if value not in (None, ""))
    # Some legacy .xls reports keep the title in a merged cell that is not in
    # the three rows immediately before the detected summary header.  The
    # file name is a safe secondary signal for campus/group, never for field
    # meaning.
    title = f"{title} {source.stem}"
    group = next((token for token in ("数学组", "理化组", "英语组", "综合组") if token in title), "")
    campus = next((token for token in ("宣城一校", "宣城二校", "宣城三校") if token in title), "")
    # A report may have multiple sheets with the same header.  The flattened
    # rows are only accepted from the header's sheet; this avoids duplicate
    # interpretation of a summary copied into another sheet.
    for absolute_index, (row_sheet, row) in enumerate(rows[header_index + 1:], start=header_index + 1):
        if row_sheet != sheet or teacher_col >= len(row):
            continue
        row_number = _sheet_row_number(rows, absolute_index, sheet)
        teacher = _text(row[teacher_col])
        if not teacher or _is_summary_teacher(teacher):
            continue
        one = _number(row[one_students]) if one_students < len(row) else None
        classes = _number(row[class_students]) if class_students < len(row) else None
        reported_total = _number(row[total_students]) if total_students is not None and total_students < len(row) else None
        # Never use a differently defined “单科总数” as the canonical
        # population when the two explicit populations are available.
        total = (one + classes) if one is not None and classes is not None else None
        weekly = _number(row[weekly_average]) if weekly_average is not None and weekly_average < len(row) else None
        avg_frequency = _number(row[frequency]) if frequency is not None and frequency < len(row) else None
        if avg_frequency is None and course_count is not None and total not in (None, 0) and course_count < len(row):
            courses = _number(row[course_count])
            avg_frequency = courses / total if courses is not None else None
        first = next((index for index, value in enumerate(row) if _text(value)), 0)
        last = max(first, len(row) - 1)
        result.records.append(WeeklyReportRecord(
            period=_period_from_value(None, period), campus=campus, group=group,
            teacher=teacher,
            one_to_one_students=one, class_students=classes, total_students=total,
            one_to_one_weekly_average=weekly, average_class_frequency=avg_frequency,
            source_file=source.name, sheet=sheet, cell=f"{sheet}!{_excel_col(teacher_col + 1)}{row_number}",
            range=f"{sheet}!{_excel_col(first + 1)}{row_number}:{_excel_col(last + 1)}{row_number}",
            evidence={"teacher": teacher, "row": row_number, "headers": headers, "reported_total_students": reported_total},
        ))
    if not result.records:
        result.warnings.append(AdapterIssue("WEEKLY_NO_DATA_ROWS", "找到周报表头，但没有可识别教师数据行。", sheet=sheet))
    result.coverage = {"records": len(result.records), "teachers": len(result.records), "recognized_fields": sum(value is not None for value in (one_students, class_students, weekly_average, frequency))}
    return result


def _excel_col(number: int) -> str:
    value = number
    output = ""
    while value:
        value, remainder = divmod(value - 1, 26)
        output = chr(65 + remainder) + output
    return output or "A"


def _sheet_row_number(rows: list[tuple[str, list[Any]]], absolute_index: int, sheet: str) -> int:
    """Translate the flattened snapshot index back to a workbook row."""
    return sum(1 for index in range(absolute_index + 1) if rows[index][0] == sheet)


read_weekly_report_excel = read_weekly_report
