"""Read-only structural audit for formula regions in payroll workbooks."""
from __future__ import annotations

import re
from collections import Counter
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

CELL = re.compile(r"(?<![A-Z0-9_])(?P<col>\$?[A-Z]{1,3})(?P<row>\$?\d+)")


@dataclass(frozen=True)
class FormulaAuditResult:
    workbook: str
    sheet: str
    cell: str
    field: str
    formula: str | None
    normalized_formula: str | None
    expected_pattern: str | None
    status: str
    evidence: str
    severity: str


def normalize_formula(formula: str, current_row: int) -> str:
    def replace(match: re.Match[str]) -> str:
        col, raw_row = match.group("col"), match.group("row")
        row = int(raw_row.replace("$", ""))
        row_token = raw_row if raw_row.startswith("$") else "{row}" if row == current_row else f"{{row{row - current_row:+d}}}"
        return f"{col}{row_token}"
    return CELL.sub(replace, formula.upper().replace(" ", ""))


def audit_payroll_formulas(path: str | Path, fields: dict[str, str] | None = None) -> list[FormulaAuditResult]:
    """Audit comparable teacher rows, without calculating any business value."""
    fields = fields or {"AE": "AE 该档每小时金额", "AF": "AF 总课时费"}
    book = load_workbook(path, data_only=False, read_only=False)
    output: list[FormulaAuditResult] = []
    for sheet in book.worksheets:
        header_row = next((row for row in range(1, min(sheet.max_row, 8) + 1) if any(sheet.cell(row, col).value == "姓名" for col in range(1, min(sheet.max_column, 20) + 1))), None)
        if header_row is None:
            continue
        name_col = next(col for col in range(1, min(sheet.max_column, 20) + 1) if sheet.cell(header_row, col).value == "姓名")
        teacher_rows = [row for row in range(header_row + 1, sheet.max_row + 1) if str(sheet.cell(row, name_col).value or "").strip()]
        for column, field in fields.items():
            formulas = {row: sheet[f"{column}{row}"].value for row in teacher_rows if isinstance(sheet[f"{column}{row}"].value, str) and sheet[f"{column}{row}"].value.startswith("=")}
            if len(formulas) < 2:
                continue  # insufficient context: do not fabricate an expected formula region
            patterns = {row: normalize_formula(value, row) for row, value in formulas.items()}
            expected, _ = Counter(patterns.values()).most_common(1)[0]
            expected_formula = next(value for row, value in formulas.items() if patterns[row] == expected)
            for row in teacher_rows:
                cell = sheet[f"{column}{row}"]
                value = cell.value
                if value is None or value == "":
                    status, evidence, severity = "FORMULA_MISSING", "同一教师公式区域存在空白。", "CRITICAL"
                elif not (isinstance(value, str) and value.startswith("=")):
                    status, evidence, severity = "FORMULA_REPLACED_BY_VALUE", "同一教师公式区域被固定值或文本替代。", "CRITICAL"
                else:
                    actual = normalize_formula(value, row)
                    if "#REF!" in value.upper() or "#VALUE!" in value.upper():
                        status, evidence, severity = "FORMULA_REGION_BREAK", "公式含有明确错误引用或错误值。", "CRITICAL"
                    elif actual == expected:
                        output.append(FormulaAuditResult(str(path), sheet.title, cell.coordinate, field, value, actual, expected, "FORMULA_MATCH", "公式结构与同列多数教师一致。", "INFO"))
                        continue
                    else:
                        references = CELL.findall(value.upper())
                        status = "FORMULA_PATTERN_MISMATCH"
                        evidence = "公式结构与同列多数教师不同。"
                        severity = "HIGH"
                        if any(not raw.startswith("$") and int(raw) != row for _, raw in references):
                            status, evidence = "ROW_REFERENCE_SHIFT", "公式引用了其他教师所在行。"
                        expected_cols = {match.group("col").replace("$", "") for match in CELL.finditer(expected_formula)}
                        actual_cols = {col.replace("$", "") for col, _ in references}
                        if actual_cols != expected_cols:
                            status, evidence = "COLUMN_REFERENCE_SHIFT", "公式引用列与同列多数公式不同。"
                output.append(FormulaAuditResult(str(path), sheet.title, cell.coordinate, field, value if isinstance(value, str) else None, normalize_formula(value, row) if isinstance(value, str) and value.startswith("=") else None, expected, status, evidence, severity))
    return output
