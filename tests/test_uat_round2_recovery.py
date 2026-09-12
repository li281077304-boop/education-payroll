"""Second-round UAT sweep: can a normal user recover from a mistake?

All cases are black box through the service / HTTP layer. Nothing here changes
business rules; it only proves that a wrong action is never a dead end.
"""
from __future__ import annotations

import json
from pathlib import Path
from threading import Thread
from urllib.request import Request, urlopen

import pytest
from openpyxl import Workbook, load_workbook

from payroll_ui.server import PayrollHttpServer
from payroll_ui.service import PayrollService


HEADERS = ["任课老师", "年级", "学科", "课程所属班型", "实到人数", "上课状态", "上课时间"]


def _schedule(path: Path, rows, *, extra_sheet: bool = False, merged: bool = False) -> Path:
    book = Workbook()
    sheet = book.active
    sheet.title = "课表"
    sheet.append(HEADERS)
    for row in rows:
        sheet.append(row)
    if merged:
        sheet.merge_cells("A9:B9")
        sheet.append([])          # 空行
        sheet.append([None] * 7)  # 全空行
    if extra_sheet:
        other = book.create_sheet("另一个工作表")
        other.append(["无关", "内容"])
        hidden = book.create_sheet("隐藏工作表")
        hidden.append(["隐藏", "内容"])
        hidden.sheet_state = "hidden"
    book.save(path)
    return path


def _lesson(teacher: str, day: str, class_type: str = "1对1", attended: int = 2):
    return [teacher, "八年级", "数学", class_type, attended, "已上课", f"{day} 10:00"]


def _service(tmp_path: Path):
    return PayrollService(tmp_path / "app")


# ------------------------------------------------------------------ 2.1 恢复
def test_wrong_upload_can_be_replaced(tmp_path):
    """传错文件后，重新选择同一个类别即可恢复，不需要重新建核算。"""
    service = _service(tmp_path)
    run = service.create("2026-08", "GENERATE")
    wrong = _schedule(tmp_path / "错的文件.xlsx", [_lesson("张三", "2026-08-05")])
    right = _schedule(tmp_path / "对的排课.xlsx", [
        _lesson("李四", "2026-08-06"), _lesson("李四", "2026-08-31"),
    ])

    service.import_file(run["id"], "schedule", str(wrong))
    recovered = service.import_file(run["id"], "schedule", str(right))

    assert recovered["files"]["schedule"]["name"] == "对的排课.xlsx"
    assert recovered["files"]["schedule"]["records"] == 2
    assert recovered["status"] == "FILES_READY", "替换后应能继续，不能卡死"


def test_replacement_does_not_destroy_run_state(tmp_path):
    """替换文件后：月份、已确认的其他材料、手工决定都不应被清空。"""
    service = _service(tmp_path)
    run = service.create("2026-08", "GENERATE")
    first = _schedule(tmp_path / "第一次.xlsx", [_lesson("张三", "2026-08-05")])
    second = _schedule(tmp_path / "第二次.xlsx", [_lesson("张三", "2026-08-20"), _lesson("张三", "2026-08-31")])

    service.import_file(run["id"], "schedule", str(first))
    service.change_period(run["id"], "2026-09")   # 制造一个非默认状态
    replaced = service.import_file(run["id"], "schedule", str(second))

    assert replaced["period"] == "2026-09", "替换文件不能悄悄改掉核算月份"
    assert replaced["files"]["schedule"]["name"] == "第二次.xlsx"
    assert replaced["period_check"]["run_month"] == "2026-09"


# ------------------------------------------------------------- 2.2 重复上传
def test_uploading_the_same_file_twice_is_recoverable(tmp_path):
    service = _service(tmp_path)
    run = service.create("2026-08", "GENERATE")
    path = _schedule(tmp_path / "排课.xlsx", [_lesson("张三", "2026-08-05")])

    first = service.import_file(run["id"], "schedule", str(path))
    second = service.import_file(run["id"], "schedule", str(path))

    assert second["files"]["schedule"]["records"] == first["files"]["schedule"]["records"]
    assert second["status"] == "FILES_READY"


def test_same_file_cannot_serve_two_roles_but_the_run_stays_usable(tmp_path):
    service = _service(tmp_path)
    run = service.create("2026-08", "GENERATE")
    path = _schedule(tmp_path / "排课.xlsx", [_lesson("张三", "2026-08-05")])
    service.import_file(run["id"], "schedule", str(path))

    with pytest.raises(ValueError) as error:
        service.import_file(run["id"], "math", str(path))

    message = str(error.value)
    assert "两个材料类别" in message
    assert "Traceback" not in message
    # 出错后原核算仍然可用
    assert service.get(run["id"])["files"]["schedule"]["records"] == 1


# --------------------------------------------------------------- 2.4 错误文案
@pytest.mark.parametrize("bad", ["", "  ", "2026-13", "2026-00", "202608", "abc-def"])
def test_invalid_period_errors_speak_chinese(tmp_path, bad):
    service = _service(tmp_path)
    run = service.create("2026-08", "GENERATE")
    with pytest.raises(ValueError) as error:
        service.change_period(run["id"], bad)
    message = str(error.value)
    for forbidden in ("Traceback", "NoneType", "database", ".py", "/Users", "/private", "internal", "KeyError"):
        assert forbidden not in message


def test_missing_file_errors_speak_chinese(tmp_path):
    service = _service(tmp_path)
    run = service.create("2026-08", "GENERATE")
    with pytest.raises(ValueError) as error:
        service.import_file(run["id"], "schedule", str(tmp_path / "不存在的文件.xlsx"))
    message = str(error.value)
    assert "找不到" in message
    for forbidden in ("Traceback", "NoneType", "database", ".py"):
        assert forbidden not in message


def test_generating_without_materials_speaks_chinese(tmp_path):
    service = _service(tmp_path)
    run = service.create("2026-08", "GENERATE")
    with pytest.raises(ValueError) as error:
        service.generate_payroll(run["id"], str(tmp_path / "工资表.xlsx"))
    message = str(error.value)
    assert "请先导入" in message
    assert "Traceback" not in message


# ------------------------------------------------- 2.5 Excel / WPS 健壮性
def test_chinese_filename_with_spaces_works(tmp_path):
    service = _service(tmp_path)
    run = service.create("2026-08", "GENERATE")
    path = _schedule(tmp_path / "8 月 排 课 表（最终版）.xlsx", [_lesson("张三", "2026-08-05")])

    imported = service.import_file(run["id"], "schedule", str(path))

    assert imported["files"]["schedule"]["name"] == "8 月 排 课 表（最终版）.xlsx"
    assert imported["files"]["schedule"]["records"] == 1


def test_merged_cells_empty_rows_and_extra_sheets_do_not_crash(tmp_path):
    service = _service(tmp_path)
    run = service.create("2026-08", "GENERATE")
    path = _schedule(
        tmp_path / "复杂表.xlsx",
        [_lesson("张三", "2026-08-05"), _lesson("张三", "2026-08-31")],
        extra_sheet=True, merged=True,
    )

    imported = service.import_file(run["id"], "schedule", str(path))

    sheets = imported["files"]["schedule"]["sheets"]
    assert imported["files"]["schedule"]["records"] == 2, "合并单元格/空行/多余工作表不能把记录吃掉"
    assert "课表" in sheets and "另一个工作表" in sheets, "多余工作表不能导致选错表"


def test_generated_workbook_reopens(tmp_path):
    """导出后必须能重新打开，且保留表头与状态说明。"""
    service = _service(tmp_path)
    run = service.create("2026-08", "GENERATE")
    path = _schedule(tmp_path / "排课.xlsx", [
        _lesson("张三", "2026-08-05"), _lesson("张三", "2026-08-31"),
    ])
    service.import_file(run["id"], "schedule", str(path))
    output = tmp_path / "工资表.xlsx"

    service.generate_payroll(run["id"], str(output))
    book = load_workbook(output)
    sheet = book["标准工资表"]

    assert sheet["A1"].value
    assert sheet["B3"].value == "AA 一对一折算小时"
    book.close()


# ------------------------------------------------------------- 2.3 导航稳定
def test_navigation_endpoints_return_expected_targets(tmp_path):
    """关键入口都必须可达：主页、静态资源、核算列表、单个核算。"""
    service = _service(tmp_path)
    service.create("2026-08", "GENERATE")
    static = Path(__file__).parents[1] / "payroll_ui" / "static"
    server = PayrollHttpServer(("127.0.0.1", 0), service, static)
    worker = Thread(target=server.serve_forever, daemon=True)
    worker.start()
    base = f"http://127.0.0.1:{server.server_port}"
    try:
        token = json.loads(urlopen(base + "/api/bootstrap").read())["token"]

        def call(path):
            request = Request(base + path, headers={"X-Payroll-Token": token})
            return urlopen(request).read()

        assert b"<!doctype html" in call("/").lower() or b"<html" in call("/").lower()
        assert b"function" in call("/static/app.js")
        runs = json.loads(call("/api/runs"))
        assert isinstance(runs, list) and runs
        detail = json.loads(call(f"/api/runs/{runs[0]['id']}"))
        assert detail["id"] == runs[0]["id"]
        assert "materials" in detail and "health" in detail, "返回后状态必须完整，不能只剩半页"
    finally:
        server.shutdown()
