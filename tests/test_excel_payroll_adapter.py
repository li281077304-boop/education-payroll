from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment

from payroll_core.excel.payroll import PAYROLL_LABELS, read_payroll_excel


FIXTURE = Path(__file__).parent / "fixtures" / "excel" / "fake_payroll.xlsx"


def _payroll_workbook(path: Path, *, formula_one_to_one: bool = False, comment: bool = False) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    for column, label in ((3, "姓名"), (27, "折算小时数"), (29, "班课折算小时数"), (31, "最终授课小时数据"), (32, "该档每小时金额"), (33, "总课时费"), (48, "总工资数")):
        sheet.cell(3, column, label)
    sheet.cell(5, 3, "张三")
    sheet.cell(5, 27, "=1+1" if formula_one_to_one else 36)
    sheet.cell(5, 29, 12)
    sheet.cell(5, 32, 100)
    sheet.cell(5, 33, 1200)
    sheet.cell(5, 48, 1500)
    if comment:
        sheet.cell(5, 29).comment = Comment("人工确认的班课值", "审核人")
    workbook.save(path)


def test_payroll_adapter_reads_semantic_columns_and_provenance():
    result = read_payroll_excel(FIXTURE, period="2026-08")

    assert result.ok
    record = result.records[0]
    assert record.teacher == "张三"
    assert record.one_to_one == 36
    assert record.provenance["one_to_one"].coordinate == "AA5"
    assert record.provenance["one_to_one"].source_field == PAYROLL_LABELS["one_to_one"]


def test_formula_without_cache_is_not_converted_to_zero(tmp_path):
    path = tmp_path / "missing-cache.xlsx"
    _payroll_workbook(path, formula_one_to_one=True)

    result = read_payroll_excel(path, period="2026-08")

    assert result.ok
    assert result.records[0].one_to_one is None
    assert any(issue.code == "MISSING_CACHE" for issue in result.warnings)


def test_payroll_adapter_extracts_comments_read_only(tmp_path):
    path = tmp_path / "comment.xlsx"
    _payroll_workbook(path, comment=True)

    result = read_payroll_excel(path, period="2026-08")

    assert result.ok
    assert len(result.comments) == 1
    assert result.comments[0].field == "class_value"
    assert result.comments[0].text == "人工确认的班课值"
