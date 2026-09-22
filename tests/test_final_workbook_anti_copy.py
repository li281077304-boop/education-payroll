from dataclasses import replace
import csv

from openpyxl import load_workbook

from payroll_core.excel.final_workbook_contract import create_empty_template
from payroll_core.excel.standard_payroll_render import render_generated_payroll
from payroll_core.final_fields import FINAL_FIELD_CODES
from payroll_ui.service import PayrollService
from tests.test_standard_payroll_output import _business_inputs, _generated
from tests.test_standard_payroll_output import _sanitized_template
from tests.test_payroll_modes_and_class_rules import _schedule


def test_company_renderer_is_current_run_driven_and_does_not_copy_old_rows(tmp_path):
    inputs = _business_inputs(complete=True)
    payroll = _generated(business_inputs=inputs)
    output = tmp_path / "current.xlsx"
    render_generated_payroll(payroll, output, template_path=_sanitized_template(tmp_path), manual_adjustments=inputs)
    book = load_workbook(output, data_only=False)
    sheet = book.active

    assert book.sheetnames[0] == "Sheet1"
    assert sheet["C5"].value == "教师甲"
    assert sheet["C6"].value is None
    assert sheet["AN5"].value == -60
    assert sheet["AN5"].comment is not None
    # A historical teacher/value/comment cannot survive an empty current roster.
    assert sheet["C16"].value is None
    assert sheet["AN16"].value is None
    assert sheet["AN16"].comment is None
    assert sheet["AK5"].value == "=AH5*1+AI5*1.5+AJ5*0.75"


def test_empty_template_has_no_historical_data_or_formula_cache(tmp_path):
    empty = tmp_path / "empty.xlsx"
    create_empty_template(_sanitized_template(tmp_path), empty)
    formula_book = load_workbook(empty, data_only=False)
    value_book = load_workbook(empty, data_only=True)
    sheet = formula_book.active
    values = value_book.active
    assert sheet["C5"].value is None
    assert sheet["AN5"].comment is None
    assert values["AK5"].value is None


def test_unresolved_m_or_renewal_components_do_not_restore_derived_formulas(tmp_path):
    payroll = _generated()
    row = payroll.rows[0]
    fields = {key: dict(value) for key, value in row.final_fields.items()}
    fields["M"] = {**fields.get("M", {}), "state": "HUMAN_REQUIRED", "value": None}
    fields["AV"] = {**fields.get("AV", {}), "state": "HUMAN_REQUIRED", "value": None}
    fields["AK"] = {**fields.get("AK", {}), "state": "HUMAN_REQUIRED", "value": None}
    blocked = replace(payroll, rows=(replace(row, final_fields=fields),))
    output = tmp_path / "blocked-derived-fields.xlsx"
    render_generated_payroll(blocked, output, template_path=_sanitized_template(tmp_path))

    sheet = load_workbook(output, data_only=False).active
    assert sheet["M5"].value is None
    assert sheet["AV5"].value is None
    assert sheet["AK5"].value is None


def test_mutated_current_source_changes_final_workbook(tmp_path):
    first_inputs = _business_inputs(complete=True)
    first = _generated(business_inputs=first_inputs)
    first_output = tmp_path / "first.xlsx"
    render_generated_payroll(first, first_output, template_path=_sanitized_template(tmp_path), manual_adjustments=first_inputs)

    mutated = [dict(item) for item in first_inputs]
    mutated[0] = {**mutated[0], "payload": {**mutated[0]["payload"], "one_to_one_hours": 91}}
    second = _generated(business_inputs=mutated)
    second_output = tmp_path / "second.xlsx"
    render_generated_payroll(second, second_output, template_path=_sanitized_template(tmp_path), manual_adjustments=mutated)

    first_sheet = load_workbook(first_output, data_only=False).active
    second_sheet = load_workbook(second_output, data_only=False).active
    assert first_sheet["AH5"].value == 2
    assert second_sheet["AH5"].value == 91
    assert second_sheet["AK5"].value == "=AH5*1+AI5*1.5+AJ5*0.75"

    refund_mutated = [dict(item) for item in first_inputs]
    refund_mutated[1] = {**refund_mutated[1], "payload": {**refund_mutated[1]["payload"], "headcount_amount": -80}}
    third = _generated(business_inputs=refund_mutated)
    third_output = tmp_path / "third.xlsx"
    render_generated_payroll(third, third_output, template_path=_sanitized_template(tmp_path), manual_adjustments=refund_mutated)
    assert load_workbook(third_output, data_only=False).active["AN5"].value == -90


def test_next_period_uses_current_roster_and_does_not_retain_august_data(tmp_path):
    inputs = _business_inputs(complete=True)
    payroll = replace(_generated(business_inputs=inputs), period="2026-09")
    payroll = replace(payroll, rows=(replace(payroll.rows[0], teacher="九月教师"),))
    output = tmp_path / "september.xlsx"
    render_generated_payroll(payroll, output, template_path=_sanitized_template(tmp_path), manual_adjustments=[])
    sheet = load_workbook(output, data_only=False).active
    assert sheet["C5"].value == "九月教师"
    assert sheet["C6"].value is None
    assert sheet["C16"].value is None
    assert sheet["AN5"].value == -60
    assert sheet["AN5"].comment is None
    assert sheet["H2"].value == "2026-09"


def test_explicit_manual_adjustment_overrides_only_current_run_field(tmp_path):
    inputs = _business_inputs(complete=True)
    payroll = _generated(business_inputs=inputs)
    adjustment = {
        "input_type": "MANUAL_ADJUSTMENT",
        "id": "adj-an-1",
        "run_id": "run-1",
        "period": "2026-08",
        "teacher_id": "教师甲",
        "teacher_name": "教师甲",
        "field": "AN",
        "raw_calculated": -60,
        "adjustment_value": -50,
        "final_value": -50,
        "reason": "已确认的人工退费调整",
        "evidence": "approved-refund.xlsx!A1",
        "actor": "审核人",
        "confirmed_at": "2026-08-31T00:00:00Z",
        "status": "APPROVED",
    }
    output = tmp_path / "adjusted.xlsx"
    render_generated_payroll(payroll, output, template_path=_sanitized_template(tmp_path), manual_adjustments=[*inputs, adjustment])
    sheet = load_workbook(output, data_only=False).active
    assert sheet["AN5"].value == -50
    assert sheet["C5"].value == "教师甲"


def test_manual_adjustment_production_path_reopen_and_revoke(tmp_path):
    service = PayrollService(tmp_path / "app-data")
    run = service.create("2026-08", "GENERATE")
    schedule = _schedule(tmp_path / "schedule.xlsx", [["教师甲", "九年级", "数学", "1对1", 1, "已上课"]] * 20)
    service.import_file(run["id"], "schedule", str(schedule))
    run = service.store.get(run["id"])
    run["template_path"] = str(_sanitized_template(tmp_path))
    service.store.save(run)

    refund = tmp_path / "refund.csv"
    with refund.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["教师", "headcount_amount", "performance_amount", "student"])
        writer.writerow(["教师甲", -100, 0, "脱敏学生"])
    imported = service.import_business_results("REFUND_RESULT", "2026-08", str(refund), "脱敏审核员")[0]
    imported = service.review_business_input(imported["id"], "START_REVIEW", "脱敏审核员")
    imported = service.review_business_input(imported["id"], "APPROVE", "脱敏审核员")
    service.bind_business_input(run["id"], imported["id"])

    adjustment = service.create_manual_adjustment(
        run["id"], teacher_id="教师甲", field="AN", raw_calculated=-100,
        adjustment_value=50, final_value=-50,
        reason="脱敏人工调整回归。", evidence="脱敏退费结果.csv:2", actor="脱敏审核员",
    )
    assert adjustment["status"] == "FINAL_CONFIRMED"
    adjusted_output = tmp_path / "adjusted.xlsx"
    service.generate_payroll(run["id"], str(adjusted_output))
    adjusted_sheet = load_workbook(adjusted_output, data_only=False).worksheets[0]
    an_column = 33 + FINAL_FIELD_CODES.index("AN")
    assert adjusted_sheet.cell(5, an_column).value == -50

    revoked = dict(adjustment, status="REVOKED")
    service.store.save_business_input(revoked)
    run = service.store.get(run["id"])
    run["business_input_bindings"] = [binding for binding in run.get("business_input_bindings", []) if binding["input_id"] != adjustment["id"]]
    service.store.save(run)
    raw_output = tmp_path / "raw.xlsx"
    service.generate_payroll(run["id"], str(raw_output))
    raw_sheet = load_workbook(raw_output, data_only=False).worksheets[0]
    assert raw_sheet.cell(5, an_column).value == -100
