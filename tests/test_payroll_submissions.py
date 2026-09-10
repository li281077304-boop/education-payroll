"""Teacher personal payroll sheets: recognize, standardize, merge, publish."""
from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook, Workbook

from payroll_ui.service import PayrollService


HEADERS_FULL = ["教师", "一对一折算", "班课折算", "课时生产", "AE", "AF", "AV"]


def _csv(path: Path, text: str) -> Path:
    path.write_text(text, encoding="utf-8")
    return path


def _sheet(path: Path, headers: list[str], rows: list[list[object]]) -> Path:
    book = Workbook()
    sheet = book.active
    sheet.title = "工资表"
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    book.save(path)
    return path


def _service(tmp_path: Path) -> PayrollService:
    return PayrollService(tmp_path / "data")


def test_single_teacher_sheet_becomes_a_standard_submission(tmp_path):
    service = _service(tmp_path)
    path = _csv(tmp_path / "teacher-a.csv", "教师,1对1折算,班课折算,课时生产,AE,AF,AV\n张三,36,10,3.24,1,2,3\n")

    batch = service.import_payroll_sheets([str(path)], "2026-08", "张三")

    merged = service.preview_payroll_merge(batch["id"])["merged"]
    assert merged["张三"]["one_to_one"] == 36.0
    assert merged["张三"]["class_value"] == 10.0
    assert merged["张三"]["av"] == 3.0
    stored = service.store.list_submissions(batch["id"])[0]
    assert stored["source_kind"] == "SINGLE"
    assert stored["source_file"] == "teacher-a.csv"


def test_many_personal_sheets_merge_into_one_table(tmp_path):
    service = _service(tmp_path)
    first = _csv(tmp_path / "a.csv", "教师,一对一折算,班课折算\n张三,36,10\n")
    second = _csv(tmp_path / "b.csv", "老师,1对1折算,班课折算小时数\n李四,40,20\n")

    batch = service.import_payroll_sheets([str(first), str(second)], "2026-08", "教务")
    preview = service.preview_payroll_merge(batch["id"])

    assert set(preview["merged"]) == {"张三", "李四"}
    assert preview["merged"]["李四"]["class_value"] == 20.0
    assert not preview["has_errors"]


def test_summary_workbook_is_the_same_standard_shape(tmp_path):
    service = _service(tmp_path)
    summary = _sheet(tmp_path / "summary.xlsx", HEADERS_FULL, [["张三", 36, 10, 1, 1, 2, 3], ["李四", 40, 20, 2, 1, 2, 3]])

    batch = service.import_payroll_sheets([str(summary)], "2026-08", "教务")
    preview = service.preview_payroll_merge(batch["id"])

    kinds = {item["source_kind"] for item in service.store.list_submissions(batch["id"])}
    assert kinds == {"SUMMARY"}
    assert set(preview["merged"]) == {"张三", "李四"}


def test_duplicate_conflict_and_missing_teachers_are_blocking_findings(tmp_path):
    service = _service(tmp_path)
    first = _csv(tmp_path / "a.csv", "教师,一对一折算,班课折算\n张三,36,10\n")
    second = _csv(tmp_path / "b.csv", "教师,一对一折算,班课折算\n张三,36,99\n")

    batch = service.import_payroll_sheets([str(first), str(second)], "2026-08", "教务")
    preview = service.preview_payroll_merge(batch["id"], expected_teachers=["张三", "王五"])

    codes = {item["code"] for item in preview["findings"]}
    assert {"DUPLICATE_TEACHER", "FIELD_CONFLICT", "MISSING_TEACHER"} <= codes
    assert preview["duplicate_teachers"] == ["张三"]
    assert preview["missing_teachers"] == ["王五"]
    with pytest.raises(ValueError, match="必须处理"):
        service.confirm_payroll_merge(batch["id"], str(tmp_path / "out.xlsx"), "教务")


def test_ambiguous_layout_waits_for_confirmation_then_reuses_the_profile(tmp_path):
    service = _service(tmp_path)
    path = _sheet(tmp_path / "ambiguous.xlsx", ["教师", "老师", "班课折算"], [["张三", "张三", 10]])

    batch = service.import_payroll_sheets([str(path)], "2026-08", "教务")
    assert batch["status"] == "NEEDS_CONFIRMATION"
    assert batch["pending_layouts"][0]["ambiguous"]["teacher"] == ["教师", "老师"]
    assert service.store.list_submissions(batch["id"]) == []

    fingerprint = batch["pending_layouts"][0]["fingerprint"]
    confirmed = service.confirm_payroll_layout(batch["id"], fingerprint, {"teacher": "教师", "class_value": "班课折算"}, "教务")
    assert confirmed["status"] == "DRAFT"
    assert service.preview_payroll_merge(batch["id"])["merged"]["张三"]["class_value"] == 10.0

    again = service.import_payroll_sheets([str(path)], "2026-08", "教务")
    assert again["pending_layouts"] == []
    assert service.store.list_submissions(again["id"])[0]["layout_profile_id"]


def test_layout_variants_between_subject_groups_are_not_rejected(tmp_path):
    service = _service(tmp_path)
    math = _csv(tmp_path / "math.csv", "教师,一对一折算,班课折算,课时生产\n张三,36,10,1\n")
    english = _csv(tmp_path / "english.csv", "老师,1对1折算,班课折算小时数,生产课时\n李四,40,20,2\n")

    batch = service.import_payroll_sheets([str(math), str(english)], "2026-08", "教务")

    assert batch["pending_layouts"] == []
    merged = service.preview_payroll_merge(batch["id"])["merged"]
    assert merged["张三"]["production"] == 1.0
    assert merged["李四"]["production"] == 2.0


def test_merge_publishes_a_new_workbook_from_the_model_without_overwriting(tmp_path):
    service = _service(tmp_path)
    source = _csv(tmp_path / "a.csv", "教师,一对一折算,班课折算\n张三,36,10\n")
    summary = _sheet(tmp_path / "s.xlsx", HEADERS_FULL, [["李四", 40, 20, 2, 1, 2, 3]])
    batch = service.import_payroll_sheets([str(source), str(summary)], "2026-08", "教务")
    source_bytes = source.read_bytes()

    target = tmp_path / "标准工资表.xlsx"
    result = service.confirm_payroll_merge(batch["id"], str(target), "教务")

    assert result["teacher_count"] == 2
    assert source.read_bytes() == source_bytes
    book = load_workbook(target)
    sheet = book["标准工资表"]
    assert sheet["A2"].value == "教师"
    assert [sheet.cell(row=row, column=1).value for row in (3, 4)] == ["张三", "李四"]
    with pytest.raises(ValueError, match="不能覆盖"):
        service.confirm_payroll_merge(batch["id"], str(target), "教务")


def test_payroll_sheets_never_create_a_management_assessment(tmp_path):
    service = _service(tmp_path)
    path = _csv(tmp_path / "a.csv", "教师,一对一折算,班课折算\n张三,36,10\n")

    service.import_payroll_sheets([str(path)], "2026-08", "张三")

    assert service.assessments.records("2026-08") == []
    assert service.assessments.results("2026-08") == []
