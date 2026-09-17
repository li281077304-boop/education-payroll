"""Regression tests for the generated standard payroll workbook contract."""
from __future__ import annotations

from pathlib import Path
from dataclasses import replace

from openpyxl import load_workbook
import pytest

from payroll_core.calculation import CompensationProfile, RatingAuthority, calculate_payroll
from payroll_core.config.core_rules import load_core_rules
from payroll_core.final_fields import DETERMINED, NOT_APPLICABLE, resolve_final_fields
from payroll_core.excel.standard_payroll_render import (
    OUT_OF_SCOPE_FINAL_FIELDS,
    HUMAN_REQUIRED,
    HEADERS,
    render_generated_payroll,
    validate_standard_payroll_workbook,
)
from payroll_core.models.records import ScheduleRecord
from payroll_core.payroll_generation import GeneratedPayroll, base_salary_field, build_generated_payroll, generated_from_calculation
from payroll_ui.service import PayrollService
from tests.payroll_test_helpers import TEMPLATE_PATH, bind_template


def _generated(*, reference_rating: bool = False, business_inputs=()):
    schedule = [
        ScheduleRecord(
            period="2026-08", teacher="教师甲", grade="九年级", subject="数学",
            class_type="1对1", attended=1, lesson_status="已上课",
            source="脱敏排课.xlsx",
        )
        for _ in range(16)
    ]
    kwargs = {
        "ratings": () if reference_rating else (RatingAuthority("教师甲", 4, "2026-08", "2026-09", "脱敏星级权威", "stars-v1"),),
        "reference_ratings": {"教师甲": 4} if reference_rating else {},
        "profiles": (CompensationProfile("教师甲", 30, "2026-08", "2026-09", "脱敏个人政策", True),),
    }
    result = calculate_payroll("2026-08", schedule, load_core_rules(), **kwargs).as_dict()
    names = {"aa": "AA", "ac": "AC", "ad": "AD", "ae": "AE", "af": "AF", "part_time_fee": "PART_TIME"}
    rows = [{
        "teacher": row["teacher"],
        "fields": {code: {**row[key], "value": row[key]["value"]} for key, code in names.items()},
    } for row in result["rows"]]
    return generated_from_calculation({
        "period": result["period"], "rows": rows,
        "rule_versions": {"core": result["rule_version_id"], "rating": "stars-v1"},
        "source_records": [{"record_key": result["course_contributions"][0]["record_key"], "teacher": "教师甲", "source": "脱敏排课.xlsx", "provenance": {}}],
    }, business_inputs=business_inputs)


def _business_inputs(*, complete: bool = False):
    values = [
        {"id": "renew-1", "teacher_id": "教师甲", "input_type": "RENEWAL_RESULT", "status": "APPROVED", "source_ref": "脱敏续费结果.csv", "source_row": "2", "payload": {"one_to_one_hours": 2, "class_hours": 3, "mentor_hours": 4}},
        {"id": "refund-1", "teacher_id": "教师甲", "input_type": "REFUND_RESULT", "status": "APPROVED", "source_ref": "脱敏退费结果.csv", "source_row": "5", "payload": {"headcount_amount": -50, "performance_amount": -10}},
    ]
    if complete:
        for index, code in enumerate(("AG", "AL", "AM", "AO", "AP", "AQ", "AR", "AS", "AT", "AU"), start=1):
            values.append({"id": f"other-{code}", "teacher_id": "教师甲", "input_type": "OTHER", "status": "APPROVED", "source_ref": "脱敏工资补充结果.csv", "source_row": str(index), "payload": {"target_field": code, "amount": index}})
    return values


def test_generated_workbook_contains_static_values_evidence_and_boundary(tmp_path: Path):
    payroll = _generated()
    output = tmp_path / "标准工资表.xlsx"

    assert payroll.rows[0].status == "NEEDS_CONFIRMATION"
    assert payroll.rows[0].blockers
    path = render_generated_payroll(payroll, output, template_path=TEMPLATE_PATH)
    book = load_workbook(path, data_only=False)
    sheet = book["Sheet1"]

    assert sheet["C5"].value == "教师甲"
    assert sheet["AA5"].value.startswith("=")
    assert sheet["AD5"].value.startswith("=")
    assert sheet["AF5"].value.startswith("=")
    assert "外围字段状态" in book.sheetnames
    assert {book["外围字段状态"].cell(row, 1).value for row in range(4, 4 + len(OUT_OF_SCOPE_FINAL_FIELDS))} == set(OUT_OF_SCOPE_FINAL_FIELDS)
    assert book["外围字段状态"]["B4"].value == HUMAN_REQUIRED
    evidence_text = "\n".join(str(cell.value) for row in book["核验与来源"].iter_rows() for cell in row if cell.value is not None)
    assert "脱敏排课.xlsx" in evidence_text
    assert '"rating": "4"' in evidence_text
    assert validate_standard_payroll_workbook(path, payroll)["ok"] is True


def test_output_validation_detects_post_generation_tampering(tmp_path: Path):
    payroll = _generated()
    output = tmp_path / "output.xlsx"
    render_generated_payroll(payroll, output, template_path=TEMPLATE_PATH)
    book = load_workbook(output)
    book["Sheet1"]["AE5"] = 999
    book.save(output)

    validation = validate_standard_payroll_workbook(output, payroll)
    assert validation["ok"] is False
    assert any("教师甲" in error and "ae" in error for error in validation["errors"])


def test_output_validation_detects_evidence_tampering(tmp_path: Path):
    payroll = _generated()
    output = tmp_path / "tampered-evidence.xlsx"
    render_generated_payroll(payroll, output, template_path=TEMPLATE_PATH)
    book = load_workbook(output)
    evidence = book["核验与来源"]
    for row in evidence.iter_rows():
        if row[0].value == "教师甲" and row[1].value == "AA":
            row[4].value = "被篡改"
            break
    book.save(output)
    validation = validate_standard_payroll_workbook(output, payroll)
    assert validation["ok"] is False
    assert any("核验与来源第" in error for error in validation["errors"])


def test_preview_uses_final_fields_without_writing_until_export(tmp_path: Path, monkeypatch):
    """Preview and export share final-field calculation; only export writes XLSX."""
    service = PayrollService(tmp_path / "app-data")
    run = service.create("2026-08", mode="GENERATE")
    run["calculation_engine"] = "CONFIGURED_V1"
    service.store.save(run)
    inputs = _business_inputs(complete=True)
    for item in inputs:
        item = {**item, "created_at": "2026-08-01T00:00:00+00:00"}
        service.store.save_business_input(item)
    checked = {
        "period": "2026-08",
        "business_input_bindings": [{"input_id": item["id"], "source_file_hash": ""} for item in inputs],
        "core_calculation": {
            "period": "2026-08",
            "rule_versions": {"core": "core-v1"},
            "rows": [{
                "teacher": "教师甲", "employment_type": "FULL_TIME",
                "fields": {
                    "AA": {"value": 20, "state": DETERMINED, "reason": "测试核心规则"},
                    "AC": {"value": 2.4, "state": DETERMINED, "reason": "测试核心规则"},
                    "AD": {"value": 22.4, "state": DETERMINED, "reason": "测试核心规则"},
                    "AE": {"value": 40, "state": DETERMINED, "reason": "测试核心规则"},
                    "AF": {"value": 80, "state": DETERMINED, "reason": "测试核心规则"},
                    "PART_TIME": {"value": None, "state": NOT_APPLICABLE, "reason": "测试核心规则"},
                },
            }],
        },
    }
    monkeypatch.setattr(service, "check", lambda _run_id: checked)

    preview = service.preview_payroll(run["id"])
    row = preview["generated_payroll"]["rows"][0]
    assert row["final_fields"]["AK"]["value"] == pytest.approx(9.5)
    assert row["final_fields"]["AV"]["value"] == pytest.approx(80 + 1 + 2 + 3 + 4 + 5 + 6 + 7 + 8 + 9 + 10 + 9.5 - 60)
    assert not list(tmp_path.rglob("*.xlsx"))

    output = tmp_path / "工资表.xlsx"
    bind_template(service, run)
    result = service.generate_payroll(run["id"], str(output))
    assert Path(result["path"]).is_file()
    reopened = load_workbook(result["path"], data_only=False)
    columns = {header: index + 1 for index, header in enumerate(HEADERS)}
    assert reopened["标准工资表"].cell(5, 37).value.startswith("=")
    assert reopened["标准工资表"]["M5"].value is None
    second = service.generate_payroll(run["id"], str(output))
    assert Path(second["path"]).is_file()
    assert Path(second["path"]) != output


def test_reference_rating_is_visible_and_determined(tmp_path: Path):
    payroll = _generated(reference_rating=True)
    output = tmp_path / "estimated.xlsx"
    render_generated_payroll(payroll, output, template_path=TEMPLATE_PATH)
    sheet = load_workbook(output)["Sheet1"]

    assert sheet["F5"].value == "已确定/上传资料/4星"
    assert payroll.status == "NEEDS_CONFIRMATION"


def test_template_output_keeps_auditable_excel_formulas(tmp_path: Path):
    payroll = _generated()
    payroll = replace(payroll, formula_inputs={"one_to_one_counts": {"教师甲": {"九年级": 16}}})
    template = Path("/Users/macos/Desktop/payroll_read_test/薪资表模板.xlsx")
    output = tmp_path / "formula-output.xlsx"
    render_generated_payroll(payroll, output, template_path=template)
    sheet = load_workbook(output, data_only=False)["Sheet1"]
    assert sheet["H2"].value == payroll.period
    assert sheet["AA5"].value.startswith("=(N5+O5+P5+Q5+R5+S5)")
    assert sheet["AD5"].value == "=AA5+AC5"
    assert sheet["AF5"].value.startswith("=MAX(0,(AD5-")
    assert sheet["AK5"].value == "=AH5*1+AI5*1.5+AJ5*0.75"
    assert sheet["AV5"].value.startswith("=M5+AF5+AG5+AK5")


def test_template_output_writes_base_salary_inputs_and_m_formula(tmp_path: Path):
    payroll = _generated()
    payroll = generated_from_calculation(
        {
            "period": "2026-08",
            "rows": [{"teacher": "教师甲", "fields": {
                "AA": {"value": 2, "state": DETERMINED},
                "AC": {"value": 0, "state": DETERMINED},
                "AD": {"value": 2, "state": DETERMINED},
                "AE": {"value": 0, "state": DETERMINED},
                "AF": {"value": 0, "state": DETERMINED},
                "PART_TIME": {"value": None, "state": NOT_APPLICABLE},
            }}],
        },
        base_salary_inputs={"教师甲": {"source": "工资资料包", "fields": {code: {"value": value} for code, value in {"G": 6000, "H": 500, "I": 200, "J": 300, "K": 20, "L": 18}.items()}}},
    )
    output = tmp_path / "base-salary-formula.xlsx"
    render_generated_payroll(payroll, output, template_path=Path("/Users/macos/Desktop/payroll_read_test/薪资表模板.xlsx"))
    sheet = load_workbook(output, data_only=False)["Sheet1"]
    assert [sheet[f"{column}5"].value for column in "GHIJKL"] == [6000, 500, 200, 300, 20, 18]
    assert sheet["M5"].value == "=(G5+H5+I5+J5)/K5*L5"


def test_calculation_derives_template_grade_inputs_from_accepted_one_to_one_courses(tmp_path: Path):
    """Template AA inputs must be populated from the canonical Core result."""
    schedule = [
        ScheduleRecord(
            period="2026-08", teacher="教师甲", grade="九年级", subject="数学",
            class_type="1对1", attended=1, lesson_status="已上课", source="脱敏排课.xlsx",
        )
        for _ in range(2)
    ]
    generated = build_generated_payroll(
        period="2026-08", schedule_records=schedule, rules=load_core_rules(),
        ratings=(RatingAuthority("教师甲", 4, "2026-08", "2026-09", "脱敏星级权威", "stars-v1"),),
        profiles=(CompensationProfile("教师甲", 30, "2026-08", "2026-09", "脱敏个人政策", True),),
    )
    assert generated.formula_inputs["one_to_one_counts"]["教师甲"]["九年级"] == 6
    output = tmp_path / "formula-inputs.xlsx"
    render_generated_payroll(generated, output, template_path=Path("/Users/macos/Desktop/payroll_read_test/薪资表模板.xlsx"))
    sheet = load_workbook(output, data_only=False)["Sheet1"]
    # 九年级 maps to V; two accepted lessons become six count units (2 * 3).
    assert sheet["V5"].value == 6


def test_approved_star_override_is_carried_even_when_zero_tier_needs_no_star_for_rate():
    schedule = [
        ScheduleRecord(
            period="2026-08", teacher="教师甲", grade="九年级", subject="数学",
            class_type="1对1", attended=1, lesson_status="已上课",
            source="脱敏排课.xlsx",
        )
        for _ in range(10)
    ]
    payroll = build_generated_payroll(
        period="2026-08",
        schedule_records=schedule,
        profiles=(CompensationProfile(
            "教师甲", 30, "2026-08", "2026-09", "脱敏个人政策", True,
            6, "审批人", "2026-08-01",
        ),),
    )

    row = payroll.rows[0]
    assert row.teaching_hours == 20
    assert row.star == 6
    assert row.ae == 0
    assert row.af == 0


def test_management_award_is_not_applicable_only_for_explicitly_non_management_rows():
    ordinary = resolve_final_fields(
        teacher="教师甲",
        core_fields={"AF": {"value": 80, "state": DETERMINED}, "PART_TIME": {"value": None, "state": NOT_APPLICABLE}},
        employment_type="FULL_TIME",
    )
    management = resolve_final_fields(
        teacher="组长甲",
        core_fields={"AF": {"value": 80, "state": DETERMINED}, "PART_TIME": {"value": None, "state": NOT_APPLICABLE}},
        employment_type="MANAGEMENT",
    )

    assert ordinary["AM"]["state"] == NOT_APPLICABLE
    assert management["AM"]["state"] == HUMAN_REQUIRED


def test_base_salary_m_uses_run_snapshot_and_binds_into_av():
    inputs = {
        "教师甲": {
            "teacher_id": "t-1",
            "display_name": "教师甲",
            "source": "工资资料包/薪资输入.xlsx",
            "provenance": {"source_row": 8},
            "fields": {code: {"value": value, "source": "工资资料包"} for code, value in {
                "G": 6000, "H": 500, "I": 200, "J": 300, "K": 20, "L": 18,
            }.items()},
        }
    }
    result = base_salary_field("教师甲", inputs)
    assert result["state"] == DETERMINED
    assert result["value"] == 6300.0
    assert result["evidence"][0]["formula"] == "M = (G + H + I + J) / K × L"

    payroll = _generated()
    result = calculate_payroll(
        "2026-08", [
            ScheduleRecord(period="2026-08", teacher="教师甲", grade="九年级", subject="数学", class_type="1对1", attended=1, lesson_status="已上课", source="排课.xlsx")
            for _ in range(16)
        ], load_core_rules(), ratings=(RatingAuthority("教师甲", 4, "2026-08", "2026-09", "星级.xlsx", "stars-v1"),),
    ).as_dict()
    generated = generated_from_calculation({"period": result["period"], "rows": [{"teacher": "教师甲", "fields": {code: {**row[code_key], "value": row[code_key]["value"]} for code_key, code in {"aa": "AA", "ac": "AC", "ad": "AD", "ae": "AE", "af": "AF", "part_time_fee": "PART_TIME"}.items()}} for row in result["rows"]]}, base_salary_inputs=inputs)
    assert generated.rows[0].final_fields["M"]["value"] == 6300.0
    assert generated.rows[0].final_fields["AV"]["state"] == HUMAN_REQUIRED


def test_base_salary_m_missing_input_is_blocked_never_zero():
    result = base_salary_field("教师甲", {"教师甲": {"fields": {"G": {"value": 1}}}})
    assert result["value"] is None
    assert result["state"] == "BLOCKED_BY_INPUT"
    assert all(code in result["reason"] for code in ("H", "I", "J", "K", "L"))


def test_empty_core_result_cannot_be_published_as_a_standard_payroll(tmp_path: Path):
    empty = GeneratedPayroll(period="2026-08", rows=(), status="NEEDS_CONFIRMATION")
    with pytest.raises(ValueError, match="没有可生成的教师记录"):
        render_generated_payroll(empty, tmp_path / "empty.xlsx")


def test_output_order_is_stable_and_duplicate_teachers_are_rejected(tmp_path: Path):
    payroll = _generated()
    row = payroll.rows[0]
    unsorted = GeneratedPayroll(
        period=payroll.period,
        rows=(replace(row, teacher="教师乙"), row),
        status=payroll.status,
        blockers=payroll.blockers,
        rule_versions=payroll.rule_versions,
        source_records=payroll.source_records,
    )
    output = tmp_path / "sorted.xlsx"
    render_generated_payroll(unsorted, output, template_path=TEMPLATE_PATH)
    sheet = load_workbook(output)["Sheet1"]
    assert [sheet.cell(index, 3).value for index in (5, 6)] == ["教师乙", "教师甲"]

    duplicate = GeneratedPayroll(
        period=payroll.period,
        rows=(row, row),
        status=payroll.status,
    )
    with pytest.raises(ValueError, match="教师不能重复"):
        render_generated_payroll(duplicate, tmp_path / "duplicate.xlsx", template_path=TEMPLATE_PATH)


def test_export_without_company_template_fails_closed(tmp_path: Path):
    with pytest.raises(ValueError, match="未绑定公司工资模板"):
        render_generated_payroll(_generated(), tmp_path / "no-template.xlsx")


def test_approved_business_results_feed_known_fields_and_ak_without_making_av_zero(tmp_path: Path):
    payroll = _generated(business_inputs=_business_inputs())
    output = tmp_path / "business-results.xlsx"
    render_generated_payroll(payroll, output, template_path=TEMPLATE_PATH)
    sheet = load_workbook(output)["Sheet1"]
    assert sheet["C5"].value == "教师甲"
    boundary = load_workbook(output)["外围字段状态"]
    boundary_rows = {boundary.cell(row, 1).value: row for row in range(4, 4 + len(OUT_OF_SCOPE_FINAL_FIELDS))}
    assert boundary.cell(boundary_rows["AH"], 2).value == "DETERMINED"
    assert boundary.cell(boundary_rows["AV"], 2).value == HUMAN_REQUIRED


def test_explicitly_named_approved_components_allow_av_formula(tmp_path: Path):
    payroll = _generated(business_inputs=_business_inputs(complete=True))
    assert payroll.final is True
    assert payroll.rows[0].final_fields["AV"]["value"] == pytest.approx(80 + 1 + (2 + 3 * 1.5 + 4 * 0.75) + 2 + 3 - 50 - 10 + 4 + 5 + 6 + 7 + 8 + 9 + 10)
    output = tmp_path / "complete-components.xlsx"
    render_generated_payroll(payroll, output, template_path=TEMPLATE_PATH)
    sheet = load_workbook(output)["Sheet1"]
    assert sheet["C5"].value == "教师甲"
    assert sheet["AV5"].value.startswith("=")


def test_reopened_validation_catches_boundary_state_tampering(tmp_path: Path):
    payroll = _generated(business_inputs=_business_inputs())
    output = tmp_path / "tampered-boundary.xlsx"
    render_generated_payroll(payroll, output, template_path=TEMPLATE_PATH)
    book = load_workbook(output)
    book["外围字段状态"]["B5"] = "HUMAN_REQUIRED"  # AH is actually determined by the bound result.
    book.save(output)

    validation = validate_standard_payroll_workbook(output, payroll)
    assert validation["ok"] is False
    assert any("外围字段 AH 状态不一致" in error for error in validation["errors"])


def test_part_time_pay_enters_final_af_and_standard_workbook(tmp_path: Path):
    schedule = [
        ScheduleRecord(
            period="2026-08", teacher="兼职甲", grade="九年级", subject="数学",
            class_type="小班", attended=6, lesson_status="已上课",
            source="脱敏排课.xlsx",
        )
        for _ in range(2)
    ]
    payroll = build_generated_payroll(
        period="2026-08",
        schedule_records=schedule,
        teacher_contexts=({"teacher": "兼职甲", "employment_type": "PART_TIME", "effective_from": "2026-08", "effective_to": "2026-09"},),
        part_time_rates=({"teacher": "兼职甲", "grade": "九年级", "rate_per_lesson": 123, "effective_from": "2026-08", "effective_to": "2026-09", "approved_by": "审核员", "approved_at": "2026-08-01", "source": "脱敏兼职单价"},),
    )

    row = payroll.rows[0]
    assert row.af is None
    assert row.final_fields["AF"]["value"] == 246.0
    assert row.final_fields["AF"]["state"] == DETERMINED
    assert row.final_fields["AV"]["state"] == HUMAN_REQUIRED

    output = tmp_path / "part-time.xlsx"
    render_generated_payroll(payroll, output, template_path=TEMPLATE_PATH)
    sheet = load_workbook(output)["Sheet1"]
    assert sheet["C5"].value == "兼职甲"
    assert sheet["AF5"].value.startswith("=")
