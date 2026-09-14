from pathlib import Path

from openpyxl import Workbook

from payroll_core.excel.schedule import read_schedule_excel
from payroll_ui.service import PayrollService


FIXTURE = Path(__file__).parent / "fixtures" / "excel" / "fake_schedule.xlsx"
REQUIRED_HEADERS = ["上课班级", "教学形式", "上课时间", "上课状态", "实到", "上课学员", "上课科目", "任课老师"]


def _schedule_workbook(path: Path, headers: list[str]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "排课记录-08月03日到08月30日"
    sheet.append(headers)
    sheet.append([{"上课班级": "初一数学一对一", "教学形式": "一对一", "上课时间": "2026-08-03 09:00", "上课状态": "已上课", "实到": 1, "上课学员": "学生甲", "上课科目": "数学", "任课老师": "张三"}.get(header, "") for header in headers])
    workbook.save(path)


def test_schedule_adapter_reads_fake_current_layout_with_normalized_fields():
    result = read_schedule_excel(FIXTURE, period="2026-08")

    assert result.ok
    assert len(result.records) == 2
    record = result.records[0]
    assert (record.grade, record.subject, record.class_type, record.attended) == ("七年级", "数学", "1对1", 1)
    assert record.lesson_date == "2026-08-10"
    assert record.provenance["teacher"].source_field == "任课老师"
    assert record.provenance["teacher"].coordinate == "O2"


def test_schedule_adapter_reports_missing_required_column(tmp_path):
    path = tmp_path / "missing-attendance.xlsx"
    headers = [header for header in REQUIRED_HEADERS if header != "实到"]
    _schedule_workbook(path, headers)

    result = read_schedule_excel(path, period="2026-08")

    assert not result.ok
    assert result.errors[0].code == "MISSING_REQUIRED_COLUMN"


def test_schedule_adapter_uses_headers_not_fixed_column_positions(tmp_path):
    path = tmp_path / "moved-columns.xlsx"
    headers = list(reversed(REQUIRED_HEADERS))
    _schedule_workbook(path, headers)

    result = read_schedule_excel(path, period="2026-08")

    assert result.ok
    assert result.records[0].teacher == "张三"
    assert result.records[0].attended == 1


def test_schedule_adapter_keeps_zero_attendance_and_excludes_out_of_period_rows(tmp_path):
    path = tmp_path / "attendance-and-period.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(REQUIRED_HEADERS)
    sheet.append([{"上课班级": "八年级数学一对一", "教学形式": "一对一", "上课时间": "2026-08-03 09:00", "上课状态": "已上课", "实到": 0, "上课学员": "学生甲", "上课科目": "数学", "任课老师": "张三"}.get(header, "") for header in REQUIRED_HEADERS])
    sheet.append([{"上课班级": "八年级数学一对一", "教学形式": "一对一", "上课时间": "2026-09-01 09:00", "上课状态": "已上课", "实到": 1, "上课学员": "学生乙", "上课科目": "数学", "任课老师": "张三"}.get(header, "") for header in REQUIRED_HEADERS])
    workbook.save(path)

    result = read_schedule_excel(path, period="2026-08")

    assert len(result.records) == 1
    assert result.records[0].attended == 0
    assert any(issue.code == "OUT_OF_PERIOD_ROWS_EXCLUDED" for issue in result.warnings)


def test_schedule_adapter_accepts_explicit_cross_month_accounting_window(tmp_path):
    path = tmp_path / "cross-month.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(REQUIRED_HEADERS)
    for day in ("2026-06-29", "2026-07-15", "2026-08-02", "2026-08-03"):
        sheet.append([{"上课班级": "八年级数学一对一", "教学形式": "一对一", "上课时间": f"{day} 09:00", "上课状态": "已上课", "实到": 1, "上课学员": "学生甲", "上课科目": "数学", "任课老师": "张三"}.get(header, "") for header in REQUIRED_HEADERS])
    workbook.save(path)

    result = read_schedule_excel(path, period="2026-07", period_start="2026-06-29", period_end="2026-08-02")

    assert [record.lesson_date for record in result.records] == ["2026-06-29", "2026-07-15", "2026-08-02"]
    assert any(issue.code == "OUT_OF_PERIOD_ROWS_EXCLUDED" for issue in result.warnings)


def test_schedule_adapter_uses_explicit_student_grade_authority_only_after_direct_extraction(tmp_path):
    path = tmp_path / "gift-course.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(REQUIRED_HEADERS)
    sheet.append([{"上课班级": "赠送课程", "教学形式": "一对一", "上课时间": "2026-08-03 09:00", "上课状态": "已上课", "实到": 1, "上课学员": "学生甲", "上课科目": "数学", "任课老师": "张三"}.get(header, "") for header in REQUIRED_HEADERS])
    workbook.save(path)

    unresolved = read_schedule_excel(path, period="2026-08")
    resolved = read_schedule_excel(path, period="2026-08", student_grades={"学生甲": "初二"})

    assert unresolved.records[0].grade == ""
    assert resolved.records[0].grade == "八年级"


def test_schedule_adapter_derives_special_type_only_from_explicit_course_context(tmp_path):
    path = tmp_path / "special-type.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(REQUIRED_HEADERS + ["上课课程"])
    base = {"上课班级": "高一特殊课程", "教学形式": "一对多", "上课时间": "2026-08-03 09:00", "上课状态": "已上课", "实到": 2, "上课学员": "学生甲", "上课科目": "数学", "任课老师": "张三", "上课课程": "高一 1v2"}
    sheet.append([base.get(header, "") for header in REQUIRED_HEADERS + ["上课课程"]])
    base["上课课程"] = "高一多人课"
    sheet.append([base.get(header, "") for header in REQUIRED_HEADERS + ["上课课程"]])
    workbook.save(path)

    result = read_schedule_excel(path, period="2026-08")

    assert [record.class_type for record in result.records] == ["1对2", "一对多"]


def test_service_prefers_local_student_grade_authority_over_legacy_skill(tmp_path):
    root = tmp_path / "app-data"
    config = root / "config"
    config.mkdir(parents=True)
    (config / "student_grade_lookup.csv").write_text("姓名,年级\n学生甲,高二\n", encoding="utf-8")

    assert PayrollService(root)._student_grade_lookup() == {"学生甲": "高二"}
