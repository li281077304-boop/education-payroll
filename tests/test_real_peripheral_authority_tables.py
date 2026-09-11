from pathlib import Path
from shutil import copyfile

from openpyxl import Workbook, load_workbook

from payroll_core.adapters import read_refund_2026, read_renewal_2026
from payroll_ui.service import PayrollService
from tests.test_payroll_ui_service import _prepared_run


TEMPLATES = Path(__file__).parents[1] / "payroll_ui" / "static" / "templates"


def _renewal_template(tmp_path: Path) -> Path:
    output = tmp_path / "续费推荐数据模板.xlsx"
    copyfile(TEMPLATES / "续费推荐数据模板.xlsx", output)
    book = load_workbook(output)
    sheet = book["8月"]
    sheet["C3"] = "教师甲"
    sheet["D3"] = 2
    sheet["J3"] = 3
    sheet["Z3"] = 4
    book.save(output)
    return output


def _july_renewal(tmp_path: Path) -> Path:
    output = tmp_path / "续费推荐数据_7月.xlsx"
    book = Workbook(); sheet = book.active; sheet.title = "7月"
    sheet.append(["序号", "学科组", "教师", "1V1课时", "", "", "", "", "", "班课", "", "", "", "", "", "", "", "", "", "小班领航伴学课次", "", "", "", "总计"])
    sheet.append(["", "", "", 1, 2, 3, 4, 5, "合计", 1, 2, 3, 4, 5, 6, 7, 8, 9, "合计", 1, 2, 3, "合计", ""])
    sheet.append([1, "示例", "教师甲", 2, 0, 0, 0, 0, "=SUM(D3:H3)", 3, 0, 0, 0, 0, 0, 0, 0, 0, "=SUM(J3:R3)", 4, 0, 0, "=SUM(T3:V3)", "=I3+S3*1.5+W3*.75"])
    book.save(output)
    return output


def _refund_template(tmp_path: Path) -> Path:
    output = tmp_path / "退费绩效确认模板.xlsx"
    copyfile(TEMPLATES / "退费绩效确认模板.xlsx", output)
    return output


def test_formal_templates_are_sanitized_and_present():
    assert (TEMPLATES / "续费推荐数据模板.xlsx").is_file()
    assert (TEMPLATES / "退费绩效确认模板.xlsx").is_file()
    assert "Demo" not in (TEMPLATES / "续费推荐数据模板.xlsx").name
    assert "Demo" not in (TEMPLATES / "退费绩效确认模板.xlsx").name


def test_renewal_template_and_august_semantics_are_importable_without_column_numbers(tmp_path):
    records = read_renewal_2026(_renewal_template(tmp_path), "2026-08")
    first = records[0]
    assert first.teacher_id == "教师甲"
    assert first.payload["payroll_fields"] == {"AH": 2.0, "AI": 3.0, "AJ": 4.0, "AK": 9.5}


def test_july_renewal_layout_is_semantically_equivalent_when_totals_move(tmp_path):
    record = read_renewal_2026(_july_renewal(tmp_path), "2026-07")[0]
    assert record.payload["payroll_fields"] == {"AH": 2.0, "AI": 3.0, "AJ": 4.0, "AK": 9.5}


def test_refund_template_keeps_one_fact_and_expands_three_teacher_groups(tmp_path):
    source = _refund_template(tmp_path)
    book = load_workbook(source); sheet = book["8月退费"]
    sheet["D5"] = "学生甲"
    sheet["P5"] = "教师甲"; sheet["Q5"] = 1; sheet["R5"] = -100
    sheet["S5"] = "教师乙"; sheet["T5"] = 0; sheet["U5"] = -50
    sheet["V5"] = "教师丙"; sheet["W5"] = 0; sheet["X5"] = -30
    book.save(source)
    records = read_refund_2026(source, "2026-08")
    same_fact = [record for record in records if record.payload.get("student") == "学生甲"]
    assert len(same_fact) == 3
    assert {record.payload["refund_group_index"] for record in same_fact} == {1, 2, 3}
    assert len({record.payload["refund_fact_id"] for record in same_fact}) == 1
    assert [record.payload["refund_performance"] for record in same_fact] == [-100, -50, -30]


def test_refund_duplicate_headcount_is_warned_but_not_changed(tmp_path):
    service, _, _ = _prepared_run(tmp_path)
    source = _refund_template(tmp_path)
    book = load_workbook(source); sheet = book["8月退费"]
    sheet["D5"] = "学生甲"; sheet["P5"] = "教师甲"; sheet["Q5"] = 1; sheet["R5"] = -100
    sheet["S5"] = "教师乙"; sheet["T5"] = 1; sheet["U5"] = -50
    book.save(source)
    imported = service.import_business_results("REFUND_RESULT", "2026-08", str(source), "审核员")
    assert len(imported) >= 2
    assert all(item["payload"]["refund_headcount"] == 1 for item in imported[:2])
    assert all(item["evidence"].get("warnings") for item in imported[:2])


def test_refund_facts_never_guess_a_payroll_an_amount(tmp_path):
    service, run, _ = _prepared_run(tmp_path)
    source = _refund_template(tmp_path)
    book = load_workbook(source); sheet = book["8月退费"]
    sheet["D5"] = "学生甲"; sheet["P5"] = "教师甲"; sheet["Q5"] = 1; sheet["R5"] = -100
    book.save(source)
    item = service.import_business_results("REFUND_RESULT", "2026-08", str(source), "审核员")[0]
    assert "payroll_amount" not in item["payload"]
    service.review_business_input(item["id"], "START_REVIEW", "审核员")
    item = service.review_business_input(item["id"], "APPROVE", "审核员")
    service.bind_business_input(run["id"], item["id"])
    result = service.peripheral_payroll(run["id"])
    assert result["teachers"]["教师甲"]["status"] == "SOURCE_READY"
    assert result["teachers"]["教师甲"]["total"] == "0"


def test_renewal_authority_uses_ah_ai_aj_and_not_a_generic_amount(tmp_path):
    service, run, _ = _prepared_run(tmp_path)
    item = service.import_business_results("RENEWAL_RESULT", "2026-08", str(_renewal_template(tmp_path)), "业务员")[0]
    service.review_business_input(item["id"], "START_REVIEW", "审核员")
    item = service.review_business_input(item["id"], "APPROVE", "审核员")
    service.bind_business_input(run["id"], item["id"])
    result = service.peripheral_payroll(run["id"])
    component = result["teachers"]["教师甲"]["components"][0]
    assert component["code"] == "RENEWAL_AWARD"
    assert component["payroll_fields"] == {"AH": "2.0", "AI": "3.0", "AJ": "4.0", "AK": "9.5"}
