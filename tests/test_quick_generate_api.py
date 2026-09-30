"""Quick-mode tests use only private temporary payroll stores and synthetic data."""

from __future__ import annotations

import csv
import hashlib
import json
import shutil
from pathlib import Path
from threading import Thread
from urllib.request import Request, urlopen

import pytest
from openpyxl import load_workbook

import payroll_ui.service as service_module
from payroll_ui.server import PayrollHttpServer
from payroll_ui.service import MONTHLY_FLOW_SUBMISSION_FIRST, PayrollService
from tests.test_standard_payroll_output import _sanitized_template
from tests.test_teacher_identity_and_support_source import _support_workbook, _teacher


def _schedule(path: Path, teachers: list[str], *, day: int = 5) -> Path:
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["teacher_id", "teacher", "grade", "subject", "class_type", "attended", "lesson_status", "time"])
        for index, teacher in enumerate(teachers):
            writer.writerow([f"t-{index}", teacher, "九年级", "数学", "1对1", 1, "已上课", f"2026-08-{day:02d} 10:00"])
    return path


def _period_schedule(path: Path, period: str, teacher: str) -> Path:
    year, month = map(int, period.split("-"))
    from calendar import monthrange
    days = monthrange(year, month)[1]
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["teacher_id", "teacher", "grade", "subject", "class_type", "attended", "lesson_status", "time"])
        for day in range(1, days + 1):
            writer.writerow(["demo-teacher", teacher, "九年级", "数学", "1对1", 1, "已上课", f"{year}-{month:02d}-{day:02d} 09:00"])
    return path


def _group(path: Path, teachers: list[str], *, base_salary: float = 1800) -> Path:
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=["teacher", "b", "d", "e", "f", "g", "h", "i", "j", "k", "l"])
        writer.writeheader()
        for teacher in teachers:
            writer.writerow({"teacher": teacher, "b": "数学组", "d": f"{teacher}@example.com", "e": "2022-01-01",
                             "f": "TR/三星级", "g": base_salary, "h": 100, "i": 20, "j": 30, "k": 26, "l": 26})
    return path


def _service(tmp_path: Path, monkeypatch: pytest.MonkeyPatch, teachers: list[str], *, rating: bool = True) -> PayrollService:
    service = PayrollService(tmp_path / "private-payroll-data")
    service.record_period_authority("2026-08", "2026-08-05", "2026-08-05", "USER_CONFIRMED",
                                   confirmed_by="脱敏 UAT", reason="合成测试周期。")
    if rating:
        service.save_rating_version("2026-08", "2026-08", "合成星级 authority", "quick-test-v1",
                                    [{"teacher": teacher, "rating": 3} for teacher in teachers])
    service.register_company_template(str(_sanitized_template(tmp_path)), "脱敏 UAT")
    monkeypatch.setattr(service_module, "default_output_dir",
                        lambda filename=None: (tmp_path / "exports" / str(filename)) if filename else (tmp_path / "exports"))
    return service


def _http_post(base: str, token: str, path: str, payload: dict) -> dict:
    request = Request(base + path, data=json.dumps(payload).encode(), method="POST",
                      headers={"Content-Type": "application/json", "X-Payroll-Token": token})
    return json.loads(urlopen(request).read())


def test_quick_http_creates_regular_run_imports_materials_generates_and_downloads_same_run(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    service = _service(tmp_path, monkeypatch, ["教师甲"])
    server = PayrollHttpServer(("127.0.0.1", 0), service, Path(__file__).parents[1] / "payroll_ui" / "static")
    worker = Thread(target=server.serve_forever, daemon=True)
    worker.start()
    base = f"http://127.0.0.1:{server.server_port}"
    try:
        token = json.loads(urlopen(base + "/api/bootstrap").read())["token"]
        run = _http_post(base, token, "/api/quick-runs", {"period": "2026-08"})["run"]
        assert run["mode"] == "GENERATE"
        assert run["monthly_flow_version"] == MONTHLY_FLOW_SUBMISSION_FIRST
        assert run["ui_mode"] == "QUICK"
        assert run["rating_version_id"] == service.store.list_rating_versions()[0]["id"]
        schedule = _schedule(tmp_path / "schedule.csv", ["教师甲"])
        group = _group(tmp_path / "group.csv", ["教师甲"])
        imported = _http_post(base, token, f"/api/quick-runs/{run['id']}/materials", {"paths": [str(group), str(schedule)]})
        assert imported["run"]["id"] == run["id"]
        assert imported["ok"] is True
        assert [item["detected_kind"] for item in imported["materials"]] == ["subject_group", "schedule"]
        result = _http_post(base, token, f"/api/quick-runs/{run['id']}/generate", {})
        assert result["ok"] is True
        assert result["teacher_count"] == 1
        assert result["download_url"] == f"/api/quick-runs/{run['id']}/download"
        request = Request(base + result["download_url"], headers={"X-Payroll-Token": token})
        assert urlopen(request).read().startswith(b"PK")
        pro = _http_post(base, token, f"/api/quick-runs/{run['id']}/professional", {})
        assert pro["run"]["id"] == run["id"]
        assert pro["run"]["ui_mode"] == "PROFESSIONAL"
        assert pro["run"]["monthly_flow_version"] == MONTHLY_FLOW_SUBMISSION_FIRST
    finally:
        server.shutdown(); server.server_close(); worker.join()


def test_quick_http_does_not_route_missing_upload_copy_to_professional_mapping(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    service = _service(tmp_path, monkeypatch, ["教师甲"])
    server = PayrollHttpServer(("127.0.0.1", 0), service, Path(__file__).parents[1] / "payroll_ui" / "static")
    worker = Thread(target=server.serve_forever, daemon=True)
    worker.start()
    base = f"http://127.0.0.1:{server.server_port}"
    try:
        token = json.loads(urlopen(base + "/api/bootstrap").read())["token"]
        run = _http_post(base, token, "/api/quick-runs", {"period": "2026-08"})["run"]
        result = _http_post(base, token, f"/api/quick-runs/{run['id']}/materials",
                            {"paths": [{"path": str(tmp_path / "missing-server-copy.xlsx"), "sha256": "0" * 64}]})
        assert result["ok"] is False
        assert result["materials"][0]["error_kind"] == "MATERIAL_READ_FAILED"
        assert result["materials"][0]["error_kind"] != "PROFESSIONAL_REQUIRED"
        assert "找不到这份材料" in result["materials"][0]["message"]
    finally:
        server.shutdown(); server.server_close(); worker.join()


def test_quick_september_august_schedule_is_not_reclassified_as_subject_group(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    service = _service(tmp_path, monkeypatch, ["教师甲"])
    run = service.create_quick_run("2026-09")["run"]
    source = _schedule(tmp_path / "august-only.csv", ["教师甲"])

    result = service.stage_quick_material(run["id"], str(source), display_name="八月排课表.csv")
    stored = service.get(run["id"])

    assert result["ok"] is False
    assert result["error_kind"] == "SCHEDULE_PERIOD_MISMATCH"
    assert result["detected_kind"] == "schedule-period-mismatch"
    assert "2026-09" in result["message"] and "2026-08" in result["message"]
    assert stored["files"].get("schedule") is None
    assert not stored.get("subject_group_materials")
    assert stored["quick_materials"][0]["kind"] == "SCHEDULE_PERIOD_MISMATCH"
    assert stored["quick_materials"][0]["path"] == str(source.resolve())
    assert service.quick_generate(run["id"])["error_kind"] == "MISSING_SCHEDULE"


def test_quick_august_run_accepts_same_schedule_csv(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    service = _service(tmp_path, monkeypatch, ["教师甲"])
    run = service.create_quick_run("2026-08")["run"]
    source = _schedule(tmp_path / "august-only.csv", ["教师甲"])

    result = service.stage_quick_material(run["id"], str(source))

    assert result["detected_kind"] == "schedule"
    assert service.get(run["id"])["quick_materials"][0]["kind"] == "schedule"


def test_quick_real_group_csv_remains_subject_group(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    service = _service(tmp_path, monkeypatch, ["教师甲"])
    run = service.create_quick_run("2026-08")["run"]
    source = _group(tmp_path / "group.csv", ["教师甲"])

    result = service.stage_quick_material(run["id"], str(source))

    assert result["detected_kind"] == "subject_group"
    assert service.get(run["id"])["quick_materials"][0]["kind"] == "subject_group"


def test_professional_auto_import_does_not_fallback_from_out_of_month_schedule_csv(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    service = _service(tmp_path, monkeypatch, ["教师甲"])
    quick = service.create_quick_run("2026-09")["run"]
    run = service.quick_to_professional(quick["id"])["run"]
    source = _schedule(tmp_path / "august-only.csv", ["教师甲"])

    with pytest.raises(ValueError, match="这是一份排课表，但没有找到 2026-09"):
        service.import_material_file(run["id"], "auto", str(source))
    stored = service.get(run["id"])
    assert stored["files"].get("schedule") is None
    assert not stored.get("subject_group_materials")


def test_quick_known_schedule_layout_outranks_incidental_support_title(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    service = _service(tmp_path, monkeypatch, ["张三", "李四"])
    run = service.create_quick_run("2026-08")["run"]
    source = tmp_path / "schedule-with-incidental-title.xlsx"
    shutil.copyfile(Path(__file__).parent / "fixtures" / "excel" / "fake_schedule.xlsx", source)
    workbook = load_workbook(source)
    workbook.active.insert_rows(1)
    workbook.active.cell(1, 1, "支持部提供：这是后续导出备注，不代表文件类型")
    workbook.save(source)

    result = service.stage_quick_material(run["id"], str(source))

    assert result["detected_kind"] == "schedule"
    assert result.get("professional_kind") != "support"
    assert service.get(run["id"])["quick_materials"][0]["kind"] == "schedule"


def test_quick_http_preserves_original_uploaded_name_and_hash(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    service = _service(tmp_path, monkeypatch, ["教师甲"])
    server = PayrollHttpServer(("127.0.0.1", 0), service, Path(__file__).parents[1] / "payroll_ui" / "static")
    worker = Thread(target=server.serve_forever, daemon=True)
    worker.start()
    base = f"http://127.0.0.1:{server.server_port}"
    source = _schedule(tmp_path / "opaque-upload-id.csv", ["教师甲"])
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    try:
        token = json.loads(urlopen(base + "/api/bootstrap").read())["token"]
        run = _http_post(base, token, "/api/quick-runs", {"period": "2026-08"})["run"]
        result = _http_post(base, token, f"/api/quick-runs/{run['id']}/materials", {
            "paths": [{"path": str(source), "sha256": digest, "name": "八月排课表.csv"}],
        })
        assert result["ok"] is True
        quick_material = result["run"]["quick_materials"][0]
        assert quick_material["name"] == "八月排课表.csv"
        assert quick_material["sha256"] == digest
        assert result["materials"][0]["material"]["name"] == "八月排课表.csv"
    finally:
        server.shutdown(); server.server_close(); worker.join()


def test_legacy_group_files_route_rejects_mismatched_upload_hash(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    from urllib.error import HTTPError

    service = _service(tmp_path, monkeypatch, ["教师甲"])
    quick = service.create_quick_run("2026-08")["run"]
    run = service.quick_to_professional(quick["id"])["run"]
    group = _group(tmp_path / "opaque-group-copy.csv", ["教师甲"])
    server = PayrollHttpServer(("127.0.0.1", 0), service, Path(__file__).parents[1] / "payroll_ui" / "static")
    worker = Thread(target=server.serve_forever, daemon=True)
    worker.start()
    base = f"http://127.0.0.1:{server.server_port}"
    try:
        token = json.loads(urlopen(base + "/api/bootstrap").read())["token"]
        try:
            _http_post(base, token, f"/api/runs/{run['id']}/files", {
                "role": "math", "path": str(group), "sha256": "0" * 64,
                "name": "original-group.csv",
            })
        except HTTPError as exc:
            assert exc.code == 400
            assert "与本次选择不一致" in exc.read().decode("utf-8")
        else:
            pytest.fail("/files math compatibility route accepted an upload copy with the wrong hash")
    finally:
        server.shutdown(); server.server_close(); worker.join()


def test_quick_historical_period_uses_its_period_and_rating_authorities(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    service = PayrollService(tmp_path / "private-payroll-data")
    service.record_period_authority("2026-08", "2026-08-03", "2026-08-30", "USER_CONFIRMED",
                                   confirmed_by="脱敏 UAT", reason="8 月合成周期。")
    service.record_period_authority("2026-09", "2026-09-01", "2026-09-30", "USER_CONFIRMED",
                                   confirmed_by="脱敏 UAT", reason="9 月合成周期。")
    aug_rating = service.save_rating_version("2026-08", "2026-08", "8 月合成星级", "aug-v1",
                                             [{"teacher": "教师甲", "rating": 3}])[0]
    service.save_rating_version("2026-09", "2026-09", "9 月合成星级", "sep-v1",
                                [{"teacher": "教师甲", "rating": 4}])
    service.register_company_template(str(_sanitized_template(tmp_path)), "脱敏 UAT")
    monkeypatch.setattr(service_module, "default_output_dir",
                        lambda filename=None: (tmp_path / "exports" / str(filename)) if filename else (tmp_path / "exports"))

    # 模拟当前日期已进入 9 月，但用户从极速首页明确选择补算 8 月。
    run = service.create_quick_run("2026-08")["run"]
    assert run["period"] == "2026-08"
    assert run["period_start"] == "2026-08-03"
    assert run["period_end"] == "2026-08-30"
    assert run["rating_version_id"] == aug_rating["id"]

    staged = service.stage_quick_material(run["id"], str(_schedule(tmp_path / "august.csv", ["教师甲"])))
    assert staged["detected_kind"] == "schedule"
    after_upload = service.get(run["id"])
    assert after_upload["period"] == "2026-08"
    assert after_upload["rating_version_id"] == aug_rating["id"]
    assert after_upload["period_authority"]["period_start"] == "2026-08-03"


def test_new_monthly_flow_run_reuses_long_term_default_af_policy_without_exceptions(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    service = _service(tmp_path, monkeypatch, ["教师甲"])
    first = service.create_quick_run("2026-08")["run"]
    service.stage_quick_material(first["id"], str(_schedule(tmp_path / "first-month.csv", ["教师甲"])))
    service.quick_to_professional(first["id"])
    service.confirm_af_policy(first["id"], "合成管理员", default_obligation_hours=30,
                              exceptions={"教师甲": {"obligation_hours": 0, "reason": "仅首 Run 例外"}})

    next_run = service.create_quick_run("2026-08")["run"]
    policy = next_run["af_policy_confirmation"]
    assert policy["confirmed"] is True
    assert policy["default_obligation_hours"] == 30
    assert policy["exceptions"] == {}


def test_two_month_product_chain_reuses_authorities_and_only_needs_next_schedule(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    service = PayrollService(tmp_path / "private-payroll-data")
    service.record_period_authority("2026-08", "2026-08-01", "2026-08-31", "USER_CONFIRMED",
                                   confirmed_by="合成 UAT", reason="8 月测试周期。")
    service.record_period_authority("2026-09", "2026-09-01", "2026-09-30", "USER_CONFIRMED",
                                   confirmed_by="合成 UAT", reason="9 月测试周期。")
    service.save_rating_version("2026-08", "2026-08", "合成星级 8 月", "aug-v1", [{"teacher": "教师甲", "rating": 3}])
    september_rating = service.save_rating_version("2026-09", "2026-09", "合成星级 9 月", "sep-v1", [{"teacher": "教师甲", "rating": 3}])
    september_rating_id = next(item["id"] for item in september_rating if item["effective_from"] == "2026-09")
    service.register_company_template(str(_sanitized_template(tmp_path)), "合成 UAT")
    monkeypatch.setattr(service_module, "default_output_dir",
                        lambda filename=None: (tmp_path / "exports" / str(filename)) if filename else (tmp_path / "exports"))

    august = service.create_quick_run("2026-08")["run"]
    service.stage_quick_material(august["id"], str(_period_schedule(tmp_path / "august.csv", "2026-08", "教师甲")))
    aug_result = service.quick_generate(august["id"])
    assert aug_result["ok"] is True

    september = service.create_quick_run("2026-09")["run"]
    assert september["period_start"] == "2026-09-01" and september["period_end"] == "2026-09-30"
    assert september["period_authority"]["is_fallback"] is False
    assert september["rating_version_id"] == september_rating_id
    assert september["af_policy_confirmation"]["confirmed"] is True
    assert september["af_policy_confirmation"]["default_obligation_hours"] == 30
    assert september["af_policy_confirmation"]["exceptions"] == {}
    assert september["template"]["source"] == "已登记公司工资模板"

    only_schedule = _period_schedule(tmp_path / "september.csv", "2026-09", "教师甲")
    service.stage_quick_material(september["id"], str(only_schedule))
    output = service.quick_generate(september["id"])
    refreshed = service.get(september["id"])
    records = service._read_for_run("schedule", only_schedule, service.store.get(september["id"])).records

    assert output["ok"] is True
    assert len(records) == 30
    assert all(record.lesson_date.startswith("2026-09-") for record in records)
    assert refreshed["rating_version_id"] == september["rating_version_id"]
    assert refreshed["af_policy_confirmation"]["default_obligation_hours"] == 30
    assert refreshed["files"]["schedule"]["name"] == only_schedule.name
    assert refreshed["template"]["source"] == "已登记公司工资模板"
    assert not refreshed.get("subject_group_materials")
    assert Path(output["path"]).is_file()


def test_professional_second_month_schedule_only_reuses_long_term_base_salary(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    service = PayrollService(tmp_path / "private-payroll-data")
    service.record_period_authority("2026-08", "2026-08-01", "2026-08-31", "USER_CONFIRMED",
                                   confirmed_by="合成 UAT", reason="8 月测试周期。")
    service.record_period_authority("2026-09", "2026-09-01", "2026-09-30", "USER_CONFIRMED",
                                   confirmed_by="合成 UAT", reason="9 月测试周期。")
    service.save_rating_version("2026-08", "2026-08", "合成星级 8 月", "aug-v1", [{"teacher": "教师甲", "rating": 3}])
    service.save_rating_version("2026-09", "2026-09", "合成星级 9 月", "sep-v1", [{"teacher": "教师甲", "rating": 3}])
    service.register_company_template(str(_sanitized_template(tmp_path)), "合成 UAT")
    monkeypatch.setattr(service_module, "default_output_dir",
                        lambda filename=None: (tmp_path / "exports" / str(filename)) if filename else (tmp_path / "exports"))
    salary_fields = {code: {"value": value, "source": "合成长期基本工资资料"} for code, value in {
        "G": 6000, "H": 500, "I": 200, "J": 300, "K": 26, "L": 26,
    }.items()}

    august = service.create("2026-08", mode="GENERATE", operator_role="DOS",
                            monthly_flow_version=MONTHLY_FLOW_SUBMISSION_FIRST)
    service.import_file(august["id"], "schedule", str(_period_schedule(tmp_path / "august-base.csv", "2026-08", "教师甲")))
    service.save_base_salary_inputs(august["id"], [{"teacher_id": "demo-teacher", "teacher": "教师甲", "fields": salary_fields}], "合成 UAT", source="合成长期基本工资资料")
    service.confirm_af_policy(august["id"], "合成 UAT", default_obligation_hours=30, exceptions={})
    service.check(august["id"])
    first_export = service.generate_payroll(august["id"], str(tmp_path / "exports" / "august.xlsx"), production=True)
    assert Path(first_export["path"]).is_file()

    september = service.create("2026-09", mode="GENERATE", operator_role="DOS",
                               monthly_flow_version=MONTHLY_FLOW_SUBMISSION_FIRST)
    assert september["base_salary_input_snapshot"]
    assert september["base_salary_inputs"]["教师甲"]["fields"]["G"]["value"] == 6000
    assert september["af_policy_confirmation"]["confirmed"] is True
    assert september["af_policy_confirmation"]["default_obligation_hours"] == 30
    assert september["af_policy_confirmation"]["exceptions"] == {}
    assert september["template"]["source"] == "已登记公司工资模板"
    service.import_file(september["id"], "schedule", str(_period_schedule(tmp_path / "september-base.csv", "2026-09", "教师甲")))

    checked = service.check(september["id"])
    preview = service.preview_payroll(september["id"])
    exported = service.generate_payroll(september["id"], str(tmp_path / "exports" / "september.xlsx"), production=True)
    refreshed = service.get(september["id"])
    assert checked["generated_payroll"]["rows"][0]["final_fields"]["M"]["state"] == "DETERMINED"
    assert preview["generated_payroll"]["rows"][0]["final_fields"]["M"]["value"] == 7000
    assert refreshed["base_salary_missing_teachers"] == []
    assert Path(exported["path"]).is_file()


def test_quick_schedule_only_generates_full_schedule_roster_with_blank_b_to_l(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    service = _service(tmp_path, monkeypatch, ["教师甲", "教师乙"])
    run = service.create_quick_run("2026-08")["run"]
    source = _schedule(tmp_path / "schedule.csv", ["教师甲", "教师乙"])
    service.stage_quick_material(run["id"], str(source))
    result = service.quick_generate(run["id"])
    assert result["ok"] is True
    sheet = load_workbook(result["path"]).active
    assert {sheet[f"C{row}"].value for row in (5, 6)} == {"教师甲", "教师乙"}
    assert all(sheet.cell(row, column).value is None for row in (5, 6) for column in range(2, 13) if column != 3)
    assert service.get(run["id"])["af_policy_confirmation"]["default_obligation_hours"] == 30
    assert service.get(run["id"])["af_policy_confirmation"]["confirmed_by"] == "QUICK_GENERATE"


def test_monthly_flow_exposes_missing_base_salary_teachers_for_inline_actions(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    service = _service(tmp_path, monkeypatch, ["教师甲", "教师乙"])
    run = service.create_quick_run("2026-08")["run"]
    service.stage_quick_material(run["id"], str(_schedule(tmp_path / "schedule.csv", ["教师甲", "教师乙"])))
    service.quick_generate(run["id"])
    rendered = service.get(run["id"])
    assert set(rendered["base_salary_missing_teachers"]) == {"教师甲", "教师乙"}

    service = _service(tmp_path / "with-group", monkeypatch, ["教师甲", "教师乙"])
    with_group = service.create_quick_run("2026-08")["run"]
    service.stage_quick_material(with_group["id"], str(_schedule(tmp_path / "with-group-schedule.csv", ["教师甲", "教师乙"])))
    service.stage_quick_material(with_group["id"], str(_group(tmp_path / "with-group.csv", ["教师甲", "教师乙"])))
    service.quick_generate(with_group["id"])
    assert service.get(with_group["id"])["base_salary_missing_teachers"] == []


def test_quick_group_scope_is_submission_union_and_conflicts_stop_export(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    service = _service(tmp_path, monkeypatch, ["教师甲", "教师乙"])
    run = service.create_quick_run("2026-08")["run"]
    sources = [
        _schedule(tmp_path / "schedule.csv", ["教师甲", "教师乙"]),
        _group(tmp_path / "math.csv", ["教师甲"], base_salary=1800),
        _group(tmp_path / "science.csv", ["教师乙"], base_salary=1800),
        _group(tmp_path / "conflict.csv", ["教师甲"], base_salary=1900),
    ]
    # Stage groups first, in arbitrary order; imports are sorted on generation.
    for source in (sources[2], sources[1], sources[3], sources[0]):
        service.stage_quick_material(run["id"], str(source))
    result = service.quick_generate(run["id"])
    assert result["ok"] is False
    assert result["error_kind"] == "NEEDS_PROFESSIONAL_REVIEW"
    assert any("教师甲" in issue for issue in result["issues"])
    run_after = service.quick_to_professional(run["id"])["run"]
    assert run_after["id"] == run["id"]
    assert run_after["ui_mode"] == "PROFESSIONAL"
    assert {item["source_name"] for item in run_after["subject_group_materials"]} == {"math.csv", "science.csv", "conflict.csv"}


def test_quick_group_union_deduplicates_same_teacher_with_identical_values(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    teachers = ["教师甲", "教师乙", "教师丙"]
    service = _service(tmp_path, monkeypatch, teachers)
    run = service.create_quick_run("2026-08")["run"]
    sources = [
        _schedule(tmp_path / "schedule.csv", teachers),
        _group(tmp_path / "group-a.csv", ["教师甲", "教师乙"]),
        _group(tmp_path / "group-b.csv", ["教师乙", "教师丙"]),
    ]
    for path in sources:
        service.stage_quick_material(run["id"], str(path))
    result = service.quick_generate(run["id"])
    assert result["ok"] is True
    assert result["teacher_count"] == 3
    assert {row["teacher"] for row in service.get(run["id"])["generated_payroll"]["rows"]} == set(teachers)


def test_quick_uses_schedule_roster_when_group_submission_covers_only_fifteen_of_fifty_one(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    schedule_teachers = [f"教师{index:02d}" for index in range(51)]
    service = _service(tmp_path, monkeypatch, schedule_teachers)
    run = service.create_quick_run("2026-08")["run"]
    service.stage_quick_material(run["id"], str(_schedule(tmp_path / "schedule.csv", schedule_teachers)))
    service.stage_quick_material(run["id"], str(_group(tmp_path / "group.csv", schedule_teachers[:15])))
    result = service.quick_generate(run["id"])
    assert result["ok"] is True
    assert result["teacher_count"] == 51
    rows = service.get(run["id"])["generated_payroll"]["rows"]
    assert {row["teacher"] for row in rows} == set(schedule_teachers)
    assert all(row["final_fields"]["M"]["state"] == "BLOCKED_BY_INPUT" for row in rows[15:])


def test_quick_support_material_is_routed_to_professional_and_kept_in_same_run(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    service = _service(tmp_path, monkeypatch, ["教师甲"])
    run = service.create_quick_run("2026-08")["run"]
    schedule = _schedule(tmp_path / "schedule.csv", ["教师甲"])
    support = _support_workbook(tmp_path / "support.xlsx", [_teacher("教师甲", 房租=500, 社保=200)])
    service.stage_quick_material(run["id"], str(schedule))
    classified = service.stage_quick_material(run["id"], str(support))
    assert classified["ok"] is False
    assert classified["professional_kind"] == "support"
    pro = service.quick_to_professional(run["id"])["run"]
    assert pro["id"] == run["id"]
    assert pro["files"].get("schedule")
    staged_support = next(item for item in pro["quick_materials"] if item.get("professional_kind") == "support")
    assert pro["support_department_pending_preview"]["source_sha256"] == staged_support["sha256"]


def test_quick_renewal_and_refund_inputs_are_kept_for_professional_mode(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    service = _service(tmp_path, monkeypatch, ["教师甲"])
    run = service.create_quick_run("2026-08")["run"]
    schedule = _schedule(tmp_path / "schedule.csv", ["教师甲"])
    renewal = tmp_path / "renewal.csv"
    renewal.write_text("教师,续费人数,总学生数\n教师甲,2,4\n", encoding="utf-8-sig")
    refund = tmp_path / "refund.csv"
    refund.write_text("教师,学生,金额,状态,说明\n教师甲,学生甲,-100,已确认,合成测试\n", encoding="utf-8-sig")
    service.stage_quick_material(run["id"], str(schedule))
    assert service.stage_quick_material(run["id"], str(renewal))["professional_kind"] == "renewal"
    assert service.stage_quick_material(run["id"], str(refund))["professional_kind"] == "refund"

    pro = service.quick_to_professional(run["id"])["run"]
    assert pro["id"] == run["id"]
    assert pro["material_inputs"]["renewal"]["name"] == renewal.name
    assert pro["material_inputs"]["refund"]["name"] == refund.name
    assert not any(item.get("transfer_reason") for item in pro["quick_materials"])


def test_quick_unclassified_upload_is_preserved_for_same_run_professional_mode(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    service = _service(tmp_path, monkeypatch, ["教师甲"])
    run = service.create_quick_run("2026-08")["run"]
    unknown = tmp_path / "unknown.csv"
    unknown.write_text("标题,说明\n合成表,专业流程使用\n", encoding="utf-8-sig")
    classified = service.stage_quick_material(run["id"], str(unknown))
    assert classified["ok"] is False
    assert classified["error_kind"] == "PROFESSIONAL_REQUIRED"
    transferred = service.quick_to_professional(run["id"])["run"]
    saved = next(item for item in transferred["quick_materials"] if item["name"] == unknown.name)
    assert transferred["id"] == run["id"]
    assert transferred["ui_mode"] == "PROFESSIONAL"
    assert saved["path"] == str(unknown.resolve())
    assert saved["transferred_to_professional"] is False


def test_quick_out_of_authority_date_does_not_guess_or_export(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    service = _service(tmp_path, monkeypatch, ["教师甲"])
    run = service.create_quick_run("2026-08")["run"]
    service.stage_quick_material(run["id"], str(_schedule(tmp_path / "outside.csv", ["教师甲"], day=6)))
    result = service.quick_generate(run["id"])
    assert result["ok"] is False
    assert not (service.get(run["id"]).get("generated_payroll") or {}).get("path")


def test_quick_orchestrator_contains_no_independent_payroll_formula_engine() -> None:
    source = Path(__file__).parents[1] / "payroll_ui" / "service.py"
    code = source.read_text(encoding="utf-8")
    block = code.split("def quick_generate(", 1)[1].split("def quick_to_professional(", 1)[0]
    for formula in ("AA =", "AC =", "AD =", "AE =", "AF =", "AK =", "AV =", "* 1.5", "* 0.75", "calculate_payroll("):
        assert formula not in block
    assert "self.check(run_id)" in block
    assert "self.generate_payroll(run_id" in block


def test_quick_group_submission_teacher_outside_schedule_roster_does_not_add_payroll_row(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    service = _service(tmp_path, monkeypatch, ["排课教师"])
    run = service.create_quick_run("2026-08")["run"]
    schedule = _schedule(tmp_path / "schedule.csv", ["排课教师"])
    group = _group(tmp_path / "group.csv", ["排课教师", "组表无排课教师"])
    service.stage_quick_material(run["id"], str(schedule))
    service.stage_quick_material(run["id"], str(group))
    result = service.quick_generate(run["id"])
    assert result["ok"] is True
    current = service.get(run["id"])
    assert {row["teacher"] for row in current["core_calculation"]["rows"]} == {"排课教师"}
    assert {row["teacher"] for row in current["generated_payroll"]["rows"]} == {"排课教师"}
    assert current["id"] == run["id"]


def test_quick_no_star_authority_is_admin_error_and_missing_schedule_is_not_exported(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    service = _service(tmp_path, monkeypatch, ["教师甲"], rating=False)
    run = service.create_quick_run("2026-08")["run"]
    source = _schedule(tmp_path / "schedule.csv", ["教师甲"])
    service.stage_quick_material(run["id"], str(source))
    result = service.quick_generate(run["id"])
    assert result["ok"] is False
    assert result["admin_configuration"] is True
    assert not (service.get(run["id"]).get("generated_payroll") or {}).get("path")

    missing = service.create_quick_run("2026-08")["run"]
    no_schedule_result = service.quick_generate(missing["id"])
    assert no_schedule_result["ok"] is False
    assert no_schedule_result["error_kind"] == "MISSING_SCHEDULE"


def test_quick_overlapping_active_star_authorities_are_admin_configuration_error(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    service = _service(tmp_path, monkeypatch, ["教师甲"])
    run = service.create_quick_run("2026-08")["run"]
    service.stage_quick_material(run["id"], str(_schedule(tmp_path / "schedule.csv", ["教师甲"])))
    list_versions = service.store.list_rating_versions
    original = list_versions()
    assert len(original) == 1
    monkeypatch.setattr(service.store, "list_rating_versions", lambda: [*original, {**original[0], "id": "overlap-test"}])
    result = service.quick_generate(run["id"])
    assert result["ok"] is False
    assert result["error_kind"] == "ADMIN_CONFIGURATION_ERROR"
    assert not (service.get(run["id"]).get("generated_payroll") or {}).get("path")


def test_quick_vs_professional_same_inputs_have_equal_core_and_final_field_values(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    service = _service(tmp_path, monkeypatch, ["教师甲"])
    schedule = _schedule(tmp_path / "schedule.csv", ["教师甲"])
    group = _group(tmp_path / "group.csv", ["教师甲"])
    quick = service.create_quick_run("2026-08")["run"]
    service.stage_quick_material(quick["id"], str(group))
    service.stage_quick_material(quick["id"], str(schedule))
    quick_result = service.quick_generate(quick["id"])
    assert quick_result["ok"] is True
    quick_rows = service.get(quick["id"])["generated_payroll"]["rows"]

    professional = service.create("2026-08", "GENERATE", operator_role="DOS",
                                  monthly_flow_version=MONTHLY_FLOW_SUBMISSION_FIRST)
    service.import_file(professional["id"], "schedule", str(schedule))
    service.import_material_file(professional["id"], "subject_group", str(group))
    service.confirm_af_policy(professional["id"], "QUICK_GENERATE", default_obligation_hours=30, exceptions={})
    service.generate_payroll(professional["id"], str(tmp_path / "professional.xlsx"), production=True)
    professional_rows = service.get(professional["id"])["generated_payroll"]["rows"]

    def values(rows):
        return {row["teacher"]: {
            code: ((row.get("final_fields") or {}).get(code) or {}).get("value")
            if code not in {"AA", "AC", "AD", "AE", "AF"}
            else ((row.get("fields") or {}).get(code) or {}).get("value")
            for code in ("AA", "AC", "AD", "AE", "AF", "AH", "AI", "AJ", "AK", "AV")
        } for row in rows}
    assert values(quick_rows) == values(professional_rows)


def test_quick_generate_has_no_second_formula_engine() -> None:
    source = Path(__file__).parents[1] / "payroll_ui" / "service.py"
    code = source.read_text(encoding="utf-8")
    quick = code.split("def quick_generate(", 1)[1].split("def quick_to_professional(", 1)[0]
    for expression in ("calculate_payroll(", "base_salary_field(", "AA =", "AC =", "AD =", "AE =", "AF =", "AK =", "AV =", "* 1.5", "* 0.75"):
        assert expression not in quick
    assert "self.check(run_id)" in quick
    assert "self.generate_payroll(run_id" in quick
