"""Synthetic end-to-end coverage for the September 26 manual Payroll UAT."""
from __future__ import annotations

import hashlib

import pytest
from openpyxl import Workbook

from payroll_core.payroll_generation import generated_from_calculation
from payroll_ui.service import PayrollService


def _run_with_roster(service: PayrollService, period: str = "2026-08") -> dict:
    created = service.create(period, "GENERATE")
    stored = service.store.get(created["id"])
    stored["core_calculation"] = {"rows": [{"teacher": "教师甲", "teacher_id": "teacher-1"}]}
    service.store.save(stored)
    return stored


def _history_salary(path, value: int = 20):
    book = Workbook()
    sheet = book.active
    sheet.title = "历史工资"
    sheet.append(["历史薪资"])
    sheet.append(["教师ID", "基本工资", "岗位津贴", "工龄工资", "其他待遇", "应出勤", "实际出勤"])
    sheet.append(["teacher-1", value, 2, 3, 4, 22, 21])
    book.save(path)
    return path


def _monthly_renewal(path):
    book = Workbook()
    january = book.active
    january.title = "1月"
    august = book.create_sheet("8月")
    for sheet, teacher, amount in ((january, "教师甲", 99), (august, "教师甲", 5)):
        header = [None] * 30
        header[0:3] = ["序号", "学科组", "教师"]
        header[3], header[9], header[25] = "1V1课时", "班课", "小班领航伴学课次"
        sheet.append(header)
        detail = [None] * 30
        detail[8] = detail[24] = detail[28] = "合计"
        sheet.append(detail)
        values = [None] * 30
        values[0:3] = [1, "数学组", teacher]
        values[8], values[24], values[28] = amount, 2, 1
        sheet.append(values)
    book.save(path)
    return path


def test_history_salary_preview_then_confirm_uses_existing_atomic_run_path(tmp_path):
    service = PayrollService(tmp_path / "data")
    run = _run_with_roster(service)
    source = _history_salary(tmp_path / "薪资表.xlsx")
    original_hash = hashlib.sha256(source.read_bytes()).hexdigest()

    preview = service.preview_base_salary_import(run["id"], str(source))
    assert preview["can_import"] is True
    assert len(preview["matched"]) == 1
    assert service.store.get(run["id"]).get("base_salary_input_snapshot") is None

    saved = service.import_base_salary_from_history(run["id"], str(source), preview["source"]["sha256"], "核算负责人")
    entry = saved["run"]["base_salary_inputs"]["教师甲"]
    assert saved["imported"] == 1
    assert entry["fields"]["G"]["value"] == 20
    assert entry["fields"]["G"]["provenance"]["source_sha256"] == original_hash
    assert entry["m"]["state"] == "DETERMINED"
    assert hashlib.sha256(source.read_bytes()).hexdigest() == original_hash
    assert service.create("2026-09", "GENERATE")["base_salary_inputs"]["教师甲"]["fields"]["G"]["value"] == 20


def test_history_salary_confirmation_rejects_source_changed_after_preview(tmp_path):
    service = PayrollService(tmp_path / "data")
    run = _run_with_roster(service)
    source = _history_salary(tmp_path / "薪资表.xlsx")
    preview = service.preview_base_salary_import(run["id"], str(source))
    _history_salary(source, 21)

    with pytest.raises(ValueError, match="已变化"):
        service.import_base_salary_from_history(run["id"], str(source), preview["source"]["sha256"], "核算负责人")
    assert service.store.get(run["id"]).get("base_salary_input_snapshot") is None


def test_monthly_renewal_import_is_visible_but_wage_fields_require_confirmation(tmp_path):
    service = PayrollService(tmp_path / "data")
    run = _run_with_roster(service)
    source = _monthly_renewal(tmp_path / "年度续费.xlsx")

    imported = service.import_material_file(run["id"], "renewal", str(source))["run"]
    assert imported["renewal_material_kind"] == "MONTHLY_FINAL_CANDIDATE"
    assert len(imported["renewal_reports"]) == 1
    assert imported["renewal_reports"][0]["sheet"] == "8月"
    assert imported.get("run_renewal_result_snapshot") is None

    preview = service.preview_renewal_material(run["id"])
    assert preview["can_confirm"] is True
    assert preview["matched"][0]["AH"] == 5
    assert preview["matched"][0]["AI"] == 2
    assert preview["matched"][0]["AJ"] == 1

    confirmed = service.confirm_renewal_material(run["id"], preview["source_sha256"], "核算负责人")
    entry = confirmed["run_renewal_result_snapshot"]["entries"]["teacher-1"]
    assert entry["AH"]["value"] == 5
    assert entry["AI"]["value"] == 2
    assert entry["AJ"]["value"] == 1
    assert confirmed["renewal_source_confirmation"]["actor"] == "核算负责人"
    assert confirmed["material_inputs"]["renewal"]["records"] == 1
    generated = generated_from_calculation({"period": "2026-08", "rows": [{
        "teacher": "教师甲", "employment_type": "FULL_TIME", "fields": {
            "AA": {"value": 1, "state": "DETERMINED"}, "AC": {"value": 2, "state": "DETERMINED"},
            "AD": {"value": 3, "state": "DETERMINED"}, "AE": {"value": 40, "state": "DETERMINED"},
            "AF": {"value": 120, "state": "DETERMINED"}, "PART_TIME": {"value": None, "state": "NOT_APPLICABLE"},
        },
    }]}, renewal_snapshot=confirmed["run_renewal_result_snapshot"])
    assert generated.rows[0].final_fields["AK"]["value"] == pytest.approx(8.75)


def _part_time_run(service: PayrollService, teacher: str = "测试兼职教师") -> dict:
    run = _run_with_roster(service)
    stored = service.store.get(run["id"])
    stored["core_calculation"] = {"rows": [{"teacher": teacher, "teacher_id": teacher}]}
    service.store.save(stored)
    service.save_employment_profile(teacher, "PART_TIME", "核算负责人", reason="脱敏 UAT 用工确认", effective_from="2026-08", effective_to="2026-08")
    return service.store.get(run["id"])


def test_part_time_company_standard_is_pending_without_authority_and_durable(tmp_path):
    database = tmp_path / "data"
    service = PayrollService(database)
    run = _part_time_run(service)

    saved = service.save_part_time_pay_decision(run["id"], "测试兼职教师", "COMPANY_STANDARD", "核算负责人")
    decision = saved["part_time_pay_decisions"]["测试兼职教师"]
    assert decision["status"] == "WAITING_FOR_AUTHORITY"
    assert decision["amount"] is None
    assert decision["rate_version_id"] == ""

    reopened = PayrollService(database).store.get(run["id"])
    assert reopened["part_time_pay_decisions"]["测试兼职教师"] == decision


def test_part_time_company_standard_uses_explicitly_bound_versioned_rate(tmp_path):
    service = PayrollService(tmp_path / "data")
    run = _part_time_run(service)
    versions = service.save_part_time_rate_version(
        [{"teacher": "测试兼职教师", "grade_scope": "*", "pricing_mode": "FIXED_GRADE_RATE", "fixed_rate": 123.0}],
        "脱敏 UAT 费率资料", "2026-08", "2026-08", "核算负责人",
    )
    version = versions[0]
    service.rebind_calculation(run["id"], "part_time", version["id"])
    saved = service.save_part_time_pay_decision(run["id"], "测试兼职教师", "COMPANY_STANDARD", "核算负责人")
    decision = saved["part_time_pay_decisions"]["测试兼职教师"]
    assert decision["status"] == "CONFIRMED"
    assert decision["rate_version_id"] == version["id"]
    assert saved["run_policy_snapshot"]["part_time_rate_version_id"] == version["id"]


def test_confirmed_part_time_profile_overrides_stale_full_time_core_policy(tmp_path):
    from payroll_core.models.records import ScheduleRecord

    service = PayrollService(tmp_path / "data")
    service.save_policy_version("2026-08", "2026-08", "脱敏旧政策", [{
        "teacher": "测试兼职教师", "role": "教师", "employment_type": "FULL_TIME",
        "allow_no_teaching": False, "obligation_hours": 30,
        "obligation_hours_deduction_enabled": True,
    }])
    run = service.create("2026-08", "GENERATE")
    service.save_employment_profile("测试兼职教师", "PART_TIME", "核算负责人", reason="长期用工性质确认", effective_from="2026-08", effective_to="9999-12")
    run = service.store.get(run["id"])
    schedule = [ScheduleRecord("2026-08", "测试兼职教师", "九年级", "数学", "1对1", 1, lesson_status="已上课")]

    calculation = service._configured_calculation(run, schedule, [])
    row = calculation["rows"][0]
    assert row["employment_type"] == "PART_TIME"
    assert all(row["fields"][code]["state"] == "NOT_APPLICABLE" for code in ("AA", "AC", "AD", "AE", "AF"))
    assert row["fields"]["PART_TIME"]["state"] == "NEEDS_INPUT"


@pytest.mark.parametrize("manual_kind,amount", [("TOTAL", 1234.5), ("UNIT_RATE", 123.45)])
def test_part_time_manual_choices_are_run_scoped_and_durable(tmp_path, manual_kind, amount):
    database = tmp_path / f"data-{manual_kind}"
    service = PayrollService(database)
    run = _part_time_run(service)

    saved = service.save_part_time_pay_decision(
        run["id"], "测试兼职教师", "MANUAL", "核算负责人",
        amount=amount, manual_kind=manual_kind, reason="脱敏 UAT 本月人工确认",
    )
    decision = saved["part_time_pay_decisions"]["测试兼职教师"]
    assert decision["status"] == "CONFIRMED"
    assert decision["amount"] == amount
    assert decision["manual_kind"] == manual_kind
    assert decision["confirmed_by"] == "核算负责人"
    assert decision["confirmed_at"]
    assert decision["run_id"] == run["id"]
    assert decision["period"] == "2026-08"
    assert PayrollService(database).store.get(run["id"])["part_time_pay_decisions"]["测试兼职教师"] == decision


def test_part_time_deferred_is_blank_not_zero_after_reload_and_canonical_generation(tmp_path):
    database = tmp_path / "data"
    service = PayrollService(database)
    run = _part_time_run(service)
    saved = service.save_part_time_pay_decision(run["id"], "测试兼职教师", "DEFERRED", "核算负责人")
    decision = PayrollService(database).store.get(run["id"])["part_time_pay_decisions"]["测试兼职教师"]
    assert decision["status"] == "DEFERRED"
    assert decision["amount"] is None

    generated = generated_from_calculation({"period": "2026-08", "rows": [{
        "teacher": "测试兼职教师", "employment_type": "PART_TIME", "fields": {
            "AA": {"value": None, "state": "NOT_APPLICABLE"},
            "AC": {"value": None, "state": "NOT_APPLICABLE"},
            "AD": {"value": None, "state": "NOT_APPLICABLE"},
            "AE": {"value": None, "state": "NOT_APPLICABLE"},
            "AF": {"value": None, "state": "NOT_APPLICABLE"},
            "PART_TIME": {"value": None, "state": "NEEDS_INPUT"},
        },
    }]}, employment_types={"测试兼职教师": "PART_TIME"}, part_time_pay_decisions={"测试兼职教师": decision})
    final_af = generated.rows[0].final_fields["AF"]
    assert final_af["state"] == "DEFERRED"
    assert final_af["value"] is None


@pytest.mark.parametrize(
    "decision,expected_state,expected_value",
    [
        ({"method": "MANUAL", "manual_kind": "TOTAL", "amount": 1350.0, "status": "CONFIRMED", "source": "本月人工确认", "reason": "脱敏依据"}, "DETERMINED", 1350.0),
        ({"method": "DEFERRED", "manual_kind": "", "amount": None, "status": "DEFERRED", "source": "用户选择暂时留白", "reason": "待补充"}, "DEFERRED", None),
    ],
)
def test_final_workbook_does_not_apply_full_time_af_formula_to_part_time(tmp_path, decision, expected_state, expected_value):
    from openpyxl import load_workbook
    from payroll_core.excel.standard_payroll_render import _template_headers, render_generated_payroll
    from test_standard_payroll_output import _sanitized_template

    decision = {"run_id": "run-u", "period": "2026-08", "teacher": "测试兼职教师", "teacher_id": "测试兼职教师", "confirmed_by": "核算负责人", "confirmed_at": "2026-09-26T00:00:00+00:00", **decision}
    core = {"period": "2026-08", "rows": [{
        "teacher": "测试兼职教师", "employment_type": "PART_TIME", "fields": {
            "AA": {"value": None, "state": "NOT_APPLICABLE"},
            "AC": {"value": None, "state": "NOT_APPLICABLE"},
            "AD": {"value": None, "state": "NOT_APPLICABLE"},
            "AE": {"value": None, "state": "NOT_APPLICABLE"},
            "AF": {"value": None, "state": "NOT_APPLICABLE"},
            "PART_TIME": {"value": None, "state": "NEEDS_INPUT"},
        },
    }]}
    payroll = generated_from_calculation(core, employment_types={"测试兼职教师": "PART_TIME"}, base_salary_inputs={}, part_time_pay_decisions={"测试兼职教师": decision})
    assert payroll.rows[0].final_fields["AF"]["state"] == expected_state
    assert payroll.rows[0].final_fields["AF"]["value"] == expected_value
    template = _sanitized_template(tmp_path)
    output = render_generated_payroll(payroll, tmp_path / f"兼职-{expected_state}.xlsx", template_path=template)
    workbook = load_workbook(output, data_only=False)
    sheet = workbook[workbook.sheetnames[0]]
    columns = _template_headers(sheet)
    assert sheet.cell(5, columns["af"]).value == expected_value
    assert not (isinstance(sheet.cell(5, columns["af"]).value, str) and sheet.cell(5, columns["af"]).value.startswith("="))
    assert sheet.cell(5, columns["af"]).comment is not None
    assert "本月兼职工资" in sheet.cell(5, columns["af"]).comment.text
    if expected_state == "DEFERRED":
        assert "兼职工资待补充" in str(sheet.cell(5, sheet.max_column).value)


def test_preview_and_export_refund_annotations_use_the_same_approved_rows(tmp_path):
    from openpyxl import load_workbook
    from payroll_core.excel.standard_payroll_render import FINAL_FIELD_CODES, _template_headers, render_generated_payroll
    from test_standard_payroll_output import _generated, _sanitized_template

    refund = {
        "id": "refund-u",
        "teacher_id": "教师甲",
        "period": "2026-08",
        "input_type": "REFUND_RESULT",
        "status": "APPROVED",
        "source_ref": "/private/sources/脱敏退费结果.xlsx",
        "source_row": "8月!5",
        "payload": {"student": "学生甲", "business_type": "退费", "headcount_amount": -50, "performance_amount": -10},
    }
    payroll = _generated(business_inputs=[refund])
    template = _sanitized_template(tmp_path)
    path = render_generated_payroll(payroll, tmp_path / "工资结果.xlsx", template_path=template, manual_adjustments=[refund])
    workbook = load_workbook(path, data_only=False)
    sheet = workbook[workbook.sheetnames[0]]
    refund_cell = sheet.cell(5, 33 + FINAL_FIELD_CODES.index("AN"))
    assert sheet.cell(5, _template_headers(sheet)["af"]).comment is None
    assert refund_cell.value == -60
    assert refund_cell.comment is not None
    assert "学生甲" in refund_cell.comment.text
    assert "退费" in refund_cell.comment.text
    assert "人头金额：-50" in refund_cell.comment.text
    assert "业绩金额：-10" in refund_cell.comment.text
    assert "脱敏退费结果.xlsx" in refund_cell.comment.text
    assert "/private/" not in refund_cell.comment.text


def test_support_annotation_is_written_to_the_matching_template_cell(tmp_path):
    from openpyxl import load_workbook
    from payroll_core.excel.standard_payroll_render import render_generated_payroll
    from test_standard_payroll_output import _generated, _sanitized_template

    annotation = {
        "teacher_id": "teacher-1", "teacher": "教师甲", "field_code": "H",
        "display_label": "H｜岗位津贴", "text": "本期岗位津贴来源说明。",
        "source_type": "SUPPORT_DEPARTMENT", "source_file": "支持部.xlsx",
        "source_sheet": "教学部", "source_row": "5", "source_field": "岗位津贴",
        "generated_from": "SUPPORT_SOURCE_CELL_COMMENT",
    }
    snapshot = {"entries": {"teacher-1": {"display_name": "教师甲", "identity": {}, "annotations": [annotation, annotation]}}}
    template = _sanitized_template(tmp_path)
    path = render_generated_payroll(_generated(), tmp_path / "含支持批注.xlsx", template_path=template, support_snapshot=snapshot)
    workbook = load_workbook(path)
    comment = workbook[workbook.sheetnames[0]]["H5"].comment
    assert comment is not None
    assert comment.text.count("本期岗位津贴来源说明") == 1
    assert "支持部.xlsx" in comment.text


def test_class_annotation_explains_current_core_contributions_without_recalculating(tmp_path):
    from openpyxl import load_workbook
    from payroll_core.calculation import calculate_payroll
    from payroll_core.config.core_rules import load_core_rules
    from payroll_core.excel.standard_payroll_render import render_generated_payroll
    from payroll_core.models.records import ScheduleRecord
    from payroll_core.payroll_generation import generated_from_calculation
    from test_standard_payroll_output import _sanitized_template

    schedule = [
        ScheduleRecord("2026-08", "教师甲", "九年级", "数学", "小班", 2, lesson_status="已上课", source="脱敏排课.xlsx", class_name="脱敏班级")
    ]
    result = calculate_payroll("2026-08", schedule, load_core_rules()).as_dict()
    names = {"aa": "AA", "ac": "AC", "ad": "AD", "ae": "AE", "af": "AF", "part_time_fee": "PART_TIME"}
    rows = [{"teacher": row["teacher"], "fields": {code: row[key] for key, code in names.items()}} for row in result["rows"]]
    generated = generated_from_calculation({
        "period": "2026-08", "rows": rows, "course_contributions": result["course_contributions"],
        "source_records": [{"record_key": result["course_contributions"][0]["record_key"], "teacher": "教师甲", "source": "脱敏排课.xlsx", "provenance": {"attended": {"sheet": "排课", "coordinate": "J5"}}}],
    })
    template = _sanitized_template(tmp_path)
    path = render_generated_payroll(generated, tmp_path / "班课说明.xlsx", template_path=template)
    workbook = load_workbook(path)
    sheet = workbook[workbook.sheetnames[0]]
    assert sheet["AC5"].comment is not None
    assert "AC 合计" in sheet["AC5"].comment.text
    assert "attended=2" in sheet["AC5"].comment.text
    assert "脱敏排课.xlsx" in sheet["AC5"].comment.text
    assert "排课" in sheet["AC5"].comment.text
    assert "第 5 行" in sheet["AC5"].comment.text
    assert sheet["AC5"].value == generated.rows[0].class_value
