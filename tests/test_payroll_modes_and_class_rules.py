"""Configurable class-type rules and the two payroll modes sharing one Core."""
from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook, Workbook

from payroll_core.models.records import ScheduleRecord
from payroll_core.models.class_type_rules import CONFIRMED_SPECIAL_CLASS_COEFFICIENTS, structured_rules
from payroll_core.reconcile.payroll_scope import (
    class_value_contribution,
    schedule_field_checks,
    schedule_totals,
)
from payroll_ui.service import PayrollService
from tests.test_payroll_ui_service import FIXTURES, _prepared_run_with_values
from tests.payroll_test_helpers import bind_template


def _record(class_type: str, attended: int, grade: str = "八年级") -> ScheduleRecord:
    return ScheduleRecord(period="2026-08", teacher="张三", grade=grade, subject="数学", class_type=class_type, attended=attended, lesson_status="已上课", source="schedule.xlsx")


def _schedule(path: Path, rows: list[list[object]], title: str = "课表") -> Path:
    book = Workbook()
    sheet = book.active
    sheet.title = title
    sheet.append(["任课老师", "年级", "学科", "课程所属班型", "实到人数", "上课状态"])
    for row in rows:
        sheet.append(row)
    book.save(path)
    return path


def test_two_person_class_uses_configured_coefficient():
    record = _record("1对2", attended=2)
    rules = structured_rules(CONFIRMED_SPECIAL_CLASS_COEFFICIENTS)
    value, calculation = class_value_contribution(record, rules)

    assert value is not None
    assert "班型系数 1.2" in calculation
    # The same record under a different configuration must give a different
    # number, which proves the coefficient is read from configuration.
    doubled, _ = class_value_contribution(record, {"special_class_coefficients": {"1对2": {"1": 1.6, "2": 2.4}}})
    assert doubled == pytest.approx(value * 2)


def test_three_person_class_uses_configured_coefficient():
    record = _record("三人班", attended=3)

    value, calculation = class_value_contribution(record, structured_rules(CONFIRMED_SPECIAL_CLASS_COEFFICIENTS))
    assert value is not None and "班型系数 1.5" in calculation
    # 未经配置的班型依旧拒绝，不按小班或 1 猜。
    assert class_value_contribution(_record("四人精品班", attended=4))[0] is None


def test_special_class_uses_its_own_attendance_mapping_not_small_group_mapping():
    """1对2 的实到 1/2 人各自查专用配置，不能退回普通小班。"""
    rules = structured_rules(CONFIRMED_SPECIAL_CLASS_COEFFICIENTS)
    two_at_one = class_value_contribution(_record("1对2", attended=1), rules)
    two_at_two = class_value_contribution(_record("1对2", attended=2), rules)
    two_at_three = class_value_contribution(_record("1对2", attended=3), rules)

    assert two_at_one[0] == pytest.approx(1.44)
    assert two_at_two[0] == pytest.approx(2.16)
    assert two_at_three[0] is None and "UNKNOWN_CLASS_RULE" in two_at_three[1]


def test_ordinary_small_group_is_priced_by_attendance_only():
    """普通小班没有班型系数：只有 年级系数 × 实到人数系数 × 2。"""
    one, two, three = (class_value_contribution(_record("小班", attended=people)) for people in (1, 2, 3))

    assert one[0] == pytest.approx(1.44) and "实到人数系数 0.8" in one[1]
    assert two[0] == pytest.approx(1.8) and "实到人数系数 1" in two[1]
    assert three[0] == pytest.approx(2.16) and "实到人数系数 1.2" in three[1]
    assert "班型系数" not in three[1]


def test_a_fixed_coefficient_for_ordinary_small_group_is_rejected():
    """小班 曾经被写成单一固定系数（1.0），这是本轮回退掉的错误建模。"""
    with pytest.raises(ValueError, match="实到人数"):
        class_value_contribution(_record("小班", attended=2), {"小班": 0.8})
    with pytest.raises(ValueError, match="实到人数"):
        class_value_contribution(_record("小班", attended=2), {"small_group_headcount_coefficients": {"2": 1.0}, "special_class_coefficients": {"小班": 0.8}})


def test_unknown_class_type_blocks_calculation():
    record = _record("四人精品班", attended=4)

    value, reason = class_value_contribution(record)
    totals, blockers = schedule_totals([record])

    assert value is None and reason.startswith("UNKNOWN_CLASS_TYPE_RULE")
    assert totals["张三"]["class_value"] == 0.0
    assert [item.status for item in blockers] == ["UNKNOWN_CLASS_TYPE_RULE"]
    assert "四人精品班" in blockers[0].reason
    checks = schedule_field_checks([record], [])
    assert any(item.status == "UNKNOWN_CLASS_TYPE_RULE" for item in checks)


def test_changing_coefficient_requires_new_version(tmp_path):
    service = PayrollService(tmp_path / "data")
    seeded = service.class_type_rule_versions()
    assert {item["id"] for item in seeded} >= {"CLASS_TYPE_RULE_2020_BASELINE", "CLASS_TYPE_RULE_2026-09"}

    with pytest.raises(ValueError, match="重叠"):
        service.save_class_type_rule_version(effective_from="2026-09-15", effective_to="2026-12-31", rules={"special_class_coefficients": {"1对2": {"1": 0.8, "2": 1.3}}}, source="测试", actor="管理员")

    updated = service.save_class_type_rule_version(
        effective_from="2027-01-01", effective_to="9999-12-31",
        rules={"special_class_coefficients": {"1对2": {"1": 0.8, "2": 1.3}}},
        source="测试", actor="管理员", supersedes_version_id="CLASS_TYPE_RULE_2026-09",
    )
    ids = {item["id"] for item in updated}
    assert len(ids) == 3, "旧版本必须保留，只能新增"
    closed = next(item for item in updated if item["id"] == "CLASS_TYPE_RULE_2026-09")
    assert closed["status"] == "SUPERSEDED" and closed["effective_to"] == "2026-12-31"
    with pytest.raises(ValueError, match="大于 0"):
        service.save_class_type_rule_version(effective_from="2028-01-01", effective_to="2028-12-31", rules={"special_class_coefficients": {"1对2": {"2": 0}}}, source="测试", actor="管理员")
    with pytest.raises(ValueError, match="实到人数"):
        service.save_class_type_rule_version(effective_from="2028-01-01", effective_to="2028-12-31", rules={"special_class_coefficients": {"小班": {"2": 0.8}}}, source="测试", actor="管理员")


def test_historical_run_keeps_historical_rule(tmp_path):
    """A future coefficient change must never rewrite an already-closed month."""
    service = PayrollService(tmp_path / "data")
    august = service.create("2026-08")
    october = service.create("2026-10")

    assert august["class_type_rule_version_id"] == "CLASS_TYPE_RULE_2020_BASELINE"
    assert october["class_type_rule_version_id"] == "CLASS_TYPE_RULE_2026-09"

    august_rules, _ = service._class_type_rules_for_run(service.store.get(august["id"]))
    october_rules, _ = service._class_type_rules_for_run(service.store.get(october["id"]))
    assert "三人班" not in august_rules["special_class_coefficients"], "2026-08 不能因为 9 月新增规则而改变"
    assert "1对3" not in august_rules["special_class_coefficients"]
    assert "small_group_headcount_coefficients" not in august_rules
    assert october_rules["special_class_coefficients"]["三人班"]["3"] == 1.5
    assert october_rules["special_class_coefficients"]["1对3"]["3"] == 1.5

    service.save_class_type_rule_version(
        effective_from="2027-01-01", effective_to="9999-12-31",
        rules={"special_class_coefficients": {"1对2": {"1": 0.8, "2": 1.6}}},
        source="测试", actor="管理员", supersedes_version_id="CLASS_TYPE_RULE_2026-09",
    )
    after = service._class_type_rules_for_run(service.store.get(august["id"]))
    assert after[0] == august_rules, "历史 Run 仍然绑定当时的规则版本"


def test_class_type_rule_is_not_hardcoded(tmp_path):
    """A class type never used before can be enabled purely by configuration."""
    service = PayrollService(tmp_path / "data")
    service.save_class_type_rule_version(
        effective_from="2027-01-01", effective_to="9999-12-31",
        rules={
            "special_class_coefficients": {"1对2": {"1": 0.8, "2": 1.2}, "三人班": {"1": 0.8, "2": 1.2, "3": 1.5}, "四人精品班": {"4": 2.0}},
        },
        source="用户配置", actor="管理员", supersedes_version_id="CLASS_TYPE_RULE_2026-09",
    )
    run = service.create("2027-03")
    coefficients, version_id = service._class_type_rules_for_run(service.store.get(run["id"]))

    assert coefficients["special_class_coefficients"]["四人精品班"]["4"] == 2.0
    assert version_id != "CLASS_TYPE_RULE_2026-09"
    value, _ = class_value_contribution(_record("四人精品班", attended=4), coefficients)
    assert value is not None


def test_generated_and_audited_modes_use_same_core_result(tmp_path):
    """Case A vs Case B: identical inputs must produce identical Core numbers."""
    service, audit_run, schedule, math, science = _prepared_run_with_values(tmp_path, one_to_one=1.8, class_value=3.24)
    checked = service.check(audit_run["id"])
    audited = {item["teacher"]: item for item in checked["field_records"] if item["field"] == "class_value"}

    generated_run = bind_template(service, service.create("2026-08", "GENERATE"))
    service.import_file(generated_run["id"], "schedule", str(schedule))
    output = tmp_path / "生成工资表.xlsx"
    produced = service.generate_payroll(generated_run["id"], str(output))

    assert produced["status"] in {"FINAL", "NEEDS_CONFIRMATION"}
    for row in produced["rows"]:
        expected = audited.get(row["teacher"])
        if expected and expected["expected"] is not None:
            assert row["class_value"] == pytest.approx(expected["expected"]), "两模式不能出现算法漂移"
    book = load_workbook(output)
    assert book["标准工资表"]["C3"].value == "姓名"
    assert book["标准工资表"]["H2"].value == "2026-08"


def test_generate_mode_does_not_require_a_submitted_payroll_sheet(tmp_path):
    service = PayrollService(tmp_path / "data")
    run = bind_template(service, service.create("2026-08", "GENERATE"))
    path = _schedule(tmp_path / "schedule.xlsx", [["张三", "八年级", "数学", "1对1", 2, "已上课"]])

    imported = service.import_file(run["id"], "schedule", str(path))
    assert imported["health"]["ready"] is True
    assert imported["materials"] and all(item["required"] is False for item in imported["materials"] if item["role"] in {"math", "science"})

    produced = service.generate_payroll(run["id"], str(tmp_path / "out.xlsx"))
    assert produced["rows"][0]["one_to_one"] > 0
    row = produced["rows"][0]
    assert row["teaching_hours"] == row["one_to_one"] + row["class_value"]
    assert row["fields"]["AD"]["state"] == "DETERMINED"
    assert "AD_MISSING_SOURCE" not in row["blockers"]
    assert row["ae"] == row["af"] == 0, "已确定AD不超过门槛，课时费为0不需要猜星级"


def test_audit_mode_still_requires_a_submitted_payroll_sheet(tmp_path):
    service = PayrollService(tmp_path / "data")
    run = service.create("2026-08")
    path = _schedule(tmp_path / "schedule.xlsx", [["张三", "八年级", "数学", "1对1", 2, "已上课"]])

    imported = service.import_file(run["id"], "schedule", str(path))

    assert imported["health"]["ready"] is False
    assert "至少一张教师提交表" in imported["health"]["missing"]
    with pytest.raises(ValueError, match="至少一张教师提交表"):
        service.generate_payroll(run["id"], str(tmp_path / "out.xlsx"))


def test_generate_mode_marks_draft_and_never_overwrites(tmp_path):
    service = PayrollService(tmp_path / "data")
    run = bind_template(service, service.create("2026-08", "GENERATE"))
    service.import_file(run["id"], "schedule", str(_schedule(tmp_path / "s.xlsx", [["张三", "八年级", "数学", "1对1", 2, "已上课"]])))
    output = tmp_path / "标准工资表.xlsx"

    service.generate_payroll(run["id"], str(output))
    book = load_workbook(output)
    assert book["标准工资表"]["H2"].value == "2026-08"
    assert book["标准工资表"]["C5"].value == "张三"
    first_bytes = output.read_bytes()
    second = service.generate_payroll(run["id"], str(output))
    assert Path(second["path"]).name == "标准工资表 (2).xlsx"
    assert output.read_bytes() == first_bytes


def test_unknown_class_type_blocks_the_generated_payroll(tmp_path):
    service = PayrollService(tmp_path / "data")
    run = bind_template(service, service.create("2026-08", "GENERATE"))
    service.import_file(run["id"], "schedule", str(_schedule(tmp_path / "s.xlsx", [["张三", "八年级", "数学", "四人精品班", 4, "已上课"]])))

    produced = service.generate_payroll(run["id"], str(tmp_path / "out.xlsx"))

    assert produced["status"] == "NEEDS_CONFIRMATION"
    assert produced["rows"][0]["fields"]["AC"]["state"] == "NEEDS_INPUT"
    assert produced["rows"][0]["fields"]["AC"]["value"] is None
