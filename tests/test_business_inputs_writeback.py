from __future__ import annotations

from pathlib import Path
from shutil import copyfile
import json
from threading import Thread
from urllib.error import HTTPError
from urllib.request import Request, urlopen

import pytest
from openpyxl import load_workbook
from openpyxl.comments import Comment

from payroll_ui.service import PayrollService
from payroll_ui.server import PayrollHttpServer
import payroll_ui.business_inputs as business_inputs_module
from payroll_core.excel.writeback import write_new_workbook
from tests.test_payroll_ui_service import FIXTURES, _prepared_run


def _csv(path: Path, rows: str) -> Path:
    path.write_text(rows, encoding="utf-8")
    return path


def _approved_refund(service: PayrollService, tmp_path: Path, run: dict | None = None) -> dict:
    source = _csv(tmp_path / "refund.csv", "教师,学生,退费金额,说明\n张三,学生甲,100,已确认退费\n")
    item = service.import_business_results("REFUND_RESULT", "2026-08", str(source), "审核员")[0]
    service.review_business_input(item["id"], "START_REVIEW", "审核员")
    approved = service.review_business_input(item["id"], "APPROVE", "审核员", "上游结果已确认")
    if run:
        service.bind_business_input(run["id"], approved["id"])
    return approved


def _approve_candidate(service: PayrollService, candidate: dict) -> dict:
    preview = service.preview_comment_candidate(candidate["id"], "APPEND")
    return service.approve_comment_candidate(candidate["id"], "审核员", preview["preview_token"])


def test_teacher_can_submit_own_business_input(tmp_path):
    service = PayrollService(tmp_path / "data")
    access = service.create_teacher_access("教师甲")
    submitted = service.teacher_submit(access["access_token"], "2026-08", "班课特殊核算", "请核实本节课程")
    assert submitted["teacher_id"] == "教师甲"
    assert submitted["status"] == "SUBMITTED"


def test_teacher_cannot_view_other_teacher_submission(tmp_path):
    service = PayrollService(tmp_path / "data")
    first, second = service.create_teacher_access("教师甲"), service.create_teacher_access("教师乙")
    service.teacher_submit(first["access_token"], "2026-08", "其他", "仅教师甲可见")
    assert service.teacher_inputs(second["access_token"]) == []


def test_teacher_submission_is_not_automatically_approved(tmp_path):
    service = PayrollService(tmp_path / "data")
    access = service.create_teacher_access("教师甲")
    record = service.teacher_submit(access["access_token"], "2026-08", "排课信息有误", "请核实")
    assert record["status"] == "SUBMITTED"


def test_draft_cannot_approve_and_review_history_is_preserved(tmp_path):
    service = PayrollService(tmp_path / "data")
    access = service.create_teacher_access("教师甲")
    draft = service.teacher_save_draft(access["access_token"], "2026-08", "其他", "草稿")
    with pytest.raises(ValueError, match="草稿必须先提交"):
        service.review_business_input(draft["id"], "APPROVE", "主管")
    submitted = service.teacher_submit(access["access_token"], "2026-08", "其他", "正式提交")
    service.review_business_input(submitted["id"], "START_REVIEW", "主管")
    service.review_business_input(submitted["id"], "REQUEST_MORE_INFO", "主管", "请补充课程")
    service.review_business_input(submitted["id"], "START_REVIEW", "主管")
    service.review_business_input(submitted["id"], "APPROVE", "主管", "资料完整")
    assert [event["event_type"] for event in service.store.list_business_input_events(submitted["id"])] == [
        "SUBMITTED", "REVIEW", "REQUEST_MORE_INFO", "REVIEW", "APPROVED",
    ]


def test_teacher_can_modify_own_draft_only(tmp_path):
    service = PayrollService(tmp_path / "data")
    alpha, beta = service.create_teacher_access("教师甲"), service.create_teacher_access("教师乙")
    draft = service.teacher_save_draft(alpha["access_token"], "2026-08", "其他", "初稿")
    revised = service.teacher_save_draft(alpha["access_token"], "2026-08", "其他", "修订稿", input_id=draft["id"])
    assert revised["payload"]["note"] == "修订稿"
    with pytest.raises(PermissionError):
        service.teacher_save_draft(beta["access_token"], "2026-08", "其他", "越权", input_id=draft["id"])


def test_approved_submission_can_link_to_run(tmp_path):
    service, run, _ = _prepared_run(tmp_path)
    access = service.create_teacher_access("教师甲")
    record = service.teacher_submit(access["access_token"], "2026-08", "排课信息有误", "请核实")
    service.review_business_input(record["id"], "START_REVIEW", "主管")
    service.review_business_input(record["id"], "APPROVE", "主管")
    bound = service.bind_business_input(run["id"], record["id"])
    assert bound["business_input_bindings"][0]["input_id"] == record["id"]


def test_rejected_submission_does_not_affect_payroll(tmp_path):
    service, run, _ = _prepared_run(tmp_path)
    access = service.create_teacher_access("教师甲")
    record = service.teacher_submit(access["access_token"], "2026-08", "其他", "说明")
    service.review_business_input(record["id"], "START_REVIEW", "主管")
    service.review_business_input(record["id"], "REJECT", "主管")
    with pytest.raises(ValueError, match="已审核通过"):
        service.bind_business_input(run["id"], record["id"])


def test_renewal_import_preserves_source_provenance(tmp_path):
    service = PayrollService(tmp_path / "data")
    source = _csv(tmp_path / "renewal.csv", "教师,学生,说明\n教师甲,学生甲,最终续费\n")
    item = service.import_business_results("RENEWAL_RESULT", "2026-08", str(source), "业务员")[0]
    assert item["evidence"]["source_file"] == "renewal.csv"
    assert item["source_row"] == "2"
    assert item["source_file_hash"]


def test_renewal_new_version_does_not_delete_old_version(tmp_path):
    service = PayrollService(tmp_path / "data")
    first = _csv(tmp_path / "renewal-v1.csv", "教师,学生\n教师甲,学生甲\n")
    second = _csv(tmp_path / "renewal-v2.csv", "教师,学生\n教师甲,学生乙\n")
    old = service.import_business_results("RENEWAL_RESULT", "2026-08", str(first), "业务员")[0]
    service.review_business_input(old["id"], "START_REVIEW", "主管")
    service.review_business_input(old["id"], "APPROVE", "主管")
    fresh = service.import_business_results("RENEWAL_RESULT", "2026-08", str(second), "业务员")[0]
    assert service.store.get_business_input(old["id"])["status"] == "APPROVED"
    assert service.store.get_business_input(fresh["id"])["status"] == "SUBMITTED"


def test_import_replacement_is_explicit_and_scoped(tmp_path):
    service = PayrollService(tmp_path / "data")
    first = service.import_business_results("RENEWAL_RESULT", "2026-08", str(_csv(tmp_path / "v1.csv", "教师,学生\n教师甲,学生甲\n")), "业务员")[0]
    second = service.import_business_results("RENEWAL_RESULT", "2026-08", str(_csv(tmp_path / "v2.csv", "教师,学生\n教师乙,学生乙\n")), "业务员")[0]
    service.review_business_input(first["id"], "START_REVIEW", "主管")
    service.review_business_input(first["id"], "APPROVE", "主管")
    service.review_business_input(second["id"], "START_REVIEW", "主管")
    service.review_business_input(second["id"], "APPROVE", "主管")
    service.import_business_results(
        "RENEWAL_RESULT", "2026-08", str(_csv(tmp_path / "v3.csv", "教师,学生\n教师丙,学生丙\n")), "业务员",
        activation_scope="REPLACE_SELECTED", replace_input_ids=[first["id"]],
    )
    assert service.store.get_business_input(first["id"])["status"] == "SUPERSEDED"
    assert service.store.get_business_input(second["id"])["status"] == "APPROVED"


def test_import_rejects_file_changed_during_read(tmp_path, monkeypatch):
    service = PayrollService(tmp_path / "data")
    source = _csv(tmp_path / "renewal.csv", "教师,学生\n教师甲,学生甲\n")
    original = business_inputs_module.file_version
    calls = 0

    def unstable_version(path):
        nonlocal calls
        calls += 1
        result = original(path)
        if calls == 2:
            return {**result, "sha256": "changed"}
        return result

    monkeypatch.setattr(business_inputs_module, "file_version", unstable_version)
    with pytest.raises(ValueError, match="读取期间发生变化"):
        service.import_business_results("RENEWAL_RESULT", "2026-08", str(source), "业务员")
    assert service.business_inputs() == []


def test_renewal_result_without_rule_does_not_guess_bonus(tmp_path):
    service = PayrollService(tmp_path / "data")
    item = service.import_business_results("RENEWAL_RESULT", "2026-08", str(_csv(tmp_path / "renewal.csv", "教师,学生\n教师甲,学生甲\n")), "业务员")[0]
    assert "bonus" not in item["payload"]


def test_refund_import_preserves_source_provenance(tmp_path):
    service = PayrollService(tmp_path / "data")
    refund = _approved_refund(service, tmp_path)
    assert refund["evidence"]["sheet"] == "CSV"
    assert refund["payload"]["student"] == "学生甲"


def test_refund_result_can_generate_comment_candidate(tmp_path):
    service, run, _ = _prepared_run(tmp_path)
    refund = _approved_refund(service, tmp_path, run)
    candidate = service.create_refund_comment_candidate(run["id"], refund["id"], "math", "Sheet1", "AC5")
    assert candidate["comment_type"] == "REFUND_NOTE"
    assert candidate["status"] == "PROPOSED"
    assert "退费" in candidate["content"]


def test_superseded_refund_invalidates_comment_candidate(tmp_path):
    service, run, _ = _prepared_run(tmp_path)
    refund = _approved_refund(service, tmp_path, run)
    candidate = service.create_refund_comment_candidate(run["id"], refund["id"], "math", "Sheet1", "AC5")
    service.import_business_results(
        "REFUND_RESULT", "2026-08", str(_csv(tmp_path / "refund-v2.csv", "教师,学生\n教师甲,学生乙\n")), "业务员",
        activation_scope="REPLACE_SELECTED", replace_input_ids=[refund["id"]],
    )
    assert service.comment_candidates(run["id"])[0]["status"] == "NEEDS_RECONFIRMATION"


def test_deleted_source_invalidates_record_and_comment_candidate(tmp_path):
    service, run, _ = _prepared_run(tmp_path)
    refund = _approved_refund(service, tmp_path, run)
    candidate = service.create_refund_comment_candidate(run["id"], refund["id"], "math", "Sheet1", "AC5")
    Path(refund["source_ref"]).unlink()
    assert not service.inputs.current(refund["id"])
    assert service.comment_candidates(run["id"])[0]["status"] == "NEEDS_RECONFIRMATION"


def test_changed_source_invalidates_comment_candidate(tmp_path):
    service, run, _ = _prepared_run(tmp_path)
    refund = _approved_refund(service, tmp_path, run)
    candidate = service.create_refund_comment_candidate(run["id"], refund["id"], "math", "Sheet1", "AC5")
    Path(refund["source_ref"]).write_text("教师,学生,退费金额,说明\n张三,学生甲,101,来源已更新\n", encoding="utf-8")
    assert service.comment_candidates(run["id"])[0]["status"] == "NEEDS_RECONFIRMATION"


def test_class_course_resolution_generates_comment_candidate(tmp_path):
    service, run, _ = _prepared_run(tmp_path)
    resolution = {"id": "resolution-1", "run_id": run["id"], "period": run["period"], "kind": "SOURCE_DATA_CORRECTION", "outcome": "RESOLVED_BY_SOURCE_CORRECTION", "teacher": "张三", "course_label": "脱敏示例课程", "reason": "脱敏：升年级未同步", "corrected_value": "高二", "system_value": "3.24", "fingerprint": "resolution-fingerprint", "created_at": "2026-08-31T00:00:00+00:00"}
    service.store.save_resolution(resolution)
    candidate = service.create_class_comment_candidate(run["id"], resolution["id"], "math", "Sheet1", "AC5")
    assert candidate["comment_type"] == "SOURCE_CORRECTION_NOTE"
    assert "升年级" in candidate["content"]


def test_class_comment_reads_resolution_persisted_inside_run_payload(tmp_path):
    """The AC workflow may store resolutions on the run; the note must still trace it."""
    service, run, _ = _prepared_run(tmp_path)
    resolution = {"id": "resolution-in-run", "run_id": run["id"], "period": run["period"], "kind": "SOURCE_DATA_CORRECTION", "outcome": "RESOLVED_BY_SOURCE_CORRECTION", "teacher": "张三", "course_label": "高三6人班1节", "reason": "脱敏的年级修正", "corrected_value": "高二", "fingerprint": "run-payload-fingerprint", "created_at": "2026-08-31T00:00:00+00:00"}
    run.setdefault("resolutions", []).append(resolution)
    service.store.save(run)
    candidate = service.create_class_comment_candidate(run["id"], resolution["id"], "math", "Sheet1", "AC5")
    assert candidate["source_resolution_id"] == resolution["id"]
    assert candidate["source_resolution_fingerprint"] == resolution["fingerprint"]


def test_deferred_issue_does_not_generate_final_comment(tmp_path):
    service, run, _ = _prepared_run(tmp_path)
    service.store.save_resolution({"id": "deferred", "run_id": run["id"], "period": run["period"], "kind": "SOURCE_DATA_CORRECTION", "outcome": "DEFERRED", "teacher": "张三", "fingerprint": "x", "created_at": "2026-08-31T00:00:00+00:00"})
    with pytest.raises(ValueError, match="已确认"):
        service.create_class_comment_candidate(run["id"], "deferred", "math", "Sheet1", "AC5")


def test_class_comment_requires_server_side_persisted_resolution(tmp_path):
    service, run, _ = _prepared_run(tmp_path)
    with pytest.raises(ValueError, match="未找到指定记录"):
        service.create_class_comment_candidate(run["id"], "browser-claimed-resolution", "math", "Sheet1", "AC5")


def test_candidate_target_must_belong_to_its_teacher(tmp_path):
    service, run, _ = _prepared_run(tmp_path)
    refund = _approved_refund(service, tmp_path, run)
    with pytest.raises(ValueError, match="不属于该教师"):
        service.create_refund_comment_candidate(run["id"], refund["id"], "math", "Sheet1", "AA6")


def test_existing_comment_is_not_silently_overwritten(tmp_path):
    service, run, _ = _prepared_run(tmp_path)
    refund = _approved_refund(service, tmp_path, run)
    candidate = service.create_refund_comment_candidate(run["id"], refund["id"], "math", "Sheet1", "AC5")
    source = Path(run["files"]["math"]["path"])
    book = load_workbook(source); book["Sheet1"]["AC5"].comment = Comment("原批注", "原作者"); book.save(source)
    # The changed workbook hash forces a new candidate rather than overwriting.
    refreshed = service.import_file(run["id"], "math", str(source))
    candidate = service.create_refund_comment_candidate(refreshed["id"], refund["id"], "math", "Sheet1", "AC5")
    preview = _approve_candidate(service, candidate)
    assert preview["before_comment"] == "原批注"
    assert "原批注" in preview["after_comment"]


def test_comment_preview_required_before_write(tmp_path):
    service, run, _ = _prepared_run(tmp_path)
    refund = _approved_refund(service, tmp_path, run)
    candidate = service.create_refund_comment_candidate(run["id"], refund["id"], "math", "Sheet1", "AC5")
    with pytest.raises(ValueError, match="先在预览后确认"):
        service.writeback_comments(run["id"], run["files"]["math"]["path"], [candidate["id"]], str(tmp_path / "output.xlsx"), "审核员")


def test_preview_required_before_candidate_approval(tmp_path):
    service, run, _ = _prepared_run(tmp_path)
    refund = _approved_refund(service, tmp_path, run)
    candidate = service.create_refund_comment_candidate(run["id"], refund["id"], "math", "Sheet1", "AC5")
    with pytest.raises(ValueError, match="先预览"):
        service.approve_comment_candidate(candidate["id"], "审核员", "not-a-preview-token")


def test_writeback_creates_new_workbook_and_preserves_source_formula_cells(tmp_path):
    service, run, _ = _prepared_run(tmp_path)
    refund = _approved_refund(service, tmp_path, run)
    candidate = service.create_refund_comment_candidate(run["id"], refund["id"], "math", "Sheet1", "AC5")
    _approve_candidate(service, candidate)
    source = Path(run["files"]["math"]["path"])
    original_bytes = source.read_bytes()
    output = tmp_path / "math_系统回填.xlsx"
    result = service.writeback_comments(run["id"], str(source), [candidate["id"]], str(output), "审核员")
    assert output.is_file() and source.read_bytes() == original_bytes
    assert result["written"] == 1
    reloaded = load_workbook(output, data_only=False)
    assert reloaded["Sheet1"]["AC5"].comment is not None
    assert reloaded["Sheet1"]["AE5"].value == load_workbook(source, data_only=False)["Sheet1"]["AE5"].value


def test_written_comment_survives_reopen(tmp_path):
    service, run, _ = _prepared_run(tmp_path)
    refund = _approved_refund(service, tmp_path, run)
    candidate = service.create_refund_comment_candidate(run["id"], refund["id"], "math", "Sheet1", "AC5")
    _approve_candidate(service, candidate)
    output = tmp_path / "out.xlsx"
    service.writeback_comments(run["id"], run["files"]["math"]["path"], [candidate["id"]], str(output), "审核员")
    assert load_workbook(output)["Sheet1"]["AC5"].comment.text


def test_no_output_overwrite_and_multiple_candidates_same_cell_are_preserved(tmp_path):
    service, run, _ = _prepared_run(tmp_path)
    first = _approved_refund(service, tmp_path, run)
    second_source = _csv(tmp_path / "refund-second.csv", "教师,学生,退费金额,说明\n张三,学生乙,200,第二条退费\n")
    second = service.import_business_results("REFUND_RESULT", "2026-08", str(second_source), "审核员")[0]
    service.review_business_input(second["id"], "START_REVIEW", "审核员")
    service.review_business_input(second["id"], "APPROVE", "审核员")
    service.bind_business_input(run["id"], second["id"])
    candidates = [
        service.create_refund_comment_candidate(run["id"], first["id"], "math", "Sheet1", "AC5"),
        service.create_refund_comment_candidate(run["id"], second["id"], "math", "Sheet1", "AC5"),
    ]
    for candidate in candidates:
        _approve_candidate(service, candidate)
    output = tmp_path / "combined.xlsx"
    service.writeback_comments(run["id"], run["files"]["math"]["path"], [candidate["id"] for candidate in candidates], str(output), "审核员")
    text = load_workbook(output)["Sheet1"]["AC5"].comment.text
    assert "已确认退费" in text and "第二条退费" in text
    another = service.create_refund_comment_candidate(run["id"], first["id"], "math", "Sheet1", "AC5")
    _approve_candidate(service, another)
    # 同名输出不再报错：自动改成 (2)，并且第一次的结果一个字节都不能动。
    before_bytes = output.read_bytes()
    written = service.writeback_comments(run["id"], run["files"]["math"]["path"], [another["id"]], str(output), "审核员")
    assert Path(written["output_path"]).name == "combined (2).xlsx"
    assert output.read_bytes() == before_bytes


def test_xlsm_writeback_is_explicitly_rejected(tmp_path):
    service, run, _ = _prepared_run(tmp_path)
    refund = _approved_refund(service, tmp_path, run)
    candidate = service.create_refund_comment_candidate(run["id"], refund["id"], "math", "Sheet1", "AC5")
    _approve_candidate(service, candidate)
    macro_source = tmp_path / "macro.xlsm"
    copyfile(run["files"]["math"]["path"], macro_source)
    with pytest.raises(ValueError, match="XLSM"):
        write_new_workbook(macro_source, tmp_path / "out.xlsx", [candidate])


def test_candidate_and_audit_history_restore_after_restart(tmp_path):
    service, run, _ = _prepared_run(tmp_path)
    refund = _approved_refund(service, tmp_path, run)
    candidate = service.create_refund_comment_candidate(run["id"], refund["id"], "math", "Sheet1", "AC5")
    _approve_candidate(service, candidate)
    restarted = PayrollService(tmp_path / "app-data")
    restored = restarted.store.get_comment_candidate(candidate["id"])
    assert restored["status"] == "APPROVED"
    assert any(event["event_type"] == "BOUND_TO_RUN" for event in restarted.store.list_business_input_events(refund["id"]))


def test_changed_source_marks_binding_stale(tmp_path):
    service, run, _ = _prepared_run(tmp_path)
    refund = _approved_refund(service, tmp_path, run)
    service.bind_business_input(run["id"], refund["id"])
    Path(refund["source_ref"]).write_text("教师,学生\n教师甲,学生变更\n", encoding="utf-8")
    assert service.get(run["id"])["status"] == "STALE"


def test_teacher_access_token_is_teacher_scoped(tmp_path):
    service = PayrollService(tmp_path / "data")
    alpha, beta = service.create_teacher_access("教师甲"), service.create_teacher_access("教师乙")
    service.teacher_submit(alpha["access_token"], "2026-08", "其他", "仅本人可见")
    assert len(service.teacher_inputs(alpha["access_token"])) == 1
    assert service.teacher_inputs(beta["access_token"]) == []


def test_teacher_cannot_access_admin_api(tmp_path):
    service = PayrollService(tmp_path / "data")
    access = service.create_teacher_access("教师甲")
    static = Path(__file__).parents[1] / "payroll_ui" / "static"
    server = PayrollHttpServer(("127.0.0.1", 0), service, static)
    worker = Thread(target=server.serve_forever, daemon=True); worker.start()
    base = f"http://127.0.0.1:{server.server_port}"
    try:
        request = Request(base + "/api/runs", headers={"X-Teacher-Token": access["access_token"]})
        with pytest.raises(HTTPError) as error:
            urlopen(request)
        assert error.value.code == 403
    finally:
        server.shutdown(); server.server_close(); worker.join()


def test_sanitized_http_e2e_submission_review_refund_preview_writeback(tmp_path):
    fixture_root = tmp_path / "fixture"; fixture_root.mkdir()
    service, run, _ = _prepared_run(fixture_root)
    static = Path(__file__).parents[1] / "payroll_ui" / "static"
    server = PayrollHttpServer(("127.0.0.1", 0), service, static)
    worker = Thread(target=server.serve_forever, daemon=True); worker.start()
    base = f"http://127.0.0.1:{server.server_port}"
    try:
        admin = json.loads(urlopen(base + "/api/bootstrap").read())["token"]
        def post(path: str, payload: dict, headers: dict | None = None):
            request = Request(base + path, data=json.dumps(payload).encode(), method="POST", headers={"Content-Type": "application/json", **(headers or {})})
            return json.loads(urlopen(request).read())
        access = post("/api/teacher-access", {"teacher_id": "教师甲"}, {"X-Payroll-Token": admin})
        submission = post("/api/teacher/inputs", {"period": "2026-08", "item_type": "班课特殊核算", "note": "脱敏测试说明"}, {"X-Teacher-Token": access["access_token"]})
        post(f"/api/business-inputs/{submission['id']}/review", {"action": "START_REVIEW", "reviewer": "主管"}, {"X-Payroll-Token": admin})
        post(f"/api/business-inputs/{submission['id']}/review", {"action": "APPROVE", "reviewer": "主管"}, {"X-Payroll-Token": admin})
        post(f"/api/runs/{run['id']}/business-inputs", {"input_id": submission["id"]}, {"X-Payroll-Token": admin})
        source = _csv(tmp_path / "refund.csv", "教师,学生,说明\n张三,学生甲,脱敏退费\n")
        imported = post("/api/business-inputs/import", {"input_type": "REFUND_RESULT", "period": "2026-08", "path": str(source), "submitted_by": "业务员"}, {"X-Payroll-Token": admin})[0]
        post(f"/api/business-inputs/{imported['id']}/review", {"action": "START_REVIEW", "reviewer": "主管"}, {"X-Payroll-Token": admin})
        post(f"/api/business-inputs/{imported['id']}/review", {"action": "APPROVE", "reviewer": "主管"}, {"X-Payroll-Token": admin})
        post(f"/api/runs/{run['id']}/business-inputs", {"input_id": imported["id"]}, {"X-Payroll-Token": admin})
        candidate = post(f"/api/runs/{run['id']}/comment-candidates/refund", {"input_id": imported["id"], "target_role": "math", "sheet": "Sheet1", "cell": "AC5"}, {"X-Payroll-Token": admin})
        preview = post(f"/api/runs/{run['id']}/comment-candidates/preview", {"candidate_id": candidate["id"], "strategy": "APPEND"}, {"X-Payroll-Token": admin})
        post(f"/api/runs/{run['id']}/comment-candidates/approve", {"candidate_id": candidate["id"], "reviewer": "主管", "preview_token": preview["preview_token"]}, {"X-Payroll-Token": admin})
        output = tmp_path / "output.xlsx"
        result = post(f"/api/runs/{run['id']}/writeback", {"source_workbook": run["files"]["math"]["path"], "candidate_ids": [candidate["id"]], "output_path": str(output), "reviewer": "主管"}, {"X-Payroll-Token": admin})
        assert Path(result["output_path"]).is_file()
        assert load_workbook(output)["Sheet1"]["AC5"].comment is not None
    finally:
        server.shutdown(); server.server_close(); worker.join()
