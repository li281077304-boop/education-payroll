"""Cross-branch integration E2E for the merged payroll workflow.

Every case uses sanitized fixtures: no real teacher, no real workbook.
The chain under test is:

Teacher Submission -> Admin Review -> create resolution -> recompute
-> resolve Business Issue -> comment candidate -> preview -> approve
-> Excel writeback -> audit history
"""
from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from payroll_ui.service import PayrollService
from tests.test_business_inputs_writeback import _approve_candidate, _approved_refund, _csv
from tests.test_payroll_ui_service import _prepared_run, _prepared_run_with_values


def _class_group(service, run):
    return next(item for item in service.check(run["id"])["issue_groups"] if "class_value" in item["affected_fields"])


def _first_course(service, run, group):
    return service.evidence(run["id"], group["id"])["resolution_courses"][0]


def _correction(service, run) -> tuple[dict, dict]:
    group = _class_group(service, run)
    course = _first_course(service, run, group)
    service.create_resolution(
        run["id"], group["id"], "SOURCE_DATA_CORRECTION", course["id"],
        {"field": "grade", "corrected_value": "高三", "reason_code": "GRADE_ROLLOVER_NOT_UPDATED", "reason": "脱敏的年级修正"},
        "审核人", group["fingerprint"],
    )
    return service.store.get(run["id"])["resolutions"][0], group


def test_e2e_source_correction_closes_gap_and_writes_new_workbook(tmp_path):
    """A. SOURCE_DATA_CORRECTION -> EXACT_CAUSE -> issue closed -> note written."""
    service, run, schedule, math, science = _prepared_run_with_values(tmp_path, one_to_one=1.8, class_value=3.24)
    original_source = science.read_bytes()

    resolution, group = _correction(service, run)

    assert resolution["outcome"] == "RESOLVED_BY_SOURCE_CORRECTION"
    recomputation = service.resolution_recomputation(run["id"], resolution["id"])
    assert recomputation["assessment"]["status"] == "EXACT_CAUSE"
    assert recomputation["assessment"]["recomputed_difference"] == pytest.approx(0.0)
    assert not any("class_value" in item["affected_fields"] for item in service.check(run["id"])["issue_groups"])

    candidate = service.create_class_comment_candidate(run["id"], resolution["id"], "science", "Sheet1", "AC5")
    assert candidate["source_resolution_id"] == resolution["id"]
    _approve_candidate(service, candidate)

    output = tmp_path / "系统回填.xlsx"
    result = service.writeback_comments(run["id"], str(science), [candidate["id"]], str(output), "审核人")

    assert result["written"] == 1
    assert science.read_bytes() == original_source
    written = load_workbook(output)
    assert written["Sheet1"]["AC5"].comment is not None
    assert written["Sheet1"]["AE5"].value == load_workbook(science, data_only=False)["Sheet1"]["AE5"].value


def test_e2e_override_stays_scoped_to_its_own_run(tmp_path):
    """B. APPROVED_PAYROLL_OVERRIDE keeps the original fact and never leaks across runs."""
    service, run, *_ = _prepared_run_with_values(tmp_path, one_to_one=1.8, class_value=10.0)
    group = _class_group(service, run)
    course = _first_course(service, run, group)
    service.create_resolution(
        run["id"], group["id"], "APPROVED_PAYROLL_OVERRIDE", course["id"],
        {"approved_treatment": "脱敏批准口径", "approved_contribution": 10.0, "reason": "脱敏验收"},
        "审核人", group["fingerprint"],
    )
    resolution = service.store.get(run["id"])["resolutions"][0]
    assert resolution["outcome"] == "RESOLVED_BY_APPROVED_OVERRIDE"
    assert resolution["run_id"] == run["id"] and resolution["teacher"] == "李四"

    other = service.create("2026-08")
    assert other["resolutions"] == []
    with pytest.raises(ValueError, match="不属于当前工资核算"):
        service.create_class_comment_candidate(other["id"], resolution["id"], "science", "Sheet1", "AC5")


def test_e2e_full_chain_survives_restart_and_does_not_duplicate_on_rerun(tmp_path):
    """C. Restart restores the whole chain; rerun creates nothing twice."""
    service, run, schedule, math, science = _prepared_run_with_values(tmp_path, one_to_one=1.8, class_value=3.24)
    resolution, _ = _correction(service, run)
    refund = _approved_refund(service, tmp_path, run)
    refund_candidate = service.create_refund_comment_candidate(run["id"], refund["id"], "math", "Sheet1", "AC5")
    _approve_candidate(service, refund_candidate)

    reopened = PayrollService(tmp_path / "app-data")
    restored = reopened.get(run["id"])
    assert [item["id"] for item in restored["resolutions"]] == [resolution["id"]]
    assert restored["resolutions"][0]["outcome"] == "RESOLVED_BY_SOURCE_CORRECTION"
    assert [item["id"] for item in reopened.comment_candidates(run["id"])] == [refund_candidate["id"]]
    assert [item["id"] for item in reopened.business_inputs("2026-08")] == [refund["id"]]
    assert len(reopened.store.list_business_input_events(refund["id"])) >= 3

    reopened.check(run["id"])
    again = PayrollService(tmp_path / "app-data").get(run["id"])
    assert len(again["resolutions"]) == 1
    assert len(reopened.comment_candidates(run["id"])) == 1


def test_e2e_source_change_invalidates_resolution_and_its_comment(tmp_path):
    """D. A changed source must propagate NEEDS_RECONFIRMATION and keep history."""
    service, run, schedule, math, science = _prepared_run_with_values(tmp_path, one_to_one=1.8, class_value=3.24)
    resolution, _ = _correction(service, run)
    candidate = service.create_class_comment_candidate(run["id"], resolution["id"], "science", "Sheet1", "AC5")
    assert candidate["status"] == "PROPOSED"

    book = load_workbook(schedule)
    book.active.append(["九年级1v1_学生乙(02-数学)", "一对一课程", "一对一", "测试校区", "2026-08-20 10:00~12:00", "2小时", "已上课", "李四", 1, 1, "学生乙", "01-数学", "线下课", "测试教室", "李四"])
    book.save(schedule)
    service.import_file(run["id"], "schedule", str(schedule))

    stored = service.store.get(run["id"])["resolutions"][0]
    assert stored["status"] == "NEEDS_RECONFIRMATION"
    assert any(event.get("history_event") == "SOURCE_CHANGED" for event in service.store.get(run["id"])["resolution_history"])
    statuses = {item["id"]: item["status"] for item in service.comment_candidates(run["id"])}
    assert statuses[candidate["id"]] == "NEEDS_RECONFIRMATION"


def test_e2e_refund_chain_writes_after_restart(tmp_path):
    """E. Refund: import -> activate -> approve -> bind Run -> note -> writeback."""
    service, run, _ = _prepared_run(tmp_path)
    refund = _approved_refund(service, tmp_path, run)
    candidate = service.create_refund_comment_candidate(run["id"], refund["id"], "math", "Sheet1", "AC5")
    _approve_candidate(service, candidate)

    reopened = PayrollService(tmp_path / "app-data")
    assert reopened.comment_candidates(run["id"])[0]["status"] == "APPROVED"
    output = tmp_path / "refund-out.xlsx"
    reopened.writeback_comments(run["id"], run["files"]["math"]["path"], [candidate["id"]], str(output), "审核员")
    assert load_workbook(output)["Sheet1"]["AC5"].comment is not None

    # An unwritten candidate must lose its basis when the refund source changes.
    second_source = _csv(tmp_path / "refund-b.csv", "教师,学生,退费金额,说明\n张三,学生乙,300,第二条退费\n")
    second = service.import_business_results("REFUND_RESULT", "2026-08", str(second_source), "审核员")[0]
    service.review_business_input(second["id"], "START_REVIEW", "审核员")
    service.review_business_input(second["id"], "APPROVE", "审核员")
    service.bind_business_input(run["id"], second["id"])
    second_candidate = service.create_refund_comment_candidate(run["id"], second["id"], "math", "Sheet1", "AC5")
    second_source.write_text("教师,学生,退费金额,说明\n张三,学生乙,999,来源已更新\n", encoding="utf-8")
    statuses = {item["id"]: item["status"] for item in service.comment_candidates(run["id"])}
    assert statuses[second_candidate["id"]] == "NEEDS_RECONFIRMATION"


def test_e2e_renewal_history_survives_restart_without_guessing_a_bonus(tmp_path):
    """F. Renewal keeps versions; no confirmed rule means no invented amount."""
    service = PayrollService(tmp_path / "data")
    service.import_business_results("RENEWAL_RESULT", "2026-08", str(_csv(tmp_path / "renewal-v1.csv", "教师,学生\n教师甲,学生甲\n")), "业务员")
    service.import_business_results("RENEWAL_RESULT", "2026-08", str(_csv(tmp_path / "renewal-v2.csv", "教师,学生\n教师甲,学生乙\n")), "业务员")

    reopened = PayrollService(tmp_path / "data")
    renewals = [item for item in reopened.business_inputs("2026-08") if item["input_type"] == "RENEWAL_RESULT"]
    assert len(renewals) == 2
    assert all("bonus" not in item["payload"] for item in renewals)


def test_e2e_teacher_scope_and_forged_resolution_are_rejected(tmp_path):
    """G. Teacher isolation, forged ids and cross-run resolutions all fail closed."""
    service, run, schedule, math, science = _prepared_run_with_values(tmp_path, one_to_one=1.8, class_value=3.24)
    resolution, _ = _correction(service, run)

    first = service.create_teacher_access("教师甲")
    second = service.create_teacher_access("教师乙")
    service.teacher_submit(first["access_token"], "2026-08", "班课特殊核算", "脱敏备注")
    service.teacher_submit(second["access_token"], "2026-08", "班课特殊核算", "另一人的备注")
    assert all(item["teacher_id"] == "教师甲" for item in service.teacher_inputs(first["access_token"]))

    other = service.create("2026-08")
    with pytest.raises(ValueError, match="未找到指定记录"):
        service.create_class_comment_candidate(run["id"], "forged-resolution-id", "science", "Sheet1", "AC5")
    with pytest.raises(ValueError, match="不属于当前工资核算"):
        service.create_class_comment_candidate(other["id"], resolution["id"], "science", "Sheet1", "AC5")
    with pytest.raises(ValueError, match="该教师"):
        service.create_class_comment_candidate(run["id"], resolution["id"], "math", "Sheet1", "AC5")
