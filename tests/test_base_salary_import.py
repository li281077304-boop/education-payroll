from __future__ import annotations

from openpyxl import Workbook
from openpyxl.styles import Alignment
from zipfile import ZipFile
import xml.etree.ElementTree as ET

from payroll_ui.base_salary_import import preview_base_salary, source_changed, source_fingerprint


def _workbook(path, *, rows, formula_g: bool = False):
    book = Workbook()
    sheet = book.active
    sheet.title = "历史工资"
    sheet["A1"] = "历史工资资料"
    sheet.merge_cells("B2:C2")
    sheet["A2"] = "教师"
    sheet["B2"] = "基本工资"
    sheet["D2"] = "工龄工资/教师等级"
    sheet["E2"] = "其他待遇"
    sheet["F2"] = "出勤"
    sheet["A3"] = "教师ID"
    sheet["B3"] = "基本工资"
    sheet["C3"] = "岗位津贴"
    sheet["D3"] = "工龄工资"
    sheet["E3"] = "其他待遇"
    sheet["F3"] = "应出勤"
    sheet["G3"] = "实际出勤"
    for row in rows:
        sheet.append(row)
    if formula_g:
        sheet["B4"] = "=10+10"
    for cell in sheet[2]:
        cell.alignment = Alignment(horizontal="center")
    book.save(path)
    return path


def _roster():
    return [
        {"teacher_id": "t-1", "display_name": "教师甲"},
        {"teacher_id": "t-2", "display_name": "教师乙"},
    ]


def test_preview_reads_two_level_merged_headers_and_matches_stable_id(tmp_path):
    source = _workbook(tmp_path / "history.xlsx", rows=[["t-1", 20, 2, 3, 4, 22, 21]])
    preview = preview_base_salary(source, "2026-09", _roster())

    assert preview["can_import"] is True
    assert preview["errors"] == []
    assert preview["matched"] == [{"teacher_id": "t-1", "match_kind": "STABLE_ID", "source_row": 4}]
    assert preview["rows"][0]["fields"] == {"G": 20.0, "H": 2.0, "I": 3.0, "J": 4.0, "K": 22.0, "L": 21.0}
    assert preview["source"]["sha256"]
    assert str(source) not in repr(preview)


def test_preview_allows_partial_match_and_redacts_unmatched_name(tmp_path):
    source = _workbook(tmp_path / "history.xlsx", rows=[
        ["t-1", 20, 2, 3, 4, 22, 21],
        ["missing", 20, 2, 3, 4, 22, 21],
    ])
    preview = preview_base_salary(source, "2026-09", _roster())

    assert len(preview["rows"]) == 1
    assert preview["can_import"] is True
    assert preview["unmatched"] == [{"code": "UNMATCHED_TEACHER_ID", "sheet": "历史工资", "source_row": 5}]
    assert "missing" not in repr(preview)


def test_preview_uses_normalized_exact_name_only_when_stable_id_is_absent(tmp_path):
    source = _workbook(tmp_path / "history.xlsx", rows=[["教师甲", 20, 2, 3, 4, 22, 21]])
    book = Workbook()
    sheet = book.active
    sheet.append(["教师", "基本工资", "岗位津贴", "工龄工资", "其他待遇", "应出勤", "实际出勤"])
    sheet.append([" 教师甲 ", 20, 2, 3, 4, 22, 21])
    book.save(source)

    preview = preview_base_salary(source, "2026-09", _roster())
    assert preview["can_import"] is True
    assert preview["matched"] == [{"teacher_id": "t-1", "match_kind": "NORMALIZED_EXACT_NAME", "source_row": 2}]


def test_preview_blocks_ambiguous_normalized_teacher_name(tmp_path):
    source = _workbook(tmp_path / "history.xlsx", rows=[["教师甲", 20, 2, 3, 4, 22, 21]])
    book = Workbook()
    sheet = book.active
    sheet.append(["教师", "基本工资", "岗位津贴", "工龄工资", "其他待遇", "应出勤", "实际出勤"])
    sheet.append(["教师甲", 20, 2, 3, 4, 22, 21])
    book.save(source)

    preview = preview_base_salary(source, "2026-09", [
        {"teacher_id": "t-1", "display_name": "教师甲"},
        {"teacher_id": "t-2", "display_name": "教师甲"},
    ])
    assert preview["can_import"] is False
    assert preview["conflicts"] == [{"code": "AMBIGUOUS_TEACHER_NAME", "sheet": "Sheet", "source_row": 2}]


def test_preview_blocks_duplicate_source_teacher(tmp_path):
    source = _workbook(tmp_path / "history.xlsx", rows=[
        ["t-1", 20, 2, 3, 4, 22, 21],
        ["t-1", 20, 2, 3, 4, 22, 21],
    ])
    preview = preview_base_salary(source, "2026-09", _roster())

    assert preview["can_import"] is False
    assert preview["conflicts"] == [{"code": "DUPLICATE_SOURCE_TEACHER", "sheet": "历史工资", "source_row": 5}]


def test_preview_fails_closed_for_missing_g_to_l_column(tmp_path):
    source = _workbook(tmp_path / "history.xlsx", rows=[["t-1", 20, 2, 3, 4, 22, 21]])
    book = Workbook()
    sheet = book.active
    sheet.append(["教师ID", "基本工资", "岗位津贴", "工龄工资", "其他待遇", "应出勤"])
    sheet.append(["t-1", 20, 2, 3, 4, 22])
    book.save(source)

    preview = preview_base_salary(source, "2026-09", _roster())
    assert preview["can_import"] is False
    assert preview["errors"] == [{"code": "MISSING_REQUIRED_COLUMNS"}]


def test_preview_fails_closed_when_formula_has_no_cached_value(tmp_path):
    source = _workbook(tmp_path / "history.xlsx", rows=[["t-1", 20, 2, 3, 4, 22, 21]], formula_g=True)
    preview = preview_base_salary(source, "2026-09", _roster())

    assert preview["can_import"] is False
    assert {issue["code"] for issue in preview["errors"]} == {"FORMULA_CACHE_MISSING"}
    assert preview["errors"][0]["field"] == "G"


def test_valid_rows_can_import_while_incomplete_teacher_remains_unmatched(tmp_path):
    source = _workbook(tmp_path / "history.xlsx", rows=[
        ["t-1", 20, 2, 3, 4, 22, 21],
        ["t-2", 25, None, 3, 4, 22, 21],
    ])

    preview = preview_base_salary(source, "2026-09", _roster())

    assert preview["can_import"] is True
    assert len(preview["rows"]) == 1
    assert preview["errors"] == [{"code": "INVALID_REQUIRED_VALUE", "sheet": "历史工资", "source_row": 5, "field": "H"}]


def test_source_fingerprint_change_is_detected(tmp_path):
    source = _workbook(tmp_path / "history.xlsx", rows=[["t-1", 20, 2, 3, 4, 22, 21]])
    before = source_fingerprint(source)
    source = _workbook(source, rows=[["t-1", 21, 2, 3, 4, 22, 21]])
    after = source_fingerprint(source)

    assert source_changed(before, after) is True
    assert source_changed(after, dict(after)) is False


def _workbook_with_cached_m(path, cached_m):
    book = Workbook()
    sheet = book.active
    sheet.title = "历史工资"
    sheet["C2"] = "教师"
    sheet["G2"] = "出勤及基本工资"
    sheet.merge_cells("G2:M2")
    for column, title in {"C": "姓名", "G": "基本工资", "H": "岗位津贴", "I": "工龄工资", "J": "其他待遇", "K": "应出勤", "L": "实际出勤", "M": "实际基本工资"}.items():
        sheet[f"{column}3"] = title
    sheet["C4"] = "教师甲"
    for column, value in {"G": 20, "I": 3, "K": 22, "L": 21}.items():
        sheet[f"{column}4"] = value
    sheet["M4"] = "=(G4+H4+I4+J4)/K4*L4"
    book.save(path)
    with ZipFile(path) as source:
        files = [(item, source.read(item.filename)) for item in source.infolist()]
    ns = {"x": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    with ZipFile(path, "w") as output:
        for item, contents in files:
            if item.filename == "xl/worksheets/sheet1.xml":
                root = ET.fromstring(contents)
                cell = next(node for node in root.findall(".//x:c", ns) if node.attrib.get("r") == "M4")
                cached = cell.find("x:v", ns)
                if cached is None:
                    cached = ET.SubElement(cell, f"{{{ns['x']}}}v")
                cached.text = str(cached_m)
                contents = ET.tostring(root, encoding="utf-8")
            output.writestr(item, contents)
    return path


def test_blank_allowances_become_zero_only_when_source_m_formula_confirms_them(tmp_path):
    source = _workbook_with_cached_m(tmp_path / "history.xlsx", (20 + 3) / 22 * 21)

    preview = preview_base_salary(source, "2026-09", [{"teacher_id": "t-1", "display_name": "教师甲"}])

    assert preview["can_import"] is True
    assert preview["rows"][0]["fields"]["H"] == 0
    assert preview["rows"][0]["fields"]["J"] == 0
    assert preview["formula_verified_zero_count"] == 2
    assert preview["rows"][0]["provenance"]["formula_verified_zero_fields"] == ["H", "J"]

    unverified = _workbook_with_cached_m(tmp_path / "unverified.xlsx", 999)
    blocked = preview_base_salary(unverified, "2026-09", [{"teacher_id": "t-1", "display_name": "教师甲"}])
    assert blocked["can_import"] is False
    assert {item["field"] for item in blocked["errors"]} == {"H", "J"}
