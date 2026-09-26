"""Teacher identity, employment type and the 支持部 workbook as one source.

These cover the second round of real UAT: why the match target is the size it
is, why a part-time teacher must not be asked for a full-time base salary, and
why one reviewed 支持部 workbook has to be enough for several fields.
"""
from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment

from payroll_core.adapters.personnel import DEFAULT_PART_TIME_RATES
from payroll_core.employment import FULL_TIME, PART_TIME, UNKNOWN
from payroll_core.final_fields import DISPLAY_LABELS, FINAL_FIELD_CODES, display_label
from payroll_core.payroll_generation import generated_from_calculation
from payroll_ui.service import PayrollService

# The 支持部 workbook: title, month, parent headers, child headers, data.
FIELD_COLUMNS = {
    "序号": 1, "科组": 2, "姓名": 3, "邮箱": 4, "入职日期": 5, "教师级别": 6,
    "基本工资": 7, "岗位津贴": 8, "工龄工资/教师等级": 9, "其他待遇": 10,
    "应出勤": 11, "实际出勤": 12, "实际基本工资": 13,
    "该档每小时金额": 31, "总课时费": 32,
    "房租": 41, "社保": 42, "工装费": 43, "内部推荐奖金": 44,
    "月度激励": 45, "补发工资": 46, "考勤罚款": 47, "总工资数": 48,
    "1对1课时": 34, "班课&1对2领航伴课次": 35, "小班领航伴学课次": 36,
}
PARENT_HEADERS = {"序号": 1, "科组": 2, "姓名": 3, "邮箱": 4, "入职日期": 5, "教师级别": 6,
                  "基本工资": 7, "该档每小时金额": 31, "总课时费": 32, "其他": 41, "总工资数": 48}


def _support_workbook(path: Path, rows: list[dict]) -> Path:
    book = Workbook()
    sheet = book.active
    sheet.title = "教学部"
    sheet.cell(1, 1).value = "教学部薪资表"
    sheet.cell(2, 3).value = "中心："
    sheet.cell(2, 5).value = "宣城二校"
    sheet.cell(2, 7).value = "月份："
    sheet.cell(2, 8).value = "2026-08"
    for label, column in PARENT_HEADERS.items():
        sheet.cell(3, column).value = label
    for label, column in FIELD_COLUMNS.items():
        if label not in PARENT_HEADERS:
            sheet.cell(4, column).value = label
    for offset, item in enumerate(rows):
        row = 5 + offset
        for label, column in FIELD_COLUMNS.items():
            if label in item:
                sheet.cell(row, column).value = item[label]
    book.save(path)
    return path


def _teacher(name: str, **overrides) -> dict:
    item = {
        "序号": 1, "科组": "数学组", "姓名": name, "邮箱": f"dc_{name}@example.com",
        "入职日期": "2020.01.01", "教师级别": "TR/四星级",
        "基本工资": 2400, "岗位津贴": 0, "工龄工资/教师等级": 1000, "其他待遇": 300,
        "应出勤": 26, "实际出勤": 26,
    }
    item.update(overrides)
    return item


def _run(service: PayrollService, teachers: list[str], period: str = "2026-08", tmp_path: Path | None = None) -> dict:
    created = service.create(period, "GENERATE")
    stored = service.store.get(created["id"])
    stored["core_calculation"] = {
        "period": period,
        "rows": [
            {"teacher": name, "teacher_id": name, "employment_type": "FULL_TIME", "fields": _core_fields()}
            for name in teachers
        ],
    }
    base = tmp_path or Path(stored.get("package_root") or ".")
    schedule = base / "排课.xls"
    math = base / "数学组提交表.xlsx"
    for target in (schedule, math):
        if not target.exists():
            target.write_bytes(b"placeholder")

    def material(target: Path, **extra) -> dict:
        stat = target.stat()
        return {"name": target.name, "path": str(target), "size": stat.st_size,
                "mtime_ns": stat.st_mtime_ns, "sha256": _sha256(target), **extra}

    stored["files"] = {
        "schedule": material(schedule, teachers=51, records=2385),
        "math": material(math, label="数学组提交表", teachers=len(teachers), records=len(teachers)),
    }
    service.store.save(stored)
    return stored


def _core_fields(after: float = 120.0) -> dict:
    return {
        "AA": {"value": 1, "state": "DETERMINED"}, "AC": {"value": 2, "state": "DETERMINED"},
        "AD": {"value": 33, "state": "DETERMINED"}, "AE": {"value": 40, "state": "DETERMINED"},
        "AF": {"value": after, "state": "DETERMINED"}, "PART_TIME": {"value": None, "state": "NOT_APPLICABLE"},
    }


# --------------------------------------------------------------- 1. roster 溯源
def test_the_match_target_states_where_it_comes_from(tmp_path):
    service = PayrollService(tmp_path / "data")
    run = _run(service, ["教师甲", "教师乙"], tmp_path=tmp_path)

    facts = service._roster_facts(service.store.get(run["id"]))
    assert facts["member_count"] == 2
    assert facts["origin"] == "core_calculation.rows"
    assert "学科组提交表" in facts["origin_label"]
    assert facts["schedule_teacher_count"] == 51
    assert "不是排课表全部教师" in facts["note"]
    # The screen must be able to answer the question without reading code.
    payload = service.get(run["id"])
    assert payload["roster_facts"]["member_count"] == 2


# ------------------------------------------------- 2. 三个集合 + 身份待确认
def test_support_preview_separates_the_four_teacher_sets(tmp_path):
    service = PayrollService(tmp_path / "data")
    run = _run(service, ["教师甲", "教师乙", "教师丙"], tmp_path=tmp_path)
    source = _support_workbook(tmp_path / "支持部.xlsx", [
        _teacher("教师甲"), _teacher("教师乙"), _teacher("只有历史的丁"),
    ])

    preview = service.preview_support_salary_import(run["id"], str(source))
    counts = preview["counts"]
    assert counts["matched"] == 2                     # A 两边都有
    assert counts["month_without_history"] == 1       # B 本月缺历史资料
    assert counts["history_only"] == 1                # C 历史表有、本月未核算
    assert counts["identity_required"] == 0           # D 证据不足
    assert preview["month_without_history"][0]["teacher"] == "教师丙"
    assert preview["history_only"][0]["teacher"] == "只有历史的丁"
    # C is a scope fact, never reported as an unrecognisable person.
    assert preview["history_only"][0]["reason"] == "NOT_IN_CURRENT_CALCULATION"


def test_the_base_salary_preview_reports_the_same_sets(tmp_path):
    service = PayrollService(tmp_path / "data")
    run = _run(service, ["教师甲", "教师乙"], tmp_path=tmp_path)
    source = _support_workbook(tmp_path / "支持部.xlsx", [_teacher("教师甲"), _teacher("历史教师")])

    preview = service.preview_base_salary_import(run["id"], str(source))
    assert preview["counts"] == {
        "source_teachers": 2, "current_run_teachers": 2, "matched": 1,
        "month_without_history": 1, "history_only": 1, "identity_required": 0,
    }
    assert preview["roster"]["member_count"] == 2


# ------------------------------------------------------------ 3. 重名 fail closed
def test_duplicate_teacher_names_fail_closed(tmp_path):
    service = PayrollService(tmp_path / "data")
    run = _run(service, ["教师甲", "教师乙"], tmp_path=tmp_path)
    stored = service.store.get(run["id"])
    stored["core_calculation"]["rows"].append({"teacher": "教师甲", "teacher_id": "另一个教师甲"})
    service.store.save(stored)

    source = _support_workbook(tmp_path / "支持部.xlsx", [_teacher("教师甲")])
    preview = service.preview_support_salary_import(run["id"], str(source))
    assert preview["counts"]["identity_required"] == 1
    assert preview["identity_required"][0]["reason"] == "AMBIGUOUS_TEACHER_NAME"
    assert preview["counts"]["matched"] == 0


# -------------------------------------------------------------- 4. 兼职语义
def test_part_time_teacher_is_not_asked_for_a_full_time_base_salary(tmp_path):
    service = PayrollService(tmp_path / "data")
    run = _run(service, ["全职甲", "兼职乙"], tmp_path=tmp_path)
    service.save_employment_profile("全职甲", FULL_TIME, "核算负责人")
    service.save_employment_profile("兼职乙", PART_TIME, "核算负责人", reason="用户确认为兼职")

    overview = service.employment_overview(service.store.get(run["id"]))
    assert overview["counts"]["part_time"] == 1
    assert overview["counts"]["full_time"] == 1
    assert overview["counts"]["needs_confirmation"] == 0

    generated = service._generated_from_checked(service.store.get(run["id"]))
    rows = {row.teacher: row for row in generated.rows}
    assert rows["兼职乙"].final_fields["M"]["state"] == "NOT_APPLICABLE"
    assert rows["兼职乙"].final_fields["M"]["value"] is None
    assert "兼职" in rows["兼职乙"].final_fields["M"]["reason"]
    assert not any("G" in blocker and "缺少" in blocker for blocker in rows["兼职乙"].blockers)


def test_part_time_employment_is_not_inferred_from_legacy_default_rate_names(tmp_path):
    """Historical example prices are not a current employment authority."""
    assert DEFAULT_PART_TIME_RATES == {}
    service = PayrollService(tmp_path / "data")
    run = _run(service, ["测试教师甲"], tmp_path=tmp_path)
    overview = service.employment_overview(service.store.get(run["id"]))
    assert overview["teachers"][0]["employment_type"] == UNKNOWN
    assert overview["teachers"][0]["needs_confirmation"] is True


def test_a_support_entry_is_document_evidence_of_full_time(tmp_path):
    service = PayrollService(tmp_path / "data")
    run = _run(service, ["教师甲"], tmp_path=tmp_path)
    source = _support_workbook(tmp_path / "支持部.xlsx", [_teacher("教师甲")])
    service.import_support_department(run["id"], str(source), _sha256(source), "核算负责人")

    overview = service.employment_overview(service.store.get(run["id"]))
    item = next(entry for entry in overview["teachers"] if entry["teacher"] == "教师甲")
    assert item["employment_type"] == FULL_TIME
    assert item["source"] == "EMPLOYMENT_FROM_DOCUMENT"
    assert item["needs_confirmation"] is False


def _sha256(path: Path) -> str:
    import hashlib
    return hashlib.sha256(path.read_bytes()).hexdigest()


# ------------------------------------- 5/6. 基本工资确认 → M → AV
def test_confirmed_base_salary_makes_m_determined_and_feeds_av(tmp_path):
    service = PayrollService(tmp_path / "data")
    run = _run(service, ["教师甲"], tmp_path=tmp_path)
    source = _support_workbook(tmp_path / "支持部.xlsx", [_teacher("教师甲")])
    result = service.import_support_department(run["id"], str(source), _sha256(source), "核算负责人")

    stored = service.store.get(run["id"])
    assert result["imported_base_salary"] == 1
    assert stored["base_salary_inputs"]["教师甲"]["fields"]["G"]["value"] == 2400
    assert stored["base_salary_inputs"]["教师甲"]["m"]["state"] == "DETERMINED"
    assert stored["base_salary_input_snapshot"]["source"]

    generated = service._generated_from_checked(stored)
    row = generated.rows[0]
    assert row.final_fields["M"]["state"] == "DETERMINED"
    assert row.final_fields["M"]["value"] == pytest.approx((2400 + 0 + 1000 + 300) / 26 * 26)


# ------------------------------------- 7/8. 续费确认 → AH/AI/AJ → AK → AV
def test_renewal_confirmation_feeds_ah_ai_aj_ak_and_av(tmp_path):
    service = PayrollService(tmp_path / "data")
    run = _run(service, ["教师甲"], tmp_path=tmp_path)
    stored = service.store.get(run["id"])
    stored["run_renewal_result_snapshot"] = _renewal_snapshot("教师甲", 5, 2, 1)
    service.store.save(stored)

    generated = generated_from_calculation(
        {"period": "2026-08", "rows": [{"teacher": "教师甲", "employment_type": "FULL_TIME", "fields": _core_fields()}]},
        base_salary_inputs={"教师甲": {"fields": {code: {"value": 1} for code in ("G", "H", "I", "J", "K", "L")}, "source": "测试"}},
        renewal_snapshot=service.store.get(run["id"])["run_renewal_result_snapshot"],
    )
    fields = generated.rows[0].final_fields
    assert fields["AH"]["value"] == 5
    assert fields["AI"]["value"] == 2
    assert fields["AJ"]["value"] == 1
    assert fields["AK"]["state"] == "DETERMINED"
    assert fields["AK"]["value"] == pytest.approx(5 + 2 * 1.5 + 1 * 0.75)
    # AV still needs the other components, but AK is no longer among them.
    reason = fields["AV"].get("reason") or ""
    undetermined = reason.split("未确定组件：")[-1] if "未确定组件：" in reason else ""
    assert undetermined and "AK" not in undetermined


def _renewal_snapshot(teacher: str, ah: float, ai: float, aj: float) -> dict:
    def field(code, value):
        return {"value": value, "state": "DETERMINED", "reason": code, "evidence": [{"kind": "TEST"}]}
    return {"version": "RUN_RENEWAL_RESULT_SNAPSHOT/v1", "period_label": "2026-08", "entries": {
        teacher: {"teacher_id": teacher, "display_name": teacher, "AH": field("AH", ah), "AI": field("AI", ai), "AJ": field("AJ", aj)},
    }}


# ------------------------------------------------- 9/10. 退费 → AN → AV
def test_approved_refund_result_feeds_an_and_av(tmp_path):
    service = PayrollService(tmp_path / "data")
    run = _run(service, ["教师甲"], tmp_path=tmp_path)
    source = _refund_workbook(tmp_path / "退费.xlsx")

    imported = service.import_business_results("REFUND_RESULT", "2026-08", str(source), "客服部")
    assert len(imported) == 1
    service.review_business_input(imported[0]["id"], "START_REVIEW", "核算负责人")
    service.review_business_input(imported[0]["id"], "APPROVE", "核算负责人")
    service.bind_business_input(run["id"], imported[0]["id"])

    stored = service.store.get(run["id"])
    generated = service._generated_from_checked(stored)
    assert generated.rows[0].final_fields["AN"]["state"] == "DETERMINED"
    assert generated.rows[0].final_fields["AN"]["value"] == pytest.approx(300)


def _refund_workbook(path: Path) -> Path:
    book = Workbook()
    sheet = book.active
    sheet.title = "8月份 "
    sheet.append(["周", "退费校区", "新签/续费", "姓名", "年级", "消耗课时", "退费科目", "科目类型",
                  "退费课时", "退费总金额", "是否为转校生", "班主任", "学科教师", "教育顾问", "备注", "教师", "人头", "业绩"])
    sheet.append([1, "宣城二校", "续费", "学员甲", "初二", 10, "数学", "班课", 5, 500, "否", "班主任甲", "教师甲", "顾问甲", "", "教师甲", 100, 200])
    book.save(path)
    return path


# ------------------------------- 11/12/13. 支持部字段：一次文件，多字段
def test_one_support_file_supplies_identity_and_support_pay_items(tmp_path):
    service = PayrollService(tmp_path / "data")
    run = _run(service, ["教师甲"], tmp_path=tmp_path)
    stale = service.store.get(run["id"])
    stale["base_salary_deferred"] = True
    stale["base_salary_deferred_by"] = "核算负责人"
    stale["base_salary_deferred_at"] = "2026-09-01T00:00:00+00:00"
    service.store.save(stale)
    source = _support_workbook(tmp_path / "支持部.xlsx", [
        _teacher("教师甲", 社保=-452.66, 补发工资=1006.92, 房租=500),
    ])
    service.import_support_department(run["id"], str(source), _sha256(source), "核算负责人")

    stored = service.store.get(run["id"])
    snapshot = stored["support_department_snapshot"]
    assert stored["base_salary_deferred"] is False
    assert stored["base_salary_deferred_by"] == ""
    assert stored["base_salary_deferred_at"] is None
    entry = snapshot["entries"]["教师甲"]
    # B 科组 / D 邮箱 / E 入职日期 / F 教师级别 are inherited verbatim.
    assert entry["identity"]["group"] == "数学组"
    assert entry["identity"]["email"] == "dc_教师甲@example.com"
    assert entry["identity"]["hire_date"] == "2020.01.01"
    assert entry["identity"]["teacher_level"] == "TR/四星级"
    # Support-only pay items, and the documented mapping table.
    assert entry["items"]["AP"] == -452.66
    assert entry["items"]["AT"] == 1006.92
    assert entry["items"]["AO"] == 500

    generated = service._generated_from_checked(stored)
    fields = generated.rows[0].final_fields
    assert fields["AP"]["value"] == -452.66
    assert fields["AT"]["value"] == 1006.92
    assert fields["AO"]["value"] == 500
    assert any(item["kind"] == "SUPPORT_DEPARTMENT_SNAPSHOT" for item in fields["AP"]["evidence"])
    # A field the material has as a column but leaves blank is a stated zero.
    assert fields["AU"]["value"] == 0
    assert "空白" in fields["AU"]["reason"]

    mapping = {item["final_field"]: item for item in snapshot["field_map"]}
    assert mapping["AP"]["business_name"] == "社保"
    assert mapping["AP"]["inherited"] is True
    assert mapping["AT"]["business_name"] == "补发工资"
    # 月度激励 belongs to the subject group, not to the support department.
    assert mapping["AS"]["source"] == "SUBJECT_GROUP_MONTHLY_INCENTIVE"
    assert mapping["AS"]["inherited"] is False
    # Computed fields are never copied from the source workbook.
    for code in ("AA", "AC", "AD", "AE", "AF", "AK", "AN", "AV"):
        assert mapping[code]["source"] == "CURRENT_RUN_CALCULATION"


def test_the_same_support_file_is_not_uploaded_twice(tmp_path):
    service = PayrollService(tmp_path / "data")
    run = _run(service, ["教师甲"], tmp_path=tmp_path)
    source = _support_workbook(tmp_path / "支持部.xlsx", [_teacher("教师甲", 社保=-300)])

    result = service.import_support_department(run["id"], str(source), _sha256(source), "核算负责人")
    # One confirmation produced both authorities.
    assert result["imported_base_salary"] == 1
    assert result["entries"] == 1
    stored = service.store.get(run["id"])
    assert stored["base_salary_inputs"]["教师甲"]["fields"]["G"]["value"] == 2400
    assert stored["support_department_snapshot"]["entries"]["教师甲"]["items"]["AP"] == -300


def test_support_annotation_refresh_preserves_confirmed_payroll_values(tmp_path):
    service = PayrollService(tmp_path / "data")
    run = _run(service, ["教师甲"], tmp_path=tmp_path)
    source = _support_workbook(tmp_path / "支持部.xlsx", [_teacher("教师甲", 岗位津贴=135)])
    book = load_workbook(source)
    book["教学部"]["H5"].comment = Comment("本项为当前岗位津贴来源备注。", "支持部")
    book.save(source)
    digest = _sha256(source)
    service.import_support_department(run["id"], str(source), digest, "核算负责人")

    stored = service.store.get(run["id"])
    snapshot = stored["support_department_snapshot"]
    before = {
        teacher: {key: snapshot["entries"][teacher].get(key) for key in ("identity", "items", "base_salary")}
        for teacher in snapshot["entries"]
    }
    confirmer, created_at = snapshot["confirmed_by"], snapshot["created_at"]
    for entry in snapshot["entries"].values():
        entry.pop("annotations", None)
    snapshot.pop("comment_count", None)
    snapshot.pop("comment_fields", None)
    service.store.save(stored)

    refreshed = service.refresh_support_annotations_from_source(run["id"], str(source), digest)
    after = service.store.get(run["id"])["support_department_snapshot"]
    entry = after["entries"]["教师甲"]
    assert refreshed["comment_count"] == 1
    assert refreshed["comment_fields"] == ["H"]
    assert entry["annotations"][0]["field_code"] == "H"
    assert entry["annotations"][0]["text"] == "本项为当前岗位津贴来源备注。"
    assert {teacher: {key: after["entries"][teacher].get(key) for key in ("identity", "items", "base_salary")} for teacher in after["entries"]} == before
    assert after["confirmed_by"] == confirmer
    assert after["created_at"] == created_at


def test_support_comments_follow_mapped_fields_including_renewal_columns(tmp_path):
    from openpyxl import load_workbook
    from openpyxl.comments import Comment
    from payroll_core.excel.standard_payroll_render import render_generated_payroll
    from payroll_ui.support_department import preview_support_department
    from test_standard_payroll_output import _generated, _sanitized_template

    source = _support_workbook(tmp_path / "支持部.xlsx", [
        _teacher("教师甲", **{"社保": -100, "1对1课时": 3, "班课&1对2领航伴课次": 2, "小班领航伴学课次": 1}),
    ])
    source_book = load_workbook(source)
    source_sheet = source_book.active
    for col, text in ((8, "岗位津贴源批注\n"), (34, "支持部 AH 原批注\n"), (35, "支持部 AI 原批注"), (36, "支持部 AJ 原批注"), (42, "社保源批注")):
        source_sheet.cell(5, col).comment = Comment(text, "测试")
    source_book.save(source)

    preview = preview_support_department(source, "2026-08", [{"teacher_id": "teacher-1", "display_name": "教师甲"}])
    annotations = preview["rows"][0]["annotations"]
    assert {item["field_code"] for item in annotations} == {"H", "AH", "AI", "AJ", "AP"}
    snapshot = {"entries": {"teacher-1": {
        "display_name": "教师甲", "identity": {},
        "items": {"AP": -100}, "base_salary": {}, "annotations": annotations,
    }}}
    template = _sanitized_template(tmp_path)
    path = render_generated_payroll(
        _generated(business_inputs=[{
            "id": "renewal", "teacher_id": "教师甲", "period": "2026-08",
            "input_type": "RENEWAL_RESULT", "status": "APPROVED",
            "payload": {"one_to_one_hours": 3, "class_hours": 2, "mentor_hours": 1},
        }]),
        tmp_path / "mapped-comments.xlsx", template_path=template, support_snapshot=snapshot,
    )
    sheet = load_workbook(path, data_only=False).worksheets[0]
    assert sheet["H5"].comment.text == "岗位津贴源批注\n"
    assert sheet["AH5"].comment.text == "支持部 AH 原批注\n"
    assert sheet["AI5"].comment.text == "支持部 AI 原批注"
    assert sheet["AJ5"].comment.text == "支持部 AJ 原批注"
    assert sheet["AP5"].comment.text == "社保源批注"
    assert sheet["AH5"].value == 3
    assert sheet["AI5"].value == 2
    assert sheet["AJ5"].value == 1


def test_support_import_rejects_a_changed_file(tmp_path):
    service = PayrollService(tmp_path / "data")
    run = _run(service, ["教师甲"], tmp_path=tmp_path)
    source = _support_workbook(tmp_path / "支持部.xlsx", [_teacher("教师甲")])
    digest = _sha256(source)
    _support_workbook(source, [_teacher("教师甲", 基本工资=9999)])

    with pytest.raises(ValueError, match="已变化"):
        service.import_support_department(run["id"], str(source), digest, "核算负责人")


# ------------------------------------------- 14. preview/export 同一 canonical
def test_the_export_renders_the_same_canonical_final_fields(tmp_path):
    """The workbook renders the canonical final fields; it never re-derives them."""
    from payroll_core.excel.standard_payroll_render import render_generated_payroll

    service = PayrollService(tmp_path / "data")
    run = _run(service, ["教师甲"], tmp_path=tmp_path)
    source = _support_workbook(tmp_path / "支持部.xlsx", [_teacher("教师甲", 社保=-452.66)])
    service.import_support_department(run["id"], str(source), _sha256(source), "核算负责人")

    stored = service.store.get(run["id"])
    canonical = service._generated_from_checked(stored)
    output = tmp_path / "导出.xlsx"
    render_generated_payroll(canonical, output, template_path=_payroll_template(tmp_path / "模板.xlsx"),
                             support_snapshot=stored["support_department_snapshot"])

    # The renderer writes FINAL_FIELD_CODES starting at column 33, and the
    # canonical object is the only source of those values.
    sheet = load_workbook(output).worksheets[0]
    row = 5
    for index, code in enumerate(FINAL_FIELD_CODES):
        expected = canonical.rows[0].final_fields[code]["value"]
        written = sheet.cell(row, 33 + index).value
        assert canonical.rows[0].teacher == "教师甲"
        if expected is None:
            assert written is None, code
        else:
            assert float(written) == pytest.approx(float(expected)), code
    # M is written as the canonical formula only when it is DETERMINED.
    assert canonical.rows[0].final_fields["M"]["state"] in {"DETERMINED", "NOT_APPLICABLE"}


# ------------------------------------------------- 15. 星级导入不改 AE 算法
def test_rating_from_a_workbook_does_not_touch_the_ae_algorithm(tmp_path):
    service = PayrollService(tmp_path / "data")
    run = _run(service, ["教师甲", "教师乙"], tmp_path=tmp_path)
    source = _support_workbook(tmp_path / "支持部.xlsx", [
        _teacher("教师甲", 教师级别="TR/四星级"),
        _teacher("教师乙", 教师级别="TR/六星级"),
    ])
    preview = service.preview_rating_from_workbook(run["id"], str(source))
    assert preview["counts"]["candidates"] == 2
    assert [item["rating"] for item in preview["candidates"]] == [4, 6]

    before = len(service.rating_versions())
    result = service.save_rating_from_workbook(run["id"], str(source), preview["source"]["sha256"], "核算负责人")
    assert result["created"]["ratings"] == 2
    assert len(service.rating_versions()) == before + 1
    version = next(item for item in service.rating_versions() if item["id"] == result["versions"][0]["id"])
    assert version["ratings"][0]["rating"] == 4
    assert "教师级别" in version["source"]


def test_rating_can_be_derived_from_the_bound_support_snapshot(tmp_path):
    service = PayrollService(tmp_path / "data")
    run = _run(service, ["教师甲"], tmp_path=tmp_path)
    source = _support_workbook(tmp_path / "支持部.xlsx", [_teacher("教师甲", 教师级别="TRST/三星级")])
    service.import_support_department(run["id"], str(source), _sha256(source), "核算负责人")

    before = len(service.rating_versions())
    result = service.derive_rating_from_support(run["id"], "核算负责人")
    assert result["created"]["ratings"] == 1
    assert len(service.rating_versions()) == before + 1


# ------------------------------------------------ 16. 每个字段都有显示名称
def test_every_final_field_has_a_code_plus_chinese_name():
    for code in ("M",) + FINAL_FIELD_CODES:
        label = display_label(code)
        assert label.startswith(code), code
        assert "名称待确认" not in label, code
        assert len(label) > len(code) + 1, code
    # Unknown fields say so instead of being invented.
    assert display_label("ZZ") == "ZZ｜名称待确认"
    # The names are the company's own headers.
    assert DISPLAY_LABELS["AO"] == "AO 房租"
    assert DISPLAY_LABELS["AU"] == "AU 考勤罚款"
    assert DISPLAY_LABELS["AV"] == "AV 总工资数"


# ------------------------------------------------------- 17. AF 例外显示真实值
def test_af_exceptions_are_shown_with_their_real_values():
    source = Path("payroll_ui/static/app.js").read_text(encoding="utf-8")
    start = source.index("function afExceptionsMarkup()")
    body = source[start:start + 1400]
    assert "af_policy_confirmation" in body
    for column in ("义务课时", "是否扣减", "原因"):
        assert column in body
    assert "obligation_hours" in body and "deduction_enabled" in body and "reason" in body
    # Rendered open, not hidden behind a second click.
    assert 'class="preview-details" open' in body


# ------------------------------------------------------------- 18. anti-copy
def test_the_company_template_never_supplies_a_payroll_value(tmp_path):
    """A template with example rows must not leak its numbers into a result."""
    from payroll_core.excel.standard_payroll_render import render_generated_payroll

    service = PayrollService(tmp_path / "data")
    run = _run(service, ["教师甲"], tmp_path=tmp_path)
    source = _support_workbook(tmp_path / "支持部.xlsx", [_teacher("教师甲")])
    service.import_support_department(run["id"], str(source), _sha256(source), "核算负责人")
    stored = service.store.get(run["id"])
    generated = service._generated_from_checked(stored)

    template = _payroll_template(tmp_path / "模板.xlsx")
    output = tmp_path / "带模板导出.xlsx"
    render_generated_payroll(generated, output, template_path=template, support_snapshot=stored["support_department_snapshot"])

    sheet = load_workbook(output).worksheets[0]
    teachers = [str(sheet.cell(row, 3).value or "") for row in range(5, sheet.max_row + 1)]
    assert "示例教师" not in teachers
    assert "教师甲" in teachers
    values = [sheet.cell(row, column).value for row in range(5, sheet.max_row + 1) for column in range(1, sheet.max_column + 1)]
    assert 99999 not in values
    assert -99999 not in values
    # The template's own month must not survive either.
    assert sheet.cell(2, 8).value == "2026-08"
    # B / D / E / F come from the reviewed support workbook.
    row = 5 + teachers.index("教师甲")
    assert sheet.cell(row, 2).value == "数学组"
    assert sheet.cell(row, 4).value == "dc_教师甲@example.com"
    assert sheet.cell(row, 5).value == "2020.01.01"
    assert sheet.cell(row, 6).value == "TR/四星级"


def _payroll_template(path: Path) -> Path:
    """A company-template shaped workbook carrying stale example values."""
    book = Workbook()
    sheet = book.active
    sheet.title = "Sheet1"
    sheet.cell(1, 1).value = "教学部薪资表"
    sheet.cell(2, 3).value = "中心："
    sheet.cell(2, 7).value = "月份："
    sheet.cell(2, 8).value = "2026-07"
    headers = {2: "科组", 3: "姓名", 4: "邮箱", 5: "入职日期", 6: "教师级别",
               7: "基本工资", 8: "岗位津贴", 9: "工龄工资/教师等级", 10: "其他待遇",
               11: "应出勤", 12: "实际出勤", 13: "实际基本工资", 27: "折算小时数",
               29: "班课折算小时数", 30: "最终授课小时数据", 31: "该档每小时金额", 32: "总课时费",
               33: "领航伴学课时费", 41: "房租", 42: "社保", 43: "工装费", 44: "内部推荐奖金",
               45: "月度激励", 46: "补发工资", 47: "考勤罚款", 48: "总工资数", 49: "备注"}
    for column, label in headers.items():
        sheet.cell(4, column).value = label
    sheet.cell(3, 48).value = "总工资数"
    sheet.cell(5, 3).value = "示例教师"
    sheet.cell(5, 6).value = "TR/一星级"
    sheet.cell(5, 7).value = 99999
    sheet.cell(5, 42).value = -99999
    book.save(path)
    return path


def test_support_snapshot_is_not_written_into_computed_fields(tmp_path):
    service = PayrollService(tmp_path / "data")
    run = _run(service, ["教师甲"], tmp_path=tmp_path)
    source = _support_workbook(tmp_path / "支持部.xlsx", [
        _teacher("教师甲", 该档每小时金额=9999, 总课时费=88888),
    ])
    service.import_support_department(run["id"], str(source), _sha256(source), "核算负责人")

    stored = service.store.get(run["id"])
    generated = service._generated_from_checked(stored)
    fields = generated.rows[0].final_fields
    # AE / AF come from this Run's own calculation, never from the workbook.
    assert fields["AF"]["value"] != 88888
    assert "9999" not in json.dumps(fields.get("AE") or {}, ensure_ascii=False)


# ------------------------- 结果表：扩展名不能决定文件真实格式（P0-5 根因）
def test_a_legacy_workbook_behind_an_xlsx_name_is_read_by_content(tmp_path, monkeypatch):
    """The office saves legacy .xls as .xlsx; the final-result reader must cope.

    Without this, the refund material could never become an approved result and
    AN stayed empty no matter how many times the operator uploaded the file.
    """
    from payroll_core.adapters import business_results

    fake = tmp_path / "退费.xlsx"
    fake.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 32)
    assert business_results._is_zip_container(fake) is False

    header = ["周", "退费校区", "新签/续费", "姓名", "年级", "消耗课时", "退费科目", "科目类型",
              "退费课时", "退费总金额", "是否为转校生", "班主任", "学科教师", "教育顾问", "备注", "教师", "人头", "业绩"]
    sheets = [
        ("1月份 ", [["1月份退费"], header, [1, "宣城二校", "续费", "学员甲", "初二", 1, "数学", "班课", 1, 1, "否", "甲", "教师甲", "甲", "", "教师甲", 11, 11]]),
        ("8月份 ", [["8月份退费"], header, [1, "宣城二校", "续费", "学员乙", "初二", 1, "数学", "班课", 1, 1, "否", "甲", "教师甲", "甲", "", "教师甲", 100, 200]]),
    ]
    monkeypatch.setattr(business_results, "_is_zip_container", lambda path: False)
    monkeypatch.setattr(business_results, "_legacy_sheets", lambda path: sheets)

    rows = business_results.read_business_result(fake, period="2026-08")
    # Only the August sheet, and the header row is found below the title row.
    assert len(rows) == 1
    assert rows[0].evidence["sheet"] == "8月份 "
    assert rows[0].payload["人头"] == 100
    assert rows[0].payload["业绩"] == 200
    assert rows[0].payload["teacher"] == "教师甲"


def test_the_header_row_is_located_below_a_title_row():
    from payroll_core.adapters.business_results import _header_row_index

    assert _header_row_index([["退费统计表"], ["周", "姓名", "教师", "人头", "业绩"], [1, "甲", "教师甲", 1, 1]]) == 1
    assert _header_row_index([["教师", "人头", "业绩"], [1, 1, 1]]) == 0


# --------------------------- 导出：F 教师级别与模板工作表集合
def test_the_teacher_level_column_comes_from_the_support_workbook(tmp_path):
    from payroll_core.excel.standard_payroll_render import render_generated_payroll

    service = PayrollService(tmp_path / "data")
    run = _run(service, ["教师甲"], tmp_path=tmp_path)
    source = _support_workbook(tmp_path / "支持部.xlsx", [_teacher("教师甲", 教师级别="TRST/三星级")])
    service.import_support_department(run["id"], str(source), _sha256(source), "核算负责人")

    stored = service.store.get(run["id"])
    generated = service._generated_from_checked(stored)
    output = tmp_path / "导出.xlsx"
    render_generated_payroll(generated, output, template_path=_payroll_template(tmp_path / "模板.xlsx"),
                             support_snapshot=stored["support_department_snapshot"])
    sheet = load_workbook(output).worksheets[0]
    assert sheet.cell(5, 6).value == "TRST/三星级"


def test_the_export_does_not_add_duplicate_evidence_sheets(tmp_path):
    """The approved template already ships 核验与来源 / 外围字段状态."""
    from payroll_core.excel.standard_payroll_render import render_generated_payroll

    service = PayrollService(tmp_path / "data")
    run = _run(service, ["教师甲"], tmp_path=tmp_path)
    source = _support_workbook(tmp_path / "支持部.xlsx", [_teacher("教师甲")])
    service.import_support_department(run["id"], str(source), _sha256(source), "核算负责人")

    template = _payroll_template(tmp_path / "模板.xlsx")
    book = load_workbook(template)
    for title in ("核验与来源", "外围字段状态"):
        book.create_sheet(title)
    book.save(template)

    stored = service.store.get(run["id"])
    generated = service._generated_from_checked(stored)
    output = tmp_path / "导出.xlsx"
    render_generated_payroll(generated, output, template_path=template,
                             support_snapshot=stored["support_department_snapshot"])
    names = load_workbook(output).sheetnames
    assert "核验与来源" in names and "外围字段状态" in names
    assert "核验与来源1" not in names and "外围字段状态1" not in names, names
