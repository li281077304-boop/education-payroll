"""Regression boundaries for the submission-first monthly payroll flow.

Every test uses a private, temporary PayrollService store and synthetic source
files. The existing legacy Run tests continue to cover the older workflow.
"""

from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import load_workbook
import pytest

from payroll_core.excel.standard_payroll_render import render_generated_payroll
from payroll_core.final_fields import DETERMINED, renewal_fields_from_snapshot
from payroll_core.payroll_generation import generated_from_calculation
from payroll_ui.service import MONTHLY_FLOW_SUBMISSION_FIRST, PayrollService
from tests.test_standard_payroll_output import _sanitized_template
from tests.test_teacher_identity_and_support_source import _sha256, _support_workbook, _teacher


def _schedule(path: Path, teachers: list[str]) -> Path:
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["teacher_id", "teacher", "grade", "subject", "class_type", "attended", "lesson_status", "time"])
        for index, teacher in enumerate(teachers):
            writer.writerow([f"teacher-{index:02d}", teacher, "九年级", "数学", "1对1", 1, "已上课", "2026-08-05 10:00"])
    return path


def _submission(path: Path, teachers: list[str]) -> Path:
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["teacher", "one_to_one", "class_value", "production", "ae", "af", "av"])
        for teacher in teachers:
            writer.writerow([teacher, 1, 0, 1, 30, 40, 0])
    return path


def _submission_with_fields(path: Path, rows: list[dict]) -> Path:
    headers = ["teacher", "one_to_one", "class_value", "production", "ae", "af", "av",
               "b", "d", "e", "f", "g", "h", "i", "j", "k", "l"]
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=headers)
        writer.writeheader()
        for row in rows:
            writer.writerow({"one_to_one": 1, "class_value": 0, "production": 1,
                             "ae": 30, "af": 40, "av": 0, **row})
    return path


def _salary_row(name: str, *, group: str = "数学组", base: int = 1800) -> dict:
    return {"teacher": name, "b": group, "d": f"{name}@example.com",
            "e": "2021.01.01", "f": "TR/三星级", "g": base,
            "h": 10, "i": 20, "j": 30, "k": 26, "l": 26}


def _run(tmp_path: Path, *, teachers: list[str]) -> tuple[PayrollService, str]:
    service = PayrollService(tmp_path / "private-payroll-store")
    run = service.create("2026-08", "GENERATE", operator_role="DOS", monthly_flow_version=MONTHLY_FLOW_SUBMISSION_FIRST)
    assert run["monthly_flow_version"] == "SUBMISSION_FIRST_V1"
    service.import_file(run["id"], "schedule", str(_schedule(tmp_path / "schedule.csv", teachers)))
    service.confirm_af_policy(run["id"], "UAT负责人")
    return service, run["id"]


def _core(teacher: str) -> dict:
    return {
        "period": "2026-08",
        "rows": [{
            "teacher": teacher,
            "employment_type": "FULL_TIME",
            "fields": {
                "AA": {"value": 1, "state": DETERMINED},
                "AC": {"value": 0, "state": DETERMINED},
                "AD": {"value": 1, "state": DETERMINED},
                "AE": {"value": 40, "state": DETERMINED},
                "AF": {"value": 40, "state": DETERMINED},
                "PART_TIME": {"value": None, "state": "NOT_APPLICABLE"},
            },
        }],
    }


def test_submission_names_define_exact_output_population_even_with_51_schedule_teachers(tmp_path: Path) -> None:
    teachers = [f"教师{index:02d}" for index in range(51)]
    service, run_id = _run(tmp_path, teachers=teachers)
    source = _submission(tmp_path / "group-submission.csv", teachers[:15])

    imported = service.import_material_file(run_id, "subject_group", str(source))
    assert imported["run"]["subject_group_materials"]
    checked = service.check(run_id)
    assert checked["processing_scope_snapshot"]["canonical_roster_count"] == 51
    assert set(checked["processing_scope_snapshot"]["teachers"]) == set(teachers[:15])
    assert {row["teacher"] for row in checked["core_calculation"]["rows"]} == set(teachers[:15])

    preview = service.preview_payroll(run_id)
    assert {row["teacher"] for row in preview["generated_payroll"]["rows"]} == set(teachers[:15])


def test_submission_teacher_without_period_schedule_keeps_payroll_row(tmp_path: Path) -> None:
    service, run_id = _run(tmp_path, teachers=["有排课教师"])
    source = _submission_with_fields(tmp_path / "group-submission.csv", [_salary_row("无本月排课教师")])

    service.import_material_file(run_id, "subject_group", str(source))
    checked = service.check(run_id)

    assert checked["processing_scope_snapshot"]["teachers"] == ["无本月排课教师"]
    assert {row["teacher"] for row in checked["core_calculation"]["rows"]} == {"无本月排课教师"}


def test_without_submission_full_schedule_population_remains_and_missing_salary_is_allowed(tmp_path: Path) -> None:
    teachers = [f"教师{index:02d}" for index in range(51)]
    service, run_id = _run(tmp_path, teachers=teachers)

    checked = service.check(run_id)
    assert set(checked["processing_scope_snapshot"]["teachers"]) == set(teachers)
    assert {row["teacher"] for row in checked["core_calculation"]["rows"]} == set(teachers)
    assert not checked.get("base_salary_inputs")
    assert not any(
        item.get("field") in {"base_salary", "salary_basis", "employment"}
        for item in checked.get("field_records", [])
    )

    preview = service.preview_payroll(run_id)
    rows = preview["generated_payroll"]["rows"]
    assert len(rows) == 51
    assert all(row["final_fields"]["M"]["value"] is None for row in rows)

    output = render_generated_payroll(
        service._generated_from_checked(checked),
        tmp_path / "generated-without-group-submission.xlsx",
        template_path=_sanitized_template(tmp_path),
    )
    sheet = load_workbook(output).active
    assert {sheet.cell(row, 3).value for row in range(5, 56)} == set(teachers)
    nonblank = [(sheet.cell(row, 3).value, sheet.cell(row, column).coordinate, sheet.cell(row, column).value)
                for row in range(5, 56) for column in range(2, 13)
                if column != 3 and sheet.cell(row, column).value is not None]
    assert not nonblank, nonblank[:12]


def test_confirmed_renewal_snapshot_without_teacher_row_sets_ah_ai_aj_ak_to_zero() -> None:
    snapshot = {
        "version": "RUN_RENEWAL_RESULT_SNAPSHOT/v1",
        "run_id": "synthetic-run",
        "period": "2026-08",
        "confirmed_by": "UAT审核员",
        "entries": {
            "其他教师": {
                "teacher_id": "other-id",
                "display_name": "其他教师",
                "AH": {"value": 2, "state": DETERMINED},
                "AI": {"value": 1, "state": DETERMINED},
                "AJ": {"value": 0, "state": DETERMINED},
            }
        },
    }
    no_row = renewal_fields_from_snapshot(snapshot, "本月无续费教师")
    assert all(no_row[code]["value"] == 0 and no_row[code]["state"] == DETERMINED for code in ("AH", "AI", "AJ"))

    generated = generated_from_calculation(_core("本月无续费教师"), renewal_snapshot=snapshot)
    final = generated.rows[0].final_fields
    assert all(final[code]["value"] == 0 and final[code]["state"] == DETERMINED for code in ("AH", "AI", "AJ", "AK"))


def test_new_monthly_flow_never_creates_salary_basis_or_employment_todos(tmp_path: Path) -> None:
    service, run_id = _run(tmp_path, teachers=["无底薪资料教师"])
    checked = service.check(run_id)
    forbidden = {"base_salary", "salary_basis", "employment", "hourly_submission_result"}
    forbidden_records = [item for item in checked.get("field_records", []) if item.get("field") in forbidden]
    assert not forbidden_records, forbidden_records[:3]
    assert not any(item.get("field") in forbidden for item in checked.get("issues", []))
    assert not any(item.get("field") in forbidden for item in checked.get("issue_groups", []))


def test_af_default_can_confirm_without_exceptions_and_checked_exception_requires_reason(tmp_path: Path) -> None:
    service, run_id = _run(tmp_path, teachers=["教师甲", "教师乙"])
    with pytest.raises(ValueError, match="例外原因"):
        service.confirm_af_policy(run_id, "UAT负责人", exceptions={"教师甲": {"obligation_hours": 0, "reason": ""}})

    rendered = service.confirm_af_policy(run_id, "UAT负责人", exceptions={"教师甲": {"obligation_hours": 0, "reason": "MT 带组且无星级"}})
    confirmation = rendered["af_policy_confirmation"]
    assert confirmation["default_obligation_hours"] == 30
    assert confirmation["exceptions"]["教师甲"]["obligation_hours"] == 0
    assert confirmation["exceptions"]["教师甲"]["reason"] == "MT 带组且无星级"
    assert "教师乙" not in confirmation["exceptions"]


def test_group_submission_b_through_l_become_current_month_values_without_support(tmp_path: Path) -> None:
    service, run_id = _run(tmp_path, teachers=["教师甲"])
    source = _submission_with_fields(tmp_path / "group.csv", [_salary_row("教师甲")])
    service.import_material_file(run_id, "subject_group", str(source))

    checked = service.check(run_id)
    fields = checked["payroll_input_fields_snapshot"]["entries"]["教师甲"]
    for code, expected in {"B": "数学组", "D": "教师甲@example.com", "E": "2021.01.01",
                           "F": "TR/三星级", "G": "1800", "H": "10", "I": "20",
                           "J": "30", "K": "26", "L": "26"}.items():
        assert str(fields[code]["value"]) == expected
        assert fields[code]["source_type"] == "SUBJECT_GROUP_SUBMISSION"
        assert fields[code]["source_file"] == source.name

    output = render_generated_payroll(
        service._generated_from_checked(checked), tmp_path / "group-values.xlsx",
        template_path=_sanitized_template(tmp_path),
    )
    sheet = load_workbook(output).active
    assert sheet["C5"].value == "教师甲"
    assert sheet["B5"].value == "数学组"
    assert sheet["G5"].value == 1800
    assert sheet["L5"].value == 26


def test_support_authority_overrides_group_submission_across_b_through_l(tmp_path: Path) -> None:
    service, run_id = _run(tmp_path, teachers=["教师甲"])
    group = _submission_with_fields(tmp_path / "group.csv", [_salary_row("教师甲")])
    service.import_material_file(run_id, "subject_group", str(group))
    support = _support_workbook(tmp_path / "support.xlsx", [
        _teacher("教师甲", 科组="理化组", 邮箱="authority@example.com",
                 入职日期="2020.02.02", 教师级别="TR/四星级",
                 基本工资=2400, 岗位津贴=100, **{"工龄工资/教师等级": 200},
                 其他待遇=300, 应出勤=25, 实际出勤=24),
    ])
    service.import_support_department(run_id, str(support), _sha256(support), "UAT负责人")

    checked = service.check(run_id)
    fields = checked["payroll_input_fields_snapshot"]["entries"]["教师甲"]
    expected = {"B": "理化组", "D": "authority@example.com", "E": "2020.02.02",
                "F": "TR/四星级", "G": 2400, "H": 100, "I": 200,
                "J": 300, "K": 25, "L": 24}
    for code, value in expected.items():
        actual = fields[code]["value"]
        if isinstance(value, (int, float)):
            assert float(actual) == float(value)
        else:
            assert str(actual) == str(value)
        assert fields[code]["source_type"] == "SUPPORT_AUTHORITY"

    output = render_generated_payroll(
        service._generated_from_checked(checked), tmp_path / "support-overrides.xlsx",
        template_path=_sanitized_template(tmp_path),
        support_snapshot=checked["support_department_snapshot"],
    )
    sheet = load_workbook(output).active
    assert [sheet[f"{code}5"].value for code in ("B", "D", "E", "F", "G", "H", "I", "J", "K", "L")] == list(expected.values())


def test_two_group_files_cross_group_merge_without_group_choice_or_confirmation(tmp_path: Path) -> None:
    service, run_id = _run(tmp_path, teachers=["数学教师", "理化教师", "跨组教师"])
    first = _submission_with_fields(tmp_path / "math.csv", [
        _salary_row("数学教师", group="数学组"), _salary_row("跨组教师", group="数学组"),
    ])
    second = _submission_with_fields(tmp_path / "science.csv", [
        _salary_row("理化教师", group="理化组"), _salary_row("跨组教师", group="数学组"),
    ])
    first_import = service.import_material_file(run_id, "subject_group", str(first))
    second_import = service.import_material_file(run_id, "subject_group", str(second))
    assert len(first_import["run"]["subject_group_materials"]) == 1
    assert len(second_import["run"]["subject_group_materials"]) == 2
    assert not second_import["run"].get("pending_subject_group_imports")

    checked = service.check(run_id)
    assert set(checked["processing_scope_snapshot"]["teachers"]) == {"数学教师", "理化教师", "跨组教师"}
    assert len(checked["core_calculation"]["rows"]) == 3
    assert not checked["subject_group_source_conflicts"]


def test_conflicting_same_teacher_same_field_creates_one_actionable_conflict(tmp_path: Path) -> None:
    service, run_id = _run(tmp_path, teachers=["教师甲"])
    first = _submission_with_fields(tmp_path / "original.csv", [_salary_row("教师甲", base=1800)])
    second = _submission_with_fields(tmp_path / "revision.csv", [_salary_row("教师甲", base=1900)])
    service.import_material_file(run_id, "subject_group", str(first))
    service.import_material_file(run_id, "subject_group", str(second))

    checked = service.check(run_id)
    conflicts = [item for item in checked["subject_group_source_conflicts"] if item["teacher"] == "教师甲"]
    assert len(conflicts) == 1
    assert conflicts[0]["field"] == "G"
    assert set(str(value) for value in conflicts[0]["values"]) == {"1800", "1900"}
    assert checked["payroll_input_fields_snapshot"]["entries"]["教师甲"]["G"]["value"] is None
    assert any(item.get("teacher") == "教师甲" and item.get("status") == "CONFLICT_NEEDS_CONFIRMATION"
               for item in checked.get("field_records", []))
