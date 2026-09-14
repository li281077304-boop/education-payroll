"""Artificial files exercise the public calculation/service/HTTP chain."""
import copy
import json
from pathlib import Path
from threading import Thread
from urllib.request import Request, urlopen

import pytest
from openpyxl import load_workbook

from payroll_ui.service import PayrollService
from payroll_ui.server import PayrollHttpServer
from tests.test_payroll_modes_and_class_rules import _schedule
from tests.test_payroll_ui_service import _payroll_with_only


def prepared(tmp_path, mode="AUDIT"):
    service = PayrollService(tmp_path / "local-data")
    service.save_rating_version("2026-08", "2026-09", "人工构造星级", "test-stars", [{"teacher": "张三", "rating": 4}])
    service.save_policy_version("2026-08", "2026-09", "人工构造个人政策", [{"teacher": "张三", "role": "TRMT", "rating": 4, "obligation_hours": 30, "obligation_hours_deduction_enabled": True}])
    run = service.create("2026-08", mode)
    source = _schedule(tmp_path / "脱敏排课.xlsx", [["张三", "九年级", "数学", "1对1", 1, "已上课"]] * 20 + [["张三", "九年级", "数学", "小班", 3, "已上课"]])
    service.import_file(run["id"], "schedule", str(source))
    target = tmp_path / "脱敏工资.xlsx"
    if mode == "AUDIT":
        _payroll_with_only(target, 5)
        book = load_workbook(target)
        sheet = book.active
        for col, value in {"AA": 40, "AC": 2.4, "AD": 999, "AE": 40, "AF": 496}.items():
            sheet[f"{col}5"] = value
        book.save(target)
        service.import_file(run["id"], "math", str(target))
    return service, run, source, target


def test_submitted_ad_never_changes_independent_chain(tmp_path):
    service, run, _, target = prepared(tmp_path)
    checked = service.check(run["id"])
    fields = checked["core_calculation"]["rows"][0]["fields"]
    assert [fields[f]["value"] for f in ("AA", "AC", "AD", "AE", "AF")] == [40, 2.4, 42.4, 40, 496]
    assert all(fields[f]["state"] == "DETERMINED" for f in ("AA", "AC", "AD", "AE", "AF"))
    ad_check = next(r for r in checked["field_records"] if r["field"] == "teaching_hours")
    assert ad_check["expected"] == 42.4 and ad_check["actual"] == 999
    book = load_workbook(target)
    book.active["AD5"] = 2
    book.save(target)
    service.import_file(run["id"], "math", str(target))
    after = service.check(run["id"])
    assert after["core_calculation"]["rows"][0]["fields"] == fields


def test_generate_and_audit_share_all_five_fields(tmp_path):
    service, audit, source, _ = prepared(tmp_path)
    checked = service.check(audit["id"])
    generate = service.create("2026-08", "GENERATE")
    service.import_file(generate["id"], "schedule", str(source))
    made = service.generate_payroll(generate["id"], str(tmp_path / "generated.xlsx"))
    assert made["rows"][0]["fields"] == checked["core_calculation"]["rows"][0]["fields"]
    with pytest.raises(ValueError, match="AD"):
        service.generate_payroll(generate["id"], str(tmp_path / "forbidden.xlsx"), confirmed_hours={"张三": 999})


def test_rule_snapshot_preserved_and_explicit_rebind_invalidates_decision(tmp_path):
    service, run, _, _ = prepared(tmp_path)
    checked = service.check(run["id"])
    group = next(g for g in checked["issue_groups"] if "teaching_hours" in g["affected_fields"])
    service.decide(run["id"], group["id"], "DEFERRED", "测试审核员", "测试待核实", expected_fingerprint=group["fingerprint"])
    before = service.core_rule_catalog()["versions"][0]
    snapshot = copy.deepcopy(before)
    rules = copy.deepcopy(before["rules"])
    rules["ae"]["star_bonuses"]["4"] = "11"
    versions = service.save_core_rule_version(rules, "人工构造修正版", "测试审核员")["versions"]
    corrected = next(v for v in versions if v["id"] != before["id"])
    assert next(v for v in versions if v["id"] == before["id"]) == snapshot
    stable = service.check(run["id"])
    assert stable["core_calculation"]["rows"][0]["fields"]["AE"]["value"] == 40
    rebound = service.rebind_calculation(run["id"], "core", corrected["id"])
    assert rebound["core_calculation"]["rows"][0]["fields"]["AE"]["value"] == 41
    assert rebound["business_decisions"][0]["status"] == "NEEDS_RECONFIRMATION"
    restored = PayrollService(tmp_path / "local-data").get(run["id"])
    assert restored["core_rule_version_id"] == corrected["id"]
    assert restored["business_decisions"][0]["status"] == "NEEDS_RECONFIRMATION"
    # Ambiguous new run must ask for a version, not pick the newest silently.
    assert service.create("2026-08")["core_rule_version_id"] is None


def test_future_rule_does_not_change_old_run(tmp_path):
    service, run, _, _ = prepared(tmp_path)
    before = service.check(run["id"])["core_calculation"]
    rules = copy.deepcopy(service.core_rule_catalog()["seed"])
    rules.update(effective_from="2026-10", effective_to="2027-09")
    rules["grade_coefficients"]["九年级"] = "1.1"
    service.save_core_rule_version(rules, "假定未来测试规则", "测试审核员")
    assert service.check(run["id"])["core_calculation"] == before
    assert service.create("2026-10")["core_rule_version_id"] != run["core_rule_version_id"]


def test_new_evidence_month_can_keep_semantic_rule_id(tmp_path):
    service = PayrollService(tmp_path / "local-data")
    rules = copy.deepcopy(service.core_rule_catalog()["seed"])
    rules.update(rule_version_id="core_rules_2026_07_v1", effective_from="2026-07", effective_to="2026-07")
    versions = service.save_core_rule_version(rules, "historical July payroll reconstruction", "evidence review")["versions"]
    july = next(item for item in versions if item["id"] == "core_rules_2026_07_v1")
    assert july["rules"]["rule_version_id"] == "core_rules_2026_07_v1"


def test_sanitized_http_generate_chain_and_restore(tmp_path):
    service, run, _, _ = prepared(tmp_path, "GENERATE")
    server = PayrollHttpServer(("127.0.0.1", 0), service, Path(__file__).parents[1] / "payroll_ui" / "static")
    thread = Thread(target=server.serve_forever, daemon=True)
    thread.start()
    def request(route, data=None):
        req = Request(f"http://127.0.0.1:{server.server_port}" + route, data=None if data is None else json.dumps(data).encode(), headers={"X-Payroll-Token": server.token, "Content-Type": "application/json"})
        return json.loads(urlopen(req).read())
    try:
        assert request("/api/core-rules")["versions"]
        checked = request(f"/api/runs/{run['id']}/check", {})
        assert checked["core_calculation"]["rows"][0]["fields"]["AF"]["value"] == 496
        generated = request(f"/api/runs/{run['id']}/generate", {"output_path": str(tmp_path / "HTTP生成.xlsx")})
        assert generated["rows"][0]["teaching_hours"] == 42.4
    finally:
        server.shutdown(); server.server_close(); thread.join()
    restored = PayrollService(tmp_path / "local-data").get(run["id"])
    assert restored["generated_payroll"]["rows"][0]["af"] == 496


def test_personal_policy_and_part_time_rates_survive_reopen(tmp_path):
    service = PayrollService(tmp_path / "app")
    service.save_policy_version("2026-08", "2026-09", "测试岗位确认", [
        {"teacher": "兼职甲", "role": "教师", "employment_type": "PART_TIME", "obligation_hours": 0, "obligation_hours_deduction_enabled": False},
        {"teacher": "管理乙", "role": "组长", "allow_no_teaching": True, "obligation_hours": 0, "obligation_hours_deduction_enabled": False},
    ])
    service.save_part_time_rate_version([{"teacher": "兼职甲", "grade_scope": "*", "rate_per_session": 123}], "假单价来源", "2026-08", "2026-09", "测试审核员")
    run = service.create("2026-08", "GENERATE")
    source = _schedule(tmp_path / "兼职排课.xlsx", [["兼职甲", "九年级", "数学", "小班", 6, "已上课"]])
    service.import_file(run["id"], "schedule", str(source))
    result = service.check(run["id"])
    by_teacher = {r["teacher"]: r["fields"] for r in result["core_calculation"]["rows"]}
    assert by_teacher["兼职甲"]["PART_TIME"]["value"] == 123
    assert by_teacher["兼职甲"]["AE"]["state"] == "NOT_APPLICABLE"
    assert by_teacher["管理乙"]["AD"]["value"] == 0
    assert by_teacher["管理乙"]["AF"]["state"] == "NOT_APPLICABLE"
    assert PayrollService(tmp_path / "app").get(run["id"])["part_time_rate_version_id"] == run["part_time_rate_version_id"]
